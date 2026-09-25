#!/usr/bin/env python3
from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
import sys
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime, timezone
from difflib import SequenceMatcher
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote, urlsplit
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from common_xml import local_name
from local_evidence_output import guard_repo_safe_output
from quiz_normalization import (
    build_normalized_model,
    build_run_receipt,
    copy_resolved_assets,
    enrich_payload,
    resolve_export,
    source_choice_projection,
    utc_now,
)
from quiz_review_projection import build_quiz_review_projection
from quiz_binder_strings import match_label


SIMILARITY_THRESHOLD = 0.95
SIMILARITY_GAP = 0.03
WORKBENCH_ROOT = Path(__file__).resolve().parents[1]
HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Open Sans", color="FFFFFF", bold=True)
WRAP_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
IMAGE_REF_PATTERN = re.compile(r"\.(?:png|jpe?g|gif|svg|bmp|webp)(?:$|[?#])", flags=re.IGNORECASE)
UNRESOLVED_SHEET_TITLE = "Unresolved Pool Match"
REVIEWER_UNRESOLVED_SHEET_TITLE = "No Safe Library Match"
REVIEWER_ALL_QUESTIONS_SHEET_TITLE = "All Questions"
UNRESOLVED_SHEET_NOTE = (
    "This sheet is the exception queue for question-library / pool matching.\n\n"
    "A row appears here when the script extracted the quiz item successfully but could not make a unique, defensible match back to questiondb.xml.\n\n"
    "It does not mean the question is broken or that extraction failed. The item may still have usable text, answer-key material, IDs, and image references.\n\n"
    "Matching is attempted in this order: qmd_displayid, qmd_globalid, exact question signature (type + prompt + choices), exact question text, then similarity.\n\n"
    "Common reasons a row lands here: the item is quiz-local, imported from another shell, duplicated in the library, or only has a weak / non-unique text match.\n\n"
    "match_score is only the best similarity candidate. A score of 1.000 can still be unresolved if the best match is not unique enough to trust.\n\n"
    "Use this sheet to review uncertain bank provenance, not to judge question quality."
)
REVIEWER_STATUS_FILLS = {
    "Direct library link": PatternFill(fill_type="solid", fgColor="E2F0D9"),
    "Linked by identical content": PatternFill(fill_type="solid", fgColor="D9EAF7"),
    "Possible link based on similar content": PatternFill(fill_type="solid", fgColor="FFF2CC"),
    "No reliable library link found": PatternFill(fill_type="solid", fgColor="FCE4D6"),
    "Used in quiz — direct library link": PatternFill(fill_type="solid", fgColor="E2F0D9"),
    "Used in quiz — linked by identical content": PatternFill(fill_type="solid", fgColor="D9EAF7"),
    "Possible quiz use — similar content; review recommended": PatternFill(fill_type="solid", fgColor="FFF2CC"),
    "Library only — not used in exported quizzes": PatternFill(fill_type="solid", fgColor="E7E6E6"),
    "Quiz only — no library link found": PatternFill(fill_type="solid", fgColor="FCE4D6"),
    "resolved": PatternFill(fill_type="solid", fgColor="E2F0D9"),
    "partially resolved": PatternFill(fill_type="solid", fgColor="FFF2CC"),
    "missing": PatternFill(fill_type="solid", fgColor="FCE4D6"),
    "none": PatternFill(fill_type="solid", fgColor="E7E6E6"),
}
REVIEWER_BAND_FILL = PatternFill(fill_type="solid", fgColor="F5F8FB")
REVIEWER_EDITABLE_FILL = PatternFill(fill_type="solid", fgColor="FFF7D6")
REVIEWER_EDITABLE_HEADERS = {
    "proposed_permanent_code",
    "proposed_scoring_mode",
    "revised_question_text",
    "revised_response_options",
    "revised_answer_key",
    "revised_question_library_location",
    "revision_reason",
    "proposed_by",
    "proposed_at",
    "approval_status",
    "approved_by",
    "approved_at",
    "reviewer_note",
}
REVIEWER_SETTINGS_EDITABLE_HEADERS = {
    "proposed_value",
    "revision_reason",
    "proposed_by",
    "proposed_at",
    "approval_status",
    "approved_by",
    "approved_at",
    "reviewer_note",
}
REVIEWER_ENTITY_IDENTIFIER_GUIDANCE = {
    "display_identifier": (
        "Reviewer-facing identifier chosen by CourseCraft for display. It uses "
        "permanent_code when one exists; otherwise it uses the first preserved "
        "source alias, then falls back to question_entity_key. The displayed value "
        "may come from the package, but this selection field is part of the "
        "CourseCraft review schema."
    ),
    "question_entity_key": (
        "CourseCraft-generated normalized key for one distinct question entity. "
        "It is not a package-native D2L field. It connects this entity to its rows "
        "in Quiz Occurrences and remains stable across refreshed exports when the "
        "course lineage and selected source identity remain stable."
    ),
    "source_aliases": (
        "Package-native identifiers preserved from the export and displayed as "
        "namespace=value pairs. CourseCraft adds the namespaces and formatting so "
        "D2L global, display, quiz-item, and library identifiers remain distinct; "
        "the identifier values themselves come from the package."
    ),
}
REVIEWER_SETTINGS_HEADERS = [
    "quiz",
    "target_entity_key",
    "setting_observation_key",
    "setting",
    "observed_state",
    "observed_value",
    "proposed_value",
    "revision_reason",
    "proposed_by",
    "proposed_at",
    "approval_status",
    "approved_by",
    "approved_at",
    "reviewer_note",
]


@dataclass
class LibraryQuestion:
    bank_ident: str
    bank_title: str
    bank_path: str
    pool_ident: str
    pool_title: str
    pool_path: str
    question_ident: str
    question_label: str
    question_title: str
    question_type: str
    question_weight: str
    qmd_displayid: str
    qmd_globalid: str
    raw_html: str
    plain_text: str
    exact_text_key: str
    similarity_text_key: str
    choices_text: str
    choice_count: int
    signature_key: str
    item_element: ET.Element


@dataclass
class MatchResult:
    library_question: LibraryQuestion | None
    evidence_level: str
    match_basis: str
    match_score: float | None
    note: str
    candidates: list[dict[str, object]] = field(default_factory=list)


@dataclass
class RenderedHtml:
    text: str
    formula_present: bool
    mathml_used: bool


@dataclass
class AssetReference:
    raw_ref: str
    normalized_ref: str
    status: str
    resolution_basis: str
    file_path: str
    absolute_path: str


SUPERSCRIPT_MAP = str.maketrans({
    "0": "⁰",
    "1": "¹",
    "2": "²",
    "3": "³",
    "4": "⁴",
    "5": "⁵",
    "6": "⁶",
    "7": "⁷",
    "8": "⁸",
    "9": "⁹",
    "+": "⁺",
    "-": "⁻",
    "=": "⁼",
    "(": "⁽",
    ")": "⁾",
    "n": "ⁿ",
    "i": "ⁱ",
})
SUBSCRIPT_MAP = str.maketrans({
    "0": "₀",
    "1": "₁",
    "2": "₂",
    "3": "₃",
    "4": "₄",
    "5": "₅",
    "6": "₆",
    "7": "₇",
    "8": "₈",
    "9": "₉",
    "+": "₊",
    "-": "₋",
    "=": "₌",
    "(": "₍",
    ")": "₎",
    "a": "ₐ",
    "e": "ₑ",
    "h": "ₕ",
    "i": "ᵢ",
    "j": "ⱼ",
    "k": "ₖ",
    "l": "ₗ",
    "m": "ₘ",
    "n": "ₙ",
    "o": "ₒ",
    "p": "ₚ",
    "r": "ᵣ",
    "s": "ₛ",
    "t": "ₜ",
    "u": "ᵤ",
    "v": "ᵥ",
    "x": "ₓ",
})
VERBAL_MATH_REPLACEMENTS = [
    ("less than or equal to", "≤"),
    ("greater than or equal to", "≥"),
    ("not equal to", "≠"),
    ("cross times", "×"),
    ("multiplied by", "×"),
    ("divided by", "÷"),
    ("open parentheses", "("),
    ("close parentheses", ")"),
    ("open parenthesis", "("),
    ("close parenthesis", ")"),
    ("left parenthesis", "("),
    ("right parenthesis", ")"),
    ("open bracket", "["),
    ("close bracket", "]"),
    ("open braces", "{"),
    ("close braces", "}"),
    ("plus or minus", "±"),
    ("minus or plus", "∓"),
    ("plus", "+"),
    ("minus", "-"),
    ("equals", "="),
    ("times", "×"),
]
ORDINAL_DENOMINATORS = {
    "half": "2",
    "halves": "2",
    "third": "3",
    "thirds": "3",
    "fourth": "4",
    "fourths": "4",
    "fifth": "5",
    "fifths": "5",
    "sixth": "6",
    "sixths": "6",
    "seventh": "7",
    "sevenths": "7",
    "eighth": "8",
    "eighths": "8",
    "ninth": "9",
    "ninths": "9",
    "tenth": "10",
    "tenths": "10",
}
NUMERATOR_WORDS = {
    "one": "1",
    "two": "2",
    "three": "3",
    "four": "4",
    "five": "5",
    "six": "6",
    "seven": "7",
    "eight": "8",
    "nine": "9",
}


def maybe_translate_script(text: str, table: dict[int, str]) -> str | None:
    if text and all(ord(char) in table or not char.isascii() for char in text):
        return text.translate(table)
    return None


def bracket_if_needed(text: str) -> str:
    cleaned = text.strip()
    if not cleaned:
        return cleaned
    if re.fullmatch(r"[A-Za-z0-9_.]+", cleaned):
        return cleaned
    if cleaned.startswith("(") and cleaned.endswith(")"):
        return cleaned
    return f"({cleaned})"


def clean_inline_text(text: str) -> str:
    cleaned = html.unescape(text).replace("\xa0", " ")
    cleaned = re.sub(r"\s+", " ", cleaned)
    cleaned = re.sub(r"\s+([,.;:!?])", r"\1", cleaned)
    cleaned = re.sub(r"([(\[{])\s+", r"\1", cleaned)
    cleaned = re.sub(r"\s+([)\]}])", r"\1", cleaned)
    cleaned = re.sub(r"\s*([=+\-×÷≤≥≠±∓])\s*", r" \1 ", cleaned)
    cleaned = re.sub(r"(?:(?<=^)|(?<=[(=+\-×÷≤≥≠±∓\[]))\s*-\s+([A-Za-z0-9.])", r"-\1", cleaned)
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned.strip()


def normalize_math_operator(text: str) -> str:
    operator = clean_inline_text(text)
    replacements = {
        "*": "×",
        "·": "·",
        "−": "-",
        "–": "-",
        "—": "-",
    }
    return replacements.get(operator, operator)


def decode_d2l_mathml(raw_mathml: str) -> str:
    decoded = raw_mathml
    decoded = decoded.replace("«", "<").replace("»", ">").replace("¨", '"').replace("§#", "&#")
    decoded = decoded.replace("§lt;", "&lt;").replace("§gt;", "&gt;").replace("§amp;", "&amp;")
    return html.unescape(decoded)


def render_mathml_element(elem: ET.Element | None) -> str:
    if elem is None:
        return ""

    tag = local_name(elem.tag)
    children = list(elem)

    if tag in {"math", "mrow", "mstyle", "mpadded", "menclose", "mphantom"}:
        return clean_inline_text(" ".join(part for part in (render_mathml_element(child) for child in children) if part))
    if tag in {"semantics"}:
        for child in children:
            if local_name(child.tag) not in {"annotation", "annotation-xml"}:
                return render_mathml_element(child)
        return ""
    if tag in {"annotation", "annotation-xml"}:
        return ""
    if tag in {"mn", "mi", "mtext", "ms"}:
        return clean_inline_text("".join(elem.itertext()))
    if tag == "mo":
        return normalize_math_operator("".join(elem.itertext()))
    if tag == "mfenced":
        open_char = elem.attrib.get("open", "(")
        close_char = elem.attrib.get("close", ")")
        separators = elem.attrib.get("separators", ",")
        separator = f"{separators[0]} " if separators else ", "
        inner = separator.join(part for part in (render_mathml_element(child) for child in children) if part)
        return clean_inline_text(f"{open_char}{inner}{close_char}")
    if tag == "msup":
        base = render_mathml_element(children[0]) if len(children) > 0 else ""
        exponent = render_mathml_element(children[1]) if len(children) > 1 else ""
        superscript = maybe_translate_script(exponent, SUPERSCRIPT_MAP)
        return clean_inline_text(f"{base}{superscript}" if superscript else f"{bracket_if_needed(base)}^{bracket_if_needed(exponent)}")
    if tag == "msub":
        base = render_mathml_element(children[0]) if len(children) > 0 else ""
        subscript = render_mathml_element(children[1]) if len(children) > 1 else ""
        translated = maybe_translate_script(subscript, SUBSCRIPT_MAP)
        return clean_inline_text(f"{base}{translated}" if translated else f"{base}_{bracket_if_needed(subscript)}")
    if tag == "msubsup":
        base = render_mathml_element(children[0]) if len(children) > 0 else ""
        subscript = render_mathml_element(children[1]) if len(children) > 1 else ""
        exponent = render_mathml_element(children[2]) if len(children) > 2 else ""
        subscript_text = maybe_translate_script(subscript, SUBSCRIPT_MAP) or f"_{bracket_if_needed(subscript)}"
        superscript_text = maybe_translate_script(exponent, SUPERSCRIPT_MAP) or f"^{bracket_if_needed(exponent)}"
        return clean_inline_text(f"{base}{subscript_text}{superscript_text}")
    if tag == "mfrac":
        numerator = render_mathml_element(children[0]) if len(children) > 0 else ""
        denominator = render_mathml_element(children[1]) if len(children) > 1 else ""
        return clean_inline_text(f"{bracket_if_needed(numerator)}/{bracket_if_needed(denominator)}")
    if tag == "msqrt":
        inner = " ".join(part for part in (render_mathml_element(child) for child in children) if part)
        return clean_inline_text(f"√{bracket_if_needed(inner)}")
    if tag == "mroot":
        base = render_mathml_element(children[0]) if len(children) > 0 else ""
        index = render_mathml_element(children[1]) if len(children) > 1 else ""
        return clean_inline_text(f"{bracket_if_needed(index)}√{bracket_if_needed(base)}")
    if tag == "mover":
        base = render_mathml_element(children[0]) if len(children) > 0 else ""
        over = render_mathml_element(children[1]) if len(children) > 1 else ""
        return clean_inline_text(f"{base}^{bracket_if_needed(over)}")
    if tag == "munder":
        base = render_mathml_element(children[0]) if len(children) > 0 else ""
        under = render_mathml_element(children[1]) if len(children) > 1 else ""
        return clean_inline_text(f"{base}_{bracket_if_needed(under)}")
    if tag == "munderover":
        base = render_mathml_element(children[0]) if len(children) > 0 else ""
        under = render_mathml_element(children[1]) if len(children) > 1 else ""
        over = render_mathml_element(children[2]) if len(children) > 2 else ""
        return clean_inline_text(f"{base}_{bracket_if_needed(under)}^{bracket_if_needed(over)}")
    if tag == "mtable":
        rows = []
        for row in children:
            if local_name(row.tag) != "mtr":
                continue
            cells = [render_mathml_element(cell) for cell in list(row) if local_name(cell.tag) == "mtd"]
            rows.append(", ".join(cell for cell in cells if cell))
        return clean_inline_text(f"[{' ; '.join(rows)}]")
    if tag in {"mtr", "mtd"}:
        return clean_inline_text(" ".join(part for part in (render_mathml_element(child) for child in children) if part))
    if not children:
        return clean_inline_text("".join(elem.itertext()))
    return clean_inline_text(" ".join(part for part in (render_mathml_element(child) for child in children) if part))


def render_mathml(raw_mathml: str) -> str:
    try:
        return render_mathml_element(ET.fromstring(decode_d2l_mathml(raw_mathml)))
    except ET.ParseError:
        return ""


def verbal_math_to_symbols(text: str) -> str:
    rendered = clean_inline_text(text)

    fraction_pattern = re.compile(
        r"fraction numerator\s+(.*?)\s+over denominator\s+(.*?)\s+end fraction",
        flags=re.IGNORECASE,
    )
    while fraction_pattern.search(rendered):
        rendered = fraction_pattern.sub(
            lambda match: clean_inline_text(
                f"{bracket_if_needed(verbal_math_to_symbols(match.group(1)))}/{bracket_if_needed(verbal_math_to_symbols(match.group(2)))}"
            ),
            rendered,
        )

    for source, target in VERBAL_MATH_REPLACEMENTS:
        rendered = re.sub(re.escape(source), target, rendered, flags=re.IGNORECASE)
    rendered = re.sub(r"\bnegative\s+([A-Za-z0-9.]+)", r"-\1", rendered, flags=re.IGNORECASE)
    rendered = re.sub(r"\bend exponent\b", "", rendered, flags=re.IGNORECASE)
    rendered = re.sub(
        r"\b(one|two|three|four|five|six|seven|eight|nine|\d+)\s+(half|halves|third|thirds|fourth|fourths|fifth|fifths|sixth|sixths|seventh|sevenths|eighth|eighths|ninth|ninths|tenth|tenths)\b",
        lambda match: f"{NUMERATOR_WORDS.get(match.group(1).lower(), match.group(1))}/{ORDINAL_DENOMINATORS[match.group(2).lower()]}",
        rendered,
        flags=re.IGNORECASE,
    )
    rendered = re.sub(
        r"([A-Za-z0-9)\]])\s+to the power of\s+(-?[A-Za-z0-9]+)",
        lambda match: (
            f"{match.group(1)}{maybe_translate_script(match.group(2), SUPERSCRIPT_MAP)}"
            if maybe_translate_script(match.group(2), SUPERSCRIPT_MAP)
            else f"{match.group(1)}^{bracket_if_needed(match.group(2))}"
        ),
        rendered,
        flags=re.IGNORECASE,
    )
    rendered = re.sub(r"([A-Za-z0-9)\]])\s+squared\b", r"\1²", rendered, flags=re.IGNORECASE)
    rendered = re.sub(r"([A-Za-z0-9)\]])\s+cubed\b", r"\1³", rendered, flags=re.IGNORECASE)
    rendered = re.sub(r"(\d)\s+([A-Za-z])\b", r"\1\2", rendered)
    rendered = re.sub(r"\b([A-Za-z]{1,4})\s*\(\s*([A-Za-z0-9+\-*/= ]+)\s*\)", r"\1(\2)", rendered)
    return clean_inline_text(rendered)


def render_formula_from_image(attrs: dict[str, str | None]) -> str:
    mathml = attrs.get("data-mathml") or ""
    if mathml:
        rendered = render_mathml(mathml)
        if rendered:
            return rendered
    alt = attrs.get("alt") or attrs.get("title") or attrs.get("data-original-title") or ""
    return verbal_math_to_symbols(alt)


class TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []
        self.formula_present = False
        self.mathml_used = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attrs_dict = {key: value for key, value in attrs}
        if tag == "img":
            self.formula_present = True
            if attrs_dict.get("data-mathml"):
                self.mathml_used = True
            formula = render_formula_from_image(attrs_dict)
            if formula:
                self.parts.append(f" {formula} ")
        elif tag in {"br", "p", "div", "li", "tr", "td", "th"}:
            self.parts.append(" ")

    def handle_data(self, data: str) -> None:
        self.parts.append(data)

    def text(self) -> str:
        return clean_inline_text(" ".join(self.parts))


class AssetRefExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.refs: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attrs_dict = {key: value or "" for key, value in attrs}
        if tag == "img":
            ref = attrs_dict.get("src", "")
            if ref and IMAGE_REF_PATTERN.search(ref):
                self.refs.append(ref)
        elif tag == "a":
            ref = attrs_dict.get("href", "")
            if ref and IMAGE_REF_PATTERN.search(ref):
                self.refs.append(ref)


def child_elements(parent: ET.Element, name: str) -> list[ET.Element]:
    return [child for child in list(parent) if local_name(child.tag) == name]


def first_child(parent: ET.Element | None, name: str) -> ET.Element | None:
    if parent is None:
        return None
    for child in list(parent):
        if local_name(child.tag) == name:
            return child
    return None


def first_descendant(parent: ET.Element | None, name: str) -> ET.Element | None:
    if parent is None:
        return None
    for elem in parent.iter():
        if local_name(elem.tag) == name:
            return elem
    return None


def text_or_blank(elem: ET.Element | None) -> str:
    return (elem.text or "").strip() if elem is not None and elem.text else ""


def qti_metadata(parent: ET.Element | None) -> dict[str, str]:
    if parent is None:
        return {}
    if local_name(parent.tag) == "item":
        qtimetadata = first_descendant(first_child(parent, "itemmetadata"), "qtimetadata")
    else:
        qtimetadata = first_child(parent, "qtimetadata")
    if qtimetadata is None:
        return {}

    fields: dict[str, str] = {}
    for field in child_elements(qtimetadata, "qti_metadatafield"):
        label = text_or_blank(first_child(field, "fieldlabel"))
        entry = text_or_blank(first_child(field, "fieldentry"))
        if label:
            fields[label] = entry
    return fields


def render_html_fragment(raw_html: str) -> RenderedHtml:
    parser = TextExtractor()
    parser.feed(html.unescape(raw_html))
    return RenderedHtml(
        text=parser.text(),
        formula_present=parser.formula_present,
        mathml_used=parser.mathml_used,
    )


def plain_text_from_html(raw_html: str) -> str:
    return render_html_fragment(raw_html).text


def fully_unescape(text: str) -> str:
    current = text
    while True:
        next_text = html.unescape(current)
        if next_text == current:
            return current
        current = next_text


def normalize_exact(text: str) -> str:
    text = text.lower()
    text = text.replace("\u2018", "'").replace("\u2019", "'")
    text = text.replace("\u201c", '"').replace("\u201d", '"')
    text = text.replace("\u2013", "-").replace("\u2014", "-")
    text = text.replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def normalize_similarity(text: str) -> str:
    normalized = normalize_exact(text)
    replacements = {
        "open parentheses": "(",
        "close parentheses": ")",
        "left parenthesis": "(",
        "right parenthesis": ")",
        "x-axis": "x axis",
        "y-axis": "y axis",
        "one third": "1 third",
        "1 third": "1/3",
        "equals": "=",
        "plus": "+",
        "minus": "-",
        "times": "x",
        "multiplied by": "x",
        "squared": "^2",
        "cubed": "^3",
    }
    for source, target in replacements.items():
        normalized = normalized.replace(source, target)
    normalized = re.sub(r"[^a-z0-9]+", " ", normalized)
    normalized = re.sub(r"\s+", " ", normalized)
    return normalized.strip()


def normalize_asset_path(path: str) -> str:
    normalized = fully_unescape(path).strip().replace("\\", "/")
    if "://" in normalized:
        normalized = urlsplit(normalized).path
    normalized = unquote(normalized)
    normalized = normalized.split("#", 1)[0]
    normalized = normalized.split("?", 1)[0]
    normalized = re.sub(r"/+", "/", normalized)
    normalized = re.sub(r"^\./", "", normalized)
    return normalized


def build_asset_index(source_dir: Path) -> dict[str, dict[str, list[Path]] | dict[str, Path]]:
    relative_paths: dict[str, Path] = {}
    basenames: dict[str, list[Path]] = defaultdict(list)
    for file_path in source_dir.rglob("*"):
        if not file_path.is_file():
            continue
        relative_path = file_path.relative_to(source_dir).as_posix()
        normalized_relative = normalize_asset_path(relative_path).lstrip("/")
        relative_paths[normalized_relative] = file_path
        basenames[file_path.name].append(file_path)
    return {
        "relative_paths": relative_paths,
        "basenames": basenames,
    }


def asset_candidates(normalized_ref: str) -> list[tuple[str, str]]:
    ref = normalized_ref.lstrip("/")
    candidates: list[tuple[str, str]] = []
    if ref:
        candidates.append((ref, "normalized_ref"))
    if ref.startswith("content/enforced/"):
        parts = ref.split("/", 3)
        if len(parts) == 4 and parts[3]:
            candidates.append((parts[3], "content_enforced_tail"))
    return candidates


def resolve_asset_reference(raw_ref: str, source_dir: Path, asset_index: dict[str, dict[str, list[Path]] | dict[str, Path]]) -> AssetReference:
    normalized_ref = normalize_asset_path(raw_ref)
    relative_paths = asset_index["relative_paths"]
    basenames = asset_index["basenames"]

    for candidate, basis in asset_candidates(normalized_ref):
        resolved = relative_paths.get(candidate)
        if resolved is not None:
            return AssetReference(
                raw_ref=raw_ref,
                normalized_ref=normalized_ref,
                status="resolved",
                resolution_basis=basis,
                file_path=resolved.relative_to(source_dir).as_posix(),
                absolute_path=str(resolved.resolve()),
            )

    basename = Path(normalized_ref.lstrip("/")).name
    basename_matches = basenames.get(basename, []) if basename else []
    if len(basename_matches) == 1:
        resolved = basename_matches[0]
        return AssetReference(
            raw_ref=raw_ref,
            normalized_ref=normalized_ref,
            status="resolved",
            resolution_basis="unique_basename",
            file_path=resolved.relative_to(source_dir).as_posix(),
            absolute_path=str(resolved.resolve()),
        )

    return AssetReference(
        raw_ref=raw_ref,
        normalized_ref=normalized_ref,
        status="missing",
        resolution_basis="not_found",
        file_path="",
        absolute_path="",
    )


def item_presentation_html_fragments(item: ET.Element) -> list[str]:
    presentation = first_descendant(item, "presentation")
    if presentation is None:
        return []
    fragments: list[str] = []
    for elem in presentation.iter():
        if local_name(elem.tag) == "mattext" and elem.text:
            fragments.append(elem.text)
    return fragments


def split_qualified_name(name: str) -> tuple[str | None, str]:
    if name.startswith("{") and "}" in name:
        namespace_uri, local = name[1:].split("}", 1)
        return namespace_uri, local
    return None, name


def source_xml_attributes(elem: ET.Element) -> list[dict[str, str | None]]:
    attributes: list[dict[str, str | None]] = []
    for qualified_name, raw_value in elem.attrib.items():
        namespace_uri, name = split_qualified_name(qualified_name)
        attributes.append({
            "qualified_name": qualified_name,
            "namespace_uri": namespace_uri,
            "name": name,
            "raw_value": raw_value,
        })
    return attributes


def source_xml_node(elem: ET.Element) -> dict[str, object]:
    namespace_uri, name = split_qualified_name(elem.tag)
    return {
        "qualified_name": elem.tag,
        "namespace_uri": namespace_uri,
        "name": name,
        "attributes": source_xml_attributes(elem),
        "raw_text": elem.text,
        "raw_tail": elem.tail,
        "children": [source_xml_node(child) for child in list(elem)],
    }


def source_xml_occurrence_fact(
    item: ET.Element,
    element_name: str,
) -> dict[str, object]:
    occurrences = [
        {
            "ordinal": ordinal,
            "node": source_xml_node(node),
        }
        for ordinal, node in enumerate(
            (
                node
                for node in item.iter()
                if local_name(node.tag) == element_name
            ),
            start=1,
        )
    ]
    return {
        "state": "present" if occurrences else "absent",
        "occurrences": occurrences,
    }


def item_metadata_field_fact(
    item: ET.Element,
    field_name: str,
) -> dict[str, object]:
    matching_fields: list[ET.Element] = []
    for field in (
        node
        for node in item.iter()
        if local_name(node.tag) == "qti_metadatafield"
    ):
        labels = [
            child
            for child in list(field)
            if local_name(child.tag) == "fieldlabel"
        ]
        if any((label.text or "").strip() == field_name for label in labels):
            matching_fields.append(field)

    occurrences: list[dict[str, object]] = []
    for ordinal, field in enumerate(matching_fields, start=1):
        labels = [
            child
            for child in list(field)
            if local_name(child.tag) == "fieldlabel"
        ]
        entries = [
            child
            for child in list(field)
            if local_name(child.tag) == "fieldentry"
        ]
        raw_value = (
            entries[0].text
            if (
                len(labels) == 1
                and (labels[0].text or "").strip() == field_name
                and len(entries) == 1
            )
            else None
        )
        occurrences.append({
            "ordinal": ordinal,
            "raw_value": raw_value,
            "node": source_xml_node(field),
        })

    if not occurrences:
        return {
            "state": "absent",
            "raw_value": None,
            "occurrences": [],
        }
    if len(occurrences) == 1:
        raw_value = occurrences[0]["raw_value"]
        return {
            "state": (
                "present_empty"
                if raw_value is None or not str(raw_value).strip()
                else "present_value"
            ),
            "raw_value": raw_value,
            "occurrences": occurrences,
        }
    return {
        "state": (
            "present_value"
            if any(
                occurrence["raw_value"] is not None
                and str(occurrence["raw_value"]).strip()
                for occurrence in occurrences
            )
            else "present_empty"
        ),
        "raw_value": None,
        "occurrences": occurrences,
    }


def response_declaration_fact(
    declaration: ET.Element,
    declaration_key: str,
    ordinal: int,
) -> tuple[dict[str, object], list[dict[str, object]]]:
    diagnostics: list[dict[str, object]] = []
    declaration_kind = local_name(declaration.tag)
    source_ident = declaration.attrib.get("ident")
    source_respident = declaration.attrib.get("respident")
    expected_identifier = source_respident if declaration_kind == "response_grp" else source_ident
    if not expected_identifier:
        diagnostics.append({
            "code": "response_fact_missing_declaration_identifier",
            "severity": "warning",
            "declaration_key": declaration_key,
            "element_name": declaration_kind,
        })

    renderers = [
        child
        for child in list(declaration)
        if local_name(child.tag).startswith("render_")
    ]
    renderer_fact: dict[str, object] | None = None
    if not renderers:
        diagnostics.append({
            "code": "response_fact_missing_renderer",
            "severity": "warning",
            "declaration_key": declaration_key,
            "element_name": declaration_kind,
        })
    else:
        if len(renderers) > 1:
            diagnostics.append({
                "code": "response_fact_multiple_renderers",
                "severity": "warning",
                "declaration_key": declaration_key,
                "renderer_count": len(renderers),
            })
        renderer = renderers[0]
        renderer_name = local_name(renderer.tag)
        if renderer_name not in {"render_choice", "render_fib"}:
            diagnostics.append({
                "code": "response_fact_unknown_renderer",
                "severity": "warning",
                "declaration_key": declaration_key,
                "element_name": renderer_name,
            })
        renderer_namespace, _renderer_local = split_qualified_name(renderer.tag)
        renderer_fact = {
            "qualified_name": renderer.tag,
            "namespace_uri": renderer_namespace,
            "name": renderer_name,
            "attributes": source_xml_attributes(renderer),
            "tree": source_xml_node(renderer),
        }

    response_labels: list[dict[str, object]] = []
    for label_ordinal, label in enumerate(
        (
            node
            for node in declaration.iter()
            if local_name(node.tag) == "response_label"
        ),
        start=1,
    ):
        source_label_ident = label.attrib.get("ident")
        response_labels.append({
            "ordinal": label_ordinal,
            "source_ident": source_label_ident,
            "attributes": source_xml_attributes(label),
        })
        if not source_label_ident:
            diagnostics.append({
                "code": "response_fact_missing_label_identifier",
                "severity": "warning",
                "declaration_key": declaration_key,
                "label_ordinal": label_ordinal,
            })

    declaration_namespace, _declaration_local = split_qualified_name(declaration.tag)
    return {
        "declaration_key": declaration_key,
        "ordinal": ordinal,
        "qualified_name": declaration.tag,
        "namespace_uri": declaration_namespace,
        "element_kind": declaration_kind,
        "attributes": source_xml_attributes(declaration),
        "tree": source_xml_node(declaration),
        "source_ident": source_ident,
        "source_respident": source_respident,
        "rcardinality": declaration.attrib.get("rcardinality"),
        "material_nodes": [
            source_xml_node(child)
            for child in list(declaration)
            if local_name(child.tag) == "material"
        ],
        "renderer": renderer_fact,
        "response_labels": response_labels,
    }, diagnostics


def extract_source_response_facts(item: ET.Element | None) -> dict[str, object]:
    facts: dict[str, object] = {
        "schema": "coursecraft.quiz_source_response_facts/0",
        "source_question_type": "",
        "item_metadata": {
            "qmd_questiontype": {
                "state": "absent",
                "raw_value": None,
                "occurrences": [],
            },
            "qmd_computerscored": {
                "state": "absent",
                "raw_value": None,
                "occurrences": [],
            },
            "qmd_weighting": {
                "state": "absent",
                "raw_value": None,
                "occurrences": [],
            },
        },
        "presentation": {
            "state": "absent",
            "occurrences": [],
        },
        "presentation_sequence": [],
        "response_declarations": [],
        "grading_type": {
            "state": "absent",
            "raw_value": None,
            "occurrences": [],
        },
        "response_processing": {
            "state": "absent",
            "occurrences": [],
        },
        "outcomes": [],
        "response_conditions": [],
        "feedback_slots": [],
        "diagnostics": [],
    }
    diagnostics = facts["diagnostics"]
    assert isinstance(diagnostics, list)

    if item is None:
        diagnostics.append({
            "code": "response_fact_source_item_unresolved",
            "severity": "warning",
        })
        return facts

    question_type = qti_metadata(item).get("qmd_questiontype", "")
    facts["source_question_type"] = question_type
    if not question_type:
        diagnostics.append({
            "code": "response_fact_missing_question_type",
            "severity": "warning",
        })

    facts["item_metadata"] = {
        field_name: item_metadata_field_fact(item, field_name)
        for field_name in (
            "qmd_questiontype",
            "qmd_computerscored",
            "qmd_weighting",
        )
    }
    facts["presentation"] = source_xml_occurrence_fact(item, "presentation")
    facts["response_processing"] = source_xml_occurrence_fact(
        item,
        "resprocessing",
    )

    presentation = first_descendant(item, "presentation")
    presentation_sequence: list[dict[str, object]] = []
    declarations: list[dict[str, object]] = []
    grading_elements: list[ET.Element] = []

    def visit_presentation(container: ET.Element, container_path: list[int]) -> None:
        for child_index, child in enumerate(list(container), start=1):
            child_path = [*container_path, child_index]
            child_name = local_name(child.tag)
            if child_name == "flow":
                visit_presentation(child, child_path)
                continue
            if child_name == "material":
                presentation_sequence.append({
                    "ordinal": len(presentation_sequence) + 1,
                    "kind": "material",
                    "container_path": child_path,
                    "node": source_xml_node(child),
                })
                continue
            if child_name in {"response_str", "response_grp", "response_lid"}:
                declaration_key = f"response-{len(declarations) + 1}"
                declaration_fact, declaration_diagnostics = response_declaration_fact(
                    child,
                    declaration_key,
                    len(declarations) + 1,
                )
                declarations.append(declaration_fact)
                diagnostics.extend(declaration_diagnostics)
                presentation_sequence.append({
                    "ordinal": len(presentation_sequence) + 1,
                    "kind": "response",
                    "container_path": child_path,
                    "declaration_key": declaration_key,
                    "element_kind": child_name,
                })
                continue
            if child_name == "response_extension":
                presentation_sequence.append({
                    "ordinal": len(presentation_sequence) + 1,
                    "kind": "response_extension",
                    "container_path": child_path,
                    "node": source_xml_node(child),
                })
                grading_elements.extend(
                    node
                    for node in child.iter()
                    if local_name(node.tag) == "grading_type"
                )
                continue
            presentation_sequence.append({
                "ordinal": len(presentation_sequence) + 1,
                "kind": "unknown",
                "container_path": child_path,
                "node": source_xml_node(child),
            })
            diagnostics.append({
                "code": "response_fact_unknown_presentation_node",
                "severity": "warning",
                "element_name": child_name,
                "container_path": child_path,
            })

    if presentation is None:
        diagnostics.append({
            "code": "response_fact_missing_presentation",
            "severity": "warning",
        })
    else:
        visit_presentation(presentation, [])

    facts["presentation_sequence"] = presentation_sequence
    facts["response_declarations"] = declarations

    grading_occurrences = [
        {
            "ordinal": ordinal,
            "raw_value": elem.text,
            "node": source_xml_node(elem),
        }
        for ordinal, elem in enumerate(grading_elements, start=1)
    ]
    if grading_occurrences:
        first_value = grading_occurrences[0]["raw_value"]
        facts["grading_type"] = {
            "state": (
                "present_empty"
                if first_value is None or not str(first_value).strip()
                else "present_value"
            ),
            "raw_value": first_value,
            "occurrences": grading_occurrences,
        }
        if len(grading_occurrences) > 1:
            diagnostics.append({
                "code": "response_fact_multiple_grading_types",
                "severity": "warning",
                "occurrence_count": len(grading_occurrences),
            })

    resprocessing_nodes = [
        node for node in item.iter() if local_name(node.tag) == "resprocessing"
    ]
    outcome_facts: list[dict[str, object]] = []
    condition_facts: list[dict[str, object]] = []
    known_predicates = {
        "and",
        "conditionvar",
        "not",
        "or",
        "other",
        "unanswered",
        "varequal",
        "vargt",
        "vargte",
        "varinside",
        "varlt",
        "varlte",
    }
    for resprocessing in resprocessing_nodes:
        for outcomes in (
            child
            for child in list(resprocessing)
            if local_name(child.tag) == "outcomes"
        ):
            for decvar in (
                child
                for child in list(outcomes)
                if local_name(child.tag) == "decvar"
            ):
                outcome_facts.append({
                    "ordinal": len(outcome_facts) + 1,
                    "attributes": source_xml_attributes(decvar),
                    "raw_text": decvar.text,
                    "node": source_xml_node(decvar),
                })
        for respcondition in (
            child
            for child in list(resprocessing)
            if local_name(child.tag) == "respcondition"
        ):
            predicates = [
                child
                for child in list(respcondition)
                if local_name(child.tag) == "conditionvar"
            ]
            setvars = [
                child
                for child in list(respcondition)
                if local_name(child.tag) == "setvar"
            ]
            other_children = [
                source_xml_node(child)
                for child in list(respcondition)
                if local_name(child.tag) not in {"conditionvar", "setvar"}
            ]
            condition_ordinal = len(condition_facts) + 1
            condition_facts.append({
                "ordinal": condition_ordinal,
                "attributes": source_xml_attributes(respcondition),
                "tree": source_xml_node(respcondition),
                "predicate_trees": [source_xml_node(predicate) for predicate in predicates],
                "setvars": [
                    {
                        "ordinal": ordinal,
                        "attributes": source_xml_attributes(setvar),
                        "raw_value": setvar.text,
                        "node": source_xml_node(setvar),
                    }
                    for ordinal, setvar in enumerate(setvars, start=1)
                ],
                "other_children": other_children,
            })
            if not predicates:
                diagnostics.append({
                    "code": "response_fact_missing_predicate",
                    "severity": "warning",
                    "condition_ordinal": condition_ordinal,
                })
            for predicate in predicates:
                for predicate_node in predicate.iter():
                    predicate_name = local_name(predicate_node.tag)
                    if predicate_name not in known_predicates:
                        diagnostics.append({
                            "code": "response_fact_unknown_predicate",
                            "severity": "warning",
                            "condition_ordinal": condition_ordinal,
                            "element_name": predicate_name,
                        })

    facts["outcomes"] = outcome_facts
    facts["response_conditions"] = condition_facts

    feedback_slots: list[dict[str, object]] = []
    for feedback in (
        node for node in item.iter() if local_name(node.tag) == "itemfeedback"
    ):
        feedback_slots.append({
            "ordinal": len(feedback_slots) + 1,
            "attributes": source_xml_attributes(feedback),
            "source_ident": feedback.attrib.get("ident"),
            "has_content": any(
                bool((node.text or "").strip())
                for node in feedback.iter()
                if node is not feedback
            ),
            "node": source_xml_node(feedback),
        })
    facts["feedback_slots"] = feedback_slots

    if question_type in {"Matching", "Ordering"} and not declarations:
        diagnostics.append({
            "code": "response_shape_empty_shell",
            "severity": "warning",
            "source_question_type": question_type,
        })

    return facts


def extract_question_asset_refs(item: ET.Element, source_dir: Path, asset_index: dict[str, dict[str, list[Path]] | dict[str, Path]]) -> list[AssetReference]:
    refs: list[str] = []
    for fragment in item_presentation_html_fragments(item):
        parser = AssetRefExtractor()
        parser.feed(fully_unescape(fragment))
        refs.extend(parser.refs)
    ordered_refs = unique_preserve_order(refs)
    return [resolve_asset_reference(ref, source_dir, asset_index) for ref in ordered_refs]


def response_choice_infos(item: ET.Element) -> list[dict[str, object]]:
    infos: list[dict[str, object]] = []
    for response_label in item.iter():
        if local_name(response_label.tag) != "response_label":
            continue
        mattexts = [elem for elem in response_label.iter() if local_name(elem.tag) == "mattext"]
        rendered = " ".join(text_or_blank(elem) for elem in mattexts if text_or_blank(elem))
        info = render_html_fragment(rendered) if rendered else RenderedHtml(text="", formula_present=False, mathml_used=False)
        infos.append({
            "ident": response_label.attrib.get("ident", ""),
            "text": info.text,
            "formula_present": info.formula_present,
            "mathml_used": info.mathml_used,
        })
    return infos


def extract_matching_structure(item: ET.Element) -> dict[str, list[dict[str, str]]]:
    prompt_infos: list[dict[str, str]] = []
    option_infos: list[dict[str, str]] = []
    seen_option_ids: set[str] = set()

    for response_grp in item.iter():
        if local_name(response_grp.tag) != "response_grp":
            continue

        prompt_material = first_child(response_grp, "material")
        prompt_html = ""
        if prompt_material is not None:
            prompt_fragments = [
                text_or_blank(elem)
                for elem in prompt_material.iter()
                if local_name(elem.tag) == "mattext" and text_or_blank(elem)
            ]
            prompt_html = " ".join(prompt_fragments)
        prompt_render = render_html_fragment(prompt_html) if prompt_html else RenderedHtml(text="", formula_present=False, mathml_used=False)
        prompt_text = prompt_render.text
        if prompt_text:
            prompt_infos.append({
                "respident": response_grp.attrib.get("respident", ""),
                "text": prompt_text,
            })

        for option_info in response_choice_infos(response_grp):
            option_ident = str(option_info["ident"])
            option_text = str(option_info["text"])
            if not option_ident or not option_text or option_ident in seen_option_ids:
                continue
            seen_option_ids.add(option_ident)
            option_infos.append({
                "ident": option_ident,
                "text": option_text,
            })

    return {
        "prompts": prompt_infos,
        "options": option_infos,
    }


def response_choice_texts(item: ET.Element) -> list[str]:
    return [str(info["text"]) for info in response_choice_infos(item) if str(info["text"])]


def source_choice_display(facts: dict[str, object]) -> dict[str, object] | None:
    """Keep readable choice text separate from source identity and rich material."""
    if str(facts.get("source_question_type")) not in {
        "Multiple Choice", "True/False", "Multi-Select", "Multi Select", "Multiple Response", "Ordering"
    }:
        return None
    projection = source_choice_projection(facts)
    options = []
    for option in projection["options"]:
        content = option["content"]
        text = render_html_fragment(content["content"]).text if content["format"] == "html" else content["content"]
        image_only = not text.strip()
        if image_only:
            parser = AssetRefExtractor()
            parser.feed(content["content"])
            refs = unique_preserve_order(parser.refs)
            text = "[Image: " + ", ".join(refs) + "]" if refs else "[Source option has no readable text]"
        options.append({"option_key": option["option_key"], "source_ident": option["source_ident"],
                        "text": text, "correct": option["correct"], "position": option["position"], "image_only": image_only})
    counts = Counter(str(option["text"]) for option in options)
    show_keys = any(option["image_only"] or counts[str(option["text"])] > 1 for option in options)
    return {"options": options, "show_keys": show_keys, "recognized": projection["recognized"], "reason": projection["reason"]}


def _choice_display_text(option: dict[str, object], show_keys: bool) -> str:
    return f"{option['option_key']}. {option['text']}" if show_keys else str(option["text"])


def render_material_block(material: ET.Element | None) -> RenderedHtml:
    if material is None:
        return RenderedHtml(text="", formula_present=False, mathml_used=False)
    mattexts = [elem for elem in material.iter() if local_name(elem.tag) == "mattext"]
    rendered = " ".join(text_or_blank(elem) for elem in mattexts if text_or_blank(elem))
    return render_html_fragment(rendered) if rendered else RenderedHtml(text="", formula_present=False, mathml_used=False)


def render_prompt_container(container: ET.Element, blank_index: int = 0) -> tuple[list[str], int, bool, bool]:
    parts: list[str] = []
    formula_present = False
    mathml_used = False

    for child in list(container):
        name = local_name(child.tag)
        if name == "material":
            rendered = render_material_block(child)
            if rendered.text:
                parts.append(rendered.text)
            formula_present = formula_present or rendered.formula_present
            mathml_used = mathml_used or rendered.mathml_used
        elif name == "response_str":
            blank_index += 1
            parts.append(f"[Blank {blank_index}]")
        elif name == "flow":
            child_parts, blank_index, child_formula, child_mathml = render_prompt_container(child, blank_index)
            parts.extend(child_parts)
            formula_present = formula_present or child_formula
            mathml_used = mathml_used or child_mathml

    return parts, blank_index, formula_present, mathml_used


def render_prompt_text(item: ET.Element) -> RenderedHtml:
    presentation = first_descendant(item, "presentation")
    if presentation is None:
        return RenderedHtml(text="", formula_present=False, mathml_used=False)

    parts, _blank_index, formula_present, mathml_used = render_prompt_container(presentation, 0)
    if parts:
        return RenderedHtml(
            text=clean_inline_text(" ".join(parts)),
            formula_present=formula_present,
            mathml_used=mathml_used,
        )

    return render_html_fragment(question_raw_html(item))


def extract_fill_in_blank_structure(item: ET.Element) -> dict[str, object]:
    blank_entries: list[dict[str, object]] = []
    for response_str in item.iter():
        if local_name(response_str.tag) != "response_str":
            continue
        response_label = first_descendant(response_str, "response_label")
        blank_entries.append({
            "blank_number": len(blank_entries) + 1,
            "respident": response_label.attrib.get("ident", "") if response_label is not None else "",
        })

    blank_respidents = {str(entry["respident"]) for entry in blank_entries if str(entry["respident"])}
    blank_answers_by_respident: dict[str, list[str]] = defaultdict(list)

    for respcondition in item.iter():
        if local_name(respcondition.tag) != "respcondition":
            continue
        setvars = [elem for elem in respcondition.iter() if local_name(elem.tag) == "setvar"]
        if not any(is_positive_setvar(setvar) for setvar in setvars):
            continue
        varequals = collect_varequals(first_descendant(respcondition, "conditionvar"))
        for entry in varequals:
            respident = str(entry["respident"])
            value = str(entry["value"])
            if not value or bool(entry["negated"]) or respident not in blank_respidents:
                continue
            blank_answers_by_respident[respident].append(value)

    blank_answer_lines: list[str] = []
    for entry in blank_entries:
        respident = str(entry["respident"])
        values = unique_preserve_order(blank_answers_by_respident.get(respident, []))
        if not values:
            continue
        blank_answer_lines.append(f"Blank {entry['blank_number']}: {' / '.join(values)}")

    return {
        "blank_count": len(blank_entries),
        "blank_answers_text": " || ".join(blank_answer_lines),
    }


def question_raw_html(item: ET.Element) -> str:
    presentation = first_descendant(item, "presentation")
    if presentation is None:
        return ""
    material = first_descendant(presentation, "material")
    mattext = first_child(material, "mattext") if material is not None else None
    return mattext.text or "" if mattext is not None else ""


def answer_key_raw_html(item: ET.Element) -> str:
    answer_key = first_descendant(item, "answer_key")
    if answer_key is None:
        return ""
    mattexts = [text_or_blank(elem) for elem in answer_key.iter() if local_name(elem.tag) == "mattext" and text_or_blank(elem)]
    return " ".join(mattexts)


def is_positive_setvar(setvar: ET.Element) -> bool:
    varname = setvar.attrib.get("varname", "")
    if varname == "D2L_Incorrect":
        return False
    if varname == "D2L_Correct":
        return True
    value = text_or_blank(setvar)
    if value == "D2L_Correct":
        return True
    try:
        return float(value) > 0
    except ValueError:
        return False


def collect_varequals(elem: ET.Element | None, negated: bool = False) -> list[dict[str, str | bool]]:
    if elem is None:
        return []
    entries: list[dict[str, str | bool]] = []
    for child in list(elem):
        name = local_name(child.tag)
        if name == "not":
            entries.extend(collect_varequals(child, not negated))
        elif name == "varequal":
            entries.append({
                "respident": child.attrib.get("respident", ""),
                "value": text_or_blank(child),
                "negated": negated,
            })
        else:
            entries.extend(collect_varequals(child, negated))
    return entries


def unique_preserve_order(values: list[str]) -> list[str]:
    seen: set[str] = set()
    ordered: list[str] = []
    for value in values:
        if value and value not in seen:
            seen.add(value)
            ordered.append(value)
    return ordered


def extract_correct_answer(item: ET.Element) -> dict[str, str]:
    question_type = qti_metadata(item).get("qmd_questiontype", "")
    display = source_choice_display(extract_source_response_facts(item))
    if display is not None:
        if not display["recognized"]:
            return {"correct_answer": "", "correct_response_ids": "", "correct_answer_basis": "unresolved_source_scoring"}
        ordering = question_type == "Ordering"
        keyed = sorted(display["options"], key=lambda option: option["position"]) if ordering else [option for option in display["options"] if option["correct"] is True]
        return {
            "correct_answer": " || ".join(str(option["text"]) for option in keyed),
            "correct_response_ids": " || ".join(str(option["source_ident"]) for option in keyed),
            "correct_answer_basis": "ordering_sequence" if ordering else "multiple_selection" if question_type in {"Multi-Select", "Multi Select", "Multiple Response"} else "single_selection",
        }
    choice_infos = response_choice_infos(item)
    choice_lookup = {str(info["ident"]): str(info["text"]) for info in choice_infos if str(info["ident"])}
    positive_choice_ids: list[str] = []
    positive_literals: list[str] = []
    ordering_positions: list[tuple[int, str]] = []
    matching_pairs: dict[str, str] = {}

    for respcondition in item.iter():
        if local_name(respcondition.tag) != "respcondition":
            continue
        setvars = [elem for elem in respcondition.iter() if local_name(elem.tag) == "setvar"]
        if not any(is_positive_setvar(setvar) for setvar in setvars):
            continue
        varequals = collect_varequals(first_descendant(respcondition, "conditionvar"))
        for entry in varequals:
            respident = str(entry["respident"])
            value = str(entry["value"])
            if not value or bool(entry["negated"]):
                continue
            if question_type == "Matching" and respident:
                if value in choice_lookup:
                    matching_pairs[respident] = value
                continue
            if question_type == "Ordering" and respident in choice_lookup:
                try:
                    ordering_positions.append((int(value), respident))
                    continue
                except ValueError:
                    pass
            if value in choice_lookup:
                positive_choice_ids.append(value)
            else:
                positive_literals.append(value)

    if question_type == "Matching":
        matching_structure = extract_matching_structure(item)
        ordered_pair_texts: list[str] = []
        ordered_pair_ids: list[str] = []
        for prompt_info in matching_structure["prompts"]:
            respident = str(prompt_info["respident"])
            prompt_text = str(prompt_info["text"])
            option_id = matching_pairs.get(respident, "")
            option_text = choice_lookup.get(option_id, "")
            if not respident or not prompt_text or not option_id or not option_text:
                continue
            ordered_pair_texts.append(f"{prompt_text} -> {option_text}")
            ordered_pair_ids.append(f"{respident}:{option_id}")
        if ordered_pair_texts:
            return {
                "correct_answer": " || ".join(ordered_pair_texts),
                "correct_response_ids": " || ".join(ordered_pair_ids),
                "correct_answer_basis": "matching_pairs",
            }

    if ordering_positions:
        ordered_pairs = sorted(ordering_positions, key=lambda pair: pair[0])
        ordered_response_ids = unique_preserve_order([respident for _position, respident in ordered_pairs if respident in choice_lookup])
        ordered_texts = unique_preserve_order([choice_lookup[respident] for respident in ordered_response_ids if respident in choice_lookup])
        if ordered_texts:
            return {
                "correct_answer": " || ".join(ordered_texts),
                "correct_response_ids": " || ".join(ordered_response_ids),
                "correct_answer_basis": "ordering_sequence",
            }

    positive_choice_ids = unique_preserve_order(positive_choice_ids)
    positive_literals = unique_preserve_order(positive_literals)
    positive_choice_texts = unique_preserve_order([choice_lookup[value] for value in positive_choice_ids if value in choice_lookup])

    if positive_choice_texts:
        basis = "multiple_selection" if question_type == "Multi-Select" or len(positive_choice_texts) > 1 else "single_selection"
        return {
            "correct_answer": " || ".join(positive_choice_texts),
            "correct_response_ids": " || ".join(positive_choice_ids),
            "correct_answer_basis": basis,
        }

    if positive_literals:
        return {
            "correct_answer": " || ".join(positive_literals),
            "correct_response_ids": "",
            "correct_answer_basis": "text_entry",
        }

    answer_key_html = answer_key_raw_html(item)
    if answer_key_html:
        return {
            "correct_answer": plain_text_from_html(answer_key_html),
            "correct_response_ids": "",
            "correct_answer_basis": "answer_key_material",
        }

    if question_type == "Long Answer":
        return {
            "correct_answer": "[manual grading; no fixed keyed answer in XML]",
            "correct_response_ids": "",
            "correct_answer_basis": "manual_review",
        }

    return {
        "correct_answer": "",
        "correct_response_ids": "",
        "correct_answer_basis": "not_found",
    }


def parse_library_item(
    bank_ident: str,
    bank_title: str,
    bank_path: str,
    pool_ident: str,
    pool_title: str,
    pool_path: str,
    item: ET.Element,
) -> LibraryQuestion:
    meta = qti_metadata(item)
    raw_html = question_raw_html(item)
    plain_text = render_prompt_text(item).text
    exact_text_key = normalize_exact(plain_text)
    similarity_text_key = normalize_similarity(plain_text)
    choice_texts = response_choice_texts(item)
    choices_text = " || ".join(choice_texts)
    choice_key = normalize_exact(choices_text)
    question_type = meta.get("qmd_questiontype", "")
    signature_key = "|".join([question_type, exact_text_key, choice_key])
    return LibraryQuestion(
        bank_ident=bank_ident,
        bank_title=bank_title,
        bank_path=bank_path,
        pool_ident=pool_ident,
        pool_title=pool_title,
        pool_path=pool_path,
        question_ident=item.attrib.get("ident", ""),
        question_label=item.attrib.get("label", ""),
        question_title=item.attrib.get("title", ""),
        question_type=question_type,
        question_weight=meta.get("qmd_weighting", ""),
        qmd_displayid=meta.get("qmd_displayid", ""),
        qmd_globalid=meta.get("qmd_globalid", ""),
        raw_html=raw_html,
        plain_text=plain_text,
        exact_text_key=exact_text_key,
        similarity_text_key=similarity_text_key,
        choices_text=choices_text,
        choice_count=len(choice_texts),
        signature_key=signature_key,
        item_element=item,
    )


def library_question_payload_row(
    library: LibraryQuestion,
    export_dir: Path,
    asset_index: dict[str, list[Path]],
) -> dict[str, object]:
    """Extract a library record's answer payload with the quiz-item machinery.

    A question-library ``<item>`` is the same shape as a quiz ``<item>``, so the
    same renderers, choice reader, answer-key resolver, and asset resolver are
    used here rather than a second implementation.  Field names match the
    canonical quiz-question row so downstream normalization can build the typed
    payload through exactly one code path.

    The row is deliberately free of placement facts.  A library record has no
    quiz, section, ordinal, pool draw, or placement points, and inventing any of
    those would make a dormant record look placed.
    """
    item = library.item_element
    render = render_prompt_text(item)
    choice_infos = response_choice_infos(item)
    matching_structure = extract_matching_structure(item)
    fill_in_blank_structure = extract_fill_in_blank_structure(item)
    answer_key = extract_correct_answer(item)
    asset_refs = extract_question_asset_refs(item, export_dir, asset_index)
    primary_asset = next(
        (asset for asset in asset_refs if asset.status == "resolved" and asset.file_path),
        None,
    )
    return {
        "library_bank_ident": library.bank_ident,
        "library_bank_title": library.bank_title,
        "library_pool_ident": library.pool_ident,
        "library_pool_title": library.pool_title,
        "library_pool_path": library.pool_path,
        "question_ident": library.question_ident,
        "question_label": library.question_label,
        "question_title": library.question_title,
        "question_type": library.question_type,
        "question_weight": library.question_weight,
        "qmd_displayid": library.qmd_displayid,
        "qmd_globalid": library.qmd_globalid,
        "question_text": render.text,
        "mathml_used": "yes" if render.mathml_used else "no",
        "formula_present": "yes" if render.formula_present else "no",
        "choices_text": " || ".join(
            str(info["text"]) for info in choice_infos if str(info["text"])
        ),
        "matching_prompts_text": " || ".join(
            str(info["text"])
            for info in matching_structure["prompts"]
            if str(info["text"])
        ),
        "matching_options_text": " || ".join(
            str(info["text"])
            for info in matching_structure["options"]
            if str(info["text"])
        ),
        "fill_in_blank_count": fill_in_blank_structure["blank_count"],
        "fill_in_blank_answers_text": fill_in_blank_structure["blank_answers_text"],
        "correct_answer": answer_key["correct_answer"],
        "correct_answer_basis": answer_key["correct_answer_basis"],
        "correct_response_ids": answer_key["correct_response_ids"],
        "source_response_facts": extract_source_response_facts(item),
        "source_choice_display": source_choice_display(extract_source_response_facts(item)),
        "question_image_count": len(asset_refs),
        "question_image_refs": " || ".join(asset.raw_ref for asset in asset_refs),
        "question_image_paths": " || ".join(
            asset.file_path for asset in asset_refs if asset.file_path
        ),
        "question_primary_image_path": primary_asset.file_path if primary_asset else "",
        "question_primary_image_absolute_path": (
            primary_asset.absolute_path if primary_asset else ""
        ),
        "question_missing_image_refs": " || ".join(
            asset.raw_ref for asset in asset_refs if asset.status != "resolved"
        ),
        "payload_state": "parsed",
    }


def parse_question_library(questiondb_path: Path) -> tuple[list[dict[str, object]], list[LibraryQuestion]]:
    root = ET.parse(questiondb_path).getroot()
    objectbank = first_descendant(root, "objectbank")
    if objectbank is None:
        raise ValueError(f"Could not locate <objectbank> in {questiondb_path}")

    pools: list[dict[str, object]] = []
    questions: list[LibraryQuestion] = []

    def visit_section(section: ET.Element, ancestors: list[str], bank_ident: str, bank_title: str, bank_path: str) -> None:
        pool_ident = section.attrib.get("ident", "")
        pool_title = section.attrib.get("title", "") or "(untitled pool)"
        pool_path = " > ".join([*ancestors, pool_title]) if ancestors else pool_title
        direct_items = child_elements(section, "item")
        pools.append({
            "bank_ident": bank_ident,
            "bank_title": bank_title,
            "bank_path": bank_path,
            "pool_ident": pool_ident,
            "pool_title": pool_title,
            "pool_path": pool_path,
            "pool_level": len(ancestors) + 1,
            "question_count": len(direct_items),
        })
        for item in direct_items:
            questions.append(parse_library_item(bank_ident, bank_title, bank_path, pool_ident, pool_title, pool_path, item))
        for child_section in child_elements(section, "section"):
            visit_section(child_section, [*ancestors, pool_title], bank_ident, bank_title, bank_path)

    root_items = child_elements(objectbank, "item")
    if root_items:
        root_pool_ident = objectbank.attrib.get("ident", "") or "OBJECTBANK_ROOT"
        root_pool_title = "Question Library Root Items"
        pools.append({
            "bank_ident": root_pool_ident,
            "bank_title": root_pool_title,
            "bank_path": root_pool_title,
            "pool_ident": root_pool_ident,
            "pool_title": root_pool_title,
            "pool_path": root_pool_title,
            "pool_level": 0,
            "question_count": len(root_items),
        })
        for item in root_items:
            questions.append(parse_library_item(root_pool_ident, root_pool_title, root_pool_title, root_pool_ident, root_pool_title, root_pool_title, item))

    for section in child_elements(objectbank, "section"):
        top_ident = section.attrib.get("ident", "")
        top_title = section.attrib.get("title", "") or "(untitled pool)"
        visit_section(section, [], top_ident, top_title, top_title)

    return pools, questions


def build_library_indexes(questions: list[LibraryQuestion]) -> dict[str, dict[str, list[LibraryQuestion]]]:
    indexes: dict[str, dict[str, list[LibraryQuestion]]] = {
        "displayid": defaultdict(list),
        "globalid": defaultdict(list),
        "label": defaultdict(list),
        "ident": defaultdict(list),
        "signature": defaultdict(list),
        "exact_text": defaultdict(list),
        "similarity_text": defaultdict(list),
        "question_type": defaultdict(list),
    }
    for question in questions:
        if question.qmd_displayid:
            indexes["displayid"][question.qmd_displayid].append(question)
        if question.qmd_globalid:
            indexes["globalid"][question.qmd_globalid].append(question)
        if question.question_label:
            indexes["label"][question.question_label].append(question)
        if question.question_ident:
            indexes["ident"][question.question_ident].append(question)
        if question.signature_key:
            indexes["signature"][question.signature_key].append(question)
        if question.exact_text_key:
            indexes["exact_text"][question.exact_text_key].append(question)
        if question.similarity_text_key:
            indexes["similarity_text"][question.similarity_text_key].append(question)
        indexes["question_type"][question.question_type].append(question)
    return indexes


def unique_match(candidates: list[LibraryQuestion], evidence_level: str, match_basis: str) -> MatchResult | None:
    if len(candidates) == 1:
        return MatchResult(
            library_question=candidates[0],
            evidence_level=evidence_level,
            match_basis=match_basis,
            match_score=1.0,
            note="",
            candidates=match_candidate_rows(candidates, match_basis),
        )
    if candidates and len({candidate.pool_ident for candidate in candidates}) == 1:
        return MatchResult(
            library_question=candidates[0],
            evidence_level=evidence_level,
            match_basis=match_basis,
            match_score=1.0,
            note="Multiple question-library items share this exact content inside one pool; pool assignment is stable but item identity is not unique.",
            candidates=match_candidate_rows(candidates, match_basis),
        )
    return None


def match_candidate_rows(
    candidates: list[LibraryQuestion],
    basis: str,
    scores: dict[int, float] | None = None,
    limit: int = 5,
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for candidate in candidates[:limit]:
        rows.append({
            "basis": basis,
            "score": round(scores.get(id(candidate), 1.0), 6) if scores else 1.0,
            "bank_ident": candidate.bank_ident,
            "bank_title": candidate.bank_title,
            "pool_ident": candidate.pool_ident,
            "pool_title": candidate.pool_title,
            "question_ident": candidate.question_ident,
            "question_label": candidate.question_label,
            "qmd_displayid": candidate.qmd_displayid,
            "qmd_globalid": candidate.qmd_globalid,
        })
    return rows


def match_question(item: ET.Element, indexes: dict[str, dict[str, list[LibraryQuestion]]]) -> MatchResult:
    meta = qti_metadata(item)
    question_type = meta.get("qmd_questiontype", "")
    plain_text = render_prompt_text(item).text
    exact_text_key = normalize_exact(plain_text)
    similarity_text_key = normalize_similarity(plain_text)
    choice_texts = response_choice_texts(item)
    choices_text = " || ".join(choice_texts)
    choice_key = normalize_exact(choices_text)
    signature_key = "|".join([question_type, exact_text_key, choice_key])
    ambiguous_candidates: list[dict[str, object]] = []

    for key_name, evidence_level, basis in [
        ("qmd_displayid", "source_evidence", "qmd_displayid"),
        ("qmd_globalid", "source_evidence", "qmd_globalid"),
    ]:
        key_value = meta.get(key_name, "")
        if not key_value:
            continue
        candidates_for_key = indexes["displayid" if key_name == "qmd_displayid" else "globalid"].get(key_value, [])
        match = unique_match(candidates_for_key, evidence_level, basis)
        if match is not None:
            return match
        if candidates_for_key:
            ambiguous_candidates = match_candidate_rows(candidates_for_key, basis)

    if signature_key:
        candidates_for_signature = indexes["signature"].get(signature_key, [])
        match = unique_match(candidates_for_signature, "inferred_exact", "question_signature_exact")
        if match is not None:
            return match
        if candidates_for_signature:
            ambiguous_candidates = match_candidate_rows(candidates_for_signature, "question_signature_exact")

    if exact_text_key:
        candidates_for_text = indexes["exact_text"].get(exact_text_key, [])
        match = unique_match(candidates_for_text, "inferred_exact", "question_text_exact")
        if match is not None:
            return match
        if candidates_for_text:
            ambiguous_candidates = match_candidate_rows(candidates_for_text, "question_text_exact")

    candidates = indexes["question_type"].get(question_type, []) or [q for group in indexes["question_type"].values() for q in group]
    filtered = [candidate for candidate in candidates if candidate.choice_count == len(choice_texts) or candidate.choice_count == 0]
    if filtered:
        candidates = filtered

    scored: list[tuple[float, LibraryQuestion]] = []
    for candidate in candidates:
        if not similarity_text_key or not candidate.similarity_text_key:
            continue
        score = SequenceMatcher(None, similarity_text_key, candidate.similarity_text_key).ratio()
        scored.append((score, candidate))

    scored.sort(key=lambda pair: pair[0], reverse=True)
    if scored:
        best_score, best_candidate = scored[0]
        second_score = scored[1][0] if len(scored) > 1 else 0.0
        if best_score >= SIMILARITY_THRESHOLD and (best_score - second_score) >= SIMILARITY_GAP:
            return MatchResult(
                library_question=best_candidate,
                evidence_level="inferred_similarity",
                match_basis="question_text_similarity",
                match_score=best_score,
                note="",
                candidates=match_candidate_rows(
                    [candidate for _score, candidate in scored],
                    "question_text_similarity",
                    {id(candidate): score for score, candidate in scored},
                ),
            )

    note = "No unique question-library match found in questiondb.xml for this quiz item."
    if scored:
        note = f"{note} Best similarity score: {scored[0][0]:.3f}."
        ambiguous_candidates = match_candidate_rows(
            [candidate for _score, candidate in scored],
            "question_text_similarity",
            {id(candidate): score for score, candidate in scored},
        )
    return MatchResult(
        library_question=None,
        evidence_level="unmatched",
        match_basis="unmatched",
        match_score=scored[0][0] if scored else None,
        note=note,
        candidates=ambiguous_candidates,
    )


def match_itemref(itemref: ET.Element, indexes: dict[str, dict[str, list[LibraryQuestion]]]) -> MatchResult:
    linkrefid = itemref.attrib.get("linkrefid", "")
    file_ref = first_descendant(itemref, "file")
    href = file_ref.attrib.get("href", "") if file_ref is not None else ""
    href_name = Path(href).name if href else ""

    if href and href_name != "questiondb.xml":
        return MatchResult(
            library_question=None,
            evidence_level="unmatched",
            match_basis="itemref_unsupported_href",
            match_score=None,
            note=f"Quiz XML itemref points to `{href}`, not questiondb.xml.",
        )

    ambiguous_candidates: list[dict[str, object]] = []
    if linkrefid:
        label_candidates = indexes["label"].get(linkrefid, [])
        match = unique_match(label_candidates, "source_evidence", "itemref_linkrefid")
        if match is not None:
            return match
        if label_candidates:
            ambiguous_candidates = match_candidate_rows(label_candidates, "itemref_linkrefid")
        ident_candidates = indexes["ident"].get(linkrefid, [])
        match = unique_match(ident_candidates, "source_evidence", "itemref_ident")
        if match is not None:
            return match
        if ident_candidates:
            ambiguous_candidates = match_candidate_rows(ident_candidates, "itemref_ident")

    note = "Quiz XML itemref did not resolve to a unique question-library item in questiondb.xml."
    if linkrefid:
        note = f"{note} linkrefid: {linkrefid}."
    return MatchResult(
        library_question=None,
        evidence_level="unmatched",
        match_basis="itemref_unmatched",
        match_score=None,
        note=note,
        candidates=ambiguous_candidates,
    )


def blank_answer_key() -> dict[str, str]:
    return {
        "correct_answer": "",
        "correct_response_ids": "",
        "correct_answer_basis": "not_found",
    }


def itemref_points(itemref: ET.Element, fallback: str) -> str:
    points = first_descendant(itemref, "points")
    if points is not None:
        value = text_or_blank(points)
        if value:
            return value
    return fallback


def section_draw_count(section: ET.Element) -> str:
    return qti_metadata(section).get("qmd_numberofitems", "")


def section_weight(section: ET.Element) -> str:
    """Points a section awards per question drawn from it.

    A ``RAND_`` draw section may declare its own weight alongside
    ``qmd_numberofitems``.  On all evidence to date this is a redundant copy
    rather than an override: every observed item carries its own weight, and
    across 1,255 items sitting under a weighted section the two agree without a
    single disagreement.  It is captured so the structure that declares an award
    can be reconstructed, not because points are read from it.

    ``qti_metadata`` reads a section's direct ``qtimetadata`` child, so this
    never picks up an enclosed item's weight.
    """
    return qti_metadata(section).get("qmd_weighting", "")


def locate_export_folder(source: Path) -> Path:
    if source.is_file():
        raise ValueError("locate_export_folder accepts directories; use resolve_export for ZIP intake.")
    questiondb = source / "questiondb.xml"
    if questiondb.exists() or any(source.glob("quiz_d2l_*.xml")):
        return source
    matches = sorted({
        path.parent
        for pattern in ("questiondb.xml", "quiz_d2l_*.xml")
        for path in source.rglob(pattern)
    })
    if len(matches) == 1:
        return matches[0]
    if not matches:
        raise FileNotFoundError(f"Could not find questiondb.xml or quiz_d2l_*.xml beneath {source}")
    raise FileExistsError(f"Found multiple quiz export folders beneath {source}; point directly to one export folder.")


def parse_quiz_files(
    export_dir: Path,
    indexes: dict[str, dict[str, list[LibraryQuestion]]],
    asset_index: dict[str, dict[str, list[Path]] | dict[str, Path]],
) -> tuple[list[dict[str, object]], list[dict[str, object]], list[dict[str, object]], list[dict[str, object]], list[dict[str, object]]]:
    question_rows: list[dict[str, object]] = []
    section_rows: list[dict[str, object]] = []
    quiz_rows: list[dict[str, object]] = []
    unresolved_rows: list[dict[str, object]] = []
    image_rows: list[dict[str, object]] = []

    for quiz_path in sorted(export_dir.glob("quiz_d2l_*.xml")):
        root = ET.parse(quiz_path).getroot()
        assessment = first_descendant(root, "assessment")
        if assessment is None:
            continue

        quiz_title = assessment.attrib.get("title", "") or quiz_path.name
        item_counter = 0
        section_records: list[dict[str, object]] = []

        def visit_section(
            section: ET.Element,
            ancestors: list[str],
            parent_ident: str,
            parent_title: str,
            source_locator: str,
            parent_source_locator: str,
        ) -> None:
            nonlocal item_counter
            section_ident = section.attrib.get("ident", "")
            section_title = section.attrib.get("title", "") or ("Quiz Root Wrapper" if section_ident == "CONTAINER_SECTION" else "(untitled section)")
            current_path = " > ".join([*ancestors, section_title]) if ancestors else section_title
            draw_count = section_draw_count(section)
            is_container = section_ident == "CONTAINER_SECTION"
            direct_question_nodes = [child for child in list(section) if local_name(child.tag) in {"item", "itemref"}]

            section_record = {
                "quiz_file": quiz_path.name,
                "quiz_title": quiz_title,
                "section_ident": section_ident,
                "section_title": section_title,
                "section_path": current_path,
                "section_source_locator": source_locator,
                "parent_section_ident": parent_ident,
                "parent_section_title": parent_title,
                "parent_section_source_locator": parent_source_locator,
                "section_level": len(ancestors) + 1,
                "is_container_section": "yes" if is_container else "no",
                "is_random_draw_section": "yes" if draw_count else "no",
                "draw_count": draw_count,
                "section_weight": section_weight(section),
                "quiz_section_item_count": len(direct_question_nodes),
                "matched_question_count": 0,
                "unmatched_question_count": 0,
                "distinct_pool_count": 0,
                "pool_titles": "",
                "pool_idents": "",
                "evidence_levels": "",
                "notes": "",
            }
            section_records.append(section_record)

            for section_item_index, question_node in enumerate(direct_question_nodes, start=1):
                item_counter += 1
                if local_name(question_node.tag) == "itemref":
                    match = match_itemref(question_node, indexes)
                    effective_item = match.library_question.item_element if match.library_question is not None else None
                    quiz_item_ident = ""
                    quiz_item_label = question_node.attrib.get("linkrefid", "")
                    quiz_item_title = match.library_question.question_title if match.library_question is not None else ""
                else:
                    match = match_question(question_node, indexes)
                    effective_item = question_node
                    quiz_item_ident = question_node.attrib.get("ident", "")
                    quiz_item_label = question_node.attrib.get("label", "")
                    quiz_item_title = question_node.attrib.get("title", "")
                library = match.library_question
                meta = qti_metadata(effective_item) if effective_item is not None else {}
                question_render = render_prompt_text(effective_item) if effective_item is not None else RenderedHtml(text="", formula_present=False, mathml_used=False)
                choice_infos = response_choice_infos(effective_item) if effective_item is not None else []
                matching_structure = extract_matching_structure(effective_item) if effective_item is not None else {"prompts": [], "options": []}
                fill_in_blank_structure = extract_fill_in_blank_structure(effective_item) if effective_item is not None else {"blank_count": 0, "blank_answers_text": ""}
                source_response_facts = extract_source_response_facts(effective_item)
                plain_text = question_render.text
                choices_text = " || ".join(str(info["text"]) for info in choice_infos if str(info["text"]))
                matching_prompts_text = " || ".join(
                    str(info["text"])
                    for info in matching_structure["prompts"]
                    if str(info["text"])
                )
                matching_options_text = " || ".join(
                    str(info["text"])
                    for info in matching_structure["options"]
                    if str(info["text"])
                )
                mathml_used = "yes" if (question_render.mathml_used or any(bool(info["mathml_used"]) for info in choice_infos)) else "no"
                formula_present = "yes" if (question_render.formula_present or any(bool(info["formula_present"]) for info in choice_infos)) else "no"
                answer_key = extract_correct_answer(effective_item) if effective_item is not None else blank_answer_key()
                asset_refs = extract_question_asset_refs(effective_item, export_dir, asset_index) if effective_item is not None else []
                question_image_paths = [asset.file_path for asset in asset_refs if asset.file_path]
                missing_image_refs = [asset.raw_ref for asset in asset_refs if asset.status != "resolved"]
                primary_asset = next((asset for asset in asset_refs if asset.status == "resolved" and asset.file_path), None)
                question_weight = itemref_points(question_node, library.question_weight if library is not None else "") if local_name(question_node.tag) == "itemref" else meta.get("qmd_weighting", "")

                if library is not None:
                    section_record["matched_question_count"] = int(section_record["matched_question_count"]) + 1
                else:
                    section_record["unmatched_question_count"] = int(section_record["unmatched_question_count"]) + 1
                    unresolved_rows.append({
                        "quiz_file": quiz_path.name,
                        "quiz_title": quiz_title,
                        "quiz_item_number": item_counter,
                        "section_ident": section_ident,
                        "section_title": section_title,
                        "section_source_locator": source_locator,
                        "quiz_item_label": quiz_item_label,
                        "quiz_item_title": quiz_item_title,
                        "question_type": meta.get("qmd_questiontype", ""),
                        "qmd_displayid": meta.get("qmd_displayid", ""),
                        "qmd_globalid": meta.get("qmd_globalid", ""),
                        "formula_present": formula_present,
                        "mathml_used": mathml_used,
                        "correct_answer": answer_key["correct_answer"],
                        "correct_answer_basis": answer_key["correct_answer_basis"],
                        "correct_response_ids": answer_key["correct_response_ids"],
                        "source_choice_display": source_choice_display(source_response_facts),
                        "question_image_count": len(asset_refs),
                        "question_image_refs": " || ".join(asset.raw_ref for asset in asset_refs),
                        "question_image_paths": " || ".join(question_image_paths),
                        "question_primary_image_path": primary_asset.file_path if primary_asset else "",
                        "question_primary_image_absolute_path": primary_asset.absolute_path if primary_asset else "",
                        "question_missing_image_refs": " || ".join(missing_image_refs),
                        "match_basis": match.match_basis,
                        "match_score": round(match.match_score, 3) if match.match_score is not None else "",
                        "match_candidates_json": json.dumps(match.candidates, ensure_ascii=False),
                        "note": match.note,
                        "choices_text": choices_text,
                        "matching_prompts_text": matching_prompts_text,
                        "matching_options_text": matching_options_text,
                        "fill_in_blank_count": fill_in_blank_structure["blank_count"],
                        "fill_in_blank_answers_text": fill_in_blank_structure["blank_answers_text"],
                        "question_text": plain_text,
                    })

                question_rows.append({
                    "quiz_file": quiz_path.name,
                    "quiz_title": quiz_title,
                    "quiz_item_number": item_counter,
                    "quiz_section_item_number": section_item_index,
                    "section_ident": section_ident,
                    "section_title": section_title,
                    "section_path": current_path,
                    "section_source_locator": source_locator,
                    "is_random_draw_section": "yes" if draw_count else "no",
                    "draw_count": draw_count,
                    "quiz_item_ident": quiz_item_ident,
                    "quiz_item_label": quiz_item_label,
                    "quiz_item_title": quiz_item_title,
                    "question_type": meta.get("qmd_questiontype", ""),
                    "question_weight": question_weight,
                    "qmd_displayid": meta.get("qmd_displayid", ""),
                    "qmd_globalid": meta.get("qmd_globalid", ""),
                    "formula_present": formula_present,
                    "mathml_used": mathml_used,
                    "question_image_count": len(asset_refs),
                    "question_image_refs": " || ".join(asset.raw_ref for asset in asset_refs),
                    "question_image_paths": " || ".join(question_image_paths),
                    "question_primary_image_path": primary_asset.file_path if primary_asset else "",
                    "question_primary_image_absolute_path": primary_asset.absolute_path if primary_asset else "",
                    "question_missing_image_refs": " || ".join(missing_image_refs),
                    "choices_text": choices_text,
                    "matching_prompts_text": matching_prompts_text,
                    "matching_options_text": matching_options_text,
                    "fill_in_blank_count": fill_in_blank_structure["blank_count"],
                    "fill_in_blank_answers_text": fill_in_blank_structure["blank_answers_text"],
                    "question_text": plain_text,
                    "correct_answer": answer_key["correct_answer"],
                    "correct_answer_basis": answer_key["correct_answer_basis"],
                    "correct_response_ids": answer_key["correct_response_ids"],
                    "source_response_facts": source_response_facts,
                    "source_choice_display": source_choice_display(source_response_facts),
                    "pool_ident": library.pool_ident if library else "",
                    "pool_title": library.pool_title if library else "",
                    "pool_path": library.pool_path if library else "",
                    "pool_question_ident": library.question_ident if library else "",
                    "pool_question_label": library.question_label if library else "",
                    "pool_question_title": library.question_title if library else "",
                    "pool_question_displayid": library.qmd_displayid if library else "",
                    "pool_question_globalid": library.qmd_globalid if library else "",
                    "evidence_level": match.evidence_level,
                    "match_basis": match.match_basis,
                    "match_score": round(match.match_score, 3) if match.match_score is not None else "",
                    "match_candidates_json": json.dumps(match.candidates, ensure_ascii=False),
                    "match_note": match.note,
                })
                for image_number, asset in enumerate(asset_refs, start=1):
                    image_rows.append({
                        "quiz_file": quiz_path.name,
                        "quiz_title": quiz_title,
                        "quiz_item_number": item_counter,
                        "quiz_item_label": quiz_item_label,
                        "quiz_item_title": quiz_item_title,
                        "section_ident": section_ident,
                        "section_title": section_title,
                        "section_path": current_path,
                        "section_source_locator": source_locator,
                        "question_type": meta.get("qmd_questiontype", ""),
                        "qmd_displayid": meta.get("qmd_displayid", ""),
                        "qmd_globalid": meta.get("qmd_globalid", ""),
                        "pool_title": library.pool_title if library else "",
                        "pool_path": library.pool_path if library else "",
                        "evidence_level": match.evidence_level,
                        "image_number": image_number,
                        "image_ref_original": asset.raw_ref,
                        "image_ref_normalized": asset.normalized_ref,
                        "image_status": asset.status,
                        "image_resolution_basis": asset.resolution_basis,
                        "image_file_path": asset.file_path,
                        "image_file_absolute_path": asset.absolute_path,
                    })

            for child_index, child_section in enumerate(
                child_elements(section, "section"), start=1
            ):
                visit_section(
                    child_section,
                    [*ancestors, section_title],
                    section_ident,
                    section_title,
                    f"{source_locator}/section[{child_index}]",
                    source_locator,
                )

        for section_index, section in enumerate(
            child_elements(assessment, "section"), start=1
        ):
            visit_section(
                section,
                [],
                "",
                "",
                f"section[{section_index}]",
                "",
            )

        quiz_item_rows = [row for row in question_rows if row["quiz_file"] == quiz_path.name]
        quiz_sections = [row for row in section_records if not (row["is_container_section"] == "yes" and int(row["quiz_section_item_count"]) == 0)]
        for record in quiz_sections:
            section_question_rows = [
                row for row in quiz_item_rows
                if row["section_source_locator"]
                == record["section_source_locator"]
            ]
            pool_titles = sorted({row["pool_title"] for row in section_question_rows if row["pool_title"]})
            pool_idents = sorted({row["pool_ident"] for row in section_question_rows if row["pool_ident"]})
            evidence_levels = Counter(row["evidence_level"] for row in section_question_rows if row["evidence_level"])
            record["distinct_pool_count"] = len(pool_titles)
            record["pool_titles"] = " | ".join(pool_titles)
            record["pool_idents"] = " | ".join(pool_idents)
            record["evidence_levels"] = " | ".join(f"{level}:{count}" for level, count in sorted(evidence_levels.items()))
            if record["is_random_draw_section"] == "yes" and len(pool_titles) != 1:
                record["notes"] = "Random-draw section does not resolve cleanly to a single pool."
            elif record["unmatched_question_count"]:
                record["notes"] = "One or more quiz items in this section did not match questiondb.xml."
            section_rows.append(record)

        evidence_levels = Counter(row["evidence_level"] for row in quiz_item_rows)
        pool_titles = sorted({row["pool_title"] for row in quiz_item_rows if row["pool_title"]})
        quiz_rows.append({
            "quiz_file": quiz_path.name,
            "quiz_title": quiz_title,
            "quiz_question_count": len(quiz_item_rows),
            "quiz_section_count": len(quiz_sections),
            "random_draw_section_count": sum(1 for row in quiz_sections if row["is_random_draw_section"] == "yes"),
            "distinct_pool_count": len(pool_titles),
            "pool_titles": " | ".join(pool_titles),
            "source_evidence_count": evidence_levels.get("source_evidence", 0),
            "inferred_exact_count": evidence_levels.get("inferred_exact", 0),
            "inferred_similarity_count": evidence_levels.get("inferred_similarity", 0),
            "unmatched_count": evidence_levels.get("unmatched", 0),
        })

    return quiz_rows, section_rows, question_rows, unresolved_rows, image_rows


def build_bank_rows(pool_rows: list[dict[str, object]], question_rows: list[dict[str, object]]) -> list[dict[str, object]]:
    bank_rows_by_key: dict[tuple[str, str, str], dict[str, object]] = {}
    pool_to_bank_key: dict[tuple[str, str], tuple[str, str, str]] = {}

    for pool in pool_rows:
        bank_ident = str(pool.get("bank_ident", "") or "")
        bank_title = str(pool.get("bank_title", "") or "")
        bank_path = str(pool.get("bank_path", "") or bank_title)
        key = (bank_ident, bank_title, bank_path)
        row = bank_rows_by_key.setdefault(key, {
            "bank_ident": bank_ident,
            "bank_title": bank_title,
            "bank_path": bank_path,
            "bank_pool_count": 0,
            "bank_library_question_count": 0,
            "bank_pool_titles": [],
            "bank_pool_idents": [],
            "used_pool_count": 0,
            "used_pool_titles": [],
            "mapped_quiz_question_count": 0,
            "distinct_quiz_count": 0,
            "used_by_quizzes": [],
            "source_evidence_count": 0,
            "inferred_exact_count": 0,
            "inferred_similarity_count": 0,
            "notes": "",
        })
        row["bank_pool_count"] = int(row["bank_pool_count"]) + 1
        row["bank_library_question_count"] = int(row["bank_library_question_count"]) + int(pool.get("question_count", 0) or 0)
        row["bank_pool_titles"].append(str(pool.get("pool_title", "") or ""))
        row["bank_pool_idents"].append(str(pool.get("pool_ident", "") or ""))
        pool_to_bank_key[(str(pool.get("pool_ident", "") or ""), str(pool.get("pool_path", "") or ""))] = key
        if bank_title == "Question Library Root Items":
            row["notes"] = "Library items are stored directly under <objectbank> with no enclosing question-library section."

    for question in question_rows:
        pool_ident = str(question.get("pool_ident", "") or "")
        pool_path = str(question.get("pool_path", "") or "")
        if not pool_ident:
            continue
        bank_key = pool_to_bank_key.get((pool_ident, pool_path))
        if bank_key is None:
            continue
        row = bank_rows_by_key[bank_key]
        row["mapped_quiz_question_count"] = int(row["mapped_quiz_question_count"]) + 1
        row["used_pool_titles"].append(str(question.get("pool_title", "") or ""))
        row["used_by_quizzes"].append(str(question.get("quiz_title", "") or ""))
        evidence_level = str(question.get("evidence_level", "") or "")
        if evidence_level == "source_evidence":
            row["source_evidence_count"] = int(row["source_evidence_count"]) + 1
        elif evidence_level == "inferred_exact":
            row["inferred_exact_count"] = int(row["inferred_exact_count"]) + 1
        elif evidence_level == "inferred_similarity":
            row["inferred_similarity_count"] = int(row["inferred_similarity_count"]) + 1

    bank_rows: list[dict[str, object]] = []
    for key in sorted(bank_rows_by_key.keys(), key=lambda item: (item[1].lower(), item[0].lower())):
        row = bank_rows_by_key[key]
        pool_titles = unique_preserve_order(sorted(title for title in row["bank_pool_titles"] if title))
        pool_idents = unique_preserve_order(sorted(ident for ident in row["bank_pool_idents"] if ident))
        used_pool_titles = unique_preserve_order(sorted(title for title in row["used_pool_titles"] if title))
        used_by_quizzes = unique_preserve_order(sorted(title for title in row["used_by_quizzes"] if title))
        bank_rows.append({
            "bank_ident": row["bank_ident"],
            "bank_title": row["bank_title"],
            "bank_path": row["bank_path"],
            "bank_pool_count": row["bank_pool_count"],
            "bank_library_question_count": row["bank_library_question_count"],
            "bank_pool_titles": " | ".join(pool_titles),
            "bank_pool_idents": " | ".join(pool_idents),
            "used_pool_count": len(used_pool_titles),
            "used_pool_titles": " | ".join(used_pool_titles),
            "mapped_quiz_question_count": row["mapped_quiz_question_count"],
            "distinct_quiz_count": len(used_by_quizzes),
            "used_by_quizzes": " | ".join(used_by_quizzes),
            "source_evidence_count": row["source_evidence_count"],
            "inferred_exact_count": row["inferred_exact_count"],
            "inferred_similarity_count": row["inferred_similarity_count"],
            "notes": row["notes"],
        })

    return bank_rows


def pool_lookup_rows(pool_rows: list[dict[str, object]]) -> dict[tuple[str, str], dict[str, object]]:
    lookup: dict[tuple[str, str], dict[str, object]] = {}
    for pool in pool_rows:
        key = (str(pool.get("pool_ident", "") or ""), str(pool.get("pool_path", "") or ""))
        lookup[key] = pool
    return lookup


def library_link_status_label(evidence_level: str) -> str:
    labels = {
        "source_evidence": "Direct library link",
        "inferred_exact": "Linked by identical content",
        "inferred_similarity": "Possible link based on similar content",
        "unmatched": "No reliable library link found",
    }
    return labels.get(evidence_level, evidence_level or "")


def reviewer_summary_unmatched_label() -> str:
    return "No reliable library link found"


def reviewer_library_location(question_row: dict[str, object], pool_lookup: dict[tuple[str, str], dict[str, object]]) -> str:
    pool_ident = str(question_row.get("pool_ident", "") or "")
    pool_path = str(question_row.get("pool_path", "") or "")
    if not pool_ident:
        return ""
    pool = pool_lookup.get((pool_ident, pool_path), {})
    bank_title = str(pool.get("bank_title", "") or "")
    pool_title = str(question_row.get("pool_title", "") or "")
    if bank_title and pool_title and bank_title != pool_title:
        location = f"{bank_title} > {pool_title}"
    else:
        location = bank_title or pool_title
    return "Question Library root-level items" if location == "Question Library Root Items" else location


def reviewer_section_label(section_title: str) -> str:
    return "" if section_title == "Quiz Root Wrapper" else section_title


def reviewer_note(question_row: dict[str, object]) -> str:
    notes: list[str] = []
    evidence_level = str(question_row.get("evidence_level", "") or "")
    match_note = str(question_row.get("match_note", "") or "").strip()
    if evidence_level == "unmatched":
        notes.append("No safe question-library location was found automatically.")
    elif evidence_level == "inferred_similarity":
        notes.append("The content matched strongly, but not by a direct library link.")
    elif "Multiple question-library items share this exact content inside one pool" in match_note:
        notes.append("This question matched one question-library location, but more than one library item in that location shares the same content.")
    elif match_note:
        notes.append(match_note)
    missing_image_refs = str(question_row.get("question_missing_image_refs", "") or "").strip()
    if missing_image_refs:
        notes.append(f"Missing image refs: {missing_image_refs}")
    return " ".join(notes)


def split_pipe_values(value: object) -> list[str]:
    text = str(value or "").strip()
    if not text:
        return []
    return [part.strip() for part in text.split(" || ") if part.strip()]


def reviewer_response_options(question_row: dict[str, object]) -> str:
    question_type = str(question_row.get("question_type", "") or "")
    display = question_row.get("source_choice_display")
    if isinstance(display, dict) and display.get("options"):
        choices = [_choice_display_text(option, bool(display["show_keys"])) for option in display["options"]]
        return "Items to order:\n" + "\n".join(f"- {choice}" for choice in choices) if question_type == "Ordering" else "\n".join(choices)
    if question_type == "Matching":
        prompts = split_pipe_values(question_row.get("matching_prompts_text", ""))
        options = split_pipe_values(question_row.get("matching_options_text", ""))
        blocks: list[str] = []
        if prompts:
            blocks.append("Prompt set:\n" + "\n".join(f"- {prompt}" for prompt in prompts))
        if options:
            blocks.append("Match options:\n" + "\n".join(f"- {option}" for option in options))
        return "\n\n".join(blocks)
    if question_type == "Ordering":
        choices = split_pipe_values(question_row.get("choices_text", ""))
        if not choices:
            return ""
        return "Items to order:\n" + "\n".join(f"- {choice}" for choice in choices)

    return "\n".join(split_pipe_values(question_row.get("choices_text", "")))


def reviewer_answer_key(question_row: dict[str, object]) -> str:
    display = question_row.get("source_choice_display")
    if isinstance(display, dict):
        if not display.get("recognized"):
            return "[Unresolved source scoring — review preserved XML facts]"
        if str(question_row.get("question_type")) == "Ordering":
            return "\n".join(f"{option['position']}. {_choice_display_text(option, bool(display['show_keys']))}" for option in sorted(display["options"], key=lambda option: option["position"]))
        return " || ".join(_choice_display_text(option, bool(display["show_keys"])) for option in display["options"] if option["correct"] is True)
    answer = str(question_row.get("correct_answer", "") or "").strip()
    if not answer:
        return ""
    if str(question_row.get("correct_answer_basis", "") or "") == "matching_pairs":
        return "\n".join(split_pipe_values(answer))
    if str(question_row.get("correct_answer_basis", "") or "") == "ordering_sequence":
        return "\n".join(f"{index}. {part}" for index, part in enumerate(split_pipe_values(answer), start=1))
    if str(question_row.get("question_type", "") or "") == "Fill in the Blanks":
        blank_answers = str(question_row.get("fill_in_blank_answers_text", "") or "").strip()
        if blank_answers:
            return "\n".join(split_pipe_values(blank_answers))
    return answer


def build_reviewer_quiz_summary_rows(payload: dict[str, object]) -> list[dict[str, object]]:
    return [
        {
            "quiz": row["quiz_title"],
            "question_count": row["quiz_question_count"],
            "question_library_locations": row["distinct_pool_count"],
            "direct_library_links": row["source_evidence_count"],
            "exact_content_matches": row["inferred_exact_count"],
            "probable_content_matches": row["inferred_similarity_count"],
            "no_safe_library_matches": row["unmatched_count"],
        }
        for row in payload["quiz_summary_rows"]
    ]


def _projection_occurrence_lookup(
    projection: dict[str, object] | None,
) -> dict[tuple[str, int], dict[str, object]]:
    candidates: dict[tuple[str, int], list[dict[str, object]]] = defaultdict(list)
    if not projection:
        return {}
    for occurrence in projection.get("question_occurrences", []):
        if not isinstance(occurrence, dict):
            continue
        quiz_file = str(occurrence.get("quiz_file") or "")
        item_number = occurrence.get("quiz_occurrence_order")
        if quiz_file and isinstance(item_number, int):
            candidates[(quiz_file, item_number)].append(occurrence)
    return {
        key: rows[0]
        for key, rows in candidates.items()
        if len(rows) == 1
    }


def _projection_entity_lookup(
    projection: dict[str, object] | None,
) -> dict[str, dict[str, object]]:
    if not projection:
        return {}
    return {
        str(row.get("question_key")): row
        for row in projection.get("question_entities", [])
        if isinstance(row, dict) and row.get("question_key")
    }


def _occurrence_question_key(occurrence: dict[str, object]) -> str:
    return str(
        occurrence.get("referenced_question_key")
        or occurrence.get("observed_question_entity_key")
        or ""
    )


def _occurrence_pool_text(occurrence: dict[str, object]) -> str:
    values: list[str] = []
    for membership in occurrence.get("pool_memberships", []):
        if not isinstance(membership, dict):
            continue
        value = str(
            membership.get("pool_path")
            or membership.get("pool_title")
            or membership.get("pool_key")
            or ""
        ).strip()
        if value and value not in values:
            values.append(value)
    return " | ".join(values)


def _first_canonical_question_row(entity: dict[str, object]) -> dict[str, object]:
    question = entity.get("question", {})
    if not isinstance(question, dict):
        return {}
    type_payload = question.get("type_payload", {})
    if not isinstance(type_payload, dict):
        return {}
    for raw in type_payload.get("raw_response_models", []):
        if not isinstance(raw, dict) or not str(raw.get("source_kind") or "").startswith(
            "canonical-extractor-row:"
        ):
            continue
        payload = raw.get("payload")
        if isinstance(payload, dict):
            return payload
    return {}


def _content_text(value: object) -> str:
    if isinstance(value, dict):
        return str(value.get("content") or "")
    return str(value or "")


def _model_response_options(question: dict[str, object]) -> str:
    payload = question.get("type_payload", {})
    if not isinstance(payload, dict):
        return ""
    lines: list[str] = []
    for option in payload.get("options", []):
        if not isinstance(option, dict):
            continue
        marker = "correct" if option.get("correct") is True else "option"
        lines.append(
            f"{option.get('option_key') or '?'} [{marker}]: "
            f"{_content_text(option.get('content'))}"
        )
    accepted = [
        str(row.get("value") or "")
        for row in payload.get("accepted_responses", [])
        if isinstance(row, dict) and row.get("value") not in (None, "")
    ]
    if accepted:
        lines.append("Accepted responses: " + " | ".join(accepted))
    pairs = []
    for row in payload.get("match_pairs", []):
        if not isinstance(row, dict):
            continue
        left = _content_text(row.get("prompt") or row.get("left") or row.get("source"))
        right = _content_text(row.get("option") or row.get("right") or row.get("target"))
        if left or right:
            pairs.append(f"{left} -> {right}".strip())
    if pairs:
        lines.append("Match pairs: " + " | ".join(pairs))
    order = [
        _content_text(value)
        for value in payload.get("correct_order", [])
        if _content_text(value)
    ]
    if order:
        lines.append("Order: " + " -> ".join(order))
    return "\n".join(lines)


def _model_answer_key(question: dict[str, object]) -> str:
    payload = question.get("type_payload", {})
    if not isinstance(payload, dict):
        return ""
    values: list[str] = []
    correct_options = [
        _content_text(row.get("content"))
        for row in payload.get("options", [])
        if isinstance(row, dict) and row.get("correct") is True
    ]
    if correct_options:
        values.append("Correct: " + " | ".join(correct_options))
    accepted = [
        str(row.get("value") or "")
        for row in payload.get("accepted_responses", [])
        if isinstance(row, dict) and row.get("value") not in (None, "")
    ]
    if accepted:
        values.append("Accepted: " + " | ".join(accepted))
    manual = _content_text(payload.get("manual_answer_key"))
    if manual:
        values.append(manual)
    return "\n".join(values)


def _question_feedback_text(question: dict[str, object]) -> str:
    rows: list[str] = []
    for feedback in question.get("feedback", []):
        if not isinstance(feedback, dict):
            continue
        content = _content_text(feedback.get("content"))
        if content:
            rows.append(f"{feedback.get('channel') or 'feedback'}: {content}")
    return "\n".join(rows)


def _question_identity_fields(question: dict[str, object]) -> dict[str, str]:
    identity = question.get("identity", {})
    if not isinstance(identity, dict):
        identity = {}
    aliases = [
        row
        for row in identity.get("source_aliases", [])
        if isinstance(row, dict) and row.get("value")
    ]
    permanent_code = str(identity.get("permanent_code") or "")
    alias_text = " | ".join(
        f"{row.get('namespace') or 'source'}={row.get('value')}"
        for row in aliases
    )
    display_identifier = (
        permanent_code
        or (str(aliases[0].get("value")) if aliases else "")
        or str(question.get("entity_key") or "")
    )
    return {
        "display_identifier": display_identifier,
        "permanent_code": permanent_code,
        "source_aliases": alias_text,
    }


def build_reviewer_entity_rows(
    projection: dict[str, object],
) -> list[dict[str, object]]:
    occurrences = {
        str(row.get("occurrence_key")): row
        for row in projection.get("question_occurrences", [])
        if isinstance(row, dict) and row.get("occurrence_key")
    }
    from quiz_normalization import authored_variant_row_digest

    rows: list[dict[str, object]] = []
    for entity in projection.get("question_entities", []):
        if not isinstance(entity, dict):
            continue
        question = entity.get("question", {})
        if not isinstance(question, dict):
            continue
        raw = _first_canonical_question_row(entity)
        # A collapsed entity would otherwise show one authoritative answer key
        # for several different authored questions.
        variant_digests = {
            authored_variant_row_digest(record["payload"])
            for record in (question.get("type_payload") or {}).get(
                "raw_response_models"
            )
            or []
            if str(record.get("source_kind", "")).startswith(
                "canonical-extractor-row:"
            )
            and isinstance(record.get("payload"), dict)
        }
        linked_occurrences = [
            occurrences[key]
            for key in entity.get("occurrence_keys", [])
            if key in occurrences
        ]
        quiz_labels = unique_preserve_order(
            str(row.get("quiz_title") or row.get("quiz_key") or "")
            for row in linked_occurrences
            if row.get("quiz_title") or row.get("quiz_key")
        )
        location_labels = unique_preserve_order(
            " / ".join(
                part
                for part in (
                    str(row.get("quiz_title") or ""),
                    str(row.get("section_path") or "quiz root"),
                    f"item {row.get('quiz_occurrence_order')}"
                    if row.get("quiz_occurrence_order") is not None
                    else "order unresolved",
                )
                if part
            )
            for row in linked_occurrences
        )
        pools = unique_preserve_order(
            _occurrence_pool_text(row)
            for row in linked_occurrences
            if _occurrence_pool_text(row)
        )
        identity_fields = _question_identity_fields(question)
        scoring = question.get("scoring", {})
        support = question.get("build_support", {})
        choice_unresolved = question.get("type_payload", {}).get("extensions", {}).get("coursecraft.source_choice_projection", {}).get("state") == "unresolved"
        prompt = str(raw.get("question_text") or "") or _content_text(
            question.get("prompt")
        )
        rows.append(
            {
                "display_identifier": identity_fields["display_identifier"],
                "title": question.get("title"),
                "question_type": question.get("source_kind") or question.get("kind"),
                "normalized_kind": question.get("kind"),
                "occurrence_count": entity.get("occurrence_count"),
                "question_entity_key": entity.get("question_key"),
                "permanent_code": identity_fields["permanent_code"],
                "source_aliases": identity_fields["source_aliases"],
                "resolved_occurrence_count": entity.get("resolved_occurrence_count"),
                "used_in_quizzes": " | ".join(quiz_labels),
                "occurrence_locations": "\n".join(location_labels),
                "question_library_locations": " | ".join(pools),
                "question_text": prompt,
                "response_options": (
                    reviewer_response_options(raw)
                    if raw
                    else _model_response_options(question)
                ),
                "answer_key": (
                    "[Unresolved source choice evidence — review individual occurrences]"
                    if choice_unresolved
                    else reviewer_answer_key(raw) if raw else _model_answer_key(question)
                ),
                "answer_key_state": (
                    "DO NOT APPROVE - source choice evidence is unresolved"
                    if choice_unresolved
                    else "authoritative"
                    if len(variant_digests) <= 1
                    else (
                        f"DO NOT APPROVE - {len(variant_digests)} authored variants "
                        "were folded into this record"
                    )
                ),
                "placement_points": (
                    " / ".join(
                        str(value)
                        for value in scoring.get("extensions", {}).get(
                            "observed_maximum_points", []
                        )
                    )
                    if isinstance(scoring, dict)
                    else ""
                ),
                "feedback": _question_feedback_text(question),
                "scoring_state": scoring.get("state")
                if isinstance(scoring, dict)
                else None,
                "scoring_mode": scoring.get("mode")
                if isinstance(scoring, dict)
                else None,
                "maximum_points": scoring.get("maximum_points")
                if isinstance(scoring, dict)
                else None,
                "extraction_storage": question.get("extensions", {}).get("storage")
                if isinstance(question.get("extensions"), dict)
                else None,
                "build_support": support.get("level")
                if isinstance(support, dict)
                else None,
                "diagnostic_ids": " | ".join(
                    str(value) for value in question.get("diagnostic_ids", [])
                ),
            }
        )
    return rows


def build_reviewer_occurrence_rows(
    projection: dict[str, object],
) -> list[dict[str, object]]:
    entities = _projection_entity_lookup(projection)
    rows: list[dict[str, object]] = []
    for occurrence in projection.get("question_occurrences", []):
        if not isinstance(occurrence, dict):
            continue
        question_key = _occurrence_question_key(occurrence)
        entity = entities.get(question_key, {})
        question = entity.get("question", {}) if isinstance(entity, dict) else {}
        if not isinstance(question, dict):
            question = {}
        raw = _first_canonical_question_row(entity) if entity else {}
        selection = occurrence.get("selection_context", {})
        if not isinstance(selection, dict):
            selection = {}
        support = question.get("build_support", {})
        identity_fields = _question_identity_fields(question)
        match_code = str(occurrence.get("library_match_status") or "unknown")
        memberships = [
            row
            for row in occurrence.get("pool_memberships", [])
            if isinstance(row, dict)
        ]
        pool_candidate_labels = unique_preserve_order(
            str(
                row.get("pool_title")
                or row.get("pool_path")
                or row.get("pool_key")
                or ""
            )
            for row in memberships
            if row.get("pool_title") or row.get("pool_path") or row.get("pool_key")
        )
        candidate_rows = [
            candidate
            for membership in memberships
            for candidate in membership.get("candidates", [])
            if isinstance(candidate, dict)
        ]
        question_candidate_labels = unique_preserve_order(
            "{}{}".format(
                candidate.get("qmd_displayid")
                or candidate.get("question_label")
                or candidate.get("qmd_globalid")
                or candidate.get("question_ident")
                or "Unnamed library candidate",
                (
                    f" (score {candidate.get('score')})"
                    if candidate.get("score") is not None
                    else ""
                ),
            )
            for candidate in candidate_rows
        )
        candidate_labels = question_candidate_labels or pool_candidate_labels
        numeric_scores = [
            float(candidate["score"])
            for candidate in candidate_rows
            if isinstance(candidate.get("score"), (int, float))
        ]
        best_score_tie = bool(numeric_scores) and numeric_scores.count(
            max(numeric_scores)
        ) > 1
        ambiguity_codes = {
            str(row.get("code") or "")
            for row in occurrence.get("projection_diagnostics", [])
            if isinstance(row, dict)
        }
        is_ambiguous = match_code == "ambiguous" or best_score_tie or bool(
            ambiguity_codes
            & {
                "occurrence_pool_ambiguous",
                "draw_pool_context_ambiguous",
            }
        )
        match_notes = unique_preserve_order(
            str(row.get("note") or "")
            for row in memberships
            if row.get("note")
        )
        if is_ambiguous:
            match_explanation = (
                "More than one library candidate remains plausible. A person "
                "must choose or decline; the Binder will not guess."
            )
        elif match_code == "unresolved":
            match_explanation = (
                "No defensible question-library location was resolved; the "
                "question remains visible for review."
            )
        elif match_code == "inferred_similarity":
            match_explanation = (
                "This is a content-based proposal rather than a direct library "
                "link. A person must confirm or decline it."
            )
        else:
            match_explanation = " | ".join(match_notes)
        rows.append(
            {
                "quiz_order": occurrence.get("quiz_order"),
                "quiz": occurrence.get("quiz_title") or occurrence.get("quiz_key"),
                "question_number": occurrence.get("quiz_occurrence_order"),
                "section_order": occurrence.get("section_order"),
                "section": occurrence.get("section_path")
                or occurrence.get("section_title")
                or "quiz root",
                "section_question_number": occurrence.get("section_occurrence_order"),
                "occurrence_key": occurrence.get("occurrence_key"),
                "question_entity_key": question_key,
                "display_identifier": identity_fields["display_identifier"],
                "source_aliases": identity_fields["source_aliases"],
                "placement": occurrence.get("placement_kind"),
                "relationship_status": occurrence.get("relationship_status"),
                "selection_mode": selection.get("mode"),
                "draw_count": selection.get("draw_count"),
                "candidate_pool_size": selection.get("candidate_pool_size"),
                "pool_context_basis": selection.get("pool_context_basis"),
                "question_library_location": _occurrence_pool_text(occurrence),
                "library_match_status": match_label(match_code),
                "library_match_code": match_code,
                "library_match_candidates": " | ".join(candidate_labels),
                "library_match_explanation": match_explanation,
                "points": occurrence.get("points"),
                "question_type": question.get("source_kind") or question.get("kind"),
                "question_text": str(raw.get("question_text") or "")
                or _content_text(question.get("prompt")),
                "response_options": (
                    reviewer_response_options(raw)
                    if raw
                    else _model_response_options(question)
                ),
                "answer_key": (
                    reviewer_answer_key(raw) if raw else _model_answer_key(question)
                ),
                "feedback": _question_feedback_text(question),
                "build_support": support.get("level")
                if isinstance(support, dict)
                else None,
                "diagnostic_ids": " | ".join(
                    str(value) for value in occurrence.get("diagnostic_ids", [])
                ),
            }
        )
    return rows


def build_reviewer_question_rows(
    payload: dict[str, object],
    projection: dict[str, object] | None = None,
) -> list[dict[str, object]]:
    pool_lookup = pool_lookup_rows(payload["pool_rows"])
    occurrence_lookup = _projection_occurrence_lookup(projection)
    entity_lookup = _projection_entity_lookup(projection)
    rows: list[dict[str, object]] = []
    for question in payload["quiz_question_rows"]:
        occurrence = occurrence_lookup.get(
            (str(question.get("quiz_file") or ""), int(question["quiz_item_number"]))
        )
        question_key = _occurrence_question_key(occurrence) if occurrence else ""
        row = {
            "quiz": question["quiz_title"],
            "question_number": question["quiz_item_number"],
            "quiz_section (blank = quiz root)": reviewer_section_label(str(question["section_title"])),
            "question_type": question["question_type"],
            "proposed_permanent_code": "",
            "proposed_scoring_mode": "",
            "question_text": question["question_text"],
            "revised_question_text": "",
            "image_link": question["question_primary_image_path"],
            "image_link_absolute_path": question.get(
                "question_primary_image_review_copy_path"
            )
            or question["question_primary_image_absolute_path"],
            "response_options": reviewer_response_options(question),
            "revised_response_options": "",
            "answer_key": reviewer_answer_key(question),
            "revised_answer_key": "",
            "question_library_location": reviewer_library_location(question, pool_lookup),
            "revised_question_library_location": "",
            "library_link_status": library_link_status_label(str(question["evidence_level"])),
            "image_count": question["question_image_count"],
            "note": reviewer_note(question),
            "revision_reason": "",
            "proposed_by": "",
            "proposed_at": "",
            "approval_status": "",
            "approved_by": "",
            "approved_at": "",
            "reviewer_note": "",
        }
        if projection is not None:
            entity = entity_lookup.get(question_key, {})
            selection = occurrence.get("selection_context", {}) if occurrence else {}
            row = {
                "quiz": row["quiz"],
                "question_number": row["question_number"],
                "quiz_section (blank = quiz root)": row[
                    "quiz_section (blank = quiz root)"
                ],
                "occurrence_key": occurrence.get("occurrence_key") if occurrence else "",
                "question_entity_key": question_key,
                "question_occurrence_count": entity.get("occurrence_count", ""),
                "placement": occurrence.get("placement_kind") if occurrence else "",
                "relationship_status": occurrence.get("relationship_status")
                if occurrence
                else "",
                "selection_mode": selection.get("mode")
                if isinstance(selection, dict)
                else "",
                "draw_count": selection.get("draw_count")
                if isinstance(selection, dict)
                else "",
                "candidate_pool_size": selection.get("candidate_pool_size")
                if isinstance(selection, dict)
                else "",
                "pool_context_basis": selection.get("pool_context_basis")
                if isinstance(selection, dict)
                else "",
                **{
                    key: value
                    for key, value in row.items()
                    if key
                    not in {
                        "quiz",
                        "question_number",
                        "quiz_section (blank = quiz root)",
                    }
                },
            }
        rows.append(row)
    return rows


def build_reviewer_image_rows(payload: dict[str, object]) -> list[dict[str, object]]:
    pool_lookup = pool_lookup_rows(payload["pool_rows"])
    question_lookup = {
        (str(question["quiz_title"]), str(question["quiz_item_number"])): question
        for question in payload["quiz_question_rows"]
    }
    rows: list[dict[str, object]] = []
    for image in payload["question_image_rows"]:
        question = question_lookup.get((str(image["quiz_title"]), str(image["quiz_item_number"])), {})
        rows.append({
            "quiz": image["quiz_title"],
            "question_number": image["quiz_item_number"],
            "question_text": str(question.get("question_text", "") or ""),
            "question_library_location": reviewer_library_location(question, pool_lookup) if question else "",
            "image_file": image["image_file_path"],
            "image_file_absolute_path": image.get("review_copy_path")
            or image["image_file_absolute_path"],
            "image_status": image["image_status"],
        })
    return rows


def build_reviewer_library_rows(payload: dict[str, object]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for bank in payload["bank_rows"]:
        rows.append({
            "question_library_location": "Question Library root-level items" if bank["bank_title"] == "Question Library Root Items" else bank["bank_title"],
            "library_sections": bank["bank_pool_count"],
            "library_questions": bank["bank_library_question_count"],
            "mapped_quiz_questions": bank["mapped_quiz_question_count"],
            "quiz_count": bank["distinct_quiz_count"],
            "used_in_quizzes": bank["used_by_quizzes"],
            "direct_library_links": bank["source_evidence_count"],
            "exact_content_matches": bank["inferred_exact_count"],
            "probable_content_matches": bank["inferred_similarity_count"],
            "note": bank["notes"],
        })
    return rows


def _library_question_identity(
    row: dict[str, object], *, matched_prefix: str = ""
) -> list[tuple[str, str]]:
    """Return strongest-first source identifiers for a question-library row."""
    if matched_prefix:
        field_pairs = (
            ("globalid", f"{matched_prefix}globalid"),
            ("displayid", f"{matched_prefix}displayid"),
            ("ident", f"{matched_prefix}ident"),
            ("label", f"{matched_prefix}label"),
        )
    else:
        field_pairs = (
            ("globalid", "qmd_globalid"),
            ("displayid", "qmd_displayid"),
            ("ident", "question_ident"),
            ("label", "question_label"),
        )
    return [
        (namespace, str(row.get(field_name) or "").strip())
        for namespace, field_name in field_pairs
        if str(row.get(field_name) or "").strip()
    ]


def _all_questions_feedback(row: dict[str, object]) -> str:
    labels = (
        ("General", "general_feedback"),
        ("Correct", "correct_feedback"),
        ("Incorrect", "incorrect_feedback"),
        ("Answer-specific", "answer_specific_feedback"),
        ("Notes", "feedback_notes"),
    )
    return "\n".join(
        f"{label}: {str(row.get(field_name) or '').strip()}"
        for label, field_name in labels
        if str(row.get(field_name) or "").strip()
    )


def _all_questions_image_fields(row: dict[str, object]) -> dict[str, object]:
    """Project source image evidence into filterable reviewer columns."""
    image_references = split_pipe_values(row.get("question_image_refs"))
    resolved_image_paths = split_pipe_values(row.get("question_image_paths"))
    included_image_paths = split_pipe_values(
        row.get("question_image_review_copy_paths")
    )
    missing_image_references = split_pipe_values(
        row.get("question_missing_image_refs")
    )
    try:
        image_reference_count = int(row.get("question_image_count") or 0)
    except (TypeError, ValueError):
        image_reference_count = 0
    if image_reference_count == 0 and image_references:
        image_reference_count = len(image_references)

    if image_reference_count == 0:
        image_status = "none"
    elif missing_image_references and resolved_image_paths:
        image_status = "partially resolved"
    elif missing_image_references:
        image_status = "missing"
    else:
        image_status = "resolved"

    if image_reference_count == 0:
        image_included = "none"
    elif not included_image_paths:
        image_included = "no"
    elif len(included_image_paths) < image_reference_count:
        image_included = "partially"
    else:
        image_included = "yes"

    primary_source_path = str(
        row.get("question_primary_image_path") or ""
    ).strip()
    primary_review_copy_path = included_image_paths[0] if included_image_paths else ""
    primary_image = primary_review_copy_path or primary_source_path
    primary_image_link_target = primary_review_copy_path or str(
        row.get("question_primary_image_absolute_path") or ""
    ).strip()

    return {
        "has_image": "yes" if image_reference_count else "no",
        "image_reference_count": image_reference_count,
        "image_references": " || ".join(image_references),
        "image_status": image_status,
        "image_included": image_included,
        "primary_image": primary_image,
        "primary_image_link_target": primary_image_link_target,
        "resolved_image_paths": " || ".join(resolved_image_paths),
        "included_image_paths": " || ".join(included_image_paths),
        "missing_image_references": " || ".join(missing_image_references),
    }


def _all_questions_quiz_locations(rows: list[dict[str, object]]) -> list[str]:
    return unique_preserve_order(
        " / ".join(
            part
            for part in (
                str(row.get("quiz_title") or ""),
                reviewer_section_label(str(row.get("section_path") or ""))
                or "quiz root",
                (
                    f"item {row.get('quiz_item_number')}"
                    if row.get("quiz_item_number") not in (None, "")
                    else ""
                ),
            )
            if part
        )
        for row in rows
    )


def build_reviewer_all_question_rows(
    payload: dict[str, object],
) -> list[dict[str, object]]:
    """Build one complete source inventory: library records plus quiz-native items."""
    library_rows = [
        row
        for row in payload.get("library_question_rows", [])
        if isinstance(row, dict)
    ]
    quiz_rows = [
        row
        for row in payload.get("quiz_question_rows", [])
        if isinstance(row, dict)
    ]

    library_indexes: dict[tuple[str, str], list[int]] = defaultdict(list)
    for index, row in enumerate(library_rows):
        for identity in _library_question_identity(row):
            library_indexes[identity].append(index)

    library_uses: dict[int, list[dict[str, object]]] = defaultdict(list)
    native_uses: dict[tuple[str, str], list[dict[str, object]]] = defaultdict(list)
    for row in quiz_rows:
        identities = _library_question_identity(row, matched_prefix="pool_question_")
        matched_index: int | None = None
        for identity in identities:
            candidates = library_indexes.get(identity, [])
            if len(candidates) == 1:
                matched_index = candidates[0]
                break
        if matched_index is not None:
            library_uses[matched_index].append(row)
            continue

        native_identities = _library_question_identity(row)
        if native_identities:
            native_key = native_identities[0]
        else:
            native_key = (
                "quiz_locator",
                "{}#{}".format(
                    row.get("quiz_file") or row.get("quiz_title") or "quiz",
                    row.get("quiz_item_number") or row.get("quiz_item_label") or "item",
                ),
            )
        native_uses[native_key].append(row)

    inventory: list[dict[str, object]] = []
    for index, library_row in enumerate(library_rows):
        uses = library_uses.get(index, [])
        evidence_levels = {
            str(row.get("evidence_level") or "") for row in uses
        }
        if "source_evidence" in evidence_levels:
            classification = "Used in quiz — direct library link"
            confidence = "direct source evidence"
        elif "inferred_exact" in evidence_levels:
            classification = "Used in quiz — linked by identical content"
            confidence = "inferred exact content match"
        elif "inferred_similarity" in evidence_levels:
            classification = (
                "Possible quiz use — similar content; review recommended"
            )
            confidence = "inferred similarity; review before treating as linked"
        else:
            classification = "Library only — not used in exported quizzes"
            confidence = "no quiz placement found in this export"

        quiz_labels = unique_preserve_order(
            str(row.get("quiz_title") or row.get("quiz_file") or "")
            for row in uses
            if row.get("quiz_title") or row.get("quiz_file")
        )
        library_location = str(library_row.get("pool_path") or "").strip()
        if not library_location:
            library_location = str(
                library_row.get("bank_path")
                or library_row.get("bank_title")
                or library_row.get("pool_title")
                or "Question Library root-level items"
            ).strip()
        inventory.append(
            {
                "classification": classification,
                "source_type": "question library",
                "link_confidence": confidence,
                "quiz_occurrence_count": len(uses),
                "direct_placement_count": sum(
                    1
                    for row in uses
                    if str(row.get("is_random_draw_section") or "") != "yes"
                ),
                "random_draw_candidate_count": sum(
                    1
                    for row in uses
                    if str(row.get("is_random_draw_section") or "") == "yes"
                ),
                "used_in_quizzes": " | ".join(quiz_labels),
                "quiz_locations": "\n".join(_all_questions_quiz_locations(uses)),
                "library_location": library_location,
                "library_bank": library_row.get("bank_title")
                or library_row.get("library_bank_title"),
                "library_pool": library_row.get("pool_title")
                or library_row.get("library_pool_title"),
                "display_identifier": library_row.get("qmd_displayid"),
                "question_label": library_row.get("question_label"),
                "question_title": library_row.get("question_title"),
                "question_type": library_row.get("question_type"),
                "question_text": library_row.get("question_text"),
                "response_options": reviewer_response_options(library_row),
                "answer_key": reviewer_answer_key(library_row),
                "answer_key_basis": library_row.get("correct_answer_basis"),
                "points": library_row.get("question_weight"),
                "feedback": _all_questions_feedback(library_row),
                **_all_questions_image_fields(library_row),
                "payload_state": library_row.get("payload_state"),
                "source_question_ident": library_row.get("question_ident"),
                "source_global_id": library_row.get("qmd_globalid"),
            }
        )

    for native_rows in native_uses.values():
        first = native_rows[0]
        quiz_labels = unique_preserve_order(
            str(row.get("quiz_title") or row.get("quiz_file") or "")
            for row in native_rows
            if row.get("quiz_title") or row.get("quiz_file")
        )
        inventory.append(
            {
                "classification": "Quiz only — no library link found",
                "source_type": "quiz native / unlinked",
                "link_confidence": "no safe question-library match",
                "quiz_occurrence_count": len(native_rows),
                "direct_placement_count": sum(
                    1
                    for row in native_rows
                    if str(row.get("is_random_draw_section") or "") != "yes"
                ),
                "random_draw_candidate_count": sum(
                    1
                    for row in native_rows
                    if str(row.get("is_random_draw_section") or "") == "yes"
                ),
                "used_in_quizzes": " | ".join(quiz_labels),
                "quiz_locations": "\n".join(
                    _all_questions_quiz_locations(native_rows)
                ),
                "library_location": "",
                "library_bank": "",
                "library_pool": "",
                "display_identifier": first.get("qmd_displayid"),
                "question_label": first.get("quiz_item_label"),
                "question_title": first.get("quiz_item_title"),
                "question_type": first.get("question_type"),
                "question_text": first.get("question_text"),
                "response_options": reviewer_response_options(first),
                "answer_key": reviewer_answer_key(first),
                "answer_key_basis": first.get("correct_answer_basis"),
                "points": " | ".join(
                    unique_preserve_order(
                        str(row.get("question_weight") or "")
                        for row in native_rows
                        if row.get("question_weight") not in (None, "")
                    )
                ),
                "feedback": _all_questions_feedback(first),
                **_all_questions_image_fields(first),
                "payload_state": "parsed from quiz XML",
                "source_question_ident": first.get("quiz_item_ident"),
                "source_global_id": first.get("qmd_globalid"),
            }
        )

    return inventory


def build_reviewer_unresolved_rows(
    payload: dict[str, object],
    projection: dict[str, object] | None = None,
) -> list[dict[str, object]]:
    occurrence_lookup = _projection_occurrence_lookup(projection)
    rows: list[dict[str, object]] = []
    for question in payload["unresolved_rows"]:
        occurrence = occurrence_lookup.get(
            (str(question.get("quiz_file") or ""), int(question["quiz_item_number"]))
        )
        row = {
            "quiz": question["quiz_title"],
            "question_number": question["quiz_item_number"],
            "quiz_section (blank = quiz root)": reviewer_section_label(str(question["section_title"])),
            "question_type": question["question_type"],
            "proposed_permanent_code": "",
            "proposed_scoring_mode": "",
            "question_text": question["question_text"],
            "revised_question_text": "",
            "response_options": reviewer_response_options(question),
            "revised_response_options": "",
            "answer_key": reviewer_answer_key(question),
            "revised_answer_key": "",
            "primary_image": question["question_primary_image_path"],
            "primary_image_absolute_path": question.get(
                "question_primary_image_review_copy_path"
            )
            or question["question_primary_image_absolute_path"],
            "why_it_needs_review": "No safe question-library location was found automatically.",
            "revision_reason": "",
            "proposed_by": "",
            "proposed_at": "",
            "approval_status": "",
            "approved_by": "",
            "approved_at": "",
            "reviewer_note": "",
        }
        if projection is not None:
            row = {
                "quiz": row["quiz"],
                "question_number": row["question_number"],
                "quiz_section (blank = quiz root)": row[
                    "quiz_section (blank = quiz root)"
                ],
                "occurrence_key": occurrence.get("occurrence_key") if occurrence else "",
                "question_entity_key": _occurrence_question_key(occurrence)
                if occurrence
                else "",
                **{
                    key: value
                    for key, value in row.items()
                    if key
                    not in {
                        "quiz",
                        "question_number",
                        "quiz_section (blank = quiz root)",
                    }
                },
            }
        rows.append(row)
    return rows


def build_reviewer_settings_rows(model: dict[str, object] | None) -> list[dict[str, object]]:
    if not model:
        return []
    quiz_titles = {
        str(row["entity_key"]): str(row.get("title") or "Untitled quiz")
        for row in model.get("quizzes", [])
    }
    rows: list[dict[str, object]] = []
    duplicate_ordinals: Counter[str] = Counter()
    for observation in model.get("settings_observations", []):
        target = str(observation["target_entity_key"])
        identity_payload = {
            "target_entity_key": target,
            "name": observation.get("name"),
            "state": observation.get("state"),
            "raw_value": observation.get("raw_value"),
            "normalized_value": observation.get("normalized_value"),
            "source_evidence_keys": observation.get("source_evidence_keys", []),
            "extensions": observation.get("extensions", {}),
        }
        identity_digest = hashlib.sha256(
            json.dumps(
                identity_payload,
                ensure_ascii=False,
                sort_keys=True,
                separators=(",", ":"),
                default=str,
            ).encode("utf-8")
        ).hexdigest()[:24]
        duplicate_ordinals[identity_digest] += 1
        observation_key = (
            f"setting.observation.{identity_digest}."
            f"{duplicate_ordinals[identity_digest]:03d}"
        )
        value = observation.get("normalized_value")
        if isinstance(value, (dict, list)):
            value = json.dumps(value, ensure_ascii=False, sort_keys=True)
        rows.append(
            {
                "quiz": quiz_titles.get(target, target),
                "target_entity_key": target,
                "setting_observation_key": observation_key,
                "setting": observation["name"],
                "observed_state": observation["state"],
                "observed_value": value,
                "proposed_value": "",
                "revision_reason": "",
                "proposed_by": "",
                "proposed_at": "",
                "approval_status": "",
                "approved_by": "",
                "approved_at": "",
                "reviewer_note": "",
            }
        )
    return rows


def write_json(output_path: Path, payload: dict[str, object]) -> None:
    output_path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def markdown_report(
    source_dir: Path,
    payload: dict[str, object],
    created_at: str,
    script_path: Path,
) -> str:
    summary = payload["summary"]
    quiz_rows = payload["quiz_summary_rows"]
    unresolved_rows = payload["unresolved_rows"]
    lines = [
        f"# Quiz Pool Review — {payload['export_label']}",
        "",
        "## Scope",
        "",
        f"- Date: `{created_at}`",
        "- Mode: `extraction + review`",
        "- Goal: extract quiz questions by quiz and indicate question-library pool sources for reviewer inspection.",
        f"- Source export: `{source_dir}`",
        "",
        "## Inputs",
        "",
        f"- `{source_dir / 'questiondb.xml'}`",
        f"- `{source_dir / 'imsmanifest.xml'}`",
        f"- `{source_dir}` / `quiz_d2l_*.xml`",
        f"- `{script_path}`",
        "",
        "## Evidence",
        "",
        f"- Quiz files parsed: **{summary['quiz_count']}**",
        f"- Question-library banks parsed: **{summary['bank_count']}**",
        f"- Question-library pools parsed: **{summary['pool_count']}**",
        f"- Quiz question rows extracted: **{summary['quiz_question_count']}**",
        f"- Pool mappings by direct source evidence: **{summary['source_evidence_count']}**",
        f"- Pool mappings by exact inferred question match: **{summary['inferred_exact_count']}**",
        f"- Pool mappings by similarity fallback: **{summary['inferred_similarity_count']}**",
        f"- Unmatched quiz items: **{summary['unmatched_count']}**",
        f"- Quiz question rows with image refs: **{summary['question_rows_with_images']}**",
        f"- Image refs resolved to packaged files: **{summary['resolved_image_reference_count']} / {summary['image_reference_count']}**",
        "- Observed quiz pattern: populated quiz sections often contain inline item copies or generic random-draw containers; some quizzes also use direct `<itemref>` references into `questiondb.xml`, and a smaller number of inline items still preserve direct `qmd_displayid` / `qmd_globalid` evidence.",
        "",
        "## Inference",
        "",
        "- Pool mappings labeled `inferred_exact` or `inferred_similarity` are inferred from quiz item content matching against `questiondb.xml`; `source_evidence` rows come from direct `<itemref>` references or direct item-ID joins via `qmd_displayid` / `qmd_globalid`.",
        "- In this export, most course quiz sections resolve cleanly to a single question-library pool; unmatched items cluster in standalone quizzes that do not appear in `questiondb.xml`.",
        "",
        "## Tooling And Outputs",
        "",
        f"- Extractor script: `{script_path}`",
        "- reviewer-facing outputs written to `workspace/review/`",
        "",
        "## Unresolved Questions",
        "",
    ]

    if unresolved_rows:
        for row in unresolved_rows[:12]:
            lines.append(
                f"- `{row['quiz_title']}` / `{row['quiz_item_label']}` — {row['note']}"
            )
        if len(unresolved_rows) > 12:
            lines.append(f"- Additional unresolved rows in workbook/json: `{len(unresolved_rows) - 12}` more")
    else:
        lines.append("- none")

    lines.extend([
        "",
        "## Next Step",
        "",
        f"- Review the workbook `Quiz Sections` sheet first to confirm each quiz section aligns with the expected pool, then use `{UNRESOLVED_SHEET_TITLE}` for any items whose bank match remains uncertain.",
        "",
        "## Quiz Summary",
        "",
    ])

    for row in quiz_rows:
        lines.append(
            f"- `{row['quiz_title']}` — questions: {row['quiz_question_count']}, pools: {row['distinct_pool_count']}, unmatched: {row['unmatched_count']}"
        )

    return "\n".join(lines) + "\n"


def apply_sheet_style(sheet, header_row: int = 1) -> None:
    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    for cell in sheet[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP_ALIGNMENT
    max_sample_row = min(sheet.max_row, header_row + 49)
    for column_cells in sheet.iter_cols(min_row=header_row, max_row=max_sample_row, min_col=1, max_col=sheet.max_column):
        values = ["" if cell.value is None else str(cell.value) for cell in column_cells]
        width = min(max(len(value) for value in values) + 2, 60)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = max(width, 12)
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
        for cell in row:
            if cell.row == header_row:
                continue
            cell.alignment = WRAP_ALIGNMENT


def write_sheet(
    workbook: Workbook,
    title: str,
    rows: list[dict[str, object]],
    hyperlink_columns: dict[str, str] | None = None,
    exclude_headers: set[str] | None = None,
    headers: list[str] | None = None,
) -> None:
    hyperlink_columns = hyperlink_columns or {}
    exclude_headers = exclude_headers or set()
    sheet = workbook.create_sheet(title=title)
    if not rows:
        if headers:
            sheet.append(headers)
            apply_sheet_style(sheet)
            return
        sheet.append(["note"])
        sheet.append(["No rows"])
        apply_sheet_style(sheet)
        return
    headers = headers or [header for header in rows[0].keys() if header not in exclude_headers]
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
        row_number = sheet.max_row
        for display_header, target_header in hyperlink_columns.items():
            if display_header not in headers:
                continue
            display_index = headers.index(display_header) + 1
            target_value = str(row.get(target_header, "") or "").strip()
            if not target_value:
                continue
            cell = sheet.cell(row=row_number, column=display_index)
            target_path = Path(target_value)
            cell.hyperlink = (
                target_path.resolve().as_uri()
                if target_path.is_absolute()
                else target_value.replace("\\", "/")
            )
            cell.style = "Hyperlink"
    apply_sheet_style(sheet)


def write_unresolved_sheet(
    workbook: Workbook,
    rows: list[dict[str, object]],
    hyperlink_columns: dict[str, str] | None = None,
) -> None:
    hyperlink_columns = hyperlink_columns or {}
    sheet = workbook.create_sheet(title=UNRESOLVED_SHEET_TITLE)
    headers = [key for key in rows[0] if key != "source_choice_display"] if rows else ["note"]
    end_column = max(len(headers), 2)

    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    sheet.cell(row=1, column=1, value="Reviewer Note")
    sheet.cell(row=1, column=1).fill = HEADER_FILL
    sheet.cell(row=1, column=1).font = HEADER_FONT
    sheet.cell(row=1, column=1).alignment = WRAP_ALIGNMENT

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    sheet.cell(row=2, column=1, value=UNRESOLVED_SHEET_NOTE)
    sheet.cell(row=2, column=1).alignment = WRAP_ALIGNMENT
    sheet.row_dimensions[2].height = 150

    header_row = 4
    for column_index, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=column_index, value=header)

    if rows:
        for row in rows:
            sheet.append([row.get(header, "") for header in headers])
            row_number = sheet.max_row
            for display_header, target_header in hyperlink_columns.items():
                if display_header not in headers:
                    continue
                display_index = headers.index(display_header) + 1
                target_value = str(row.get(target_header, "") or "").strip()
                if not target_value:
                    continue
                cell = sheet.cell(row=row_number, column=display_index)
                cell.hyperlink = Path(target_value).resolve().as_uri()
                cell.style = "Hyperlink"
    else:
        sheet.append(["No rows"])

    apply_sheet_style(sheet, header_row=header_row)


def write_overview_sheet(
    workbook: Workbook,
    payload: dict[str, object],
    source_dir: Path,
    created_at: str,
    script_path: Path,
) -> None:
    sheet = workbook.create_sheet(title="Overview")
    overview_rows = [
        ("Export label", payload["export_label"]),
        ("Source folder", str(source_dir)),
        ("Created at", created_at),
        ("Created by", str(script_path)),
        ("Quiz files parsed", payload["summary"]["quiz_count"]),
        ("Question-library banks parsed", payload["summary"]["bank_count"]),
        ("Question-library pools parsed", payload["summary"]["pool_count"]),
        ("Quiz question rows extracted", payload["summary"]["quiz_question_count"]),
        ("Pool mapping by direct source evidence", payload["summary"]["source_evidence_count"]),
        ("Pool mapping by exact inference", payload["summary"]["inferred_exact_count"]),
        ("Pool mapping by similarity inference", payload["summary"]["inferred_similarity_count"]),
        ("Unmatched quiz items", payload["summary"]["unmatched_count"]),
        ("Quiz question rows with image refs", payload["summary"]["question_rows_with_images"]),
        ("Image refs observed", payload["summary"]["image_reference_count"]),
        ("Image refs resolved to packaged files", payload["summary"]["resolved_image_reference_count"]),
        ("Missing image refs", payload["summary"]["missing_image_reference_count"]),
        ("Evidence note", "Most quiz rows in this export do not declare an explicit source pool section in quiz XML. Pool rows labeled inferred are content matches against questiondb.xml; `source_evidence` rows come from direct `<itemref>` references or item-ID joins via `qmd_displayid` / `qmd_globalid`."),
        ("Suggested review order", f"1) Overview 2) Quiz Sections 3) Quiz Questions 4) Question Images 5) Question Banks 6) Question Pools 7) {UNRESOLVED_SHEET_TITLE}"),
    ]
    readme_rows = [
        ("How to use this workbook", f"Start on Overview for scope and counts. Then use Quiz Sections to review quiz-side pools and section behavior, Quiz Questions for item-level review, Question Images for full image references, Question Banks for question-library bank usage, Question Pools for raw question-library section detail, and {UNRESOLVED_SHEET_TITLE} for items whose bank provenance remains uncertain."),
        ("How this was created", f"Generated by {script_path.name} from the unpacked export folder. The script parses questiondb.xml and quiz_d2l_*.xml, extracts inline quiz items and direct questiondb itemrefs, compares quiz items back to the question library by IDs and content, and writes reviewer-facing XLSX, JSON, and Markdown artifacts."),
        ("Quiz Summary", "One row per quiz with question counts, pool counts, and evidence totals."),
        ("Quiz Sections", "One row per quiz section with random-draw information, matched pool distribution, and section-level notes."),
        ("Quiz Questions", "One row per quiz item with extracted prompt text, answer-key material when present, evidence level, likely pool match, and the primary image link when available."),
        ("Question Images", "One row per question image reference with original XML path, normalized path, resolution status, and clickable local file link when the packaged asset was found."),
        ("Question Banks", "One row per question-library bank with aggregated pool counts, mapped quiz usage, evidence totals, and a note when the bank is made of root-level objectbank items."),
        ("Question Pools", "One row per question-library pool / section from questiondb.xml with bank context, pool path, level, and direct question count."),
        (UNRESOLVED_SHEET_TITLE, "Exception queue for rows whose question-library / pool match was not unique enough to trust. These items were still extracted successfully; only their bank match remains uncertain."),
        ("Evidence labels", "`source_evidence` means a direct question-library reference or item-ID join, `inferred_exact` means a unique exact content match, `inferred_similarity` means a high-confidence similarity guess, and `unmatched` means no defensible bank match was found."),
    ]

    sheet.append(["field", "value"])
    for row in overview_rows:
        sheet.append(list(row))

    summary_end_row = sheet.max_row
    sheet.auto_filter.ref = f"A1:B{summary_end_row}"
    sheet.freeze_panes = "A2"

    readme_header_row = summary_end_row + 2
    sheet.merge_cells(start_row=readme_header_row, start_column=1, end_row=readme_header_row, end_column=2)
    sheet.cell(row=readme_header_row, column=1, value="Workbook README")

    current_row = readme_header_row + 1
    for label, value in readme_rows:
        sheet.cell(row=current_row, column=1, value=label)
        sheet.cell(row=current_row, column=2, value=value)
        current_row += 1

    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 100

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP_ALIGNMENT

    readme_header = sheet.cell(row=readme_header_row, column=1)
    readme_header.fill = HEADER_FILL
    readme_header.font = HEADER_FONT
    readme_header.alignment = WRAP_ALIGNMENT

    for row in range(2, sheet.max_row + 1):
        for column in range(1, 3):
            cell = sheet.cell(row=row, column=column)
            cell.alignment = WRAP_ALIGNMENT
        if row > readme_header_row:
            sheet.cell(row=row, column=1).font = Font(bold=True)


def write_workbook(output_path: Path, payload: dict[str, object], source_dir: Path, created_at: str, script_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    write_overview_sheet(workbook, payload, source_dir, created_at, script_path)

    write_sheet(workbook, "Quiz Summary", payload["quiz_summary_rows"])
    write_sheet(workbook, "Quiz Sections", payload["quiz_section_rows"])
    write_sheet(
        workbook,
        "Quiz Questions",
        payload["quiz_question_rows"],
        hyperlink_columns={"question_primary_image_path": "question_primary_image_absolute_path"},
        exclude_headers={"source_response_facts", "source_choice_display"},
    )
    write_sheet(
        workbook,
        "Question Images",
        payload["question_image_rows"],
        hyperlink_columns={"image_file_path": "image_file_absolute_path"},
    )
    write_sheet(workbook, "Question Banks", payload["bank_rows"])
    write_sheet(workbook, "Question Pools", payload["pool_rows"])
    write_unresolved_sheet(
        workbook,
        payload["unresolved_rows"],
        hyperlink_columns={"question_primary_image_path": "question_primary_image_absolute_path"},
    )

    workbook.save(output_path)


def write_reviewer_overview_sheet(
    workbook: Workbook,
    payload: dict[str, object],
    source_dir: Path,
    created_at: str,
    script_path: Path,
    projection: dict[str, object] | None = None,
    all_question_rows: list[dict[str, object]] | None = None,
) -> None:
    sheet = workbook.create_sheet(title="Overview")
    summary = payload["summary"]
    overview_rows = [
        ("Export label", payload["export_label"]),
        ("Source label", payload["export_label"]),
        ("Created at", created_at),
        ("Created by", script_path.name),
        ("Quizzes reviewed", summary["quiz_count"]),
        ("Questions reviewed", summary["quiz_question_count"]),
        ("Quiz questions with direct library links", summary["source_evidence_count"]),
        (
            "Additional quiz–library links inferred from identical content",
            summary["inferred_exact_count"],
        ),
        (
            "Possible quiz–library links suggested by similar content",
            summary["inferred_similarity_count"],
        ),
        (
            "Quiz questions with no reliable library link found",
            summary["unmatched_count"],
        ),
        (
            "About these library-link counts",
            "They describe question provenance and relationships—not errors or "
            "judgments about question quality.",
        ),
        ("Questions with images", summary["question_rows_with_images"]),
    ]
    if projection is not None:
        projection_summary = projection["summary"]
        overview_rows.extend(
            [
                ("Unique question entities", projection_summary["unique_question_count"]),
                ("Question occurrences", projection_summary["question_occurrence_count"]),
                ("Resolved occurrences", projection_summary["resolved_occurrence_count"]),
                ("Unresolved occurrences", projection_summary["unresolved_occurrence_count"]),
                (
                    "Random-draw candidate occurrences",
                    projection_summary["random_draw_occurrence_count"],
                ),
            ]
        )
    if all_question_rows is not None:
        inventory_counts = Counter(
            str(row.get("classification") or "") for row in all_question_rows
        )
        overview_rows.extend(
            [
                ("All source questions", len(all_question_rows)),
                (
                    "Question-library records",
                    sum(
                        1
                        for row in all_question_rows
                        if row.get("source_type") == "question library"
                    ),
                ),
                (
                    "Used in quiz — direct library link",
                    inventory_counts["Used in quiz — direct library link"],
                ),
                (
                    "Used in quiz — linked by identical content",
                    inventory_counts["Used in quiz — linked by identical content"],
                ),
                (
                    "Possible quiz use — similar content; review recommended",
                    inventory_counts[
                        "Possible quiz use — similar content; review recommended"
                    ],
                ),
                (
                    "Library only — not used in exported quizzes",
                    inventory_counts["Library only — not used in exported quizzes"],
                ),
                (
                    "Quiz only — no library link found",
                    inventory_counts["Quiz only — no library link found"],
                ),
                (
                    "All source questions with image refs",
                    sum(
                        1
                        for row in all_question_rows
                        if row.get("has_image") == "yes"
                    ),
                ),
                (
                    "Image-bearing questions with included files",
                    sum(
                        1
                        for row in all_question_rows
                        if row.get("image_included") in {"yes", "partially"}
                    ),
                ),
                (
                    "Image-bearing questions with missing files",
                    sum(
                        1
                        for row in all_question_rows
                        if row.get("image_status")
                        in {"missing", "partially resolved"}
                    ),
                ),
            ]
        )
    readme_rows = [
        (
            "What this workbook is",
            "This workbook presents the exported quizzes and complete Question "
            "Library for review, including question content, answer data, quiz "
            "placement, library relationships, settings, and image references.",
        ),
        (
            "How to use it",
            f"Start on Overview, then use {REVIEWER_ALL_QUESTIONS_SHEET_TITLE} "
            "for the complete source inventory: every question-library record "
            "plus quiz-only or unlinked questions. Quiz Occurrences shows each "
            "individual ordered placement, including pool and draw context, with "
            "the full question content beside it. Use Quiz Questions for marginalia, "
            "Quiz Settings for open setting decisions, Question Images for media, "
            "Question Library for source-location totals, and "
            f"{REVIEWER_UNRESOLVED_SHEET_TITLE} for uncertain library links.",
        ),
        (
            "How to propose revisions",
            "Yellow proposal cells are editable; extracted source cells are locked to preserve the export.\n"
            "1. To suggest a question change, enter replacement text, responses, answer key, or library location in the matching revised_* column.\n"
            "2. Explain the change in revision_reason and add your name and date in proposed_by and proposed_at.\n"
            "3. Use proposed_permanent_code only when assigning a stable code. Use proposed_scoring_mode only when the extracted scoring mode is unresolved.\n"
            "4. Leave approval_status, approved_by, and approved_at for the person approving the change.\n"
            "Use reviewer_note for context only; it does not replace an extracted value or submit a change.",
        ),
        (
            "How quiz–library links are identified",
            "These labels describe how the extractor associated quiz questions "
            "with Question Library records; they are not error counts or judgments "
            "about question quality. `Direct library link` means the quiz XML "
            "explicitly identifies a library item. `Linked by identical content` "
            "means no direct link was available, but the question and response "
            "content matched a library item exactly. `Possible link based on similar "
            "content` means the strongest candidate is similar but not identical "
            "and requires review. `No reliable library link found` means the "
            "question stays visible as quiz-only or unlinked instead of being "
            "assigned a questionable library location.",
        ),
        ("Quiz Summary", "One row per quiz with question counts and a simple breakdown of library-link status."),
        (REVIEWER_ALL_QUESTIONS_SHEET_TITLE, "The complete source-question inventory. It contains every record parsed from questiondb.xml and adds one row per unique quiz-only/unlinked question. The classification column distinguishes direct library links, links inferred from identical content, possible links based on similar content that require review, unused library records, and quiz-only questions."),
        ("Images in All Questions", "Filter has_image to yes. image_status says whether each XML reference resolved; image_included says whether a stable reviewer copy was packaged. Click a blue/underlined primary_image value to open the file. A blank primary_image with image_status=missing means the export referenced an image but did not include the file. When sharing or moving the workbook, keep its generated *_assets folder beside it so relative image links continue to work."),
        ("Quiz Questions", "The editable occurrence-level marginalia surface. Stable occurrence and question keys bind decisions to the model; extracted columns remain locked."),
        ("Quiz Occurrences", "One self-contained row for every quiz placement or pool candidate, with quiz and section order, inline/itemref status, draw context, content, answers, diagnostics, and build-support state."),
        ("Quiz Settings", "Observed quiz settings with an additive proposal/approval wing. setting_observation_key is a hidden, CourseCraft-generated row identity rather than a package-native D2L field; it keeps repeated source observations distinct. Extracted setting columns remain locked; accepted decisions materialize as quiz_decision inputs during local re-ingest."),
        ("Question Images", "One row per question image with a direct file link when the image was bundled in the export."),
        ("Question Library", "One row per question-library location with counts showing how much of the reviewed quiz content appears to come from that location."),
        (REVIEWER_UNRESOLVED_SHEET_TITLE, "Questions that remain visible even though the script could not make a safe library-location call; response options and answer-key material are still shown when present in the quiz XML."),
    ]

    sheet.append(["field", "value"])
    for row in overview_rows:
        sheet.append(list(row))

    summary_end_row = sheet.max_row
    sheet.auto_filter.ref = f"A1:B{summary_end_row}"
    sheet.freeze_panes = "A2"

    readme_header_row = summary_end_row + 2
    sheet.merge_cells(start_row=readme_header_row, start_column=1, end_row=readme_header_row, end_column=2)
    sheet.cell(row=readme_header_row, column=1, value="Workbook README")

    current_row = readme_header_row + 1
    for label, value in readme_rows:
        sheet.cell(row=current_row, column=1, value=label)
        sheet.cell(row=current_row, column=2, value=value)
        current_row += 1

    sheet.column_dimensions["A"].width = 48
    sheet.column_dimensions["B"].width = 100

    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP_ALIGNMENT

    readme_header = sheet.cell(row=readme_header_row, column=1)
    readme_header.fill = HEADER_FILL
    readme_header.font = HEADER_FONT
    readme_header.alignment = WRAP_ALIGNMENT
    sheet.row_dimensions[readme_header_row].height = 28

    for row in range(2, sheet.max_row + 1):
        for column in range(1, 3):
            cell = sheet.cell(row=row, column=column)
            cell.alignment = WRAP_ALIGNMENT
        if row > readme_header_row:
            sheet.cell(row=row, column=1).font = Font(bold=True)
            readme_text = str(sheet.cell(row=row, column=2).value or "")
            estimated_lines = max(1, (len(readme_text) // 70) + 1)
            explicit_lines = readme_text.count("\n") + 1
            sheet.row_dimensions[row].height = min(
                260,
                max(36, 20 * (max(estimated_lines, explicit_lines) + 1)),
            )


def set_column_width_by_header(sheet, header: str, width: float) -> None:
    headers = [cell.value for cell in sheet[1]]
    if header not in headers:
        return
    column_index = headers.index(header) + 1
    sheet.column_dimensions[get_column_letter(column_index)].width = width


def add_header_guidance(sheet, guidance: dict[str, str]) -> None:
    headers = [cell.value for cell in sheet[1]]
    for header, text in guidance.items():
        if header not in headers:
            continue
        cell = sheet.cell(row=1, column=headers.index(header) + 1)
        cell.comment = Comment(text, "CourseCraft")


def apply_status_fill(sheet, header: str) -> None:
    headers = [cell.value for cell in sheet[1]]
    if header not in headers:
        return
    column_index = headers.index(header) + 1
    for row in range(2, sheet.max_row + 1):
        cell = sheet.cell(row=row, column=column_index)
        fill = REVIEWER_STATUS_FILLS.get(str(cell.value or ""))
        if fill is not None:
            cell.fill = fill


def apply_font_size(sheet, size: int) -> None:
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, min_col=1, max_col=sheet.max_column):
        for cell in row:
            font = copy(cell.font)
            font.sz = size
            font.name = "Open Sans"
            cell.font = font


def apply_banded_rows(sheet, header_row: int = 1) -> None:
    for row_index in range(header_row + 1, sheet.max_row + 1):
        if (row_index - (header_row + 1)) % 2 != 0:
            continue
        for column_index in range(1, sheet.max_column + 1):
            sheet.cell(row=row_index, column=column_index).fill = REVIEWER_BAND_FILL


def configure_reviewer_marginalia(
    sheet, editable_headers: set[str] | None = None
) -> None:
    editable_headers = editable_headers or REVIEWER_EDITABLE_HEADERS
    headers = [str(cell.value or "") for cell in sheet[1]]
    editable_columns: list[int] = []
    for column_index, header in enumerate(headers, start=1):
        if header not in editable_headers:
            continue
        editable_columns.append(column_index)
        sheet.column_dimensions[get_column_letter(column_index)].width = 28
        for row_index in range(2, sheet.max_row + 1):
            cell = sheet.cell(row=row_index, column=column_index)
            cell.protection = Protection(locked=False)
            cell.fill = REVIEWER_EDITABLE_FILL

    if "approval_status" in headers and sheet.max_row >= 2:
        status_column = get_column_letter(headers.index("approval_status") + 1)
        validation = DataValidation(
            type="list",
            formula1='"open,accepted,rejected"',
            allow_blank=True,
        )
        validation.error = "Use open, accepted, or rejected."
        validation.errorTitle = "Unknown approval status"
        validation.prompt = "Choose the review disposition for this row."
        validation.promptTitle = "Approval status"
        validation.showErrorMessage = True
        validation.showInputMessage = True
        sheet.add_data_validation(validation)
        validation.add(f"{status_column}2:{status_column}{sheet.max_row}")

    if editable_columns:
        sheet.protection.sheet = True
        # Keep normal review mechanics available while extracted cells remain locked.
        sheet.protection.autoFilter = False
        sheet.protection.sort = False
        sheet.protection.selectLockedCells = False
        sheet.protection.selectUnlockedCells = False


def protect_reviewer_source_sheet(sheet) -> None:
    """Keep projection-only sheets immutable while retaining review mechanics."""
    sheet.protection.sheet = True
    sheet.protection.autoFilter = False
    sheet.protection.sort = False
    sheet.protection.selectLockedCells = False


def configure_reviewer_print_sheet(
    sheet, *, fit_width: int = 1, repeat_columns: str | None = None
) -> None:
    """Set predictable print/PDF defaults without changing the live grid."""
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = "1"  # US Letter
    sheet.page_setup.fitToWidth = fit_width
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    if repeat_columns:
        sheet.print_title_cols = repeat_columns
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.4
    sheet.page_margins.bottom = 0.4


def apply_reviewer_workbook_tweaks(workbook: Workbook) -> None:
    for sheet_name in workbook.sheetnames:
        apply_font_size(workbook[sheet_name], 14)
        configure_reviewer_print_sheet(workbook[sheet_name])

    quiz_summary = workbook["Quiz Summary"]
    set_column_width_by_header(quiz_summary, "quiz", 48)
    set_column_width_by_header(quiz_summary, "no_safe_library_matches", 24)
    apply_banded_rows(quiz_summary)

    if REVIEWER_ALL_QUESTIONS_SHEET_TITLE in workbook.sheetnames:
        all_questions = workbook[REVIEWER_ALL_QUESTIONS_SHEET_TITLE]
        all_questions.freeze_panes = "F2"
        all_questions.sheet_view.zoomScale = 80
        configure_reviewer_print_sheet(
            all_questions, fit_width=3, repeat_columns="A:E"
        )
        set_column_width_by_header(all_questions, "classification", 42)
        set_column_width_by_header(all_questions, "source_type", 24)
        set_column_width_by_header(all_questions, "link_confidence", 40)
        set_column_width_by_header(all_questions, "used_in_quizzes", 52)
        set_column_width_by_header(all_questions, "quiz_locations", 72)
        set_column_width_by_header(all_questions, "library_location", 64)
        set_column_width_by_header(all_questions, "display_identifier", 32)
        set_column_width_by_header(all_questions, "question_text", 72)
        set_column_width_by_header(all_questions, "has_image", 12)
        set_column_width_by_header(all_questions, "image_reference_count", 18)
        set_column_width_by_header(all_questions, "image_references", 52)
        set_column_width_by_header(all_questions, "image_status", 20)
        set_column_width_by_header(all_questions, "image_included", 16)
        set_column_width_by_header(all_questions, "primary_image", 52)
        set_column_width_by_header(all_questions, "resolved_image_paths", 52)
        set_column_width_by_header(all_questions, "included_image_paths", 52)
        set_column_width_by_header(all_questions, "missing_image_references", 52)
        set_column_width_by_header(all_questions, "response_options", 56)
        set_column_width_by_header(all_questions, "answer_key", 48)
        set_column_width_by_header(all_questions, "feedback", 56)
        all_questions.row_dimensions[1].height = 30
        for row_index in range(2, all_questions.max_row + 1):
            all_questions.row_dimensions[row_index].height = 60
        apply_banded_rows(all_questions)
        apply_status_fill(all_questions, "classification")
        apply_status_fill(all_questions, "image_status")
        protect_reviewer_source_sheet(all_questions)

    quiz_questions = workbook["Quiz Questions"]
    # Source type is reviewer context, while stable technical joins stay intact.
    for type_sheet in [quiz_questions, *([workbook["All Questions"]] if "All Questions" in workbook.sheetnames else [])]:
        type_headers = [cell.value for cell in type_sheet[1]]
        if "question_type" in type_headers:
            type_column = get_column_letter(type_headers.index("question_type") + 1)
            type_sheet.column_dimensions[type_column].hidden = False
            type_sheet.column_dimensions[type_column].width = 24
    quiz_questions.freeze_panes = "D2"
    quiz_questions.sheet_view.zoomScale = 80
    configure_reviewer_print_sheet(
        quiz_questions, fit_width=3, repeat_columns="A:C"
    )
    set_column_width_by_header(quiz_questions, "question_library_location", 42)
    set_column_width_by_header(quiz_questions, "library_link_status", 24)
    set_column_width_by_header(quiz_questions, "note", 68)
    apply_banded_rows(quiz_questions)
    apply_status_fill(quiz_questions, "library_link_status")
    configure_reviewer_marginalia(quiz_questions)

    if "Question Entities" in workbook.sheetnames:
        entities = workbook["Question Entities"]
        entities.freeze_panes = "D2"
        entities.sheet_view.zoomScale = 85
        configure_reviewer_print_sheet(entities, fit_width=3, repeat_columns="A:C")
        set_column_width_by_header(entities, "display_identifier", 32)
        set_column_width_by_header(entities, "question_entity_key", 42)
        set_column_width_by_header(entities, "source_aliases", 52)
        set_column_width_by_header(entities, "used_in_quizzes", 48)
        set_column_width_by_header(entities, "occurrence_locations", 72)
        set_column_width_by_header(entities, "question_library_locations", 52)
        set_column_width_by_header(entities, "question_text", 72)
        set_column_width_by_header(entities, "response_options", 56)
        set_column_width_by_header(entities, "answer_key", 48)
        set_column_width_by_header(entities, "feedback", 48)
        add_header_guidance(entities, REVIEWER_ENTITY_IDENTIFIER_GUIDANCE)
        apply_banded_rows(entities)
        protect_reviewer_source_sheet(entities)
        entities.sheet_state = "hidden"

    if "Quiz Occurrences" in workbook.sheetnames:
        occurrences = workbook["Quiz Occurrences"]
        occurrences.freeze_panes = "G2"
        occurrences.sheet_view.zoomScale = 80
        configure_reviewer_print_sheet(
            occurrences, fit_width=3, repeat_columns="A:F"
        )
        set_column_width_by_header(occurrences, "occurrence_key", 42)
        set_column_width_by_header(occurrences, "question_entity_key", 42)
        set_column_width_by_header(occurrences, "display_identifier", 32)
        set_column_width_by_header(occurrences, "source_aliases", 52)
        set_column_width_by_header(occurrences, "quiz", 42)
        set_column_width_by_header(occurrences, "section", 56)
        set_column_width_by_header(occurrences, "question_library_location", 52)
        set_column_width_by_header(occurrences, "library_match_status", 28)
        set_column_width_by_header(occurrences, "library_match_code", 24)
        set_column_width_by_header(occurrences, "library_match_candidates", 52)
        set_column_width_by_header(occurrences, "library_match_explanation", 72)
        set_column_width_by_header(occurrences, "question_text", 72)
        set_column_width_by_header(occurrences, "response_options", 56)
        set_column_width_by_header(occurrences, "answer_key", 48)
        set_column_width_by_header(occurrences, "feedback", 48)
        apply_banded_rows(occurrences)
        apply_status_fill(occurrences, "library_match_status")
        protect_reviewer_source_sheet(occurrences)

    quiz_settings = workbook["Quiz Settings"]
    quiz_settings.freeze_panes = "C2"
    set_column_width_by_header(quiz_settings, "quiz", 42)
    set_column_width_by_header(quiz_settings, "setting", 28)
    set_column_width_by_header(quiz_settings, "observed_value", 32)
    settings_headers = [cell.value for cell in quiz_settings[1]]
    observation_key_column = settings_headers.index("setting_observation_key") + 1
    quiz_settings.column_dimensions[
        get_column_letter(observation_key_column)
    ].hidden = True
    apply_banded_rows(quiz_settings)
    configure_reviewer_marginalia(quiz_settings, REVIEWER_SETTINGS_EDITABLE_HEADERS)

    question_images = workbook["Question Images"]
    apply_banded_rows(question_images)

    question_library = workbook["Question Library"]
    set_column_width_by_header(question_library, "question_library_location", 52)
    set_column_width_by_header(question_library, "quiz_count", 14)
    set_column_width_by_header(question_library, "used_in_quizzes", 72)
    set_column_width_by_header(question_library, "note", 68)
    apply_banded_rows(question_library)

    unresolved = workbook[REVIEWER_UNRESOLVED_SHEET_TITLE]
    unresolved.freeze_panes = "D2"
    unresolved.sheet_view.zoomScale = 80
    configure_reviewer_print_sheet(unresolved, fit_width=3, repeat_columns="A:C")
    set_column_width_by_header(unresolved, "why_it_needs_review", 68)
    apply_banded_rows(unresolved)
    configure_reviewer_marginalia(unresolved)


def write_reviewer_workbook(
    output_path: Path,
    payload: dict[str, object],
    source_dir: Path,
    created_at: str,
    script_path: Path,
    model: dict[str, object] | None = None,
) -> None:
    projection = (
        build_quiz_review_projection(model, include_content=True)
        if model is not None
        else None
    )
    all_question_rows = build_reviewer_all_question_rows(payload)
    workbook = Workbook()
    workbook.remove(workbook.active)
    write_reviewer_overview_sheet(
        workbook,
        payload,
        source_dir,
        created_at,
        script_path,
        projection,
        all_question_rows,
    )
    write_sheet(workbook, "Quiz Summary", build_reviewer_quiz_summary_rows(payload))
    write_sheet(
        workbook,
        REVIEWER_ALL_QUESTIONS_SHEET_TITLE,
        all_question_rows,
        hyperlink_columns={"primary_image": "primary_image_link_target"},
        exclude_headers={"primary_image_link_target"},
    )
    if projection is not None:
        write_sheet(
            workbook,
            "Quiz Occurrences",
            build_reviewer_occurrence_rows(projection),
        )
    write_sheet(
        workbook,
        "Quiz Questions",
        build_reviewer_question_rows(payload, projection),
        hyperlink_columns={"image_link": "image_link_absolute_path"},
        exclude_headers={"image_link_absolute_path"},
    )
    write_sheet(
        workbook,
        "Quiz Settings",
        build_reviewer_settings_rows(model),
        headers=REVIEWER_SETTINGS_HEADERS,
    )
    write_sheet(
        workbook,
        "Question Images",
        build_reviewer_image_rows(payload),
        hyperlink_columns={"image_file": "image_file_absolute_path"},
        exclude_headers={"image_file_absolute_path"},
    )
    write_sheet(workbook, "Question Library", build_reviewer_library_rows(payload))
    write_sheet(
        workbook,
        REVIEWER_UNRESOLVED_SHEET_TITLE,
        build_reviewer_unresolved_rows(payload, projection),
        hyperlink_columns={"primary_image": "primary_image_absolute_path"},
        exclude_headers={"primary_image_absolute_path"},
    )
    if projection is not None:
        write_sheet(
            workbook,
            "Question Entities",
            build_reviewer_entity_rows(projection),
        )
    apply_reviewer_workbook_tweaks(workbook)
    workbook.save(output_path)


def build_payload(source_dir: Path, export_label: str | None = None) -> dict[str, object]:
    questiondb_path = source_dir / "questiondb.xml"
    if questiondb_path.exists():
        pools, library_questions = parse_question_library(questiondb_path)
    else:
        pools, library_questions = [], []
    indexes = build_library_indexes(library_questions)
    asset_index = build_asset_index(source_dir)
    quiz_rows, section_rows, question_rows, unresolved_rows, image_rows = parse_quiz_files(source_dir, indexes, asset_index)
    bank_rows = build_bank_rows(pools, question_rows)

    summary = {
        "quiz_count": len(quiz_rows),
        "bank_count": len(bank_rows),
        "pool_count": len(pools),
        "quiz_question_count": len(question_rows),
        "source_evidence_count": sum(1 for row in question_rows if row["evidence_level"] == "source_evidence"),
        "inferred_exact_count": sum(1 for row in question_rows if row["evidence_level"] == "inferred_exact"),
        "inferred_similarity_count": sum(1 for row in question_rows if row["evidence_level"] == "inferred_similarity"),
        "unmatched_count": sum(1 for row in question_rows if row["evidence_level"] == "unmatched"),
        "question_rows_with_images": sum(1 for row in question_rows if int(row["question_image_count"]) > 0),
        "image_reference_count": len(image_rows),
        "resolved_image_reference_count": sum(1 for row in image_rows if row["image_status"] == "resolved"),
        "missing_image_reference_count": sum(1 for row in image_rows if row["image_status"] != "resolved"),
    }

    return {
        "export_label": export_label or source_dir.name,
        "summary": summary,
        "bank_rows": bank_rows,
        "pool_rows": pools,
        "library_question_rows": [
            {
                **library_question_payload_row(question, source_dir, asset_index),
                "bank_ident": question.bank_ident,
                "bank_title": question.bank_title,
                "bank_path": question.bank_path,
                "pool_ident": question.pool_ident,
                "pool_title": question.pool_title,
                "pool_path": question.pool_path,
                "plain_text": question.plain_text,
                "choice_count": question.choice_count,
            }
            for question in library_questions
        ],
        "quiz_summary_rows": quiz_rows,
        "quiz_section_rows": section_rows,
        "quiz_question_rows": question_rows,
        "question_image_rows": image_rows,
        "unresolved_rows": unresolved_rows,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        description="Extract quiz questions by quiz and map them back to question-library pools for review."
    )
    parser.add_argument("source", help="Path to a Brightspace export ZIP or unpacked folder")
    parser.add_argument(
        "--output-dir",
        required=True,
        help=(
            "Directory for local review artifacts, normalized model, and run "
            "receipt. Inside the Workbench this must be git-ignored."
        ),
    )
    parser.add_argument(
        "--allow-tracked-output",
        action="store_true",
        help=(
            "Explicitly permit output inside a Workbench path that git does not "
            "ignore. Course evidence may include answer keys and large artifacts."
        ),
    )
    parser.add_argument(
        "--source-lineage-key",
        default="",
        help="Optional stable cc:lineage key for refresh-stable entity identities when the export has no course identity.",
    )
    parser.add_argument(
        "--copy-images-to-assets",
        action="store_true",
        help=(
            "Legacy alias for --asset-mode copy. Copy resolved question images "
            "beside the review outputs."
        ),
    )
    parser.add_argument(
        "--asset-mode",
        choices=("auto", "copy", "reference"),
        default="auto",
        help=(
            "Reviewer asset handling: auto preserves the historical behavior "
            "(copy for ZIP intake, reference for folders); copy always emits a "
            "self-contained assets folder; reference never copies source assets."
        ),
    )
    args = parser.parse_args(argv)

    try:
        started_at = utc_now()
        source_arg = Path(args.source).expanduser().resolve()
        output_dir = Path(args.output_dir).expanduser().resolve()
        guard_repo_safe_output(
            output_dir,
            repo_root=WORKBENCH_ROOT,
            allow_tracked=args.allow_tracked_output,
            ignored_lane_example=(
                "workspace/review/quiz_binder_runs/<label>__extractor"
            ),
        )
        output_dir.mkdir(parents=True, exist_ok=True)
        with resolve_export(source_arg) as resolved:
            source_dir = resolved.source_dir
            payload = build_payload(source_dir, resolved.export_label)
            enrich_payload(payload, source_dir)
            base_name = f"{resolved.export_label}__quiz_pool_review"
            created_at = datetime.now(timezone.utc).astimezone().strftime("%Y-%m-%d %H:%M:%S %Z")
            script_path = Path(__file__).resolve()

            xlsx_path = output_dir / f"{base_name}.xlsx"
            reviewer_xlsx_path = output_dir / f"{base_name}_reviewer.xlsx"
            json_path = output_dir / f"{base_name}.json"
            md_path = output_dir / f"{base_name}.md"
            model_path = output_dir / f"{base_name}.model.json"
            run_path = output_dir / f"{base_name}.run.json"

            if args.copy_images_to_assets and args.asset_mode == "reference":
                raise ValueError(
                    "--copy-images-to-assets cannot be combined with --asset-mode reference"
                )
            should_copy_assets = (
                args.copy_images_to_assets
                or args.asset_mode == "copy"
                or (args.asset_mode == "auto" and resolved.source_kind == "zip")
            )
            copied_asset_count = 0
            if should_copy_assets:
                copied_asset_count = copy_resolved_assets(
                    payload,
                    source_dir,
                    output_dir,
                    f"{base_name}_assets",
                )

            model, source_meta = build_normalized_model(
                payload,
                source_dir,
                source_arg,
                resolved.source_kind,
                source_lineage_key=args.source_lineage_key or None,
            )
            write_workbook(xlsx_path, payload, source_dir, created_at, script_path)
            write_reviewer_workbook(
                reviewer_xlsx_path,
                payload,
                source_dir,
                created_at,
                script_path,
                model,
            )
            write_json(json_path, payload)
            md_path.write_text(markdown_report(source_dir, payload, created_at, script_path), encoding="utf-8")
            write_json(model_path, model)

            finished_at = utc_now()
            emitted_artifacts = [
                (xlsx_path, "review"),
                (reviewer_xlsx_path, "review"),
                (json_path, "review"),
                (md_path, "review"),
                (model_path, "model"),
            ]
            copied_asset_dir = output_dir / f"{base_name}_assets"
            if copied_asset_dir.exists():
                emitted_artifacts.extend(
                    (path, "review")
                    for path in sorted(copied_asset_dir.rglob("*"))
                    if path.is_file()
                )
            receipt = build_run_receipt(
                model,
                source_meta,
                source_arg,
                output_dir,
                emitted_artifacts,
                started_at,
                finished_at,
                f"python scripts/extract_quiz_pool_review.py {source_arg.name} --output-dir {output_dir.name}",
            )
            write_json(run_path, receipt)

            print(json.dumps({
                "source": str(source_arg),
                "source_kind": resolved.source_kind,
                "xlsx": str(xlsx_path),
                "reviewer_xlsx": str(reviewer_xlsx_path),
                "json": str(json_path),
                "markdown": str(md_path),
                "model": str(model_path),
                "run_receipt": str(run_path),
                "copied_asset_count": copied_asset_count,
                "summary": payload["summary"],
            }, indent=2))
        return 0
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
