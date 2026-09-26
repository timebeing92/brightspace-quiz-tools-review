#!/usr/bin/env python3
from __future__ import annotations

import argparse
import html
import hashlib
import json
import re
import sys
import zipfile
from dataclasses import dataclass, field
from pathlib import Path, PurePosixPath
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

from common_xml import clean, local_name
from quiz_build_support import load_authoring_projection, package_member_path
from quiz_contracts import validate_contract
from extract_quiz_pool_review import extract_source_response_facts
from quiz_normalization import source_choice_projection


REQUIRED_FILES = [
    "imsmanifest.xml",
    "questiondb.xml",
    "orgunitconfig/orgunitconfig.xml",
]


def normalized_header(value: object) -> str:
    return clean(value).lower().replace(" ", "_").replace("-", "_")


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


@dataclass
class ValidationReport:
    package_dir: Path
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)

    def error(self, message: str) -> None:
        self.errors.append(message)

    def warning(self, message: str) -> None:
        self.warnings.append(message)

    def note(self, message: str) -> None:
        self.notes.append(message)


def parse_xml(path: Path, report: ValidationReport) -> ET.Element | None:
    try:
        return ET.parse(path).getroot()
    except ET.ParseError as exc:
        report.error(f"{path.relative_to(report.package_dir)} is not well-formed XML: {exc}")
    except FileNotFoundError:
        report.error(f"Missing XML file: {path.relative_to(report.package_dir)}")
    return None


def metadata_value(elem: ET.Element, label: str) -> str:
    for field in elem.iter():
        if local_name(field.tag) != "qti_metadatafield":
            continue
        field_label = ""
        field_entry = ""
        for child in field:
            if local_name(child.tag) == "fieldlabel":
                field_label = clean(child.text)
            elif local_name(child.tag) == "fieldentry":
                field_entry = clean(child.text)
        if field_label == label:
            return field_entry
    return ""


def direct_metadata_value(elem: ET.Element, label: str) -> str:
    for child in elem:
        if local_name(child.tag) != "qtimetadata":
            continue
        for field in child:
            if local_name(field.tag) != "qti_metadatafield":
                continue
            field_label = ""
            field_entry = ""
            for field_child in field:
                if local_name(field_child.tag) == "fieldlabel":
                    field_label = clean(field_child.text)
                elif local_name(field_child.tag) == "fieldentry":
                    field_entry = clean(field_child.text)
            if field_label == label:
                return field_entry
    return ""


def direct_children_by_local_name(elem: ET.Element, name: str) -> list[ET.Element]:
    return [child for child in list(elem) if local_name(child.tag) == name]


def find_quiz_files(package_dir: Path) -> list[Path]:
    return sorted(package_dir.glob("quiz_d2l_*.xml"))


def validate_required_files(package_dir: Path, report: ValidationReport, *, library_only: bool = False) -> list[Path]:
    for relative in REQUIRED_FILES:
        if not (package_dir / relative).exists():
            report.error(f"Missing required package file: {relative}")
    quiz_files = find_quiz_files(package_dir)
    if not quiz_files and not library_only:
        report.error("Missing quiz_d2l_*.xml payload.")
    return quiz_files


def validate_manifest(
    package_dir: Path,
    manifest_root: ET.Element | None,
    report: ValidationReport,
) -> set[str]:
    if manifest_root is None:
        return set()
    hrefs: list[str] = []
    for elem in manifest_root.iter():
        if local_name(elem.tag) not in {"resource", "file"}:
            continue
        href = elem.attrib.get("href", "")
        if href:
            hrefs.append(href)
            # Quicklink/contentlink and external hrefs are LMS URLs (e.g.
            # /d2l/common/dialogs/quickLink/...), not files in the package.
            if href.startswith("/") or href.startswith(("http://", "https://", "mailto:", "data:", "#")):
                continue
            normalized_href = href.replace("\\", "/")
            if ".." in PurePosixPath(normalized_href).parts:
                report.error(f"Unsafe manifest href escapes the package: {href}")
                continue
            try:
                normalized = package_member_path(href).as_posix()
            except (UnicodeDecodeError, ValueError):
                report.error(f"Unsafe manifest href escapes or ambiguously encodes the package: {href}")
                continue
            if not (package_dir / normalized).is_file():
                report.error(f"Manifest href does not exist in package: {href}")
    if not hrefs:
        report.warning("Manifest has no resource hrefs.")
    else:
        report.note(f"Manifest resource/file hrefs checked: {len(hrefs)}")
    local_hrefs: set[str] = set()
    for href in hrefs:
        if not href or href.startswith("/") or href.startswith(("http://", "https://", "mailto:", "data:", "#")):
            continue
        try:
            local_hrefs.add(package_member_path(href).as_posix())
        except (UnicodeDecodeError, ValueError):
            continue
    return local_hrefs


def validate_package_closure(
    package_dir: Path,
    manifest_hrefs: set[str],
    report: ValidationReport,
    *,
    strict: bool,
) -> None:
    package_files = {
        path.relative_to(package_dir).as_posix()
        for path in package_dir.rglob("*")
        if path.is_file()
    }
    structural_files = {
        path
        for path in package_files
        if path == "questiondb.xml"
        or path == "grades_d2l.xml"
        or path == "orgunitconfig/orgunitconfig.xml"
        or path.startswith("quiz_d2l_") and path.endswith(".xml")
    }
    undeclared_structural = sorted(structural_files - manifest_hrefs)
    for path in undeclared_structural:
        report.error(f"Package structural file is not declared by the manifest: {path}")

    allowed_supplemental = {"imsmanifest.xml", "README.md"}
    unlisted = sorted(package_files - manifest_hrefs - allowed_supplemental)
    unlisted_nonstructural = [path for path in unlisted if path not in structural_files]
    for path in unlisted_nonstructural:
        message = f"Package file is not declared by the manifest and may be stale or nonportable: {path}"
        if strict:
            report.error(message)
        else:
            report.warning(message)
    if not undeclared_structural and not unlisted_nonstructural:
        report.note(f"Package closure checked: {len(package_files)} files")


def validate_local_asset_references(package_dir: Path, report: ValidationReport) -> None:
    referenced: set[str] = set()
    for xml_path in sorted(package_dir.glob("*.xml")):
        root = parse_xml(xml_path, report)
        if root is None:
            continue
        for elem in root.iter():
            if local_name(elem.tag) != "mattext" or not elem.text:
                continue
            for ref in re.findall(r'(?:src|href)=["\']([^"\']+)', html.unescape(elem.text), re.IGNORECASE):
                value = ref.split("?", 1)[0].split("#", 1)[0].replace("\\", "/")
                if value and not value.startswith(("/", "http://", "https://", "data:", "mailto:", "#", "//")):
                    referenced.add(value)
    missing = []
    for ref in sorted(referenced):
        try:
            member_path = package_member_path(ref)
        except (UnicodeDecodeError, ValueError):
            report.error(f"Question HTML has an unsafe or ambiguously encoded local asset reference: {ref}")
            continue
        if not package_dir.joinpath(*member_path.parts).is_file():
            missing.append(ref)
    for ref in missing:
        report.error(f"Question HTML references a missing local asset: {ref}")
    if referenced and not missing:
        report.note(f"Question HTML asset references checked: {len(referenced)}")


def collect_library_items(questiondb_root: ET.Element | None, report: ValidationReport) -> dict[str, dict[str, str]]:
    items: dict[str, dict[str, str]] = {}
    if questiondb_root is None:
        return items
    for item in questiondb_root.iter():
        if local_name(item.tag) != "item":
            continue
        label = item.attrib.get("label", "")
        ident = item.attrib.get("ident", "")
        displayid = metadata_value(item, "qmd_displayid")
        qtype = metadata_value(item, "qmd_questiontype")
        if not label:
            report.error(f"Question library item {ident or '[no ident]'} has no label.")
        if not displayid:
            report.warning(f"Question library item {label or ident or '[no label]'} has no qmd_displayid.")
        if not qtype:
            report.warning(f"Question library item {label or ident or '[no label]'} has no qmd_questiontype.")
        if label:
            items[label] = {"ident": ident, "displayid": displayid, "qtype": qtype}
        if ident:
            items.setdefault(ident, {"ident": ident, "displayid": displayid, "qtype": qtype})
    if not items:
        report.error("questiondb.xml contains no library items.")
    else:
        unique_labels = {key for key, value in items.items() if key != value.get("ident")}
        report.note(f"Question library items checked: {len(unique_labels) or len(items)}")
    return items


def validate_truefalse_answer_labels(
    questiondb_root: ET.Element | None,
    report: ValidationReport,
) -> None:
    """Require the answer-label encoding proven by the 2026-07-20 tenant probe."""
    if questiondb_root is None:
        return
    checked = 0
    for item in questiondb_root.iter():
        if local_name(item.tag) != "item" or metadata_value(item, "qmd_questiontype") != "True/False":
            continue
        checked += 1
        identity = (
            metadata_value(item, "qmd_displayid")
            or item.attrib.get("label", "")
            or item.attrib.get("ident", "")
            or "[unidentified]"
        )
        labels: list[tuple[str, str]] = []
        for response_label in item.iter():
            if local_name(response_label.tag) != "response_label":
                continue
            mattext = next(
                (node for node in response_label.iter() if local_name(node.tag) == "mattext"),
                None,
            )
            if mattext is not None:
                labels.append((mattext.attrib.get("texttype", ""), clean(mattext.text)))
        if [text for _texttype, text in labels] != ["True", "False"]:
            report.error(
                f"True/False question {identity} must contain exactly the bare answer labels "
                f"True and False in that order; found {labels}."
            )
            continue
        invalid_texttypes = [texttype for texttype, _text in labels if texttype != "text/plain"]
        if invalid_texttypes:
            report.error(
                f"True/False question {identity} answer labels must use mattext "
                f"texttype='text/plain'; found {invalid_texttypes}. The HTML form caused "
                "broken localization labels in the 2026-07-20 Brightspace tenant probe."
            )
    if checked and not any(error.startswith("True/False question") for error in report.errors):
        report.note(f"Tenant-verified True/False answer labels checked: {checked}")


def section_draw_count(section: ET.Element) -> int | None:
    value = direct_metadata_value(section, "qmd_numberofitems")
    if not value:
        return None
    try:
        return int(float(value))
    except ValueError:
        return -1


def validate_multiselect_keys(root: ET.Element | None, report: ValidationReport) -> None:
    """Require the native scoring profile, not merely interpretable generic QTI."""
    if root is None:
        return
    for item in root.iter('item'):
        if metadata_value(item, 'qmd_questiontype') != 'Multi-Select':
            continue
        identity = metadata_value(item, 'qmd_displayid') or item.get('ident', '[unidentified]')
        projection = source_choice_projection(extract_source_response_facts(item))
        conditions = item.findall('resprocessing/respcondition')
        setters = [c.find('setvar') for c in conditions]
        native = (len(conditions) == 2 and all(s is not None for s in setters)
            and [(s.get('action'), clean(s.text)) for s in setters] == [('Add', '1'), ('Add', '0')]
            and conditions[0].find('conditionvar') is not None
            and all(e.tag in {'varequal', 'not'} for e in conditions[0].find('conditionvar'))
            and [clean(e.text) for e in item.iter() if local_name(e.tag) == 'grading_type'] == ['0'])
        if not projection['recognized'] or not native:
            report.error(f"Multi-Select question {identity} lacks a supported native answer key: "
                         f"{projection['reason'] or 'expected native Add 1 / complementary Add 0 profile' }.")


def itemref_file_href(itemref: ET.Element) -> str:
    for child in itemref:
        if local_name(child.tag) == "file":
            return child.attrib.get("href", "")
    return ""


def validate_quiz_payload(
    package_dir: Path,
    quiz_path: Path,
    quiz_root: ET.Element | None,
    library_items: dict[str, dict[str, str]],
    report: ValidationReport,
) -> None:
    if quiz_root is None:
        return

    itemref_count = 0
    for itemref in quiz_root.iter():
        if local_name(itemref.tag) != "itemref":
            continue
        itemref_count += 1
        linkrefid = itemref.attrib.get("linkrefid", "")
        if not linkrefid:
            report.error(f"{quiz_path.name} has an itemref without linkrefid.")
        elif linkrefid not in library_items:
            report.error(f"{quiz_path.name} itemref linkrefid does not resolve in questiondb.xml: {linkrefid}")
        href = itemref_file_href(itemref)
        if not href:
            report.warning(f"{quiz_path.name} itemref {linkrefid or '[no linkrefid]'} has no file href.")
        elif not (package_dir / href.replace("\\", "/")).exists():
            report.error(f"{quiz_path.name} itemref {linkrefid} points to missing file: {href}")

    for section in quiz_root.iter():
        if local_name(section.tag) != "section":
            continue
        draw_count = section_draw_count(section)
        if draw_count is None:
            continue
        if draw_count < 0:
            report.error(f"{quiz_path.name} section {section.attrib.get('title', '')!r} has invalid qmd_numberofitems.")
            continue
        candidate_count = len(direct_children_by_local_name(section, "itemref")) + len(direct_children_by_local_name(section, "item"))
        if draw_count > candidate_count:
            report.error(
                f"{quiz_path.name} section {section.attrib.get('title', '')!r} draws {draw_count} "
                f"from {candidate_count} candidate questions."
            )

    if itemref_count:
        report.note(f"{quiz_path.name} itemrefs checked: {itemref_count}")
    else:
        report.warning(f"{quiz_path.name} has no itemrefs; inline quiz items may be intentional, but itemref joins were not tested.")


def workbook_question_codes(workbook_path: Path, report: ValidationReport) -> set[str]:
    try:
        wb = load_workbook(workbook_path, data_only=True)
    except Exception as exc:  # noqa: BLE001
        report.error(f"Could not read workbook {workbook_path}: {exc}")
        return set()

    codes: set[str] = set()
    skipped = {"README", "BANK_INDEX", "QUIZ_STRUCTURE"}
    for sheet_name in wb.sheetnames:
        if sheet_name.upper() in skipped:
            continue
        ws = wb[sheet_name]
        headers = [normalized_header(ws.cell(1, column).value) for column in range(1, ws.max_column + 1)]
        if "question_code" not in headers:
            continue
        code_col = headers.index("question_code") + 1
        for row_index in range(2, ws.max_row + 1):
            code = clean(ws.cell(row_index, code_col).value)
            if code:
                codes.add(code)
    return codes


def validate_workbook_coverage(
    workbook_path: Path | None,
    library_items: dict[str, dict[str, str]],
    report: ValidationReport,
) -> None:
    if workbook_path is None:
        return
    codes = workbook_question_codes(workbook_path, report)
    if not codes:
        report.warning("Workbook coverage check found no question_code values.")
        return
    package_displayids = {value["displayid"] for value in library_items.values() if value.get("displayid")}
    missing = sorted(code for code in codes if code not in package_displayids)
    if missing:
        report.error(f"Workbook question_code values missing from questiondb.xml qmd_displayid: {', '.join(missing)}")
    else:
        report.note(f"Workbook question_code coverage checked: {len(codes)}")


def _assessment_settings(root: ET.Element) -> dict[str, str]:
    settings: dict[str, str] = {}
    assessment = next((elem for elem in root.iter() if local_name(elem.tag) == "assessment"), None)
    if assessment is None:
        return settings
    for child in assessment:
        if local_name(child.tag) != "assess_procextension":
            continue
        for setting in child:
            if not list(setting):
                settings[local_name(setting.tag)] = clean(setting.text)
    return settings


def validate_authoring_projection(
    args: argparse.Namespace,
    package_dir: Path,
    library_items: dict[str, dict[str, str]],
    quiz_roots: list[tuple[Path, ET.Element | None]],
    report: ValidationReport,
) -> None:
    if not args.model:
        return
    try:
        projection = load_authoring_projection(
            Path(args.model).expanduser().resolve(),
            quiz_entity_key=args.quiz_entity_key,
            settings_path=Path(args.settings).expanduser().resolve() if args.settings else None,
            asset_root=[Path(root).expanduser().resolve() for root in args.asset_root] or None,
            promotion_receipt_path=(
                Path(args.promotion_receipt).expanduser().resolve()
                if args.promotion_receipt
                else None
            ),
            trial_authorization_path=(
                Path(args.phase5_candidate_authorization).expanduser().resolve()
                if args.phase5_candidate_authorization
                else None
            ),
        )
    except Exception as exc:  # noqa: BLE001
        report.error(f"Authoring contract projection failed: {exc}")
        return

    expected_codes = {row["question_code"] for row in projection["questions"]}
    actual_codes = {value["displayid"] for value in library_items.values() if value.get("displayid")}
    missing = sorted(expected_codes - actual_codes)
    extra = sorted(actual_codes - expected_codes)
    if missing:
        report.error(f"Authoring model question codes missing from package: {', '.join(missing)}")
    if extra:
        report.error(f"Package contains question codes outside selected authoring projection: {', '.join(extra)}")
    if not missing and not extra:
        report.note(f"Authoring model question identity coverage checked: {len(expected_codes)}")

    questiondb_root = parse_xml(package_dir / "questiondb.xml", report)
    actual_shuffle: dict[str, bool] = {}
    actual_keys: dict[str, list[str] | None] = {}
    if questiondb_root is not None:
        for item in questiondb_root.iter():
            if local_name(item.tag) != "item":
                continue
            code = metadata_value(item, "qmd_displayid")
            render_choice = next(
                (elem for elem in item.iter() if local_name(elem.tag) == "render_choice"),
                None,
            )
            actual_shuffle[code] = bool(render_choice is not None and render_choice.attrib.get("shuffle") == "yes")
            if metadata_value(item, 'qmd_questiontype') in {'Multiple Choice', 'True/False', 'Multi-Select'}:
                key = source_choice_projection(extract_source_response_facts(item))
                actual_keys[code] = ([o['option_key'] for o in key['options'] if o['correct']]
                                     if key['recognized'] else None)
    for row in projection['questions']:
        if row['target_question_type'] in {'MULTICHOICE', 'TRUEFALSE', 'MULTISELECT'}:
            expected = row['correct_option_key'].split(';')
            if row['target_question_type'] == 'TRUEFALSE':
                expected = [{'T': 'A', 'F': 'B'}.get(key, key) for key in expected]
            if actual_keys.get(row['question_code']) != expected:
                report.error(f"Question answer-key projection mismatch: {row['question_code']}; "
                             f"expected {expected}, found {actual_keys.get(row['question_code'])}.")
    expected_shuffle = {
        row["question_code"]: bool(row["randomize_answers"])
        for row in projection["questions"]
        if row["target_question_type"] != "WRITTEN_RESPONSE"
    }
    shuffle_differences = {
        code: {"expected": value, "actual": actual_shuffle.get(code)}
        for code, value in expected_shuffle.items()
        if actual_shuffle.get(code) != value
    }
    if shuffle_differences:
        report.error(f"Question answer-randomization projection mismatch: {shuffle_differences}")
    else:
        report.note(f"Question answer-randomization values checked: {len(expected_shuffle)}")

    if not args.library_only and len(quiz_roots) == 1 and quiz_roots[0][1] is not None:
        quiz_root = quiz_roots[0][1]
        actual_draws = {
            elem.attrib.get("title", ""): section_draw_count(elem)
            for elem in quiz_root.iter()
            if local_name(elem.tag) == "section" and section_draw_count(elem) is not None
        }
        expected_draws = {
            row["section_title"]: row["recommended_draw_count"] for row in projection["sections"]
        }
        if actual_draws != expected_draws:
            report.error(f"Authoring draw projection mismatch: expected {expected_draws}, found {actual_draws}")
        else:
            report.note(f"Authoring draw relationships checked: {len(expected_draws)}")

        expected_settings = projection["quiz_settings"]
        actual_settings = _assessment_settings(quiz_root)
        rendered_expected = {
            "is_active": "yes" if expected_settings["is_active"] else "no",
            "attempts_allowed": str(expected_settings["attempts_allowed"]),
            "time_limit": str(expected_settings["time_limit"]),
            "show_clock": "yes" if expected_settings["show_clock"] else "no",
            "enforce_time_limit": "yes" if expected_settings["enforce_time_limit"] else "no",
            "is_forward_only": "yes" if expected_settings["is_forward_only"] else "no",
        }
        differences = {
            name: {"expected": value, "actual": actual_settings.get(name)}
            for name, value in rendered_expected.items()
            if actual_settings.get(name) != value
        }
        if differences:
            report.error(f"Settings receipt projection mismatch: {differences}")
        else:
            report.note(f"Settings receipt values checked: {len(rendered_expected)}")

    for asset in projection["assets"]:
        package_path = package_dir.joinpath(*asset["archive_path"].parts)
        if not package_path.is_file():
            report.error(f"Projected asset missing from package: {asset['archive_path']}")
        elif file_hash(package_path) != asset["sha256"]:
            report.error(f"Projected asset checksum differs in package: {asset['archive_path']}")
    if projection["assets"] and not any("Projected asset" in error for error in report.errors):
        report.note(f"Authoring asset checksums checked: {len(projection['assets'])}")


def validate_run_receipt(path: Path | None, package_dir: Path, report: ValidationReport) -> None:
    if path is None:
        return
    try:
        receipt = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:  # noqa: BLE001
        report.error(f"Could not read run receipt {path}: {exc}")
        return
    issues = validate_contract(receipt, mode="transform")
    for issue in issues:
        if issue.severity == "error":
            report.error(f"Run receipt {issue.render()}")
        else:
            report.warning(f"Run receipt {issue.render()}")
    for row in [*receipt.get("inputs", []), *receipt.get("artifacts", [])]:
        if row.get("status") not in {"read", "emitted"}:
            continue
        artifact_path = Path(row["path"]).expanduser()
        if not artifact_path.is_file():
            report.error(f"Run receipt path is missing: {artifact_path}")
            continue
        if row.get("sha256") and file_hash(artifact_path) != row["sha256"]:
            report.error(f"Run receipt checksum mismatch: {artifact_path}")
    settings_path_value = receipt.get("extensions", {}).get("coursecraft.settings_receipt")
    if settings_path_value:
        settings_path = Path(settings_path_value).expanduser()
        try:
            settings_receipt = json.loads(settings_path.read_text(encoding="utf-8"))
            for issue in validate_contract(settings_receipt, mode="transform"):
                if issue.severity == "error":
                    report.error(f"Settings receipt {issue.render()}")
                else:
                    report.warning(f"Settings receipt {issue.render()}")
        except Exception as exc:  # noqa: BLE001
            report.error(f"Could not validate settings receipt {settings_path}: {exc}")
    recorded = {str(Path(row["path"]).resolve()) for row in receipt.get("artifacts", [])}
    unrecorded = [
        str(path.resolve())
        for path in sorted(item for item in package_dir.rglob("*") if item.is_file())
        if str(path.resolve()) not in recorded
    ]
    if unrecorded:
        report.error(f"Package files absent from run receipt: {', '.join(unrecorded)}")
    elif not any(error.startswith("Run receipt") or "run receipt" in error for error in report.errors):
        report.note(f"Run receipt checksums checked: {len(receipt.get('artifacts', []))} artifacts")


def validate_zip(package_dir: Path, zip_path: Path | None, report: ValidationReport) -> None:
    if zip_path is None:
        return
    if not zip_path.exists():
        report.error(f"Zip file does not exist: {zip_path}")
        return
    expected = {
        file_path.relative_to(package_dir).as_posix(): file_hash(file_path)
        for file_path in package_dir.rglob("*")
        if file_path.is_file()
    }
    try:
        with zipfile.ZipFile(zip_path) as archive:
            names = {name for name in archive.namelist() if not name.endswith("/")}
            missing = sorted(set(expected) - names)
            extra = sorted(names - set(expected))
            if missing:
                report.error(f"Zip is missing package files: {', '.join(missing)}")
            if extra:
                report.warning(f"Zip contains extra files not present in package folder: {', '.join(extra)}")
            for name in sorted(set(expected) & names):
                digest = hashlib.sha256(archive.read(name)).hexdigest()
                if digest != expected[name]:
                    report.error(f"Zip file is stale or differs from package folder: {name}")
    except zipfile.BadZipFile:
        report.error(f"Not a readable zip file: {zip_path}")
        return
    report.note(f"Zip root files checked: {len(expected)}")


def render_report(report: ValidationReport) -> str:
    lines = [
        "# Quiz Package Validation",
        "",
        f"- Package: `{report.package_dir}`",
        f"- Errors: {len(report.errors)}",
        f"- Warnings: {len(report.warnings)}",
        "",
        "## Notes",
        "",
    ]
    if report.notes:
        lines.extend(f"- {note}" for note in report.notes)
    else:
        lines.append("- No successful checks recorded.")
    lines.extend(["", "## Errors", ""])
    if report.errors:
        lines.extend(f"- {error}" for error in report.errors)
    else:
        lines.append("- None.")
    lines.extend(["", "## Warnings", ""])
    if report.warnings:
        lines.extend(f"- {warning}" for warning in report.warnings)
    else:
        lines.append("- None.")
    lines.append("")
    return "\n".join(lines)


def validate_package(args: argparse.Namespace) -> ValidationReport:
    package_dir = Path(args.package_dir).expanduser().resolve()
    report = ValidationReport(package_dir=package_dir)
    if not package_dir.exists():
        report.error(f"Package directory does not exist: {package_dir}")
        return report

    quiz_files = validate_required_files(package_dir, report, library_only=args.library_only)
    manifest_root = parse_xml(package_dir / "imsmanifest.xml", report)
    questiondb_root = parse_xml(package_dir / "questiondb.xml", report)
    parse_xml(package_dir / "orgunitconfig" / "orgunitconfig.xml", report)
    quiz_roots = [(quiz_file, parse_xml(quiz_file, report)) for quiz_file in quiz_files]

    manifest_hrefs = validate_manifest(package_dir, manifest_root, report)
    validate_package_closure(package_dir, manifest_hrefs, report, strict=args.strict_closure)
    validate_local_asset_references(package_dir, report)
    library_items = collect_library_items(questiondb_root, report)
    validate_truefalse_answer_labels(questiondb_root, report)
    validate_multiselect_keys(questiondb_root, report)
    for quiz_file, quiz_root in quiz_roots:
        validate_quiz_payload(package_dir, quiz_file, quiz_root, library_items, report)
        validate_multiselect_keys(quiz_root, report)
    validate_workbook_coverage(Path(args.workbook).expanduser().resolve() if args.workbook else None, library_items, report)
    validate_authoring_projection(args, package_dir, library_items, quiz_roots, report)
    validate_zip(package_dir, Path(args.zip).expanduser().resolve() if args.zip else None, report)
    validate_run_receipt(
        Path(args.run_receipt).expanduser().resolve() if args.run_receipt else None,
        package_dir,
        report,
    )

    return report


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Validate a generated Brightspace quiz package folder and optional import zip.")
    parser.add_argument("package_dir", help="Path to the generated package folder.")
    parser.add_argument("--workbook", default="", help="Optional normalized workbook used to generate the package.")
    parser.add_argument("--model", default="", help="Optional coursecraft.quiz/1 authoring model used to generate the package.")
    parser.add_argument("--settings", default="", help="Optional coursecraft.quiz_settings/1 receipt used by the model projection.")
    parser.add_argument("--quiz-entity-key", default="", help="Selected quiz key when the authoring model contains multiple quizzes.")
    parser.add_argument("--asset-root", action="append", default=[], help="Root for relative model asset source_path values; repeat for multiple folders.")
    parser.add_argument("--promotion-receipt", default="", help="Optional verified Quiz Binder promotion receipt for a strict-model build.")
    parser.add_argument("--phase5-candidate-authorization", default="", help="Optional exact local-only Phase 5 candidate authorization; requires --promotion-receipt.")
    parser.add_argument("--run-receipt", default="", help="Optional coursecraft.quiz_run/1 build receipt to validate against emitted files.")
    parser.add_argument("--library-only", action="store_true", help="Validate a question-library-only package without requiring quiz_d2l XML.")
    parser.add_argument(
        "--strict-closure",
        action="store_true",
        help="Treat non-manifest package files other than imsmanifest.xml/README.md as errors.",
    )
    parser.add_argument("--zip", default="", help="Optional import zip to compare against the package folder.")
    parser.add_argument("--output", default="", help="Optional markdown validation note path.")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    report = validate_package(args)
    rendered = render_report(report)
    if args.output:
        output_path = Path(args.output).expanduser()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(rendered, encoding="utf-8")
    print(rendered)
    return 2 if report.errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
