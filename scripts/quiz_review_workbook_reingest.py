#!/usr/bin/env python3
"""Re-ingest only the additive marginalia wing of a reviewer workbook.

The baseline workbook, edited workbook, and normalized source model are bound by
hash. Extracted/source-cell edits fail closed. Emitted annotations use the
existing ``coursecraft.quiz/1`` annotation shape; the surrounding overlay is a
local transport envelope, not a fourth CourseCraft contract.
"""

from __future__ import annotations

import argparse
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import sys
from typing import Any
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

from quiz_contracts import validate_contract
from quiz_review_projection import build_quiz_review_projection


FORMAT = "quiz-binder-local-decision-overlay-v0"
SHEET_NAME = "Quiz Questions"
UNRESOLVED_SHEET_NAME = "No Safe Library Match"
SETTINGS_SHEET_NAME = "Quiz Settings"
ROW_KEY_HEADERS = ("quiz", "question_number")
OCCURRENCE_ROW_KEY_HEADERS = ("occurrence_key",)
EDITABLE_HEADERS = {
    "proposed_permanent_code",
    "proposed_scoring_mode",
    "revised_question_text",
    "revised_response_options",
    "revised_answer_key",
    "revised_question_library_location",
    "revision_reason",
    "proposed_by",
    "proposed_at",
    "approval_status",
    "approved_by",
    "approved_at",
    "reviewer_note",
}
UNRESOLVED_EDITABLE_HEADERS = EDITABLE_HEADERS - {
    "revised_question_library_location"
}
REVISION_FIELDS = {
    "proposed_permanent_code": "/identity/permanent_code",
    "proposed_scoring_mode": "/scoring/mode",
    "revised_question_text": "/prompt/content",
    "revised_response_options": "/type_payload/options",
    "revised_answer_key": "/type_payload/answer_key",
    "revised_question_library_location": "/relationships/member_of",
}
SETTINGS_OBSERVATION_KEY_HEADER = "setting_observation_key"
SETTINGS_ROW_KEY_HEADERS = (SETTINGS_OBSERVATION_KEY_HEADER,)
LEGACY_SETTINGS_ROW_KEY_HEADERS = ("target_entity_key", "setting")
SETTINGS_EDITABLE_HEADERS = {
    "proposed_value",
    "revision_reason",
    "proposed_by",
    "proposed_at",
    "approval_status",
    "approved_by",
    "approved_at",
    "reviewer_note",
}


class ReingestError(ValueError):
    pass


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _canonical_digest(value: Any) -> str:
    encoded = json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _cell_value(value: Any) -> Any:
    if isinstance(value, datetime):
        if value.tzinfo is None:
            value = value.replace(tzinfo=timezone.utc)
        return value.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")
    return value


def _read_sheet(
    path: Path,
    *,
    sheet_name: str = SHEET_NAME,
    row_key_headers: tuple[str, ...] = ROW_KEY_HEADERS,
    editable_headers: set[str] = EDITABLE_HEADERS,
) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    if sheet_name not in workbook.sheetnames:
        raise ReingestError(f"Workbook is missing required sheet {sheet_name!r}: {path}")
    sheet = workbook[sheet_name]
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    if not headers or any(not header for header in headers):
        raise ReingestError(f"Workbook has a blank header in {sheet_name!r}: {path}")
    duplicates = sorted({header for header in headers if headers.count(header) > 1})
    if duplicates:
        raise ReingestError(f"Workbook has duplicate headers: {', '.join(duplicates)}")
    missing = sorted((set(row_key_headers) | editable_headers) - set(headers))
    if missing:
        raise ReingestError(f"Workbook is missing required headers: {', '.join(missing)}")

    rows: list[dict[str, Any]] = []
    for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(headers)), start=2):
        values = [_cell_value(cell.value) for cell in cells]
        if not any(value not in (None, "") for value in values):
            continue
        row = dict(zip(headers, values))
        row["__row_number"] = row_number
        for header, value in zip(headers, values):
            if isinstance(value, str) and value.startswith("="):
                raise ReingestError(
                    f"Formula cells are not accepted ({sheet_name}!{header}{row_number})."
                )
        rows.append(row)
    return headers, rows


def _row_key(
    row: dict[str, Any], headers: tuple[str, ...] = ROW_KEY_HEADERS
) -> tuple[str, ...]:
    return tuple(str(row.get(header) or "").strip() for header in headers)


def _index_rows(
    rows: list[dict[str, Any]],
    label: str,
    row_key_headers: tuple[str, ...] = ROW_KEY_HEADERS,
) -> dict[tuple[str, ...], dict[str, Any]]:
    result: dict[tuple[str, ...], dict[str, Any]] = {}
    for row in rows:
        key = _row_key(row, row_key_headers)
        if not all(key):
            raise ReingestError(f"{label} workbook row {row['__row_number']} has a blank row key.")
        if key in result:
            raise ReingestError(f"{label} workbook repeats row key {key!r}.")
        result[key] = row
    return result


def _model_row_index(model: dict[str, Any]) -> dict[tuple[str, str], str]:
    candidates: dict[tuple[str, str], set[str]] = {}
    for question in model.get("questions", []):
        entity_key = question["entity_key"]
        reviewer_row = question.get("extensions", {}).get("coursecraft.binder.reviewer_row")
        payloads: list[dict[str, Any]] = []
        if isinstance(reviewer_row, dict):
            payloads.append(reviewer_row)
        for raw in question.get("type_payload", {}).get("raw_response_models", []):
            payload = raw.get("payload")
            if isinstance(payload, dict):
                payloads.append(payload)
        for payload in payloads:
            quiz = payload.get("quiz", payload.get("quiz_title"))
            number = payload.get("question_number", payload.get("quiz_item_number"))
            if quiz in (None, "") or number in (None, ""):
                continue
            key = (str(quiz).strip(), str(number).strip())
            candidates.setdefault(key, set()).add(entity_key)
    ambiguous = {key: values for key, values in candidates.items() if len(values) != 1}
    if ambiguous:
        rendered = ", ".join(repr(key) for key in sorted(ambiguous))
        raise ReingestError(f"Model has ambiguous reviewer row identities: {rendered}")
    return {key: next(iter(values)) for key, values in candidates.items()}


def _model_occurrence_index(model: dict[str, Any]) -> dict[tuple[str], str]:
    """Map stable occurrence keys to their observed question entities.

    Resolved placements use ``referenced_question_key``.  An unresolved
    placement may still carry an observed question entity from the canonical
    extractor row; that entity remains a valid target for reviewer marginalia
    without pretending the source relationship itself was resolved.
    """

    candidates: dict[tuple[str], set[str]] = {}
    projection = build_quiz_review_projection(model)
    for occurrence in projection.get("question_occurrences", []):
        occurrence_key = str(occurrence.get("occurrence_key") or "").strip()
        entity_key = str(
            occurrence.get("referenced_question_key")
            or occurrence.get("observed_question_entity_key")
            or ""
        ).strip()
        if not occurrence_key or not entity_key:
            continue
        candidates.setdefault((occurrence_key,), set()).add(entity_key)
    ambiguous = {key: values for key, values in candidates.items() if len(values) != 1}
    if ambiguous:
        rendered = ", ".join(repr(key[0]) for key in sorted(ambiguous))
        raise ReingestError(f"Model has ambiguous occurrence identities: {rendered}")
    return {key: next(iter(values)) for key, values in candidates.items()}


def _timestamp(
    value: Any, field: str, row_number: int, *, required: bool = True
) -> str | None:
    text = str(value or "").strip()
    if not text:
        if not required:
            return None
        raise ReingestError(f"Row {row_number} requires {field}.")
    normalized = text.replace("Z", "+00:00")
    try:
        parsed = datetime.fromisoformat(normalized)
    except ValueError as exc:
        raise ReingestError(f"Row {row_number} has invalid {field}: {text!r}.") from exc
    if parsed.tzinfo is None:
        raise ReingestError(f"Row {row_number} {field} must include a timezone.")
    return parsed.astimezone(timezone.utc).isoformat().replace("+00:00", "Z")


def _review_metadata(
    row: dict[str, Any], row_number: int, decision: str, metadata_policy: str
) -> tuple[str | None, str | None, str | None, str | None, str | None]:
    """Preserve supplied attribution; optional blanks do not invent review events."""
    required = metadata_policy == "required"
    reason = str(row.get("revision_reason") or "").strip() or None
    proposer = str(row.get("proposed_by") or "").strip() or None
    for field, value in (("revision_reason", reason), ("proposed_by", proposer)):
        if required and value is None:
            raise ReingestError(f"Row {row_number} requires {field}.")
    proposed_at = _timestamp(
        row.get("proposed_at"), "proposed_at", row_number, required=required
    )
    approver = approved_at = None
    if decision in {"accepted", "rejected"}:
        approver = str(row.get("approved_by") or "").strip() or None
        if required and approver is None:
            raise ReingestError(f"Row {row_number} requires approved_by.")
        approved_at = _timestamp(
            row.get("approved_at"), "approved_at", row_number, required=required
        )
    return reason, proposer, proposed_at, approver, approved_at


def _stable_id(prefix: str, payload: Any) -> str:
    return f"{prefix}.{_canonical_digest(payload)[:24]}"


def _validate_source_cells(
    headers: list[str],
    baseline: dict[str, Any],
    edited: dict[str, Any],
    editable_headers: set[str] = EDITABLE_HEADERS,
    sheet_name: str = SHEET_NAME,
) -> None:
    row_number = edited["__row_number"]
    for header in headers:
        if header in editable_headers:
            continue
        if baseline.get(header) != edited.get(header):
            raise ReingestError(
                f"Source field edit rejected at {sheet_name} row {row_number}, "
                f"column {header!r}."
            )


def _workbook_has_sheet(path: Path, sheet_name: str) -> bool:
    return sheet_name in load_workbook(path, read_only=True).sheetnames


def _placeholder_sheet_values(path: Path, sheet_name: str) -> list[list[Any]] | None:
    workbook = load_workbook(path, data_only=False, read_only=True)
    if sheet_name not in workbook.sheetnames:
        return None
    sheet = workbook[sheet_name]
    values = [[cell.value for cell in row] for row in sheet.iter_rows()]
    if not values or [str(value or "").strip() for value in values[0]] != ["note"]:
        return None
    return values


def _setting_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    text = value.strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        return text


def _materialize_setting_decisions(
    baseline_path: Path, edited_path: Path, *, metadata_policy: str = "required"
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    baseline_has = _workbook_has_sheet(baseline_path, SETTINGS_SHEET_NAME)
    edited_has = _workbook_has_sheet(edited_path, SETTINGS_SHEET_NAME)
    if baseline_has != edited_has:
        raise ReingestError("Edited and baseline workbooks disagree on Quiz Settings presence.")
    if not baseline_has:
        return [], [], []
    baseline_headers, baseline_rows = _read_sheet(
        baseline_path,
        sheet_name=SETTINGS_SHEET_NAME,
        row_key_headers=LEGACY_SETTINGS_ROW_KEY_HEADERS,
        editable_headers=SETTINGS_EDITABLE_HEADERS,
    )
    edited_headers, edited_rows = _read_sheet(
        edited_path,
        sheet_name=SETTINGS_SHEET_NAME,
        row_key_headers=LEGACY_SETTINGS_ROW_KEY_HEADERS,
        editable_headers=SETTINGS_EDITABLE_HEADERS,
    )
    if baseline_headers != edited_headers:
        raise ReingestError("Edited Quiz Settings headers do not match the baseline workbook.")
    row_key_headers = (
        SETTINGS_ROW_KEY_HEADERS
        if SETTINGS_OBSERVATION_KEY_HEADER in baseline_headers
        else LEGACY_SETTINGS_ROW_KEY_HEADERS
    )
    baseline_index = _index_rows(
        baseline_rows, "Baseline Quiz Settings", row_key_headers
    )
    edited_index = _index_rows(edited_rows, "Edited Quiz Settings", row_key_headers)
    if set(baseline_index) != set(edited_index):
        raise ReingestError("Edited Quiz Settings row identities do not match the baseline.")

    decisions: list[dict[str, Any]] = []
    inputs: list[dict[str, Any]] = []
    diffs: list[dict[str, Any]] = []
    proposed_settings: dict[tuple[str, str], int] = {}
    for row_key in sorted(baseline_index):
        baseline = baseline_index[row_key]
        edited = edited_index[row_key]
        _validate_source_cells(
            baseline_headers,
            baseline,
            edited,
            SETTINGS_EDITABLE_HEADERS,
            SETTINGS_SHEET_NAME,
        )
        changed = [
            header
            for header in SETTINGS_EDITABLE_HEADERS
            if baseline.get(header) != edited.get(header)
        ]
        if not changed:
            continue
        row_number = edited["__row_number"]
        target_entity_key = str(edited.get("target_entity_key") or "").strip()
        setting = str(edited.get("setting") or "").strip()
        observation_key = str(
            edited.get(SETTINGS_OBSERVATION_KEY_HEADER) or ""
        ).strip()
        has_proposal = str(edited.get("proposed_value") or "").strip() != ""
        decision_status = str(edited.get("approval_status") or "open").strip().lower()
        if decision_status not in {"open", "accepted", "rejected"}:
            raise ReingestError(
                f"Quiz Settings row {row_number} approval_status must be open, accepted, or rejected."
            )
        if not has_proposal and decision_status != "open":
            raise ReingestError(
                f"Quiz Settings row {row_number} has a decision but no proposed value."
            )
        if has_proposal:
            setting_target = (target_entity_key, setting)
            if setting_target in proposed_settings:
                raise ReingestError(
                    "Quiz Settings has more than one proposal for "
                    f"{setting!r} on {target_entity_key!r} (rows "
                    f"{proposed_settings[setting_target]} and {row_number}); keep "
                    "one proposal so the effective value is unambiguous."
                )
            proposed_settings[setting_target] = row_number
            reason, proposer, proposed_at, approver, approved_at = _review_metadata(
                edited, row_number, decision_status, metadata_policy
            )
            value = _setting_value(edited.get("proposed_value"))
            decision_identity = {
                "target_entity_key": target_entity_key,
                "setting": setting,
                "value": value,
                "proposer": proposer,
                "proposed_at": proposed_at,
                "setting_observation_key": observation_key or None,
            }
            decision_id = _stable_id("setting.decision", decision_identity)
            decisions.append(
                {
                    "decision_id": decision_id,
                    "target_entity_key": target_entity_key,
                    "setting": setting,
                    "value": value,
                    "status": decision_status,
                    "proposed_by": proposer,
                    "proposed_at": proposed_at,
                    "approved_by": approver,
                    "approved_at": approved_at,
                    "reason": reason,
                    "setting_observation_key": observation_key or None,
                }
            )
            if decision_status == "accepted":
                inputs.append(
                    {
                        "input_id": _stable_id("in.quiz-decision", decision_identity),
                        "setting": setting,
                        "layer": "quiz_decision",
                        "target_entity_key": target_entity_key,
                        "state": "known",
                        "value": value,
                        "source_reference": f"Quiz Binder settings marginalia row {row_number}",
                        "source_evidence_keys": [],
                        "extensions": {
                            "coursecraft.binder.decision_id": decision_id,
                            "coursecraft.binder.reason": reason,
                            "coursecraft.binder.proposed_by": proposer,
                            "coursecraft.binder.proposed_at": proposed_at,
                            "coursecraft.binder.approved_by": approver,
                            "coursecraft.binder.approved_at": approved_at,
                            "coursecraft.binder.setting_observation_key": (
                                observation_key or None
                            ),
                        },
                    }
                )
        diffs.append(
            {
                "row_number": row_number,
                "row_key_headers": list(row_key_headers),
                "row_key": list(row_key),
                "target_entity_key": target_entity_key,
                "changed_fields": sorted(changed),
                "decision": decision_status,
            }
        )
    return decisions, inputs, diffs


def _question_sheet_inputs(
    baseline_path: Path, edited_path: Path
) -> list[
    tuple[
        str,
        list[str],
        dict[tuple[str, ...], dict[str, Any]],
        dict[tuple[str, ...], dict[str, Any]],
        set[str],
        tuple[str, ...],
    ]
]:
    result = []
    for sheet_name, editable_headers, required in (
        (SHEET_NAME, EDITABLE_HEADERS, True),
        (UNRESOLVED_SHEET_NAME, UNRESOLVED_EDITABLE_HEADERS, False),
    ):
        baseline_has = _workbook_has_sheet(baseline_path, sheet_name)
        edited_has = _workbook_has_sheet(edited_path, sheet_name)
        if baseline_has != edited_has:
            raise ReingestError(
                f"Edited and baseline workbooks disagree on {sheet_name!r} presence."
            )
        if not baseline_has:
            if required:
                raise ReingestError(f"Workbook is missing required sheet {sheet_name!r}.")
            continue
        if not required:
            baseline_placeholder = _placeholder_sheet_values(baseline_path, sheet_name)
            edited_placeholder = _placeholder_sheet_values(edited_path, sheet_name)
            if baseline_placeholder is not None or edited_placeholder is not None:
                if baseline_placeholder != edited_placeholder:
                    raise ReingestError(
                        f"Source field edit rejected in placeholder sheet {sheet_name!r}."
                    )
                continue
        baseline_headers, baseline_rows = _read_sheet(
            baseline_path,
            sheet_name=sheet_name,
            row_key_headers=(),
            editable_headers=editable_headers,
        )
        edited_headers, edited_rows = _read_sheet(
            edited_path,
            sheet_name=sheet_name,
            row_key_headers=(),
            editable_headers=editable_headers,
        )
        if baseline_headers != edited_headers:
            raise ReingestError(
                f"Edited {sheet_name} headers do not exactly match the baseline workbook."
            )
        row_key_headers = (
            OCCURRENCE_ROW_KEY_HEADERS
            if all(header in baseline_headers for header in OCCURRENCE_ROW_KEY_HEADERS)
            else ROW_KEY_HEADERS
        )
        missing_row_keys = sorted(set(row_key_headers) - set(baseline_headers))
        if missing_row_keys:
            raise ReingestError(
                f"Workbook is missing required row identity headers: "
                f"{', '.join(missing_row_keys)}"
            )
        baseline_index = _index_rows(
            baseline_rows, f"Baseline {sheet_name}", row_key_headers
        )
        edited_index = _index_rows(
            edited_rows, f"Edited {sheet_name}", row_key_headers
        )
        if set(baseline_index) != set(edited_index):
            raise ReingestError(
                f"Edited {sheet_name} row identities do not exactly match the baseline."
            )
        result.append(
            (
                sheet_name,
                baseline_headers,
                baseline_index,
                edited_index,
                editable_headers,
                row_key_headers,
            )
        )
    return result


def materialize_workbook_decisions(
    model_path: Path, baseline_path: Path, edited_path: Path, *,
    metadata_policy: str = "required",
) -> dict[str, Any]:
    if metadata_policy not in {"required", "optional"}:
        raise ReingestError(f"Unknown review metadata policy: {metadata_policy!r}.")
    def assessment_layout(path: Path) -> bool:
        try:
            with zipfile.ZipFile(path) as archive:
                root = ET.fromstring(archive.read("xl/workbook.xml"))
        except (OSError, zipfile.BadZipFile, KeyError, ET.ParseError) as exc:
            raise ReingestError(f"Workbook is missing or changed (not readable XLSX): {path.name}") from exc
        return any(node.get("name") == "_Assessment Review" for node in root.iter())

    if assessment_layout(baseline_path):
        from quiz_assessment_review import materialize_assessment_decisions
        return materialize_assessment_decisions(
            model_path, baseline_path, edited_path, metadata_policy=metadata_policy,
        )
    if assessment_layout(edited_path):
        raise ReingestError("Grouped edits require their assessment review packet baseline")
    model = json.loads(model_path.read_text(encoding="utf-8"))
    model_issues = validate_contract(model, mode="transform")
    if model_issues:
        raise ReingestError(f"Source model is invalid: {model_issues[0].render()}")

    sheet_inputs = _question_sheet_inputs(baseline_path, edited_path)
    legacy_model_index = _model_row_index(model)
    occurrence_model_index = _model_occurrence_index(model)

    annotations: list[dict[str, Any]] = []
    row_diffs: list[dict[str, Any]] = []
    changed_row_keys: set[tuple[tuple[str, ...], tuple[str, ...]]] = set()
    question_rows: list[
        tuple[
            str,
            list[str],
            tuple[str, ...],
            dict[str, Any],
            dict[str, Any],
            set[str],
            tuple[str, ...],
        ]
    ] = []
    for (
        sheet_name,
        baseline_headers,
        baseline_index,
        edited_index,
        editable_headers,
        row_key_headers,
    ) in sheet_inputs:
        question_rows.extend(
            (
                sheet_name,
                baseline_headers,
                row_key,
                baseline_index[row_key],
                edited_index[row_key],
                editable_headers,
                row_key_headers,
            )
            for row_key in sorted(baseline_index)
        )

    for (
        sheet_name,
        baseline_headers,
        row_key,
        baseline,
        edited,
        editable_headers,
        row_key_headers,
    ) in question_rows:
        _validate_source_cells(
            baseline_headers, baseline, edited, editable_headers, sheet_name
        )

        changed = [
            header
            for header in editable_headers
            if baseline.get(header) != edited.get(header)
        ]
        if not changed:
            continue
        stable_row_identity = (row_key_headers, row_key)
        if stable_row_identity in changed_row_keys:
            raise ReingestError(
                f"Reviewer row {row_key!r} was edited in more than one question sheet; "
                "keep its marginalia in one place."
            )
        changed_row_keys.add(stable_row_identity)
        entity_key = (
            occurrence_model_index.get(row_key)
            if row_key_headers == OCCURRENCE_ROW_KEY_HEADERS
            else legacy_model_index.get(row_key)
        )
        if not entity_key:
            raise ReingestError(
                f"No unique model entity maps to reviewer row {row_key!r}."
            )
        row_number = edited["__row_number"]
        revisions = [
            header
            for header in REVISION_FIELDS
            if str(edited.get(header) or "").strip()
        ]
        decision = str(edited.get("approval_status") or "open").strip().lower()
        if decision not in {"open", "accepted", "rejected"}:
            raise ReingestError(
                f"Row {row_number} approval_status must be open, accepted, or rejected."
            )
        if not revisions and decision != "open":
            raise ReingestError(f"Row {row_number} has a decision but no proposed revision.")

        if revisions:
            reason, proposer, proposed_at, approver, approved_at = _review_metadata(
                edited, row_number, decision, metadata_policy
            )

            for header in revisions:
                value = edited.get(header)
                identity = {
                    "target_entity_key": entity_key,
                    "field_path": REVISION_FIELDS[header],
                    "value": value,
                    "actor": proposer,
                    "timestamp": proposed_at,
                    "reason": reason,
                }
                proposal_id = _stable_id("ann.proposal", identity)
                proposal = {
                    "annotation_id": proposal_id,
                    "target_entity_key": entity_key,
                    "field_path": REVISION_FIELDS[header],
                    "kind": "proposed_revision",
                    "value": value,
                    "actor": proposer,
                    "timestamp": proposed_at,
                    "source_evidence_keys": [],
                    "status": decision,
                    "extensions": {
                        "coursecraft.binder.reason": reason,
                        "coursecraft.binder.workbook_sheet": sheet_name,
                        "coursecraft.binder.workbook_row": row_number,
                        "coursecraft.binder.source_column": header,
                    },
                }
                annotations.append(proposal)
                if decision == "accepted":
                    approval_identity = {
                        "proposal_id": proposal_id,
                        "actor": approver,
                        "timestamp": approved_at,
                    }
                    annotations.append(
                        {
                            "annotation_id": _stable_id("ann.approval", approval_identity),
                            "target_entity_key": entity_key,
                            "field_path": REVISION_FIELDS[header],
                            "kind": "approved_change",
                            "value": value,
                            "actor": approver,
                            "timestamp": approved_at,
                            "source_evidence_keys": [],
                            "status": "accepted",
                            "extensions": {
                                "coursecraft.binder.proposal_id": proposal_id,
                                "coursecraft.binder.reason": reason,
                                "coursecraft.binder.workbook_sheet": sheet_name,
                                "coursecraft.binder.workbook_row": row_number,
                            },
                        }
                    )

        reviewer_note = str(edited.get("reviewer_note") or "").strip()
        if reviewer_note and baseline.get("reviewer_note") != edited.get("reviewer_note"):
            actor = str(edited.get("proposed_by") or "").strip() or None
            timestamp = (
                _timestamp(edited.get("proposed_at"), "proposed_at", row_number,
                           required=metadata_policy == "required")
                if edited.get("proposed_at")
                else None
            )
            note_identity = {
                "target_entity_key": entity_key,
                "value": reviewer_note,
                "actor": actor,
                "timestamp": timestamp,
            }
            annotations.append(
                {
                    "annotation_id": _stable_id("ann.note", note_identity),
                    "target_entity_key": entity_key,
                    "field_path": "/reviewer_note",
                    "kind": "reviewer_annotation",
                    "value": reviewer_note,
                    "actor": actor,
                    "timestamp": timestamp,
                    "source_evidence_keys": [],
                    "status": "informational",
                    "extensions": {
                        "coursecraft.binder.workbook_sheet": sheet_name,
                        "coursecraft.binder.workbook_row": row_number,
                    },
                }
            )

        row_diffs.append(
            {
                "sheet": sheet_name,
                "row_number": row_number,
                "row_key_headers": list(row_key_headers),
                "row_key": list(row_key),
                "target_entity_key": entity_key,
                "changed_fields": sorted(changed),
                "decision": decision,
            }
        )

    settings_decisions, settings_inputs, settings_row_diffs = _materialize_setting_decisions(
        baseline_path, edited_path, metadata_policy=metadata_policy
    )

    test_model = deepcopy(model)
    test_model["annotations"].extend(annotations)
    issues = validate_contract(test_model, mode="transform")
    if issues:
        raise ReingestError(f"Materialized annotations are invalid: {issues[0].render()}")

    for index, setting_input in enumerate(settings_inputs):
        validation_receipt = {
            "schema": "coursecraft.quiz_settings/1",
            "receipt_id": f"cc:settings-receipt:binder:{index}",
            "run_id": f"cc:run:quiz-binder:settings:{index}",
            "target_quiz_key": setting_input["target_entity_key"],
            "profile_key": None,
            "inputs": [setting_input],
            "resolutions": [],
            "diagnostics": [],
            "extensions": {"coursecraft.binder.route_status": "local_only"},
        }
        setting_issues = validate_contract(validation_receipt, mode="transform")
        if setting_issues:
            raise ReingestError(
                f"Materialized settings input is invalid: {setting_issues[0].render()}"
            )

    materialized = {
        "annotations": annotations,
        "settings_decisions": settings_decisions,
        "settings_inputs": settings_inputs,
    }

    return {
        "format": FORMAT,
        "source_model": {
            "model_id": model["model_id"],
            "sha256": sha256_file(model_path),
            "source_fingerprint": model["source"]["fingerprint"]["digest"],
        },
        "workbooks": {
            "baseline_sha256": sha256_file(baseline_path),
            "edited_sha256": sha256_file(edited_path),
        },
        "materialized_view_fingerprint": _canonical_digest(materialized),
        "row_diffs": row_diffs,
        "settings_row_diffs": settings_row_diffs,
        "annotations": annotations,
        "settings_decisions": settings_decisions,
        "settings_inputs": settings_inputs,
        "content_minimized_summary": {
            "changed_row_count": len(row_diffs) + len(settings_row_diffs),
            "annotation_count": len(annotations),
            "accepted_change_count": sum(
                row["kind"] == "approved_change" for row in annotations
            ),
            "accepted_setting_input_count": len(settings_inputs),
        },
    }


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Materialize Quiz Binder workbook marginalia into existing annotation records."
    )
    parser.add_argument("--model", required=True)
    parser.add_argument("--baseline-workbook", required=True)
    parser.add_argument("--edited-workbook", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument(
        "--metadata-policy", choices=("required", "optional"), default="required",
        help="Require review attribution/reason/dates, or preserve missing values as null."
    )
    args = parser.parse_args()
    try:
        result = materialize_workbook_decisions(
            Path(args.model), Path(args.baseline_workbook), Path(args.edited_workbook),
            metadata_policy=args.metadata_policy,
        )
    except (OSError, json.JSONDecodeError, ReingestError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2
    Path(args.output).write_text(
        json.dumps(result, indent=2, ensure_ascii=False, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(result["content_minimized_summary"], sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
