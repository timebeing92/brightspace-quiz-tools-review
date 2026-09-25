#!/usr/bin/env python3
"""Deterministic fresh-question intake. No revisions, inference, or evidence promotion."""
from __future__ import annotations

import argparse
from collections import defaultdict
from decimal import Decimal, InvalidOperation
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
from zipfile import BadZipFile, ZipFile
import xml.etree.ElementTree as ET

from jsonschema import Draft202012Validator

from quiz_contracts import validate_contract
from quiz_question_types import canonical_question_type, question_type_catalog
from quiz_build_support import CAPABILITY_REGISTRY_PATH, sha256_file

SCHEMA = 'coursecraft.quiz_draft/1'
PARSER_VERSION = 'coursecraft.quiz_draft_intake/1'
SCHEMA_ROOT = Path(__file__).resolve().parents[1] / 'workspace/reference/schemas/quiz'
INPUT_SCHEMA_PATH = SCHEMA_ROOT / 'quiz_draft_intake_schema.json'
REPORT_SCHEMA_PATH = SCHEMA_ROOT / 'quiz_draft_intake_report_schema.json'
MAX_EXACT_INTEGER = 2**53 - 1
TABLES = {
    'Questions': ['question_code', 'question_type', 'question_text', 'points', 'pool_code', 'answer_key', 'feedback', 'content_format'],
    'Responses': ['question_code', 'response_key', 'role', 'text', 'content_format', 'correct', 'case_sensitive', 'group_key', 'match_key', 'position'],
    'Pools': ['pool_code', 'title'],
    'Draws': ['draw_code', 'pool_code', 'title', 'draw_count', 'points_per_question'],
}
FORMATS = {'plain_text', 'html', 'mathml', 'latex'}
CODE = re.compile(r'^[A-Za-z0-9][A-Za-z0-9._-]*$')


class IntakeError(ValueError):
    pass


class WorkbookData(dict):
    """Input data plus physical row locators; locators are not authored fields."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.row_numbers = {}


def require(condition, message):
    if not condition:
        raise IntakeError(message)


def code(value, where):
    require(isinstance(value, str) and bool(CODE.fullmatch(value)), f'{where}: use a nonempty stable code (letters, digits, dot, underscore, hyphen).')
    return value


def number(value, where, integer=False):
    require(isinstance(value, (str, int, float, Decimal)) and not isinstance(value, bool),
            f'{where}: positive number required; booleans are not numbers.')
    try:
        exact = Decimal(str(value))
    except InvalidOperation:
        raise IntakeError(f'{where}: positive number required.') from None
    require(exact.is_finite() and 0 < exact <= MAX_EXACT_INTEGER,
            f'{where}: positive finite number no greater than {MAX_EXACT_INTEGER} required.')
    require(not integer or exact == exact.to_integral_value(), f'{where}: positive integer required.')
    # Canonical JSON and current QTI output use binary64 / nine decimal places.
    # Refuse precision loss instead of silently rounding authored point values.
    require(exact.normalize().as_tuple().exponent >= -9,
            f'{where}: at most nine fractional decimal places are supported; no rounding is performed.')
    result = float(exact)
    require(Decimal(str(result)) == exact, f'{where}: value cannot round-trip through a JSON number without precision loss.')
    return int(exact) if exact == exact.to_integral_value() else result


def populated(value):
    return value is not None and value != ''


def boolean(value, where):
    if isinstance(value, bool):
        return value
    require(isinstance(value, str) and value.lower() in ('true', 'false'), f'{where}: enter TRUE or FALSE explicitly.')
    return value.lower() == 'true'


def rich(value, fmt, where):
    require(isinstance(value, str) and bool(value.strip()), f'{where}: nonempty text required.')
    require(isinstance(fmt, str) and fmt in FORMATS, f'{where}: explicit content_format required: {sorted(FORMATS)}.')
    # The canonical formattedContent enum has no latex member. Keep exact bytes
    # and an explicit encoding extension; do not pretend LaTeX is HTML or MathML.
    return {'format': 'plain_text' if fmt == 'latex' else fmt, 'content': value,
            'extensions': {'coursecraft.source_encoding': 'latex'} if fmt == 'latex' else {}}


def identity(value):
    return {'strategy': 'permanent_code', 'permanent_code': value, 'source_aliases': [], 'content_fingerprints': [], 'extensions': {}}


def location(data, parts):
    path = '$' + ''.join(f'[{part}]' if isinstance(part, int) else f'.{part}' for part in parts)
    if isinstance(data, WorkbookData) and len(parts) >= 2:
        physical = data.row_numbers.get((parts[0], parts[1]))
        if physical is not None:
            field = '.' + str(parts[2]) if len(parts) > 2 else ''
            return f'{parts[0]} row {physical}{field} ({path})'
    return path


def validate_shape(data, schema_path=INPUT_SCHEMA_PATH):
    def finite(value, parts=()):
        if isinstance(value, (float, Decimal)):
            require(Decimal(str(value)).is_finite(), f'{location(data, parts)}: finite JSON number required.')
        elif isinstance(value, dict):
            for key, child in value.items():
                finite(child, (*parts, key))
        elif isinstance(value, list):
            for index, child in enumerate(value):
                finite(child, (*parts, index))

    finite(data)
    validator = Draft202012Validator(json.loads(schema_path.read_text(encoding='utf-8')))
    errors = sorted(validator.iter_errors(data), key=lambda err: str(list(err.absolute_path)))
    if errors:
        error = errors[0]
        raise IntakeError(f'{location(data, list(error.absolute_path))}: {error.message}')


def read_json(path):
    def pairs(values):
        out = {}
        for key, value in values:
            require(key not in out, f'JSON: duplicate object key {key!r}.')
            out[key] = value
        return out

    def constant(value):
        raise IntakeError(f'JSON: {value} is not a finite JSON number.')

    snapshot = path if isinstance(path, bytes) else Path(path).read_bytes()
    return json.loads(snapshot.decode('utf-8'), object_pairs_hook=pairs,
                      parse_float=Decimal, parse_constant=constant)


def template_spec():
    return {'schema': SCHEMA, 'tables': TABLES, 'types': question_type_catalog()}


def workbook_numeric_tokens(path):
    """Preserve stored numeric lexemes before a spreadsheet reader casts floats."""
    import posixpath
    namespace = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    relationship = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
    tokens = {}
    with ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read('xl/workbook.xml'))
        relations = {row.get('Id'): row.get('Target') for row in ET.fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
        for sheet in workbook.findall(f'.//{namespace}sheet'):
            title = sheet.get('name')
            if title not in {'Config', *TABLES}:
                continue
            target = relations.get(sheet.get(relationship))
            require(isinstance(target, str) and bool(target), f'{title}: missing workbook sheet relationship.')
            part = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
            document = ET.fromstring(archive.read(part))
            for cell in document.findall(f'.//{namespace}c'):
                value = cell.find(f'{namespace}v')
                if cell.get('t', 'n') == 'n' and value is not None and value.text is not None and cell.find(f'{namespace}f') is None:
                    try:
                        tokens[(title, cell.get('r'))] = Decimal(value.text)
                    except (InvalidOperation, TypeError):
                        raise IntakeError(f'{title} cell {cell.get("r")}: invalid stored numeric value.') from None
    return tokens


def read_workbook(path):
    import openpyxl  # Read only; authoring templates are made with Artifact Tool.
    snapshot = path if isinstance(path, bytes) else Path(path).read_bytes()
    numeric_tokens = workbook_numeric_tokens(BytesIO(snapshot))
    wb = openpyxl.load_workbook(BytesIO(snapshot), read_only=True, data_only=False)
    def value(cell, title):
        if cell.data_type == 'n' and cell.value is not None:
            return numeric_tokens.get((title, cell.coordinate), cell.value)
        return cell.value
    try:
        require(set(wb.sheetnames) == {'START HERE', 'Config', 'Types', *TABLES}, 'Workbook tabs do not match quiz_draft/1; preserve all template tabs.')
        config = {}
        result = WorkbookData(config=config)
        config_headers = [c.value for c in next(wb['Config'].iter_rows(), ())]
        while config_headers and config_headers[-1] is None:
            config_headers.pop()
        require(config_headers == ['setting', 'value'],
                'Config row 1: headers must be setting, value.')
        for physical_row, row in enumerate(wb['Config'].iter_rows(min_row=2), 2):
            if all(c.value is None for c in row):
                continue
            where = f'Config row {physical_row}'
            require(len(row) >= 2 and not any(c.data_type == 'f' for c in row), f'{where}: formulas are not accepted.')
            key, entry = row[0].value, value(row[1], 'Config')
            require(isinstance(key, str), f'{where}.setting: text key required.')
            require(key not in config and not any(c.value is not None for c in row[2:]), f'{where}: duplicate key or extra value.')
            config[key] = entry
            result.row_numbers[('config', key)] = physical_row
        for title, headers in TABLES.items():
            sheet = wb[title]
            actual = [c.value for c in next(sheet.iter_rows(), ())]
            while actual and actual[-1] is None:
                actual.pop()
            require(actual == headers, f'{title}: headers must match the versioned template.')
            rows = []
            for physical_row, row in enumerate(sheet.iter_rows(min_row=2), 2):
                if all(c.value is None for c in row):
                    continue
                for cell, header in zip(row, headers):
                    require(cell.data_type != 'f', f'{title} row {physical_row}.{header}: formulas are not accepted; enter literal content.')
                require(not any(c.value is not None for c in row[len(headers):]), f'{title} row {physical_row}: extra populated column.')
                result.row_numbers[(title, len(rows))] = physical_row
                rows.append({h: value(c, title) for h, c in zip(headers, row)})
            result[title] = rows
        return result
    finally:
        wb.close()


def ingest(data, *, source_name='draft.json', source_digest=None, input_format=None):
    validate_shape(data)
    for title, fields in {'Questions': ['points'], 'Responses': ['position'], 'Draws': ['draw_count', 'points_per_question']}.items():
        for index, row in enumerate(data[title]):
            for field in fields:
                if populated(row.get(field)):
                    number(row[field], location(data, [title, index, field]), field in ('position', 'draw_count'))
    input_format = input_format or ('xlsx' if isinstance(data, WorkbookData) else 'json')
    require(input_format in ('xlsx', 'json'), 'source.input_format: expected xlsx or json.')
    require(isinstance(source_name, str) and bool(source_name), 'source.name: nonempty text required.')
    require(source_digest is None or isinstance(source_digest, str) and bool(re.fullmatch('[0-9a-f]{64}', source_digest)),
            'source.sha256: lowercase 64-character SHA-256 required.')
    cfg = data['config']
    lineage = code(cfg['lineage_code'], 'Config.lineage_code')
    quiz_code = code(cfg['quiz_code'], 'Config.quiz_code') if populated(cfg['quiz_code']) else None
    require(bool(quiz_code) == populated(cfg['quiz_title']), 'Config.quiz_code/quiz_title: supply both or leave both blank for a library.')
    if quiz_code:
        require(bool(cfg['quiz_title'].strip()), 'Config.quiz_title: nonblank title required.')
    digest = source_digest or hashlib.sha256(json.dumps(data, sort_keys=True, ensure_ascii=False,
                                                       default=lambda value: number(value, 'source numeric value'), allow_nan=False).encode()).hexdigest()
    source = {'name': source_name, 'input_format': input_format, 'sha256': digest,
              'fingerprint_basis': 'file_bytes' if source_digest else 'canonical_json',
              'parser': PARSER_VERSION, 'input_schema_sha256': sha256_file(INPUT_SCHEMA_PATH)}
    key = lambda kind, value: f'cc:{kind}:draft:{lineage}:{value}'
    model = {'schema': 'coursecraft.quiz/1', 'model_id': key('model', digest[:24]), 'run_id': key('run', digest[:24]),
             'source': {'source_key': f'cc:source:{digest}', 'source_lineage_key': key('lineage', lineage), 'source_kind': 'workbook' if input_format == 'xlsx' else 'unknown',
                        'fingerprint': {'algorithm': 'sha256', 'digest': digest, 'scope': 'file' if source_digest else 'unknown'}, 'references': [source_name],
                        'extensions': {'coursecraft.intake_schema': SCHEMA, 'coursecraft.draft_input': source}},
             **{name: [] for name in ('quizzes', 'structures', 'questions', 'assets', 'resources', 'relationships', 'evidence', 'lineage', 'settings_observations', 'annotations', 'diagnostics')},
             'extensions': {'coursecraft.authoring_projection': 'quiz-builder-v1', 'coursecraft.intake_schema': SCHEMA,
                            'coursecraft.intake_catalog': {'registry_sha256': sha256_file(CAPABILITY_REGISTRY_PATH),
                                'catalog_sha256': hashlib.sha256(json.dumps(question_type_catalog(), sort_keys=True).encode()).hexdigest()}}}
    base = lambda: {'source_evidence_keys': [], 'diagnostic_ids': [], 'extensions': {}}
    row_locations = {id(row): location(data, [title, index]) for title in TABLES for index, row in enumerate(data[title])}
    def where(row, field):
        return f'{row_locations[id(row)]}.{field}'
    def indexed(title, field):
        out = {}
        for row in data[title]:
            value = code(row.get(field), where(row, field))
            require(value not in out, f'{where(row, field)}: duplicate code {value}.')
            out[value] = row
        return out
    pools = indexed('Pools', 'pool_code')
    questions = indexed('Questions', 'question_code')
    draws = indexed('Draws', 'draw_code')
    require(bool(questions), 'Questions: no authored questions found.')
    require(bool(pools), 'Pools: declare at least one bank/pool.')
    grouped = defaultdict(list)
    for row in data['Responses']:
        qcode = row.get('question_code')
        require(qcode in questions, f'{where(row, "question_code")}: unknown question_code {qcode!r}.')
        grouped[qcode].append(row)
    members = defaultdict(list)
    catalog = {r['kind']: r for r in question_type_catalog()}
    diagnostics = []
    def relation(kind, start, end, ordinal):
        model['relationships'].append({'relationship_key': key('relationship', hashlib.sha256(f'{kind}|{start}|{end}'.encode()).hexdigest()[:24]), 'kind': kind,
            'source_kind': 'explicit draft intake', 'from_entity_key': start, 'to_entity_key': end, 'status': 'resolved', 'ordinal': ordinal,
            'attributes': {}, 'candidates': [], **base()})
    for qcode, row in questions.items():
        try:
            kind = canonical_question_type(row.get('question_type'))
        except ValueError as exc:
            raise IntakeError(f'{where(row, "question_type")}: {exc}') from None
        fmt = row.get('content_format')
        prompt = rich(row.get('question_text'), fmt, where(row, 'question_text'))
        points = number(row.get('points'), where(row, 'points'))
        pool = row.get('pool_code')
        require(pool in pools, f'{where(row, "pool_code")}: unknown pool_code {pool!r}.')
        members[pool].append(qcode)
        payload = {name: [] for name in ('options', 'accepted_responses', 'blanks', 'match_pairs', 'correct_order', 'raw_response_models')}
        payload['extensions'] = {}
        responses, seen, blanks = grouped[qcode], set(), defaultdict(list)
        allowed_roles = {'multiple_choice': {'option'}, 'true_false': {'option'}, 'multi_select': {'option'}, 'long_answer': set(), 'short_answer': {'accepted'},
                         'multi_short_answer': {'accepted'}, 'fill_in_blanks': {'accepted'}, 'matching': {'left', 'right'}, 'ordering': {'option'}}[kind]
        for response in responses:
            rid = code(response.get('response_key'), where(response, 'response_key'))
            require(rid not in seen, f'{qcode}: duplicate response_key {rid}.')
            seen.add(rid)
            role = response.get('role')
            require(role in allowed_roles, f'{where(response, "role")}: incompatible with {qcode} type {kind}.')
            content = rich(response.get('text'), response.get('content_format'), where(response, 'text'))
            used = {'question_code', 'response_key', 'role', 'text', 'content_format'}
            if role == 'accepted':
                used.add('case_sensitive')
                answer = {'value': content['content'], 'case_sensitive': boolean(response.get('case_sensitive'), where(response, 'case_sensitive')), 'weight': 100,
                          'extensions': {'coursecraft.response_key': rid, 'coursecraft.formatted_content': content}}
                if kind == 'short_answer':
                    payload['accepted_responses'].append(answer)
                else:
                    used.add('group_key')
                    blanks[code(response.get('group_key'), where(response, 'group_key'))].append(answer)
            else:
                correct = None
                if kind in ('multiple_choice', 'true_false', 'multi_select'):
                    used.add('correct')
                    correct = boolean(response.get('correct'), where(response, 'correct'))
                payload['options'].append({'option_key': rid, 'content': content, 'correct': correct, 'weight': (100 if correct else 0) if correct is not None else None,
                                           'source_evidence_keys': [], 'extensions': {'coursecraft.match_side': role} if kind == 'matching' else {}})
                if kind == 'ordering':
                    used.add('position')
                    payload['correct_order'].append({'option_key': rid, 'position': number(response.get('position'), where(response, 'position'), True), 'source_evidence_keys': [], 'extensions': {}})
                if role == 'left':
                    used.add('match_key')
                    payload['match_pairs'].append({'left_key': rid, 'right_key': code(response.get('match_key'), where(response, 'match_key')), 'weight': 100, 'source_evidence_keys': [], 'extensions': {}})
            for name, value in response.items():
                require(name in used or not populated(value), f'{where(response, name)}: unused field contains data; review the type change.')
        ncorrect = sum(o['correct'] is True for o in payload['options'])
        if kind in ('multiple_choice', 'multi_select'):
            require(len(responses) >= 2, f'{qcode}: at least two options required.')
            require(ncorrect == 1 if kind == 'multiple_choice' else ncorrect >= 1, f'{qcode}: invalid correct-option count.')
        if kind == 'true_false':
            require([(o['option_key'], o['content']['content'], o['content']['format']) for o in payload['options']] == [('T', 'True', 'plain_text'), ('F', 'False', 'plain_text')] and ncorrect == 1,
                    f'{qcode}: True/False requires T=True, F=False in that order and exactly one correct option.')
        if kind == 'short_answer':
            require(bool(payload['accepted_responses']), f'{qcode}: accepted answer required.')
        if kind in ('multi_short_answer', 'fill_in_blanks'):
            require(len(blanks) >= (2 if kind == 'multi_short_answer' else 1), f'{qcode}: missing answer groups.')
            payload['blanks'] = [{'blank_key': b, 'accepted_responses': answers, 'source_evidence_keys': [], 'extensions': {}} for b, answers in blanks.items()]
            if kind == 'fill_in_blanks':
                placeholders = re.findall(r'\[\[blank:([A-Za-z0-9._-]+)\]\]', prompt['content'])
                require(len(placeholders) == len(set(placeholders)) and set(placeholders) == set(blanks), f'{qcode}: [[blank:code]] placeholders must match answer groups exactly once.')
        if kind == 'matching':
            left = [r for r in responses if r['role'] == 'left']
            right = {r['response_key'] for r in responses if r['role'] == 'right'}
            targets = [p['right_key'] for p in payload['match_pairs']]
            require(len(left) >= 2 and len(targets) == len(set(targets)) and set(targets) == right, f'{qcode}: v1 matching requires at least two complete one-to-one pairs.')
        if kind == 'ordering':
            require(len(responses) >= 2 and sorted(o['position'] for o in payload['correct_order']) == list(range(1, len(responses)+1)), f'{qcode}: positions must be unique and contiguous from 1.')
        require(kind == 'long_answer' or not populated(row.get('answer_key')), f'{where(row, "answer_key")}: only for Written Response; mark keyed responses in Responses.')
        if populated(row.get('answer_key')):
            payload['manual_answer_key'] = rich(row['answer_key'], fmt, where(row, 'answer_key'))
        feedback = []
        if populated(row.get('feedback')):
            feedback = [{'channel': 'general', 'source_kind': 'draft feedback', 'content': rich(row['feedback'], fmt, where(row, 'feedback')), 'source_evidence_keys': [], 'extensions': {}}]
        model['questions'].append({'entity_key': key('question', qcode), 'identity': identity(qcode), 'kind': kind, 'source_kind': row['question_type'], 'title': qcode, 'prompt': prompt,
            'type_payload': payload, 'scoring': {'state': 'known', 'mode': 'manual' if kind == 'long_answer' else 'all_or_nothing' if kind == 'multi_select' else 'exact', 'maximum_points': points, 'rules': [], 'extensions': {}},
            'feedback': feedback, 'build_support': {'level': 'extraction_only', 'receipt_refs': [], 'notes': ['Fresh draft; no question-instance import evidence.'], 'extensions': {}}, **base()})
        relation('member_of', key('question', qcode), key('structure', 'pool.'+pool), len(members[pool]))
        diagnostics.append({'question_code': qcode, 'kind': kind, 'type_capability': catalog[kind]['status'], 'instance_build_support': 'extraction_only'})
    if quiz_code:
        model['quizzes'].append({'entity_key': key('quiz', quiz_code), 'identity': identity(quiz_code), 'title': cfg['quiz_title'], **base()})
    require(quiz_code or not draws, 'Library intake cannot declare quiz draws without quiz identity.')
    for index, (pcode, row) in enumerate(pools.items(), 1):
        require(bool(members[pcode]), f'{pcode}: empty pools are not accepted.')
        require(isinstance(row.get('title'), str) and bool(row['title'].strip()), f'{pcode}: title required.')
        model['structures'].append({'entity_key': key('structure', 'pool.'+pcode), 'identity': identity('pool.'+pcode), 'kind': 'pool', 'source_kind': 'draft pool', 'title': row['title'], 'ordinal': index,
            'selection': {'mode': 'all', 'requested_count': None, 'available_count': len(members[pcode]), 'extensions': {}}, **base()})
    selected_pools = set()
    for index, (dcode, row) in enumerate(draws.items(), 1):
        pool = row.get('pool_code')
        require(pool in pools and pool not in selected_pools, f'{dcode}: pool must exist and may be drawn once in v1.')
        selected_pools.add(pool)
        count = number(row.get('draw_count'), where(row, 'draw_count'), True)
        require(count <= len(members[pool]), f'{dcode}: draw exceeds pool membership.')
        points = number(row.get('points_per_question'), where(row, 'points_per_question'))
        require(isinstance(row.get('title'), str) and bool(row['title'].strip()), f'{dcode}: title required.')
        model['structures'].append({'entity_key': key('structure', 'draw.'+dcode), 'identity': identity('draw.'+dcode), 'kind': 'draw', 'source_kind': 'draft draw', 'title': row['title'], 'ordinal': index,
            'selection': {'mode': 'random', 'requested_count': count, 'available_count': len(members[pool]), 'extensions': {'points_per_question': points}}, **base()})
        relation('contains', key('quiz', quiz_code), key('structure', 'draw.'+dcode), index)
        relation('draws_from', key('structure', 'draw.'+dcode), key('structure', 'pool.'+pool), 1)
    errors = [i.render() for i in validate_contract(model, mode='transform') if i.severity == 'error']
    require(not errors, 'Canonical contract validation failed: '+ '; '.join(errors))
    report = {'schema': 'coursecraft.quiz_draft_intake_report/1', 'valid_intake': True, 'ready_to_build': False, 'source': source, 'question_count': len(questions), 'pool_count': len(pools), 'draw_count': len(draws),
                   'question_types': diagnostics, 'notes': ['Intake does not approve content, convert equations, or create import evidence.', 'Use existing readiness and reviewed promotion/candidate workflows before package generation.']}
    validate_shape(report, REPORT_SCHEMA_PATH)
    return model, report


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('input', type=Path, nargs='?')
    p.add_argument('--output-dir', type=Path)
    p.add_argument('--template-spec', type=Path, help='Write current headers/catalog for the generic template builder, then exit.')
    args = p.parse_args()
    try:
        if args.template_spec:
            require(args.input is None and args.output_dir is None, '--template-spec cannot be combined with intake arguments.')
            with args.template_spec.open('x', encoding='utf-8') as target:
                target.write(json.dumps(template_spec(), ensure_ascii=False, indent=2)+'\n')
            return
        require(args.input is not None and args.output_dir is not None, 'input and --output-dir are required.')
        require(args.input.suffix.lower() in ('.xlsx', '.json'), 'Input must be a .xlsx or .json file.')
        snapshot = args.input.read_bytes()
        data = read_workbook(snapshot) if args.input.suffix.lower() == '.xlsx' else read_json(snapshot)
        model, report = ingest(data, source_name=args.input.name, source_digest=hashlib.sha256(snapshot).hexdigest(), input_format=args.input.suffix.lower()[1:])
        require(not args.output_dir.exists(), 'Output directory already exists; use a new run directory.')
        args.output_dir.mkdir(parents=True)
        for name, value in [('quiz.model.json', model), ('intake-report.json', report)]:
            (args.output_dir/name).write_text(json.dumps(value, ensure_ascii=False, indent=2)+'\n')
        from quiz_authoring_readiness import analyze_authoring_readiness, render_readiness_markdown
        readiness = analyze_authoring_readiness(args.output_dir/'quiz.model.json')
        (args.output_dir/'readiness.json').write_text(json.dumps(readiness, indent=2)+'\n')
        (args.output_dir/'readiness.md').write_text(render_readiness_markdown(readiness))
        print(json.dumps(report, ensure_ascii=False))
    except (ValueError, KeyError, OSError, BadZipFile, ET.ParseError) as exc:
        p.exit(2, f'Intake refused: {exc}\n')


if __name__ == '__main__':
    main()
