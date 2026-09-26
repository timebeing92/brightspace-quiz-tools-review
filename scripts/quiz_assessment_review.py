#!/usr/bin/env python3
"""Prepare and re-import portable, assessment-grouped quiz review workbooks.

This is a presentation adapter over the existing native review contract. Exact
source content stays in the copied model and native workbook. No Google service
or network access is used. Run with --help for independent commands.
"""
from __future__ import annotations

import argparse
from collections import OrderedDict
from copy import copy
import hashlib
import json
import math
from pathlib import Path
import re
import shutil
import tempfile
from urllib.parse import unquote, urlsplit
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from quiz_build_support import resolve_projection_relationships
from quiz_review_projection import build_quiz_review_projection
from quiz_review_text import readable_html
from quiz_contracts import validate_contract
from quiz_review_workbook_reingest import (
    EDITABLE_HEADERS, SETTINGS_EDITABLE_HEADERS, ReingestError,
    _canonical_digest, _review_metadata, _stable_id, materialize_workbook_decisions, sha256_file,
)

SCHEMA = "coursecraft.assessment_review_packet/2"
LEGACY_SCHEMA = "coursecraft.assessment_review_packet/1"
MARKER = "_Assessment Review"
WORKING = "reviewer_working.xlsx"
BASELINE = "reviewer_baseline_DO_NOT_EDIT.xlsx"
LAYOUT = "assessment-review.json"
EXTRA_HEADERS = ["response_options_markup", "answer_key_markup", "question_text_markup",
                 "source_points", "source_feedback", "content_preservation", "source_content_ref"]
VISIBLE = {"question_number", "Source pool / folder", "question_type", "question_text",
           "revised_question_text", "image_link", "response_options", "revised_response_options",
           "answer_key", "revised_answer_key", "revision_reason", "proposed_by", "approval_status",
           "reviewer_note", "source_points", "source_feedback", "content_preservation"}
IMAGE_REPLACEMENT_SHEET = "Image Replacements"
NEW_QUESTION_SHEET = "New Questions"
NEW_RESPONSE_SHEET = "New Responses"
IMAGE_HEADERS = ["replacement_id", "quiz_title", "question_title", "original_image",
                 "original_asset", "replacement_file", "revision_reason", "approval_status",
                 "proposed_by", "proposed_at", "approved_by", "approved_at"]
NEW_QUESTION_HEADERS = ["question_code", "question_type", "question_text", "points",
                        "target_quiz_entity_key", "target_pool_entity_key", "answer_key",
                        "feedback", "content_format", "approval_status", "revision_reason",
                        "proposed_by", "proposed_at", "approved_by", "approved_at"]
NEW_RESPONSE_HEADERS = ["question_code", "response_key", "role", "text", "content_format",
                        "correct", "case_sensitive", "group_key", "match_key", "position"]
IMAGE_SIGNATURES = {".png": (b"\x89PNG\r\n\x1a\n", "image/png"),
                    ".jpg": (b"\xff\xd8\xff", "image/jpeg"),
                    ".jpeg": (b"\xff\xd8\xff", "image/jpeg"),
                    ".gif": (b"GIF8", "image/gif")}
TEAL, PALE, YELLOW = "175C72", "E9F4F7", "FFF2CC"
RICH_MARKUP = re.compile(r"<(?:[\w.-]+:)?(?:math|img|matimage|sup|sub)\b", re.I)


def _json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _inside(root: Path, relative: str) -> Path:
    path = root / relative
    if Path(relative).is_absolute() or ".." in Path(relative).parts:
        raise ReingestError("Unsafe packet path: " + relative)
    if path.is_symlink() or not path.resolve().is_relative_to(root.resolve()):
        raise ReingestError("Packet path escapes its folder: " + relative)
    return path


def _ref(root: Path, path: Path) -> dict:
    return {"path": path.relative_to(root).as_posix(), "sha256": sha256_file(path)}


def _replacement_image(root: Path, value: str) -> tuple[Path, str]:
    rel = Path(value)
    if rel.is_absolute() or ".." in rel.parts or not rel.parts or rel.parts[0] != "Replacement Images":
        raise ReingestError("Replacement files must be placed under Replacement Images/ in this packet.")
    path = _inside(root, rel.as_posix())
    if not path.is_file() or path.stat().st_size == 0 or path.stat().st_size > 20 * 1024 * 1024:
        raise ReingestError("Replacement image must be a nonempty file no larger than 20 MB.")
    signature = IMAGE_SIGNATURES.get(path.suffix.lower())
    if signature is None or not path.read_bytes().startswith(signature[0]):
        raise ReingestError("Replacement image must be a matching PNG, JPEG, or GIF file.")
    return path, signature[1]


def _sheet_rows(workbook, title: str, expected_headers: list[str]) -> list[dict]:
    if title not in workbook.sheetnames:
        raise ReingestError(f"Workbook is missing required sheet {title!r}.")
    sheet = workbook[title]
    headers = [str(c.value or "").strip() for c in sheet[1]]
    while headers and not headers[-1]:
        headers.pop()
    if headers != expected_headers:
        raise ReingestError(f"{title} headers do not match the review packet contract.")
    rows = []
    for row_number in range(2, sheet.max_row + 1):
        cells = [sheet.cell(row_number, col) for col in range(1, len(headers) + 1)]
        if all(cell.value in (None, "") for cell in cells):
            continue
        if any(cell.data_type == "f" for cell in cells):
            raise ReingestError(f"Formula cells are not accepted ({title} row {row_number}).")
        rows.append({**dict(zip(headers, (cell.value for cell in cells))), "__row_number": row_number})
    return rows


def _annotation_pair(target: str, field: str, value: object, decision: str,
                     row: dict, row_number: int, metadata_policy: str,
                     source_sheet: str) -> list[dict]:
    reason, proposer, proposed_at, approver, approved_at = _review_metadata(
        row, row_number, decision, metadata_policy)
    identity = {"target_entity_key": target, "field_path": field, "value": value,
                "actor": proposer, "timestamp": proposed_at, "reason": reason}
    proposal_id = _stable_id("ann.proposal", identity)
    extensions = {"coursecraft.binder.reason": reason,
                  "coursecraft.binder.workbook_sheet": source_sheet,
                  "coursecraft.binder.workbook_row": row_number}
    proposal = {"annotation_id": proposal_id, "target_entity_key": target,
                "field_path": field, "kind": "proposed_revision", "value": value,
                "actor": proposer, "timestamp": proposed_at, "source_evidence_keys": [],
                "status": decision, "extensions": extensions}
    if decision != "accepted":
        return [proposal]
    approval_identity = {"proposal_id": proposal_id, "actor": approver, "timestamp": approved_at}
    approval = {"annotation_id": _stable_id("ann.approval", approval_identity),
                "target_entity_key": target, "field_path": field,
                "kind": "approved_change", "value": value, "actor": approver,
                "timestamp": approved_at, "source_evidence_keys": [], "status": "accepted",
                "extensions": {**extensions, "coursecraft.binder.proposal_id": proposal_id}}
    return [proposal, approval]


def _nodes(node: dict, name: str):
    if node.get("name") == name:
        yield node
    for child in node.get("children", []):
        yield from _nodes(child, name)


def _presentation(question: dict, native: dict, model_index: int) -> dict:
    """Keep original rich fragments addressable; render only known structure."""
    result = {h: native.get(h) for h in ("question_text", "response_options", "answer_key")}
    result.update({h: None for h in EXTRA_HEADERS})
    result["source_content_ref"] = f"model.json#/questions/{model_index}"
    result["content_preservation"] = "Source retained; review only"
    payload = question.get("type_payload", {})
    facts = [r["payload"] for r in payload.get("raw_response_models", [])
             if r.get("source_kind") == "d2l_qti_response_facts/0"]
    prompt_fragments: list[str] = []
    ambiguous_source = bool(facts) and (len(facts) != 1 or len(facts[0].get("presentation", {}).get("occurrences", [])) != 1)
    if ambiguous_source:
        result["content_preservation"] = "Source retained; multiple source presentations need display review"
        result["question_text_markup"] = "See " + result["source_content_ref"]
    if len(facts) == 1 and len(facts[0].get("presentation", {}).get("occurrences", [])) == 1:
        tree = facts[0]["presentation"]["occurrences"][0]["node"]

        def prompt_material(node):
            for child in node.get("children", []):
                if child.get("name") == "material":
                    yield from (n.get("raw_text") or "" for n in _nodes(child, "mattext"))
                elif child.get("name") == "flow":
                    yield from prompt_material(child)

        prompt_fragments = list(prompt_material(tree))
    elif question.get("prompt", {}).get("format") in {"html", "xhtml", "mathml"}:
        prompt_fragments = [question["prompt"]["content"]]

    options = payload.get("options", [])
    originals = [{"option_key": o.get("option_key"), "source_response_ident":
                  o.get("extensions", {}).get("coursecraft.source_response_ident"),
                  "content": o.get("content"), "correct": o.get("correct")}
                 for o in options]
    raw_fields = {"question_text_markup": prompt_fragments,
                  "response_options_markup": originals,
                  "answer_key_markup": [o for o in originals if o["correct"] is True]}
    for name, value in raw_fields.items():
        if value:
            encoded = json.dumps(value, ensure_ascii=False)
            result[name] = encoded if len(encoded) < 30000 else "See " + result["source_content_ref"]
    try:
        if prompt_fragments:
            result["question_text"] = readable_html(" ".join(prompt_fragments)) or native.get("question_text")
        if options and all(o.get("content", {}).get("format") in {"html", "xhtml", "plain_text", "mathml"}
                           for o in options):
            labels = []
            for option in options:
                content = option.get("content", {})
                raw = content.get("content", "")
                label = raw if content.get("format") == "plain_text" else readable_html(raw)
                if not label and re.search(r"<(?:img|matimage)\b", raw, re.I):
                    label = "[Image option — open the question images]"
                labels.append(label or "[No visible text; inspect source]")
            # Always expose option keys; duplicate and image-only choices stay distinct.
            result["response_options"] = "\n".join(f"{o['option_key']}. {label}" for o, label in zip(options, labels))
            correct = [f"{o['option_key']}. {label}" for o, label in zip(options, labels) if o.get("correct") is True]
            if correct:
                result["answer_key"] = "\n".join(correct)
        feedback = []
        if len(facts) == 1:
            for slot in facts[0].get("feedback_slots", []):
                fragments = [n.get("raw_text") or "" for n in _nodes(slot.get("node", {}), "mattext")]
                text = readable_html(" ".join(fragments))
                if text:
                    feedback.append(str(slot.get("source_ident") or "Feedback") + ": " + text)
        if not feedback:
            for item in question.get("feedback", []):
                content = item.get("content", {})
                if isinstance(content, dict) and content.get("content"):
                    feedback.append(readable_html(content["content"]))
        result["source_feedback"] = "\n".join(feedback) or None
    except (ValueError, ET.ParseError) as exc:
        result["content_preservation"] = "Source retained; display needs review: " + str(exc)
    result["rich_revision_hold"] = any(RICH_MARKUP.search(s) for s in prompt_fragments)
    result["rich_revision_hold"] |= any(o.get("content", {}).get("format") != "plain_text" and
        RICH_MARKUP.search(o.get("content", {}).get("content", "")) for o in options)
    result["rich_revision_hold"] |= native.get("image_count") not in (None, "", 0, "0")
    if ambiguous_source:
        result["rich_revision_hold"] = True
    return result


def _copy_cell(src, dst) -> None:
    dst.value = src.value
    if isinstance(src.value, str):
        dst.data_type = "s"  # source text beginning '=' stays literal
    dst._style = copy(src._style)
    if src.comment:
        dst.comment = copy(src.comment)
    if src.hyperlink:
        dst._hyperlink = copy(src.hyperlink)
        dst.hyperlink.ref = dst.coordinate


def _portable_links(book, source_dir: Path, output_dir: Path) -> list[dict]:
    assets: dict[str, dict] = {}
    for sheet in book:
        for row in sheet:
            for cell in row:
                if not cell.hyperlink or not cell.hyperlink.target:
                    continue
                old = cell.hyperlink.target
                parsed = urlsplit(old)
                if parsed.scheme in {"http", "https", "mailto"} or old.startswith("#"):
                    continue
                source = (source_dir / unquote(old)).resolve()
                if parsed.scheme or not source.is_relative_to(source_dir.resolve()) or not source.is_file():
                    raise ValueError("Image link is missing or outside the extraction folder; use self_contained extraction: " + old)
                digest = sha256_file(source)
                name = re.sub(r"[^A-Za-z0-9._-]", "_", source.name)[-100:]
                relative = "Images/" + digest + "-" + name
                target = output_dir / relative
                if relative not in assets:
                    target.parent.mkdir(parents=True, exist_ok=True)
                    shutil.copyfile(source, target)
                    assets[relative] = _ref(output_dir, target)
                cell.hyperlink = relative
    return list(assets.values())


def _sheet_title(title: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", " ", title).strip().strip("'") or "Assessment"
    candidate = base[:31]
    number = 2
    while candidate.casefold() in used:
        suffix = f" ({number})"
        candidate = base[:31-len(suffix)] + suffix
        number += 1
    used.add(candidate.casefold())
    return candidate


def prepare_review_packet(source_workbook: Path, model_path: Path, output_dir: Path) -> dict:
    """Create a new packet. Existing files and reviewer work are never replaced."""
    source_workbook, model_path, output_dir = map(Path, (source_workbook, model_path, output_dir))
    if output_dir.exists() and any(output_dir.iterdir()):
        raise ValueError("Review output must be a new or empty folder; preserve existing reviewer edits")
    model = json.loads(model_path.read_text(encoding="utf-8"))
    projection = build_quiz_review_projection(model, include_content=True)
    # Confirm the baseline really belongs to the supplied model before adapting it.
    if projection["question_occurrences"]:
        materialize_workbook_decisions(model_path, source_workbook, source_workbook, metadata_policy="optional")
    book = load_workbook(source_workbook)
    if MARKER in book.sheetnames:
        raise ValueError("This workbook is already grouped")
    output_dir.mkdir(parents=True, exist_ok=True)
    assets = _portable_links(book, source_workbook.parent, output_dir)
    shutil.copyfile(model_path, output_dir / "model.json")
    # Native baseline shares the packet's image folder and remains a separate source.
    native = output_dir / "native_source_DO_NOT_EDIT.xlsx"
    book.save(native)
    if not projection["question_occurrences"]:
        # A question-library-only export has no assessment placement to revise.
        # Preserve successful extraction and portable links without inventing one.
        guide = book.create_sheet("START HERE", 0)
        guide.append(["Question inventory", f"{len(model['questions'])} source questions; no quiz occurrences in this export."])
        guide.append(["Review", "Use All Questions and Question Images. Keep this folder together for local image links."])
        guide.append(["Author", "Use the blank drafting template for new questions/assessments. This inventory has no assessment edits to import."])
        guide.column_dimensions["A"].width = 24
        guide.column_dimensions["B"].width = 105
        for cells in guide:
            guide.row_dimensions[cells[0].row].height = 45
            for cell in cells:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for sheet in book:
            sheet.protection.sheet = True
            for cells in sheet:
                for cell in cells:
                    cell.protection = Protection(locked=True)
        book.active = 0
        baseline = output_dir / BASELINE
        book.save(baseline)
        shutil.copyfile(baseline, output_dir / WORKING)
        spec = {"schema": SCHEMA, "profile": "inventory_only", "model": _ref(output_dir, output_dir / "model.json"),
                "native_workbook": _ref(output_dir, native), "baseline": _ref(output_dir, baseline),
                "assets": assets, "sheets": [], "question_count": 0, "inventory_question_count": len(model["questions"])}
        _json(output_dir / LAYOUT, spec)
        (output_dir / "README.md").write_text("# Question inventory\n\nOpen reviewer_working.xlsx. "
            "This export has no quiz occurrences, so there are no assessment edits to import. "
            "Keep Images/ beside the workbook. Use the separate blank drafting template for new assessments.\n", encoding="utf-8")
        return spec
    src = book["Quiz Questions"]
    original_headers = [c.value for c in src[1]]
    required = {"occurrence_key", "question_entity_key", "quiz", "question_number"}
    if not required.issubset(original_headers):
        raise ValueError("Assessment review needs the current occurrence-aware native workbook")
    columns = {h: i + 1 for i, h in enumerate(original_headers)}
    headers = original_headers[:2] + ["Source pool / folder"] + original_headers[2:] + EXTRA_HEADERS
    dest_columns = {h: i + 1 for i, h in enumerate(headers)}
    occurrences = {o["occurrence_key"]: o for o in projection["question_occurrences"]}
    questions = {q["entity_key"]: (i, q) for i, q in enumerate(model["questions"])}
    assessments: OrderedDict[str, list] = OrderedDict()
    seen = set()
    for row in range(2, src.max_row + 1):
        values = {h: src.cell(row, c).value for h, c in columns.items()}
        key = values["occurrence_key"]
        if key in seen or key not in occurrences:
            raise ValueError("Duplicate or unknown source occurrence: " + str(key))
        seen.add(key)
        occurrence = occurrences[key]
        assessments.setdefault(occurrence["quiz_key"], []).append((row, values, occurrence))
    if seen != set(occurrences):
        raise ValueError("Native workbook does not cover every source occurrence")
    quiz_titles = {q["entity_key"]: q.get("title") or q["entity_key"] for q in model.get("quizzes", [])}
    guide = book.create_sheet("START HERE", 0)
    source_label = ", ".join(model["source"].get("references", [])) or model["source"]["source_key"]
    guide_rows = [
        ("Assessment review", "Edit the yellow fields on each assessment tab."),
        ("Source", source_label),
        ("Questions", f"{len(seen)} occurrences in {len(assessments)} assessments. Source sets/pools appear within each tab."),
        ("Images", "Open image_link or Question Images. Keep Images/ beside this workbook when sharing the complete folder."),
        ("Draft and accept", "Leave approval_status blank or open while drafting. Enter accepted only for an accepted change. Reviewer names, reasons and dates are optional."),
        ("Source content", "Original text and keys are protected. Readable text is paired with original markup columns and model.json."),
        ("Import edits", "Keep the complete packet, including hidden sheets and companion files. Compose collects assessment-tab edits automatically."),
        ("Quiz settings", "Quiz Settings retains its own editable proposed-value fields. Question edits belong on assessment tabs."),
        ("Replace an image", "Place a PNG, JPEG or GIF under Replacement Images/. In Image Replacements, enter its path for the matching question and source image, add a reason, and mark it accepted."),
        ("New questions", "Add question rows to New Questions and answer rows to New Responses. Use exact quiz and pool keys from Target Pools; only explicitly accepted questions are added."),
        ("Question types", "Use the type names and response fields from the blank quiz drafting template. The draft importer checks the content before accepted questions are promoted."),
        ("Math and diagrams", "Accepted revisions to questions containing equations or diagrams require a supported rich-content workflow. Drafts and notes can be collected for review."),
    ]
    for row in guide_rows:
        guide.append(row)
    guide.column_dimensions["A"].width = 24
    guide.column_dimensions["B"].width = 108
    for row in guide:
        guide.row_dimensions[row[0].row].height = 48
        for cell in row:
            cell.font = Font(name="Calibri", size=11, bold=cell.column == 1, color=TEAL if cell.column == 1 else "222222")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    guide.freeze_panes = "B2"
    spec = {"schema": SCHEMA, "model": _ref(output_dir, output_dir / "model.json"),
            "native_workbook": _ref(output_dir, native), "assets": assets, "sheets": [], "question_count": len(seen)}
    used = {name.casefold() for name in book.sheetnames} | {MARKER.casefold()}
    for index, (quiz_key, entries) in enumerate(sorted(assessments.items(), key=lambda p: -len(p[1])), 1):
        title = entries[0][2]["quiz_title"] or "Untitled assessment"
        dst = book.create_sheet(_sheet_title(title, used), index)
        dst.sheet_properties.tabColor = TEAL
        dst.sheet_properties.outlinePr.summaryBelow = False
        dst.sheet_view.showGridLines = False
        dst.freeze_panes = "D3"
        dst.merge_cells("B1:C1")
        dst.cell(1, 2, title)
        prompt_col = dest_columns["question_text"]
        dst.merge_cells(start_row=1, start_column=prompt_col, end_row=1, end_column=prompt_col + 2)
        dst.cell(1, prompt_col, "Source: " + source_label)
        dst.row_dimensions[1].height = 64
        dst.row_dimensions[2].height = 48
        for col, header in enumerate(headers, 1):
            dst.cell(2, col, header)
            letter = get_column_letter(col)
            width = 24
            if header in {"question_text", "revised_question_text", "response_options", "revised_response_options", "source_feedback", "reviewer_note"}:
                width = 56
            elif header == "Source pool / folder":
                width = 52
            elif header == "question_number":
                width = 10
            dst.column_dimensions[letter].width = width
            dst.column_dimensions[letter].hidden = header not in VISIBLE
        groups: OrderedDict[tuple, list] = OrderedDict()
        for entry in entries:
            o = entry[2]
            pools = tuple(p.get("pool_path") or p.get("pool_title") or "" for p in o.get("pool_memberships", []))
            group = (o.get("section_key"), pools)
            groups.setdefault(group, []).append(entry)
        tab = {"title": dst.title, "quiz_key": quiz_key, "rows": [], "group_count": len(groups)}
        row = 3
        for (_section_key, pools), members in groups.items():
            occurrence = members[0][2]
            local_label = "Quiz pool > " if occurrence.get("selection_context", {}).get("mode") == "random" else "Quiz section > "
            label = " | ".join(p for p in pools if p) or (
                local_label + str(occurrence.get("section_path") or occurrence.get("section_title"))
                if occurrence.get("section_key") else "Quiz root")
            dst.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
            dst.cell(row, 2, label)
            selection = occurrence.get("selection_context", {})
            count_text = f"{len(members)} questions"
            if selection.get("mode") == "random":
                count_text = f"{len(members)} candidates; {selection.get('draw_count')} drawn per attempt"
            dst.cell(row, prompt_col, count_text)
            dst.row_dimensions[row].height = max(44, math.ceil(len(label)/62)*15)
            for cell in dst[row]:
                cell.fill = PatternFill("solid", fgColor=PALE)
                cell.font = Font(name="Calibri", size=11, bold=True, color=TEAL)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            row += 1
            start = row
            for source_row, values, occurrence in members:
                qindex, question = questions[values["question_entity_key"]]
                display = _presentation(question, values, qindex)
                display["source_points"] = occurrence.get("points")
                diagnostics = occurrence.get("diagnostic_ids", [])
                if diagnostics:
                    display["content_preservation"] += f"; {len(diagnostics)} source/projection diagnostics"
                for header in original_headers:
                    _copy_cell(src.cell(source_row, columns[header]), dst.cell(row, dest_columns[header]))
                for header in ("question_text", "response_options", "answer_key", *EXTRA_HEADERS):
                    cell = dst.cell(row, dest_columns[header], display.get(header))
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
                dst.cell(row, dest_columns["Source pool / folder"], label)
                lines = 1
                for col, header in enumerate(headers, 1):
                    cell = dst.cell(row, col)
                    cell.font = Font(name="Calibri", size=11, color="0563C1" if cell.hyperlink else "222222")
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    cell.protection = Protection(locked=header not in EDITABLE_HEADERS)
                    if header in EDITABLE_HEADERS:
                        cell.fill = PatternFill("solid", fgColor=YELLOW)
                    if header in VISIBLE:
                        width = dst.column_dimensions[get_column_letter(col)].width
                        lines = max(lines, sum(max(1, math.ceil(len(s)/width)) for s in str(cell.value or "").split("\n")))
                dst.row_dimensions[row].height = min(409, max(54, lines*15 + 12))
                tab["rows"].append({"row": row, "native_row": source_row, "occurrence_key": values["occurrence_key"],
                                    "rich_revision_hold": bool(display["rich_revision_hold"])})
                row += 1
            dst.row_dimensions.group(start, row-1, outline_level=1, hidden=False)
        for r in (1, 2):
            for cell in dst[r]:
                cell.fill = PatternFill("solid", fgColor=TEAL)
                cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
        # Copy native dropdowns by header and by question row, never group headings.
        for validation in src.data_validations.dataValidation:
            new = copy(validation)
            new.sqref = ""
            for item in tab["rows"]:
                for header, col in columns.items():
                    if f"{get_column_letter(col)}{item['native_row']}" in validation:
                        new.add(dst.cell(item["row"], dest_columns[header]))
            if str(new.sqref):
                dst.add_data_validation(new)
        dst.protection.sheet = True
        dst.protection.selectLockedCells = False
        dst.protection.selectUnlockedCells = False
        dst.protection.formatRows = False
        dst.protection.formatColumns = False
        spec["sheets"].append(tab)

    # Replacement rows are pre-bound to stable occurrence, question and asset
    # identities. Reviewers add a file and decision; source image cells stay locked.
    image_sheet = book.create_sheet(IMAGE_REPLACEMENT_SHEET)
    image_sheet.append(IMAGE_HEADERS)
    image_sheet.freeze_panes = "F2"
    image_sheet.auto_filter.ref = f"A1:L1"
    image_sheet.sheet_view.showGridLines = False
    widths = [44, 34, 48, 38, 52, 48, 45, 18, 24, 28, 24, 28]
    for col, (header, width) in enumerate(zip(IMAGE_HEADERS, widths), 1):
        cell = image_sheet.cell(1, col)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        image_sheet.column_dimensions[get_column_letter(col)].width = width
        image_sheet.column_dimensions[get_column_letter(col)].hidden = header in {"replacement_id", "original_asset"}
    image_rows = []
    asset_map = {a["entity_key"]: a for a in model.get("assets", [])}
    relation_map: dict[tuple[str, str], dict] = {}
    for relation in model.get("relationships", []):
        if relation.get("kind") == "uses_asset" and relation.get("status") == "resolved":
            relation_map[(relation["from_entity_key"], relation["to_entity_key"])] = relation
    for occurrence_key, occurrence in occurrences.items():
        question_key = occurrence.get("referenced_question_key") or occurrence.get("observed_question_entity_key")
        if not question_key:
            continue
        qindex, question = questions[question_key]
        for (owner, asset_key), relation in sorted(relation_map.items()):
            if owner != question_key or asset_key not in asset_map:
                continue
            asset = asset_map[asset_key]
            if not asset.get("package_path"):
                continue
            stable = hashlib.sha256(f"{occurrence_key}\0{asset_key}".encode()).hexdigest()[:24]
            row_number = image_sheet.max_row + 1
            image_sheet.append([stable, quiz_titles.get(occurrence["quiz_key"], occurrence["quiz_key"]),
                                question.get("title") or question_key, asset["package_path"], asset_key,
                                None, None, None, None, None, None, None])
            for cell in image_sheet[row_number]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.protection = Protection(locked=cell.column not in {6, 7, 8, 9, 10, 11, 12})
                if cell.column in {6, 7, 8, 9, 10, 11, 12}:
                    cell.fill = PatternFill("solid", fgColor=YELLOW)
            image_rows.append({"row": row_number, "replacement_id": stable,
                               "occurrence_key": occurrence_key, "quiz_key": occurrence["quiz_key"],
                               "question_entity_key": question_key, "asset_entity_key": asset_key,
                               "package_path": asset["package_path"]})
    image_sheet.row_dimensions[1].height = 36
    image_sheet.protection.sheet = True
    image_sheet.protection.selectLockedCells = False
    image_sheet.protection.selectUnlockedCells = False
    image_sheet.protection.formatRows = False
    image_sheet.protection.formatColumns = False
    status_validation = DataValidation(type="list", formula1='"open,accepted,rejected"', allow_blank=True)
    image_sheet.add_data_validation(status_validation)
    for item in image_rows:
        status_validation.add(image_sheet.cell(item["row"], 8))
    spec["image_replacements"] = {"sheet": IMAGE_REPLACEMENT_SHEET, "headers": IMAGE_HEADERS, "rows": image_rows}

    # New questions use a dedicated row-addition lane that reuses the blank
    # drafting workbook's typed question/response vocabulary. Source rows remain
    # fixed; every new row needs a target quiz, existing pool and explicit status.
    target_pools = {}
    for quiz_key in assessments:
        # Use the same bounded source-evidence joins as readiness/building.
        # Do not manufacture these authoring relationships in the source model.
        relations, _derivations = resolve_projection_relationships(model, quiz_key)
        reachable = {quiz_key}
        while True:
            children = {r["to_entity_key"] for r in relations
                        if r["kind"] == "contains" and r["from_entity_key"] in reachable}
            if children <= reachable:
                break
            reachable |= children
        pools = {r["to_entity_key"] for r in relations
                 if r["kind"] == "draws_from" and r["from_entity_key"] in reachable}
        for pool in pools:
            structure = next((s for s in model["structures"] if s["entity_key"] == pool), None)
            if structure:
                target_pools[(quiz_key, pool)] = structure.get("title") or pool
    targets = book.create_sheet("Target Pools")
    targets.append(["target_quiz_entity_key", "target_pool_entity_key", "source_pool_title"])
    for (quiz_key, pool_key), pool_title in sorted(target_pools.items()):
        targets.append([quiz_key, pool_key, pool_title])
    targets.freeze_panes = "A2"
    targets.column_dimensions["A"].width = 72
    targets.column_dimensions["B"].width = 76
    targets.column_dimensions["C"].width = 52
    for cells in targets:
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor=TEAL)
                cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    targets.protection.sheet = True
    qsheet = book.create_sheet(NEW_QUESTION_SHEET)
    qsheet.append(NEW_QUESTION_HEADERS)
    qsheet.freeze_panes = "A2"
    qsheet.auto_filter.ref = f"A1:{get_column_letter(len(NEW_QUESTION_HEADERS))}1"
    qwidths = [24, 24, 58, 12, 72, 76, 48, 45, 18, 18, 42, 24, 28, 24, 28]
    for col, (header, width) in enumerate(zip(NEW_QUESTION_HEADERS, qwidths), 1):
        qsheet.column_dimensions[get_column_letter(col)].width = width
        cell = qsheet.cell(1, col)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    qsheet.protection.sheet = True
    qsheet.protection.selectLockedCells = False
    qsheet.protection.selectUnlockedCells = False
    qsheet.protection.formatRows = False
    qsheet.protection.formatColumns = False
    qsheet.protection.sheet = False
    rsheet = book.create_sheet(NEW_RESPONSE_SHEET)
    rsheet.append(NEW_RESPONSE_HEADERS)
    rsheet.freeze_panes = "A2"
    for col, header in enumerate(NEW_RESPONSE_HEADERS, 1):
        rsheet.column_dimensions[get_column_letter(col)].width = 30 if header != "text" else 58
        cell = rsheet.cell(1, col)
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    rsheet.protection.sheet = True
    rsheet.protection.selectLockedCells = False
    rsheet.protection.selectUnlockedCells = False
    rsheet.protection.formatRows = False
    rsheet.protection.formatColumns = False
    rsheet.protection.sheet = False
    qstatus = DataValidation(type="list", formula1='"open,accepted,rejected"', allow_blank=True)
    qsheet.add_data_validation(qstatus)
    qstatus.add("J2:J1048576")
    spec["question_additions"] = {"question_sheet": NEW_QUESTION_SHEET,
                                   "question_headers": NEW_QUESTION_HEADERS,
                                   "response_sheet": NEW_RESPONSE_SHEET,
                                   "response_headers": NEW_RESPONSE_HEADERS,
                                   "target_pool_sheet": "Target Pools",
                                   "target_pools": [{"quiz_key": q, "pool_key": p, "title": t}
                                                    for (q, p), t in sorted(target_pools.items())]}
    # The inventory is a readable source view too. Stable source aliases bind it
    # to model entities; ambiguous aliases are left intact instead of guessed.
    if "All Questions" in book:
        inventory = book["All Questions"]
        ih = {c.value: c.column for c in inventory[1]}
        aliases = {}
        for key, (qi, question) in questions.items():
            for alias in question.get("identity", {}).get("source_aliases", []):
                value = alias.get("value")
                aliases.setdefault(value, set()).add(key)
        for name in EXTRA_HEADERS[:3]:
            ih[name] = inventory.max_column + 1
            inventory.cell(1, ih[name], name)
            inventory.column_dimensions[get_column_letter(ih[name])].hidden = True
        for row in range(2, inventory.max_row + 1):
            candidates = set()
            for h in ("source_global_id", "source_question_ident"):
                if h in ih:
                    candidates.update(aliases.get(inventory.cell(row, ih[h]).value, set()))
            if len(candidates) != 1:
                continue
            qi, question = questions[next(iter(candidates))]
            old = {h: inventory.cell(row, col).value for h, col in ih.items()}
            display = _presentation(question, old, qi)
            for h in ("question_text", "response_options", "answer_key", *EXTRA_HEADERS[:3]):
                cell = inventory.cell(row, ih[h], display.get(h))
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        if "question_type" in ih:
            dim = inventory.column_dimensions[get_column_letter(ih["question_type"])]
            dim.hidden, dim.width = False, 24
    # Native evidence tabs are snapshots. Quiz Settings remains an active native surface.
    for sheet in book:
        if sheet.title in {"START HERE", "Quiz Settings", "Question Images", "All Questions", IMAGE_REPLACEMENT_SHEET, NEW_QUESTION_SHEET, NEW_RESPONSE_SHEET, "Target Pools"} or sheet.title in {s["title"] for s in spec["sheets"]}:
            sheet.sheet_state = "visible"
        else:
            sheet.sheet_state = "hidden"
        if sheet.title not in {s["title"] for s in spec["sheets"]} | {"Quiz Settings", NEW_QUESTION_SHEET, NEW_RESPONSE_SHEET}:
            for cells in sheet:
                for cell in cells:
                    cell.protection = Protection(locked=True)
            sheet.protection.sheet = True
    marker = book.create_sheet(MARKER)
    marker.append(["schema", SCHEMA])
    marker.sheet_state = "veryHidden"
    marker.protection.sheet = True
    book.active = 0
    baseline = output_dir / BASELINE
    book.save(baseline)
    shutil.copyfile(baseline, output_dir / WORKING)
    spec["baseline"] = _ref(output_dir, baseline)
    _json(output_dir / LAYOUT, spec)
    (output_dir / "README.md").write_text(
        "# Assessment review\n\nOpen `reviewer_working.xlsx` and start with START HERE. "
        "Edit the yellow fields on assessment tabs and use the addition sheets as needed. Keep this entire folder together; "
        "Images/ holds source diagrams and Replacement Images/ holds new files. Hidden source sheets and companion files are required for import.\n\n"
        "Use Quiz Workshop Compose, or run from the tool repository:\n\n"
        "```sh\npython scripts/quiz_assessment_review.py import --packet /path/to/this-folder "
        "--output /path/to/decisions.json\n```\n\n"
        "A downloaded or renamed edited copy can be supplied with `--edited-workbook /path/to/edited.xlsx`. "
        "Keep the baseline and companion files unchanged. Import collects decisions; use the original full "
        "extraction workspace for readiness and package building.\n", encoding="utf-8")
    return spec


def validate_packet(packet_dir: Path) -> dict:
    packet_dir = Path(packet_dir)
    spec = json.loads((packet_dir / LAYOUT).read_text(encoding="utf-8"))
    if spec.get("schema") not in {SCHEMA, LEGACY_SCHEMA}:
        raise ReingestError("Unknown assessment review packet version")
    # A copied packet must have ordinary local companions, not symlinked folders.
    if packet_dir.is_symlink() or any(p.is_symlink() for p in packet_dir.rglob("*")):
        raise ReingestError("Review packets cannot contain symlinks")
    for ref in [spec["model"], spec["native_workbook"], spec["baseline"], *spec["assets"]]:
        path = _inside(packet_dir, ref["path"])
        if not path.is_file() or sha256_file(path) != ref["sha256"]:
            raise ReingestError("Review companion missing or changed: " + ref["path"])
    if spec.get("image_replacements"):
        if spec["image_replacements"].get("headers") != IMAGE_HEADERS:
            raise ReingestError("Image replacement sheet schema is not supported.")
    if spec.get("question_additions"):
        draft = spec["question_additions"]
        if draft.get("question_headers") != NEW_QUESTION_HEADERS or draft.get("response_headers") != NEW_RESPONSE_HEADERS:
            raise ReingestError("New-question sheet schema is not supported.")
    return spec


def _image_replacement_annotations(spec: dict, edited, packet: Path, metadata_policy: str) -> list[dict]:
    layout = spec.get("image_replacements")
    if not layout:
        return []
    sheet = edited[layout["sheet"]]
    headers = {cell.value: cell.column for cell in sheet[1]}
    by_id = {row["replacement_id"]: row for row in layout["rows"]}
    if len(by_id) != len(layout["rows"]):
        raise ReingestError("Image replacement layout repeats a stable replacement ID.")
    annotations = []
    seen = set()
    for row_number in range(2, sheet.max_row + 1):
        values = {h: sheet.cell(row_number, col).value for h, col in headers.items()}
        if all(value in (None, "") for value in values.values()):
            continue
        stable_id = values.get("replacement_id")
        spec_row = by_id.get(stable_id)
        if spec_row is None or stable_id in seen or row_number != spec_row["row"]:
            raise ReingestError("Image replacement row identity changed; use the pre-bound source row.")
        seen.add(stable_id)
        status = str(values.get("approval_status") or "open").strip().lower()
        if status not in {"open", "accepted", "rejected"}:
            raise ReingestError(f"Image Replacements row {row_number} approval_status must be open, accepted, or rejected.")
        replacement_file = str(values.get("replacement_file") or "").strip()
        if not replacement_file:
            if status != "open" or values.get("revision_reason"):
                raise ReingestError(f"Image Replacements row {row_number} needs a replacement file for its decision.")
            continue
        path, media_type = _replacement_image(packet, replacement_file)
        value = {"quiz_entity_key": spec_row["quiz_key"],
                 "occurrence_key": spec_row["occurrence_key"],
                 "asset_entity_key": spec_row["asset_entity_key"],
                 "original_package_path": spec_row["package_path"],
                 "replacement_file": Path(replacement_file).as_posix(),
                 "sha256": sha256_file(path), "media_type": media_type}
        metadata = {"revision_reason": values.get("revision_reason"),
                    "approval_status": status, "proposed_by": values.get("proposed_by"),
                    "proposed_at": values.get("proposed_at"), "approved_by": values.get("approved_by"),
                    "approved_at": values.get("approved_at")}
        field_path = "/assets/replacements/" + spec_row["asset_entity_key"]
        annotations.extend(_annotation_pair(spec_row["question_entity_key"], field_path,
                                            value, status, metadata, row_number,
                                            metadata_policy, layout["sheet"]))
    if seen != set(by_id):
        raise ReingestError("Image replacement rows were removed or moved.")
    return annotations


def _new_question_annotations(spec: dict, edited, model: dict, metadata_policy: str) -> tuple[list[dict], list[dict]]:
    from quiz_draft_intake import ingest

    cfg = spec.get("question_additions")
    if not cfg:
        return [], []
    question_rows = _sheet_rows(edited, cfg["question_sheet"], NEW_QUESTION_HEADERS)
    response_rows = _sheet_rows(edited, cfg["response_sheet"], NEW_RESPONSE_HEADERS)
    pool_rows = cfg.get("target_pools", [])
    targets = {(row["quiz_key"], row["pool_key"]): row["title"] for row in pool_rows}
    if len(targets) != len(pool_rows):
        raise ReingestError("Target Pools repeats a quiz/pool identity.")
    questions_by_code = {}
    for row in question_rows:
        qcode = str(row.get("question_code") or "").strip()
        status = str(row.get("approval_status") or "open").strip().lower()
        if status not in {"open", "accepted", "rejected"}:
            raise ReingestError(f"New Questions row {row['__row_number']} has an invalid approval_status.")
        if not qcode:
            if any(value not in (None, "") for key, value in row.items() if key != "__row_number"):
                raise ReingestError(f"New Questions row {row['__row_number']} needs question_code.")
            continue
        if qcode in questions_by_code:
            raise ReingestError(f"New Questions repeats question_code {qcode!r}.")
        questions_by_code[qcode] = row
    response_codes = {str(row.get("question_code") or "").strip() for row in response_rows}
    if response_codes - set(questions_by_code):
        raise ReingestError("New Responses contains a question_code absent from New Questions.")
    accepted = [row for row in question_rows
                if str(row.get("approval_status") or "open").strip().lower() == "accepted"
                and str(row.get("question_code") or "").strip()]
    decisions = [{"question_code": code,
                  "status": str(row.get("approval_status") or "open").strip().lower()}
                 for code, row in sorted(questions_by_code.items())]
    if not accepted:
        return [], decisions
    model_codes = {q.get("identity", {}).get("permanent_code") for q in model.get("questions", [])}
    for row in accepted:
        qcode = str(row["question_code"]).strip()
        if qcode in model_codes:
            raise ReingestError(f"New question code {qcode!r} already exists in the source model.")
        target = (str(row.get("target_quiz_entity_key") or "").strip(),
                  str(row.get("target_pool_entity_key") or "").strip())
        if target not in targets:
            raise ReingestError(f"New Questions row {row['__row_number']} must select an exact quiz and pool from Target Pools.")
    lineage = "review-additions-" + hashlib.sha256(
        json.dumps([(row.get("question_code"), row.get("question_text"), row.get("target_pool_entity_key"))
                    for row in accepted], ensure_ascii=False, sort_keys=True, default=str).encode()).hexdigest()[:20]
    pool_pairs = sorted({(str(row["target_quiz_entity_key"]).strip(),
                          str(row["target_pool_entity_key"]).strip()) for row in accepted})
    pool_codes = {pair: f"TARGET{index:04d}" for index, pair in enumerate(pool_pairs, 1)}
    accepted_codes = {str(row["question_code"]).strip() for row in accepted}
    draft_data = {"config": {"schema": "coursecraft.quiz_draft/1", "lineage_code": lineage,
                             "quiz_code": None, "quiz_title": None},
                  "Questions": [], "Responses": [], "Pools": [], "Draws": []}
    for pair, pool_code in pool_codes.items():
        draft_data["Pools"].append({"pool_code": pool_code, "title": targets[pair]})
    for row in accepted:
        qcode = str(row["question_code"]).strip()
        pair = (str(row["target_quiz_entity_key"]).strip(),
                str(row["target_pool_entity_key"]).strip())
        draft_data["Questions"].append({"question_code": qcode, "question_type": row.get("question_type"),
            "question_text": row.get("question_text"), "points": row.get("points"),
            "pool_code": pool_codes[pair], "answer_key": row.get("answer_key"),
            "feedback": row.get("feedback"), "content_format": row.get("content_format")})
    for row in response_rows:
        if str(row.get("question_code") or "").strip() in accepted_codes:
            draft_data["Responses"].append({key: row.get(key) for key in NEW_RESPONSE_HEADERS})
    draft_model, report = ingest(draft_data, source_name="assessment review new-question rows", input_format="json")
    if not report.get("valid_intake"):
        raise ReingestError("New question rows did not pass draft intake validation.")
    by_code = {q["identity"]["permanent_code"]: q for q in draft_model["questions"]}
    annotations = []
    for row in accepted:
        qcode = str(row["question_code"]).strip()
        qmodel = by_code[qcode]
        addition = {"question": qmodel,
                    "target_quiz_entity_key": str(row["target_quiz_entity_key"]).strip(),
                    "target_pool_entity_key": str(row["target_pool_entity_key"]).strip(),
                    "draft_model_sha256": draft_model["source"]["fingerprint"]["digest"]}
        annotations.extend(_annotation_pair(addition["target_quiz_entity_key"], "/questions/add/" + qcode,
            addition, "accepted", row, row["__row_number"], metadata_policy, cfg["question_sheet"]))
    return annotations, decisions


def materialize_assessment_decisions(model_path: Path, baseline_path: Path, edited_path: Path, *, metadata_policy="optional") -> dict:
    """Validate the grouped source view, then collect edits through native reingest."""
    packet = baseline_path.parent
    spec = validate_packet(packet)
    if spec.get("profile") == "inventory_only":
        raise ReingestError("This library-only inventory has no assessment edits to import; use the blank drafting template for new assessments")
    if sha256_file(baseline_path) != spec["baseline"]["sha256"] or sha256_file(model_path) != spec["model"]["sha256"]:
        raise ReingestError("Assessment review model or baseline does not match the packet")
    baseline, edited = load_workbook(baseline_path), load_workbook(edited_path)
    if set(baseline.sheetnames) != set(edited.sheetnames):
        raise ReingestError("Assessment review sheets were added or removed")
    native_path = _inside(packet, spec["native_workbook"]["path"])
    normalized = load_workbook(native_path)
    native_columns = {c.value: c.column for c in normalized["Quiz Questions"][1]}
    native_sheet = normalized["Quiz Questions"]
    expected_keys = {native_sheet.cell(r, native_columns["occurrence_key"]).value
                     for r in range(2, native_sheet.max_row+1)}
    model = json.loads(model_path.read_text(encoding="utf-8"))
    model_questions = {q["entity_key"]: (i, q) for i, q in enumerate(model["questions"])}
    occurrence_quizzes = {o["occurrence_key"]: o["quiz_key"] for o in
                          build_quiz_review_projection(model)["question_occurrences"]}
    seen = set()
    asset_paths = {ref["path"] for ref in spec["assets"]}
    allowed = {}
    for tab in spec["sheets"]:
        sheet = baseline[tab["title"]]
        headers = {c.value: c.column for c in sheet[2]}
        for item in tab["rows"]:
            key = item["occurrence_key"]
            row, native_row = item["row"], item["native_row"]
            if (key in seen or key not in expected_keys or
                sheet.cell(row, headers["occurrence_key"]).value != key or
                native_sheet.cell(native_row, native_columns["occurrence_key"]).value != key or
                occurrence_quizzes.get(key) != tab["quiz_key"]):
                raise ReingestError("Assessment layout does not match source occurrence identities")
            values = {h: native_sheet.cell(native_row, c).value for h, c in native_columns.items()}
            qi, question = model_questions[values["question_entity_key"]]
            held = bool(_presentation(question, values, qi)["rich_revision_hold"])
            if item["rich_revision_hold"] is not held:
                raise ReingestError("Assessment layout rich-content policy differs from source")
            seen.add(key)
        allowed[tab["title"]] = {(r["row"], headers[h]) for r in tab["rows"] for h in EDITABLE_HEADERS}
    if seen != expected_keys or spec["question_count"] != len(expected_keys):
        raise ReingestError("Assessment layout does not cover all source occurrences")
    if "Quiz Settings" in baseline:
        sheet = baseline["Quiz Settings"]
        allowed["Quiz Settings"] = {(r, c.column) for r in range(2, sheet.max_row+1)
                                    for c in sheet[1] if c.value in SETTINGS_EDITABLE_HEADERS}
    image_layout = spec.get("image_replacements")
    if image_layout:
        sheet = baseline[image_layout["sheet"]]
        cols = {cell.value: cell.column for cell in sheet[1]}
        allowed[image_layout["sheet"]] = {
            (item["row"], cols[name])
            for item in image_layout["rows"]
            for name in ("replacement_file", "revision_reason", "approval_status",
                         "proposed_by", "proposed_at", "approved_by", "approved_at")
        }
    draft_layout = spec.get("question_additions")
    if draft_layout:
        for sheet_name in (draft_layout["question_sheet"], draft_layout["response_sheet"]):
            sheet = baseline[sheet_name]
            allowed[sheet_name] = {
                (row, col)
                for row in range(2, max(sheet.max_row, edited[sheet_name].max_row) + 1)
                for col in range(1, sheet.max_column + 1)
            }
    for before in baseline:
        after = edited[before.title]
        permitted = allowed.get(before.title, set())
        for cells in after:
            for cell in cells:
                if cell.data_type == "f":
                    raise ReingestError(f"Formula cells are not accepted ({after.title}!{cell.coordinate})")
        for r in range(1, max(before.max_row, after.max_row)+1):
            for c in range(1, max(before.max_column, after.max_column)+1):
                a, b = before.cell(r, c), after.cell(r, c)
                if (r, c) in permitted:
                    continue
                # Blank and empty-string serialize identically across spreadsheet editors.
                av = None if a.value in (None, "") else a.value
                bv = None if b.value in (None, "") else b.value
                if av != bv:
                    raise ReingestError(f"Source or layout cell changed: {before.title}!{a.coordinate}")
                if (a.hyperlink.target if a.hyperlink else None) != (b.hyperlink.target if b.hyperlink else None):
                    raise ReingestError(f"Source image link changed: {before.title}!{a.coordinate}")
                if a.hyperlink and a.hyperlink.target and a.hyperlink.target.startswith("Images/"):
                    if a.hyperlink.target not in asset_paths:
                        raise ReingestError("Image link is missing from the packet asset manifest")
    for tab in spec["sheets"]:
        sheet = edited[tab["title"]]
        headers = {c.value: c.column for c in sheet[2]}
        for item in tab["rows"]:
            row = item["row"]
            status = str(sheet.cell(row, headers["approval_status"]).value or "").lower().strip()
            if item["rich_revision_hold"] and status == "accepted" and any(
                sheet.cell(row, headers[h]).value not in (None, "") for h in
                ("revised_question_text", "revised_response_options", "revised_answer_key")):
                raise ReingestError("Accepted equation/image revisions need a supported rich-content path; retain as an open draft")
            for header in EDITABLE_HEADERS:
                normalized["Quiz Questions"].cell(item["native_row"], native_columns[header]).value = sheet.cell(row, headers[header]).value
    if "Quiz Settings" in normalized:
        for cells in edited["Quiz Settings"]:
            for cell in cells:
                normalized["Quiz Settings"].cell(cell.row, cell.column).value = cell.value
    with tempfile.TemporaryDirectory(prefix="quiz-assessment-import-") as scratch:
        normalized_path = Path(scratch) / "native-review.xlsx"
        normalized.save(normalized_path)
        overlay = materialize_workbook_decisions(model_path, native_path, normalized_path, metadata_policy=metadata_policy)
    overlay["workbooks"] = {"baseline_sha256": sha256_file(baseline_path), "edited_sha256": sha256_file(edited_path)}
    overlay["assessment_review"] = {"schema": spec["schema"], "layout_sha256": sha256_file(packet / LAYOUT), "question_count": spec["question_count"]}
    extra_annotations = _image_replacement_annotations(spec, edited, packet, metadata_policy)
    draft_annotations, draft_decisions = _new_question_annotations(spec, edited, model, metadata_policy)
    overlay["annotations"].extend(extra_annotations)
    overlay["annotations"].extend(draft_annotations)
    overlay["question_draft_decisions"] = draft_decisions
    response_row_count = (sum(1 for row in edited[draft_layout["response_sheet"]].iter_rows(min_row=2)
                              if any(cell.value not in (None, "") for cell in row))
                          if draft_layout else 0)
    overlay["content_minimized_summary"]["changed_row_count"] += len(draft_decisions) + response_row_count
    overlay["content_minimized_summary"]["annotation_count"] = len(overlay["annotations"])
    overlay["content_minimized_summary"]["accepted_change_count"] = sum(
        annotation["kind"] == "approved_change" for annotation in overlay["annotations"])
    if extra_annotations or draft_annotations:
        check_model = json.loads(model_path.read_text(encoding="utf-8"))
        check_model["annotations"].extend(extra_annotations + draft_annotations)
        issues = validate_contract(check_model, mode="transform")
        if issues:
            raise ReingestError(f"Assessment addition decisions are invalid: {issues[0].render()}")
        overlay["materialized_view_fingerprint"] = _canonical_digest({
            "annotations": overlay["annotations"],
            "settings_decisions": overlay["settings_decisions"],
            "settings_inputs": overlay["settings_inputs"],
        })
    return overlay


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    prepare = commands.add_parser("prepare", help="Create a new portable review folder from native extraction")
    prepare.add_argument("--source-workbook", type=Path, required=True)
    prepare.add_argument("--model", type=Path, required=True)
    prepare.add_argument("--output-dir", type=Path, required=True)
    collect = commands.add_parser("import", help="Collect assessment edits into a decision overlay")
    collect.add_argument("--packet", type=Path, required=True)
    collect.add_argument("--edited-workbook", type=Path)
    collect.add_argument("--output", type=Path, required=True)
    collect.add_argument("--metadata-policy", choices=("optional", "required"), default="optional")
    args = parser.parse_args(argv)
    try:
        if args.command == "prepare":
            result = prepare_review_packet(args.source_workbook, args.model, args.output_dir)
            print(f"Created {len(result['sheets'])} assessment tabs, {result['question_count']} quiz occurrences and {len(result['assets'])} image files in {args.output_dir}")
        else:
            if args.output.exists():
                raise ValueError("Choose a new output filename; existing decisions will not be overwritten")
            result = materialize_assessment_decisions(args.packet / "model.json", args.packet / BASELINE,
                args.edited_workbook or args.packet / WORKING, metadata_policy=args.metadata_policy)
            args.output.parent.mkdir(parents=True, exist_ok=True)
            _json(args.output, result)
            print(json.dumps(result["content_minimized_summary"]))
    except (OSError, ValueError, KeyError) as exc:
        parser.exit(2, f"assessment review: {exc}\n")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
