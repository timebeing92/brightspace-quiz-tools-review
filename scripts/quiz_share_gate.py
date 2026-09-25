#!/usr/bin/env python3
"""Fail closed when a proposed quiz handoff contains configured privacy leaks."""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import sys
import zipfile
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path, PurePosixPath
from typing import BinaryIO

from openpyxl import load_workbook


RESULT_FORMAT = "coursecraft.quiz_share_gate_result/1"
MAX_MEMBER_BYTES = 20 * 1024 * 1024
MAX_ARCHIVE_BYTES = 200 * 1024 * 1024
TEXT_EXTENSIONS = {
    ".cfg",
    ".css",
    ".csv",
    ".htm",
    ".html",
    ".ini",
    ".js",
    ".json",
    ".jsonl",
    ".md",
    ".mjs",
    ".py",
    ".qmd",
    ".rels",
    ".rst",
    ".svg",
    ".toml",
    ".tsv",
    ".txt",
    ".xml",
    ".yaml",
    ".yml",
}
TEXT_FILENAMES = {".gitignore", ".gitattributes"}
IMAGE_EXTENSIONS = {".bmp", ".gif", ".ico", ".jpeg", ".jpg", ".png", ".webp"}
OOXML_EXTENSIONS = {".docx", ".pptx"}
LOCAL_ABSOLUTE_PATH_PATTERNS = (
    re.compile(
        r"(?i)(?:file://)?/(?:Users|home|Volumes|tmp|private/(?:tmp|var)|var/folders)/"
        r"[^\s\"'<>]+"
    ),
    re.compile(
        r"(?i)(?:file:///)?[A-Z]:[\\/](?:Users|Documents and Settings)[\\/]"
        r"[^\s\"'<>]+"
    ),
    re.compile(r"\\\\[^\s\\/]+[\\/][^\s\"'<>]+"),
)
ABSOLUTE_PATH_FIELD_PATTERN = re.compile(
    r"(?<![A-Za-z0-9_])[A-Za-z0-9_]*_absolute_path(?![A-Za-z0-9_])",
    flags=re.IGNORECASE,
)


@dataclass(frozen=True)
class Finding:
    rule: str
    artifact: str
    location: str


@dataclass
class ScanState:
    person_patterns: list[re.Pattern[str]] = field(default_factory=list)
    lineage_patterns: list[re.Pattern[str]] = field(default_factory=list)
    findings: list[Finding] = field(default_factory=list)
    scanned_artifacts: int = 0
    scanned_values: int = 0
    image_files_name_scanned: int = 0
    image_files_embedded_text_scanned: int = 0
    unscanned_artifacts: list[dict[str, str]] = field(default_factory=list)
    _finding_keys: set[tuple[str, str, str]] = field(default_factory=set)

    def safe_reference(self, value: str) -> str:
        sensitive = (
            any(pattern.search(value) for pattern in LOCAL_ABSOLUTE_PATH_PATTERNS)
            or bool(ABSOLUTE_PATH_FIELD_PATTERN.search(value))
            or any(pattern.search(value) for pattern in self.person_patterns)
            or any(pattern.search(value) for pattern in self.lineage_patterns)
        )
        if not sensitive:
            return value
        digest = hashlib.sha256(value.encode("utf-8")).hexdigest()[:16]
        return f"redacted-reference:{digest}"

    def add_finding(self, rule: str, artifact: str, location: str) -> None:
        artifact = self.safe_reference(artifact)
        location = self.safe_reference(location)
        key = (rule, artifact, location)
        if key in self._finding_keys:
            return
        self._finding_keys.add(key)
        self.findings.append(Finding(rule=rule, artifact=artifact, location=location))

    def add_unscanned(self, artifact: str, reason: str) -> None:
        self.unscanned_artifacts.append(
            {"artifact": self.safe_reference(artifact), "reason": reason}
        )


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def load_term_file(path: Path | None) -> list[str]:
    if path is None:
        return []
    terms: list[str] = []
    for raw_line in path.read_text(encoding="utf-8").splitlines():
        term = raw_line.strip()
        if term and not term.startswith("#"):
            terms.append(term)
    return terms


def normalized_terms(values: list[str]) -> list[str]:
    return sorted({value.strip() for value in values if value.strip()}, key=str.casefold)


def term_patterns(terms: list[str]) -> list[re.Pattern[str]]:
    return [
        re.compile(rf"(?<!\w){re.escape(term)}(?!\w)", flags=re.IGNORECASE)
        for term in terms
    ]


def scan_value(
    value: object,
    *,
    artifact: str,
    location: str,
    state: ScanState,
    person_patterns: list[re.Pattern[str]],
    lineage_patterns: list[re.Pattern[str]],
) -> None:
    text = str(value or "")
    if not text:
        return
    state.scanned_values += 1
    if any(pattern.search(text) for pattern in LOCAL_ABSOLUTE_PATH_PATTERNS):
        state.add_finding("local_absolute_path", artifact, location)
    if ABSOLUTE_PATH_FIELD_PATTERN.search(text):
        state.add_finding("absolute_path_field", artifact, location)
    if any(pattern.search(text) for pattern in person_patterns):
        state.add_finding("configured_person_name", artifact, location)
    if any(pattern.search(text) for pattern in lineage_patterns):
        state.add_finding("disallowed_lineage_identifier", artifact, location)


def scan_json_value(
    value: object,
    *,
    artifact: str,
    location: str,
    state: ScanState,
    person_patterns: list[re.Pattern[str]],
    lineage_patterns: list[re.Pattern[str]],
) -> None:
    if isinstance(value, dict):
        for key, child in value.items():
            child_location = f"{location}.{key}"
            scan_value(
                key,
                artifact=artifact,
                location=f"{child_location} (key)",
                state=state,
                person_patterns=person_patterns,
                lineage_patterns=lineage_patterns,
            )
            scan_json_value(
                child,
                artifact=artifact,
                location=child_location,
                state=state,
                person_patterns=person_patterns,
                lineage_patterns=lineage_patterns,
            )
        return
    if isinstance(value, list):
        for index, child in enumerate(value):
            scan_json_value(
                child,
                artifact=artifact,
                location=f"{location}[{index}]",
                state=state,
                person_patterns=person_patterns,
                lineage_patterns=lineage_patterns,
            )
        return
    scan_value(
        value,
        artifact=artifact,
        location=location,
        state=state,
        person_patterns=person_patterns,
        lineage_patterns=lineage_patterns,
    )


def scan_text(
    text: str,
    *,
    suffix: str,
    artifact: str,
    state: ScanState,
    person_patterns: list[re.Pattern[str]],
    lineage_patterns: list[re.Pattern[str]],
) -> None:
    if suffix == ".json":
        try:
            payload = json.loads(text)
        except json.JSONDecodeError:
            pass
        else:
            scan_json_value(
                payload,
                artifact=artifact,
                location="$",
                state=state,
                person_patterns=person_patterns,
                lineage_patterns=lineage_patterns,
            )
            return

    if suffix == ".jsonl":
        all_json = True
        for line_number, line in enumerate(text.splitlines(), start=1):
            if not line.strip():
                continue
            try:
                payload = json.loads(line)
            except json.JSONDecodeError:
                all_json = False
                break
            scan_json_value(
                payload,
                artifact=artifact,
                location=f"line:{line_number}",
                state=state,
                person_patterns=person_patterns,
                lineage_patterns=lineage_patterns,
            )
        if all_json:
            return

    if suffix in {".csv", ".tsv"}:
        delimiter = "\t" if suffix == ".tsv" else ","
        for row_number, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1):
            for column_number, value in enumerate(row, start=1):
                scan_value(
                    value,
                    artifact=artifact,
                    location=f"row:{row_number},column:{column_number}",
                    state=state,
                    person_patterns=person_patterns,
                    lineage_patterns=lineage_patterns,
                )
        return

    for line_number, line in enumerate(text.splitlines(), start=1):
        scan_value(
            line,
            artifact=artifact,
            location=f"line:{line_number}",
            state=state,
            person_patterns=person_patterns,
            lineage_patterns=lineage_patterns,
        )


def scan_image_bytes(
    data: bytes,
    *,
    artifact: str,
    state: ScanState,
    person_patterns: list[re.Pattern[str]],
    lineage_patterns: list[re.Pattern[str]],
) -> None:
    if len(data) > MAX_MEMBER_BYTES:
        state.add_unscanned(artifact, "image_scan_size_limit")
        return
    state.scanned_artifacts += 1
    state.image_files_name_scanned += 1
    state.image_files_embedded_text_scanned += 1
    for encoding in ("utf-8", "utf-16-le"):
        decoded = data.decode(encoding, errors="ignore")
        if not decoded:
            continue
        scan_value(
            decoded,
            artifact=artifact,
            location=f"embedded image text ({encoding})",
            state=state,
            person_patterns=person_patterns,
            lineage_patterns=lineage_patterns,
        )


def scan_xlsx(
    source: Path | BinaryIO,
    *,
    artifact: str,
    state: ScanState,
    person_patterns: list[re.Pattern[str]],
    lineage_patterns: list[re.Pattern[str]],
) -> None:
    try:
        workbook = load_workbook(source, read_only=False, data_only=False, keep_links=True)
    except Exception:
        state.add_unscanned(artifact, "xlsx_open_failed")
        return
    state.scanned_artifacts += 1
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value not in (None, ""):
                        scan_value(
                            cell.value,
                            artifact=artifact,
                            location=f"{sheet.title}!{cell.coordinate}",
                            state=state,
                            person_patterns=person_patterns,
                            lineage_patterns=lineage_patterns,
                        )
                    if cell.hyperlink and cell.hyperlink.target:
                        scan_value(
                            cell.hyperlink.target,
                            artifact=artifact,
                            location=f"{sheet.title}!{cell.coordinate} (hyperlink)",
                            state=state,
                            person_patterns=person_patterns,
                            lineage_patterns=lineage_patterns,
                        )
        properties = workbook.properties
        for property_name in ("creator", "lastModifiedBy", "title", "subject", "description"):
            property_value = getattr(properties, property_name, None)
            if property_value:
                scan_value(
                    property_value,
                    artifact=artifact,
                    location=f"workbook.properties.{property_name}",
                    state=state,
                    person_patterns=person_patterns,
                    lineage_patterns=lineage_patterns,
                )
    finally:
        workbook.close()


def safe_member_name(name: str) -> bool:
    path = PurePosixPath(name)
    return not path.is_absolute() and ".." not in path.parts


def scan_zip(
    source: Path | BinaryIO,
    *,
    artifact: str,
    state: ScanState,
    person_patterns: list[re.Pattern[str]],
    lineage_patterns: list[re.Pattern[str]],
    depth: int = 0,
) -> None:
    if depth > 2:
        state.add_unscanned(artifact, "nested_archive_depth_exceeded")
        return
    try:
        archive = zipfile.ZipFile(source)
    except (OSError, zipfile.BadZipFile):
        state.add_unscanned(artifact, "zip_open_failed")
        return
    total_bytes = 0
    with archive:
        for info in archive.infolist():
            if info.is_dir():
                continue
            member_artifact = f"{artifact}!/{info.filename}"
            scan_value(
                info.filename,
                artifact=member_artifact,
                location="archive member name",
                state=state,
                person_patterns=person_patterns,
                lineage_patterns=lineage_patterns,
            )
            if not safe_member_name(info.filename):
                state.add_unscanned(member_artifact, "unsafe_archive_member_name")
                continue
            total_bytes += info.file_size
            if info.file_size > MAX_MEMBER_BYTES or total_bytes > MAX_ARCHIVE_BYTES:
                state.add_unscanned(member_artifact, "archive_scan_size_limit")
                continue
            data = archive.read(info)
            suffix = Path(info.filename).suffix.lower()
            if suffix == ".xlsx":
                scan_xlsx(
                    io.BytesIO(data),
                    artifact=member_artifact,
                    state=state,
                    person_patterns=person_patterns,
                    lineage_patterns=lineage_patterns,
                )
            elif suffix in OOXML_EXTENSIONS or suffix == ".zip":
                scan_zip(
                    io.BytesIO(data),
                    artifact=member_artifact,
                    state=state,
                    person_patterns=person_patterns,
                    lineage_patterns=lineage_patterns,
                    depth=depth + 1,
                )
            elif suffix in TEXT_EXTENSIONS or Path(info.filename).name.lower() in TEXT_FILENAMES:
                state.scanned_artifacts += 1
                scan_text(
                    data.decode("utf-8", errors="replace"),
                    suffix=suffix,
                    artifact=member_artifact,
                    state=state,
                    person_patterns=person_patterns,
                    lineage_patterns=lineage_patterns,
                )
            elif suffix in IMAGE_EXTENSIONS:
                scan_image_bytes(
                    data,
                    artifact=member_artifact,
                    state=state,
                    person_patterns=person_patterns,
                    lineage_patterns=lineage_patterns,
                )
            else:
                state.add_unscanned(member_artifact, "unsupported_archive_member_type")


def scan_file(
    path: Path,
    *,
    artifact: str,
    state: ScanState,
    person_patterns: list[re.Pattern[str]],
    lineage_patterns: list[re.Pattern[str]],
) -> None:
    scan_value(
        artifact,
        artifact=artifact,
        location="artifact name",
        state=state,
        person_patterns=person_patterns,
        lineage_patterns=lineage_patterns,
    )
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        scan_xlsx(
            path,
            artifact=artifact,
            state=state,
            person_patterns=person_patterns,
            lineage_patterns=lineage_patterns,
        )
    elif suffix in OOXML_EXTENSIONS or suffix == ".zip":
        scan_zip(
            path,
            artifact=artifact,
            state=state,
            person_patterns=person_patterns,
            lineage_patterns=lineage_patterns,
        )
    elif suffix in TEXT_EXTENSIONS or path.name.lower() in TEXT_FILENAMES:
        if path.stat().st_size > MAX_MEMBER_BYTES:
            state.add_unscanned(artifact, "text_scan_size_limit")
            return
        state.scanned_artifacts += 1
        scan_text(
            path.read_text(encoding="utf-8", errors="replace"),
            suffix=suffix,
            artifact=artifact,
            state=state,
            person_patterns=person_patterns,
            lineage_patterns=lineage_patterns,
        )
    elif suffix in IMAGE_EXTENSIONS:
        scan_image_bytes(
            path.read_bytes(),
            artifact=artifact,
            state=state,
            person_patterns=person_patterns,
            lineage_patterns=lineage_patterns,
        )
    else:
        state.add_unscanned(artifact, "unsupported_file_type")


def iter_target_files(target: Path) -> list[tuple[Path, str]]:
    if target.is_file():
        return [(target, target.name)]
    rows: list[tuple[Path, str]] = []
    for path in sorted(target.rglob("*")):
        if path.is_symlink():
            rows.append((path, path.relative_to(target).as_posix()))
        elif path.is_file():
            rows.append((path, path.relative_to(target).as_posix()))
    return rows


def configuration_failure(message: str) -> int:
    print(json.dumps({"result_format": RESULT_FORMAT, "status": "configuration_error", "message": message}, indent=2))
    return 2


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("target", type=Path, help="File or directory proposed for handoff")
    parser.add_argument("--output", type=Path, help="Optional path for the content-minimized JSON result")
    parser.add_argument("--person", action="append", default=[], help="Personal name that must not appear; repeat as needed")
    parser.add_argument("--person-file", type=Path, help="UTF-8 file with one disallowed personal name per line")
    parser.add_argument(
        "--allow-no-person-terms",
        action="store_true",
        help="Explicitly record that the name scan has no configured terms",
    )
    parser.add_argument(
        "--lineage-policy",
        choices=("allow", "deny"),
        required=True,
        help="Explicit disclosure decision for configured lineage identifiers",
    )
    parser.add_argument("--lineage-id", action="append", default=[], help="Lineage identifier to deny; repeat as needed")
    parser.add_argument("--lineage-file", type=Path, help="UTF-8 file with one lineage identifier per line")
    parser.add_argument(
        "--allow-unscanned",
        action="store_true",
        help="Permit unsupported/oversized artifacts while recording them in the result",
    )
    args = parser.parse_args()

    target = args.target.expanduser()
    if not target.exists():
        return configuration_failure("target does not exist")
    try:
        people = normalized_terms(args.person + load_term_file(args.person_file))
        lineage_ids = normalized_terms(args.lineage_id + load_term_file(args.lineage_file))
    except OSError:
        return configuration_failure("a configured term file could not be read")
    if not people and not args.allow_no_person_terms:
        return configuration_failure("configure --person/--person-file or explicitly pass --allow-no-person-terms")
    if args.lineage_policy == "deny" and not lineage_ids:
        return configuration_failure("lineage-policy deny requires --lineage-id or --lineage-file")

    people_patterns = term_patterns(people)
    lineage_patterns = term_patterns(lineage_ids) if args.lineage_policy == "deny" else []
    state = ScanState(
        person_patterns=people_patterns,
        lineage_patterns=lineage_patterns,
    )
    scan_value(
        target.name,
        artifact="target",
        location="target name",
        state=state,
        person_patterns=people_patterns,
        lineage_patterns=lineage_patterns,
    )
    for path, artifact in iter_target_files(target):
        if path.is_symlink():
            state.add_unscanned(artifact, "symlink_not_followed")
            continue
        scan_file(
            path,
            artifact=artifact,
            state=state,
            person_patterns=people_patterns,
            lineage_patterns=lineage_patterns,
        )

    findings = sorted(state.findings, key=lambda row: (row.rule, row.artifact, row.location))
    unscanned_blocking = bool(state.unscanned_artifacts and not args.allow_unscanned)
    status = "fail" if findings or unscanned_blocking else "pass"
    target_fingerprint = hashlib.sha256(target.name.encode("utf-8")).hexdigest()
    result = {
        "result_format": RESULT_FORMAT,
        "generated_at": utc_now(),
        "status": status,
        "target": {
            "name_sha256": target_fingerprint,
            "kind": "directory" if target.is_dir() else "file",
        },
        "policy": {
            "absolute_local_paths": "deny",
            "absolute_path_fields": "deny",
            "named_person_scan": "configured_terms" if people else "explicitly_no_terms",
            "named_person_term_count": len(people),
            "lineage_disclosure": args.lineage_policy,
            "lineage_term_count": len(lineage_ids),
            "unscanned_artifacts": "allow" if args.allow_unscanned else "deny",
        },
        "summary": {
            "finding_count": len(findings),
            "scanned_artifact_count": state.scanned_artifacts,
            "scanned_value_count": state.scanned_values,
            "image_filename_scan_count": state.image_files_name_scanned,
            "image_embedded_text_scan_count": state.image_files_embedded_text_scanned,
            "unscanned_artifact_count": len(state.unscanned_artifacts),
        },
        "findings": [asdict(row) for row in findings],
        "unscanned_artifacts": sorted(
            state.unscanned_artifacts,
            key=lambda row: (row["artifact"], row["reason"]),
        ),
        "content_minimization": "Matched values and configured terms are not copied into this result.",
    }
    rendered = json.dumps(result, indent=2) + "\n"
    if args.output:
        output_path = args.output.expanduser()
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(rendered, encoding="utf-8")
    sys.stdout.write(rendered)
    return 0 if status == "pass" else 1


if __name__ == "__main__":
    raise SystemExit(main())
