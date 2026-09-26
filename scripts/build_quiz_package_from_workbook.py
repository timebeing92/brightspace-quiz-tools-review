#!/usr/bin/env python3
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import html
import json
import shutil
import sys
import uuid
import zipfile
from dataclasses import dataclass
from pathlib import Path
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

from d2l_package_lib import build_grades_xml
from quiz_build_support import (
    CAPABILITY_REGISTRY_PATH,
    artifact_row,
    build_run_receipt,
    default_settings_receipt,
    load_authoring_projection,
    materialize_effective_settings_receipt,
    safe_ident,
    sha256_file,
    target_identifier_issues,
)


D2L_NS = "http://desire2learn.com/xsd/d2lcp_v2p0"
IMS_NS = "http://www.imsglobal.org/xsd/imscp_v1p1"
IMSMD_NS = "http://www.imsglobal.org/xsd/imsmd_rootv1p2p1"  # canonical D2L LOM ns; v1p2 is not what D2L emits
XML_LANG = "{http://www.w3.org/XML/1998/namespace}lang"
UUID_NAMESPACE = uuid.UUID("9d8758ca-6353-4d61-9e6b-4a52b82c0f24")

ET.register_namespace("d2l_2p0", D2L_NS)
ET.register_namespace("", IMS_NS)
ET.register_namespace("imsmd", IMSMD_NS)


def ns_attr(namespace: str, name: str) -> str:
    return f"{{{namespace}}}{name}"


def clean(value: object) -> str:
    if value is None:
        return ""
    text = str(value)
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text.strip()


def normalized_header(value: object) -> str:
    return clean(value).lower().replace(" ", "_").replace("-", "_")


def stable_uuid(value: str) -> str:
    return str(uuid.uuid5(UUID_NAMESPACE, value))


def as_float(value: str, default: float = 1.0) -> float:
    if not value:
        return default
    return float(value)


def as_int(value: str, default: int = 0) -> int:
    if not value:
        return default
    return int(float(value))


def truthy(value: str) -> bool:
    return clean(value).upper() in {"TRUE", "YES", "Y", "1"}


def html_fragment(value: str, content_format: str | None = None) -> str:
    # Explicit canonical formats outrank the legacy workbook HTML heuristic.
    if content_format == "plain_text":
        return f"<p>{html.escape(value, quote=False)}</p>" if value else ""
    if content_format in {"html", "xhtml"}:
        return value
    if content_format is not None:
        raise ValueError(f"Content format {content_format!r} requires a reviewed package conversion.")
    value = clean(value)
    if not value:
        return ""
    if "<" in value and ">" in value:
        return value
    return f"<p>{value}</p>"


def normalize_question_type(value: str) -> str:
    token = clean(value).upper().replace("-", "_").replace(" ", "_").replace("/", "_")
    aliases = {
        "MULTICHOICE": "MULTICHOICE",
        "MULTIPLE_CHOICE": "MULTICHOICE",
        "MULTIPLECHOICE": "MULTICHOICE",
        "TRUEFALSE": "TRUEFALSE",
        "TRUE_FALSE": "TRUEFALSE",
        "TRUE_OR_FALSE": "TRUEFALSE",
        "MULTISELECT": "MULTISELECT",
        "MULTI_SELECT": "MULTISELECT",
        "MULTIPLE_SELECT": "MULTISELECT",
        "WRITTEN_RESPONSE": "WRITTEN_RESPONSE",
        "WRITTENRESPONSE": "WRITTEN_RESPONSE",
        "LONG_ANSWER": "WRITTEN_RESPONSE",
        "LONGANSWER": "WRITTEN_RESPONSE",
    }
    if token not in aliases:
        raise ValueError(f"Unsupported target_question_type: {value!r}")
    return aliases[token]


@dataclass(frozen=True)
class QuestionRow:
    bank_id: str
    bank_title: str
    source_label: str
    question_code: str
    question_title: str
    target_question_type: str
    scoring_policy: str
    scoring_policy_value: str
    question_text: str
    points: float
    randomize_answers: bool
    options: list[str]
    correct_option_key: str
    correct_option_text: str
    evaluator_answer_key: str
    notes: str
    label: str
    question_text_format: str | None = None
    option_formats: list[str] | None = None
    evaluator_answer_key_format: str | None = None


@dataclass(frozen=True)
class QuizSection:
    section_order: int
    section_title: str
    source_bank_id: str
    recommended_draw_count: int
    points_per_question: float
    notes: str


@dataclass(frozen=True)
class QuizSettings:
    is_active: bool = False
    attempts_allowed: int = 1
    time_limit: int = 0
    show_clock: bool = False
    enforce_time_limit: bool = False
    is_forward_only: bool = False


def question_from_projection(row: dict) -> QuestionRow:
    question_code = row["question_code"]
    return QuestionRow(
        bank_id=row["bank_id"],
        bank_title=row["bank_title"],
        source_label=row["source_label"],
        question_code=question_code,
        question_title=row["question_title"],
        target_question_type=row["target_question_type"],
        scoring_policy=row["scoring_policy"],
        scoring_policy_value=row["scoring_policy_value"],
        question_text=row["question_text"],
        points=float(row["points"]),
        randomize_answers=bool(row["randomize_answers"]),
        options=[*row["options"], "", "", "", "", ""][:5],
        correct_option_key=row["correct_option_key"],
        correct_option_text=row["correct_option_text"],
        evaluator_answer_key=row["evaluator_answer_key"],
        notes=row["notes"],
        label=f"QUES_{safe_ident(question_code, 'Q')}",
        question_text_format=row.get("question_text_format"),
        option_formats=row.get("option_formats"),
        evaluator_answer_key_format=row.get("evaluator_answer_key_format"),
    )


def section_from_projection(row: dict) -> QuizSection:
    return QuizSection(
        section_order=int(row["section_order"]),
        section_title=row["section_title"],
        source_bank_id=row["source_bank_id"],
        recommended_draw_count=int(row["recommended_draw_count"]),
        points_per_question=float(row["points_per_question"]),
        notes=row["notes"],
    )


def row_dict(ws, row_index: int) -> dict[str, str]:
    headers = [normalized_header(ws.cell(1, column).value) for column in range(1, ws.max_column + 1)]
    return {
        headers[column - 1]: clean(ws.cell(row_index, column).value)
        for column in range(1, ws.max_column + 1)
        if headers[column - 1]
    }


def load_workbook_data(workbook_path: Path) -> tuple[list[QuestionRow], list[QuizSection]]:
    wb = load_workbook(workbook_path, data_only=True)
    questions: list[QuestionRow] = []

    skipped = {"README", "BANK_INDEX", "QUIZ_STRUCTURE"}
    for sheet_name in wb.sheetnames:
        if sheet_name.upper() in skipped:
            continue
        ws = wb[sheet_name]
        if ws.max_row < 2:
            continue
        headers = {normalized_header(ws.cell(1, column).value) for column in range(1, ws.max_column + 1)}
        if "question_code" not in headers:
            continue

        for row_index in range(2, ws.max_row + 1):
            row = row_dict(ws, row_index)
            if not row.get("question_code"):
                continue
            qtype = normalize_question_type(row.get("target_question_type", ""))
            scoring_policy = row.get("scoring_policy", "").upper()
            if qtype == "MULTISELECT" and scoring_policy not in {"", "ALL_OR_NOTHING"}:
                raise ValueError(
                    f"{row['question_code']} requests unsupported Multi-Select scoring_policy "
                    f"{row.get('scoring_policy')!r}; supported value is ALL_OR_NOTHING."
                )

            bank_id = row.get("bank_id") or sheet_name
            question_code = row["question_code"]
            label = f"QUES_{safe_ident(question_code, 'Q')}"
            questions.append(
                QuestionRow(
                    bank_id=bank_id,
                    bank_title=row.get("bank_title") or bank_id,
                    source_label=row.get("week") or row.get("scope_label") or row.get("source_label") or sheet_name,
                    question_code=question_code,
                    question_title=row.get("question_title") or question_code,
                    target_question_type=qtype,
                    scoring_policy=scoring_policy,
                    scoring_policy_value=row.get("scoring_policy_value", ""),
                    question_text=row.get("question_text", ""),
                    points=as_float(row.get("points", ""), 1.0),
                    randomize_answers=truthy(row.get("randomize_answers", "")),
                    options=[
                        row.get("option_a", ""),
                        row.get("option_b", ""),
                        row.get("option_c", ""),
                        row.get("option_d", ""),
                        row.get("option_e", ""),
                    ],
                    correct_option_key=row.get("correct_option_key", ""),
                    correct_option_text=row.get("correct_option_text", ""),
                    evaluator_answer_key=row.get("evaluator_answer_key", ""),
                    notes=row.get("notes", ""),
                    label=label,
                )
            )

    if not questions:
        raise ValueError(f"No question rows found in {workbook_path}.")

    sections: list[QuizSection] = []
    if "Quiz_Structure" in wb.sheetnames:
        ws = wb["Quiz_Structure"]
        for row_index in range(2, ws.max_row + 1):
            row = row_dict(ws, row_index)
            if not row.get("source_bank_id"):
                continue
            sections.append(
                QuizSection(
                    section_order=as_int(row.get("section_order", ""), len(sections) + 1),
                    section_title=row.get("section_title") or row["source_bank_id"],
                    source_bank_id=row["source_bank_id"],
                    recommended_draw_count=as_int(row.get("recommended_draw_count", ""), 1),
                    points_per_question=as_float(row.get("points_per_question", ""), 1.0),
                    notes=row.get("notes", ""),
                )
            )

    if not sections:
        bank_order: list[str] = []
        bank_titles: dict[str, str] = {}
        for question in questions:
            if question.bank_id not in bank_order:
                bank_order.append(question.bank_id)
                bank_titles[question.bank_id] = question.bank_title
        for index, bank_id in enumerate(bank_order, start=1):
            bank_questions = [question for question in questions if question.bank_id == bank_id]
            sections.append(
                QuizSection(
                    section_order=index,
                    section_title=bank_titles[bank_id],
                    source_bank_id=bank_id,
                    recommended_draw_count=len(bank_questions),
                    points_per_question=bank_questions[0].points if bank_questions else 1.0,
                    notes="Auto-created because workbook had no Quiz_Structure sheet.",
                )
            )

    bank_counts: dict[str, int] = {}
    for question in questions:
        bank_counts[question.bank_id] = bank_counts.get(question.bank_id, 0) + 1
    for section in sections:
        available = bank_counts.get(section.source_bank_id, 0)
        if not available:
            raise ValueError(f"Quiz section {section.section_title!r} points to empty bank {section.source_bank_id!r}.")
        if section.recommended_draw_count > available:
            raise ValueError(
                f"Quiz section {section.section_title!r} draws {section.recommended_draw_count} "
                f"from {available} available questions in bank {section.source_bank_id!r}."
            )

    return questions, sections


def add_qti_metadata(parent: ET.Element, metadata: dict[str, str]) -> None:
    qtimetadata = ET.SubElement(parent, "qtimetadata")
    for label, entry in metadata.items():
        field = ET.SubElement(qtimetadata, "qti_metadatafield")
        ET.SubElement(field, "fieldlabel").text = label
        ET.SubElement(field, "fieldentry").text = entry


def answer_keys(question: QuestionRow, option_count: int) -> set[str]:
    keys = {part.strip().upper() for part in question.correct_option_key.replace(",", ";").split(";") if part.strip()}
    if question.target_question_type == "TRUEFALSE":
        if keys & {"TRUE", "T", "YES"}:
            return {"A"}
        if keys & {"FALSE", "F", "NO"}:
            return {"B"}
    if keys:
        return keys

    expected = clean(question.correct_option_text).lower()
    if expected:
        letters = ["A", "B", "C", "D", "E"]
        return {
            letters[index]
            for index, option in enumerate(question.options[:option_count])
            if clean(option).lower() == expected
        }
    return set()


def item_options(question: QuestionRow) -> list[str]:
    if question.target_question_type == "TRUEFALSE":
        return ["True", "False"]
    return [option for option in question.options if clean(option)]


def add_choice_question(
    flow: ET.Element,
    item: ET.Element,
    question: QuestionRow,
    *,
    true_false_answer_texttype: str = "text/plain",
) -> None:
    response_extension = ET.SubElement(flow, "response_extension")
    ET.SubElement(response_extension, ns_attr(D2L_NS, "display_style")).text = "2"
    ET.SubElement(response_extension, ns_attr(D2L_NS, "enumeration")).text = "6"
    ET.SubElement(response_extension, ns_attr(D2L_NS, "grading_type")).text = "0"

    options = item_options(question)
    if len(options) < 2:
        raise ValueError(f"{question.question_code} needs at least two options.")

    lid_ident = f"{question.label}_LID"
    rcardinality = "Multiple" if question.target_question_type == "MULTISELECT" else "Single"
    response_lid = ET.SubElement(flow, "response_lid", {"ident": lid_ident, "rcardinality": rcardinality})
    render_choice = ET.SubElement(
        response_lid,
        "render_choice",
        {"shuffle": "yes" if question.randomize_answers and question.target_question_type != "TRUEFALSE" else "no"},
    )

    option_idents: list[tuple[str, str]] = []
    for index, option_text in enumerate(options, start=1):
        option_ident = f"{question.label}_A{index:04d}"
        option_idents.append((option_ident, option_text))
        flow_label = ET.SubElement(render_choice, "flow_label", {"class": "Block"})
        response_label = ET.SubElement(flow_label, "response_label", {"ident": option_ident})
        flow_mat = ET.SubElement(response_label, "flow_mat")
        material = ET.SubElement(flow_mat, "material")
        # True/False labels stay bare and use text/plain. A 2026-07-20 tenant
        # probe showed that text/html was reinterpreted as a localization key.
        # Multiple-choice option text remains <p>-wrapped HTML.
        option_format = question.option_formats[index - 1] if question.option_formats is not None else None
        option_markup = option_text if question.target_question_type == "TRUEFALSE" else html_fragment(option_text, option_format)
        option_texttype = (
            true_false_answer_texttype
            if question.target_question_type == "TRUEFALSE"
            else "text/html"
        )
        ET.SubElement(material, "mattext", {"texttype": option_texttype}).text = option_markup

    keys = answer_keys(question, len(options))
    if not keys:
        raise ValueError(f"{question.question_code} has no correct answer key.")

    letters = ["A", "B", "C", "D", "E"]
    correct_idents = {
        option_ident
        for index, (option_ident, _option_text) in enumerate(option_idents)
        if letters[index] in keys
    }

    resprocessing = ET.SubElement(item, "resprocessing")

    if question.target_question_type == "MULTISELECT":
        # Multi-select carries an outcomes/decvar block and an all-or-nothing
        # respcondition (per QUIZ_PACKAGE_BUILD_AND_VALIDATION_METHOD.md). MC/TF
        # do NOT — canonical D2L MC/TF resprocessing has no outcomes/decvar.
        outcomes = ET.SubElement(resprocessing, "outcomes")
        ET.SubElement(outcomes, "decvar", {"vartype": "Integer", "defaultval": "0"})
        respcondition = ET.SubElement(resprocessing, "respcondition", {"title": "All correct selections"})
        conditionvar = ET.SubElement(respcondition, "conditionvar")
        and_block = ET.SubElement(conditionvar, "and")
        for option_ident, _option_text in option_idents:
            if option_ident in correct_idents:
                ET.SubElement(and_block, "varequal", {"respident": lid_ident}).text = option_ident
            else:
                not_block = ET.SubElement(and_block, "not")
                ET.SubElement(not_block, "varequal", {"respident": lid_ident}).text = option_ident
        ET.SubElement(respcondition, "setvar", {"action": "Set"}).text = "100.000000000"
        fallback = ET.SubElement(resprocessing, "respcondition", {"title": "Incorrect selections"})
        ET.SubElement(fallback, "conditionvar")
        ET.SubElement(fallback, "setvar", {"action": "Set"}).text = "0.000000000"
        return

    # MC / TF: one respcondition per option (conditionvar -> setvar ->
    # displayfeedback), each linked to an empty itemfeedback block appended to the
    # item after resprocessing — matching canonical D2L MC/TF structure.
    feedback_idents: list[str] = []
    for index, (option_ident, _option_text) in enumerate(option_idents, start=1):
        feedback_ident = f"{question.label}_IF{index:04d}"
        feedback_idents.append(feedback_ident)
        respcondition = ET.SubElement(resprocessing, "respcondition", {"title": f"Response Condition {index}"})
        conditionvar = ET.SubElement(respcondition, "conditionvar")
        ET.SubElement(conditionvar, "varequal", {"respident": lid_ident}).text = option_ident
        score = "100.000000000" if option_ident in correct_idents else "0.000000000"
        ET.SubElement(respcondition, "setvar", {"action": "Set"}).text = score
        ET.SubElement(respcondition, "displayfeedback", {"feedbacktype": "Response", "linkrefid": feedback_ident})

    for feedback_ident in feedback_idents:
        itemfeedback = ET.SubElement(item, "itemfeedback", {"ident": feedback_ident})
        feedback_material = ET.SubElement(itemfeedback, "material")
        ET.SubElement(feedback_material, "mattext", {"texttype": "text/plain"})


def add_written_response(flow: ET.Element, item: ET.Element, question: QuestionRow) -> None:
    response_extension = ET.SubElement(flow, "response_extension")
    ET.SubElement(response_extension, ns_attr(D2L_NS, "has_signed_comments")).text = "no"
    ET.SubElement(response_extension, ns_attr(D2L_NS, "has_htmleditor")).text = "no"
    ET.SubElement(response_extension, ns_attr(D2L_NS, "has_fileupload")).text = "no"

    response_str = ET.SubElement(flow, "response_str", {"ident": f"{question.label}_STR", "rcardinality": "Multiple"})
    render_fib = ET.SubElement(
        response_str,
        "render_fib",
        {"rows": "5", "columns": "100", "prompt": "Box", "fibtype": "String"},
    )
    response_label = ET.SubElement(render_fib, "response_label", {"ident": f"{question.label}_LA"})
    material = ET.SubElement(response_label, "material")
    ET.SubElement(material, "mattext", {"texttype": "text/plain"})

    answer_text = question.evaluator_answer_key or question.notes or "Instructor-evaluated written response."
    answer_key = ET.SubElement(item, "answer_key")
    answer_key_material = ET.SubElement(answer_key, "answer_key_material")
    flow_mat = ET.SubElement(answer_key_material, "flow_mat")
    material = ET.SubElement(flow_mat, "material")
    ET.SubElement(material, "mattext", {"texttype": "text/html"}).text = html_fragment(answer_text, question.evaluator_answer_key_format)


def create_item(
    parent: ET.Element,
    question: QuestionRow,
    d2l_id: int,
    *,
    true_false_answer_texttype: str = "text/plain",
) -> None:
    rendered_title = (
        f"{question.source_label} - {question.question_title}"
        if question.source_label
        else question.question_title
    )
    item = ET.SubElement(
        parent,
        "item",
        {
            ns_attr(D2L_NS, "id"): str(d2l_id),
            ns_attr(D2L_NS, "page"): "1",
            "ident": f"OBJ_{safe_ident(question.question_code, 'Q')}",
            "label": question.label,
            "title": rendered_title,
        },
    )

    qtype_map = {
        "MULTICHOICE": "Multiple Choice",
        "MULTISELECT": "Multi-Select",
        "TRUEFALSE": "True/False",
        "WRITTEN_RESPONSE": "Long Answer",
    }
    itemmetadata = ET.SubElement(item, "itemmetadata")
    add_qti_metadata(
        itemmetadata,
        {
            "qmd_computerscored": "yes",  # D2L emits "yes" for all types, incl. Long Answer (canonical confirmed)
            "qmd_questiontype": qtype_map[question.target_question_type],
            "qmd_weighting": f"{question.points:.9f}",
            "qmd_globalid": stable_uuid(f"{question.question_code}:global"),
            "qmd_displayid": question.question_code,
            "qmd_aihumanorigin": "HumanGenerated",
        },
    )

    itemproc = ET.SubElement(item, "itemproc_extension")
    ET.SubElement(itemproc, ns_attr(D2L_NS, "difficulty")).text = "3"
    ET.SubElement(itemproc, ns_attr(D2L_NS, "isbonus")).text = "no"
    ET.SubElement(itemproc, ns_attr(D2L_NS, "ismandatory")).text = "yes"

    presentation = ET.SubElement(item, "presentation")
    flow = ET.SubElement(presentation, "flow")
    material = ET.SubElement(flow, "material")
    ET.SubElement(material, "mattext", {"texttype": "text/html"}).text = html_fragment(question.question_text, question.question_text_format)

    if question.target_question_type in {"MULTICHOICE", "MULTISELECT", "TRUEFALSE"}:
        add_choice_question(
            flow,
            item,
            question,
            true_false_answer_texttype=true_false_answer_texttype,
        )
    else:
        add_written_response(flow, item, question)


def grouped_questions(questions: list[QuestionRow]) -> dict[str, list[QuestionRow]]:
    grouped: dict[str, list[QuestionRow]] = {}
    for question in questions:
        grouped.setdefault(question.bank_id, []).append(question)
    return grouped


def create_questiondb(
    questions: list[QuestionRow],
    question_library_ident: str,
    *,
    true_false_answer_texttype: str = "text/plain",
) -> ET.ElementTree:
    root = ET.Element("questestinterop")
    objectbank = ET.SubElement(root, "objectbank", {"ident": question_library_ident})
    next_d2l_id = 1

    for bank_index, (bank_id, bank_questions) in enumerate(grouped_questions(questions).items(), start=1):
        section = ET.SubElement(
            objectbank,
            "section",
            {
                ns_attr(D2L_NS, "id"): str(next_d2l_id),
                "ident": f"SECT_{safe_ident(bank_id, 'BANK')}",
                "title": bank_questions[0].bank_title,
            },
        )
        next_d2l_id += 1
        sectionproc = ET.SubElement(section, "sectionproc_extension")
        ET.SubElement(sectionproc, ns_attr(D2L_NS, "display_section_name")).text = "no"
        ET.SubElement(sectionproc, ns_attr(D2L_NS, "display_section_line")).text = "no"
        ET.SubElement(sectionproc, ns_attr(D2L_NS, "type_display_section")).text = "0"
        for question in bank_questions:
            create_item(
                section,
                question,
                bank_index * 1000 + next_d2l_id,
                true_false_answer_texttype=true_false_answer_texttype,
            )
            next_d2l_id += 1

    return ET.ElementTree(root)


def create_quiz(
    questions: list[QuestionRow],
    sections: list[QuizSection],
    quiz_title: str,
    quiz_ident: str,
    resource_code: str,
    condition_set: str,
    grade_resource_code: str | None = None,
    settings: QuizSettings | None = None,
) -> ET.ElementTree:
    settings = settings or QuizSettings()
    root = ET.Element("questestinterop")
    assessment = ET.SubElement(
        root,
        "assessment",
        {
            "title": quiz_title,
            "ident": quiz_ident,
            "condition_set": condition_set,
            ns_attr(D2L_NS, "resource_code"): resource_code,
        },
    )

    assessmentcontrol = ET.SubElement(
        assessment,
        "assessmentcontrol",
        {
            "hide_question_pointsswitch": "no",
            "hintswitch": "no",
            "solutionswitch": "no",
            "feedbackswitch": "no",
        },
    )
    assessmentcontrol.text = ""

    assess_proc = ET.SubElement(assessment, "assess_procextension")
    if grade_resource_code:
        # Joins the quiz to its gradebook item by resource_code. The inner value is
        # the grades_d2l.xml item identifier (build_grades_xml assigns 700001 to the
        # first item); D2L reassigns ids on import, so resource_code is the real join.
        grade_item = ET.SubElement(
            assess_proc,
            "grade_item",
            {ns_attr(D2L_NS, "is_autoexport"): "yes", "resource_code": grade_resource_code},
        )
        grade_item.text = "700001"
    ET.SubElement(assess_proc, "is_active").text = "yes" if settings.is_active else "no"
    ET.SubElement(assess_proc, ns_attr(D2L_NS, "time_limit")).text = str(settings.time_limit)
    ET.SubElement(assess_proc, ns_attr(D2L_NS, "show_clock")).text = "yes" if settings.show_clock else "no"
    ET.SubElement(assess_proc, ns_attr(D2L_NS, "enforce_time_limit")).text = (
        "yes" if settings.enforce_time_limit else "no"
    )
    ET.SubElement(assess_proc, ns_attr(D2L_NS, "attempts_allowed")).text = str(settings.attempts_allowed)
    ET.SubElement(assess_proc, ns_attr(D2L_NS, "mark_calculation_type")).text = "1"
    ET.SubElement(assess_proc, ns_attr(D2L_NS, "is_forward_only")).text = (
        "yes" if settings.is_forward_only else "no"
    )

    assess_feedback = ET.SubElement(assessment, "assessfeedback", {"title": ""})
    rubric = ET.SubElement(assess_feedback, "rubric")
    flow_mat = ET.SubElement(rubric, "flow_mat")
    material = ET.SubElement(flow_mat, "material")
    ET.SubElement(material, "mattext", {"texttype": "text/html"}).text = "<p>Your quiz has been submitted.</p>"

    container = ET.SubElement(assessment, "section", {"ident": "CONTAINER_SECTION"})
    by_bank = grouped_questions(questions)
    for section in sorted(sections, key=lambda item: item.section_order):
        rand_section = ET.SubElement(
            container,
            "section",
            {
                "ident": f"RAND_{safe_ident(str(section.section_order), 'SECTION')}_{safe_ident(section.source_bank_id, 'BANK')}",
                "title": section.section_title,
                ns_attr(D2L_NS, "page"): "1",
            },
        )
        add_qti_metadata(
            rand_section,
            {
                "qmd_numberofitems": str(section.recommended_draw_count),
                "qmd_weighting": f"{section.points_per_question:.9f}",
            },
        )
        sectionproc = ET.SubElement(rand_section, "sectionproc_extension")
        ET.SubElement(sectionproc, ns_attr(D2L_NS, "display_section_name")).text = "no"
        ET.SubElement(sectionproc, ns_attr(D2L_NS, "display_section_line")).text = "no"
        ET.SubElement(sectionproc, ns_attr(D2L_NS, "type_display_section")).text = "0"

        for index, question in enumerate(by_bank[section.source_bank_id], start=1):
            itemref = ET.SubElement(
                rand_section,
                "itemref",
                {
                    "linkrefid": question.label,
                    ns_attr(D2L_NS, "page"): "1",
                },
            )
            ET.SubElement(itemref, ns_attr(D2L_NS, "file"), {"href": "questiondb.xml"})
            ET.SubElement(itemref, ns_attr(D2L_NS, "points")).text = f"{section.points_per_question:.9f}"
            ET.SubElement(itemref, ns_attr(D2L_NS, "difficulty")).text = "3"
            ET.SubElement(itemref, ns_attr(D2L_NS, "isbonus")).text = "no"
            ET.SubElement(itemref, ns_attr(D2L_NS, "ismandatory")).text = "no"
            ET.SubElement(itemref, ns_attr(D2L_NS, "order")).text = str(index)

    return ET.ElementTree(root)


def create_manifest(
    manifest_ident: str,
    quiz_file: str,
    title: str,
    resource_code: str,
    *,
    library_only: bool = False,
    grade_resource_code: str | None = None,
    module_title: str = "",
    asset_paths: list[str] | None = None,
) -> ET.ElementTree:
    manifest = ET.Element(ns_attr(IMS_NS, "manifest"), {"identifier": manifest_ident})
    metadata = ET.SubElement(manifest, ns_attr(IMS_NS, "metadata"))
    lom = ET.SubElement(metadata, ns_attr(IMSMD_NS, "lom"))
    general = ET.SubElement(lom, ns_attr(IMSMD_NS, "general"))
    title_node = ET.SubElement(general, ns_attr(IMSMD_NS, "title"))
    ET.SubElement(title_node, ns_attr(IMSMD_NS, "langstring"), {XML_LANG: "en-us"}).text = title
    ET.SubElement(general, ns_attr(IMSMD_NS, "language")).text = "en-us"

    place_quiz = bool(module_title) and not library_only
    if place_quiz:
        organizations = ET.SubElement(manifest, ns_attr(IMS_NS, "organizations"), {"default": "d2l_orgs"})
        organization = ET.SubElement(organizations, ns_attr(IMS_NS, "organization"), {"identifier": "d2l_org"})
        module_item = ET.SubElement(
            organization,
            ns_attr(IMS_NS, "item"),
            {
                "identifier": "ITEM_MODULE",
                "identifierref": "RES_MODULE",
                ns_attr(D2L_NS, "id"): "2001",
                ns_attr(D2L_NS, "resource_code"): stable_uuid(f"{title}:module"),
                "completion_type": "2",
                ns_attr(D2L_NS, "ai_human_origin"): "0",
            },
        )
        ET.SubElement(module_item, ns_attr(IMS_NS, "title")).text = module_title
        quiz_item = ET.SubElement(
            module_item,
            ns_attr(IMS_NS, "item"),
            {
                "identifier": "ITEM_QUIZ_LINK",
                "identifierref": "RES_QUIZ_LINK",
                ns_attr(D2L_NS, "id"): "2002",
                ns_attr(D2L_NS, "resource_code"): stable_uuid(f"{title}:quizlink"),
                "completion_type": "2",
                ns_attr(D2L_NS, "ai_human_origin"): "0",
                "resource_type_key": "D2L.LE.Quizzing.Quiz",
            },
        )
        ET.SubElement(quiz_item, ns_attr(IMS_NS, "title")).text = title

    resources = ET.SubElement(manifest, ns_attr(IMS_NS, "resources"))

    def add_resource(attrs: dict) -> ET.Element:
        return ET.SubElement(resources, ns_attr(IMS_NS, "resource"), attrs)

    add_resource({
        "identifier": "res_orgunitconfig", "type": "webcontent",
        ns_attr(D2L_NS, "material_type"): "orgunitconfig", ns_attr(D2L_NS, "link_target"): "",
        "href": "orgunitconfig/orgunitconfig.xml", "title": "",
    })
    if grade_resource_code and not library_only:
        add_resource({
            "identifier": "res_grades", "type": "webcontent",
            ns_attr(D2L_NS, "material_type"): "d2lgrades", ns_attr(D2L_NS, "link_target"): "",
            "href": "grades_d2l.xml", "title": "",
        })
    question_resource = add_resource({
        "identifier": "res_question_library", "type": "webcontent",
        ns_attr(D2L_NS, "material_type"): "d2lquestionlibrary", ns_attr(D2L_NS, "link_target"): "",
        "href": "questiondb.xml", "title": "Question Library",
    })
    for asset_path in sorted(asset_paths or []):
        ET.SubElement(question_resource, ns_attr(IMS_NS, "file"), {"href": asset_path})
    if not library_only:
        add_resource({
            "identifier": "res_quiz", "type": "webcontent",
            ns_attr(D2L_NS, "material_type"): "d2lquiz", ns_attr(D2L_NS, "link_target"): "",
            ns_attr(D2L_NS, "resource_code"): resource_code, "href": quiz_file, "title": title,
        })
    if place_quiz:
        add_resource({
            "identifier": "RES_MODULE", "type": "webcontent",
            ns_attr(D2L_NS, "material_type"): "contentmodule", ns_attr(D2L_NS, "link_target"): "",
            "href": "", "title": "",
        })
        add_resource({
            "identifier": "RES_QUIZ_LINK", "type": "webcontent",
            ns_attr(D2L_NS, "material_type"): "contentlink", ns_attr(D2L_NS, "link_target"): "_self",
            "href": f"/d2l/common/dialogs/quickLink/quickLink.d2l?ou={{orgUnitId}}&type=quiz&rCode={resource_code}",
            "title": "",
        })
    return ET.ElementTree(manifest)


def write_xml(path: Path, tree: ET.ElementTree) -> None:
    ET.indent(tree, space="  ")
    tree.write(path, encoding="utf-8", xml_declaration=True)


def write_orgunitconfig(output_dir: Path, orgunitconfig: Path | None) -> str:
    target_dir = output_dir / "orgunitconfig"
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / "orgunitconfig.xml"
    if orgunitconfig:
        shutil.copy2(orgunitconfig, target)
        return "Copied from supplied --orgunitconfig."
    target.write_text('<?xml version="1.0" encoding="UTF-8"?>\n<orgunitconfig />\n', encoding="utf-8")
    return "Generated minimal placeholder; replace with target-shell orgunitconfig.xml before production import."


def write_readme(
    output_dir: Path,
    source_path: Path,
    quiz_title: str,
    quiz_file: str,
    question_count: int,
    section_count: int,
    orgunit_note: str,
) -> None:
    source_kind = "CourseCraft quiz model" if source_path.suffix.lower() == ".json" else "normalized workbook"
    readme = f"""# Quiz Package Build

Generated from {source_kind}:

- `{source_path.name}`

Package files:

- `imsmanifest.xml`
- `questiondb.xml`
- `{quiz_file}`
- `orgunitconfig/orgunitconfig.xml`

Summary:

- quiz title: {quiz_title}
- question count: {question_count}
- quiz section count: {section_count}
- orgunitconfig: {orgunit_note}

Validate before import:

```bash
python3 scripts/validate_quiz_package.py .
```
"""
    (output_dir / "README.md").write_text(readme, encoding="utf-8")


def write_zip(output_dir: Path, zip_path: Path) -> None:
    zip_path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for file_path in sorted(path for path in output_dir.rglob("*") if path.is_file()):
            archive.write(file_path, file_path.relative_to(output_dir).as_posix())


def build_package(args: argparse.Namespace) -> Path:
    started_at = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
    source_path = Path(args.source).expanduser().resolve()
    if not source_path.is_file():
        raise ValueError(f"Authoring source does not exist: {source_path}")
    output_dir = Path(args.output_dir).expanduser().resolve()
    existing_output_files = (
        sorted(item.relative_to(output_dir).as_posix() for item in output_dir.rglob("*") if item.is_file())
        if output_dir.is_dir()
        else []
    )
    if existing_output_files and not args.allow_nonempty_output:
        preview = ", ".join(existing_output_files[:5])
        remainder = f" (+{len(existing_output_files) - 5} more)" if len(existing_output_files) > 5 else ""
        raise ValueError(
            f"Output directory is not empty ({preview}{remainder}). Use a new directory or pass "
            "--allow-nonempty-output and run strict package validation to guard against stale files."
        )
    authoring_projection = None
    assets: list[dict] = []
    if source_path.suffix.lower() == ".json":
        authoring_projection = load_authoring_projection(
            source_path,
            quiz_entity_key=args.quiz_entity_key,
            settings_path=Path(args.settings).expanduser().resolve() if args.settings else None,
            asset_root=[Path(root).expanduser().resolve() for root in args.asset_root] or None,
            promotion_receipt_path=(
                Path(args.promotion_receipt).expanduser().resolve()
                if args.promotion_receipt
                else None
            ),
            trial_authorization_path=(
                Path(args.phase5_candidate_authorization).expanduser().resolve()
                if args.phase5_candidate_authorization
                else None
            ),
        )
        questions = [question_from_projection(row) for row in authoring_projection["questions"]]
        sections = [section_from_projection(row) for row in authoring_projection["sections"]]
        quiz_title = args.quiz_title or authoring_projection["quiz_title"]
        quiz_key = authoring_projection["quiz_key"]
        run_id = authoring_projection["run_id"]
        settings_receipt = authoring_projection["settings_receipt"]
        quiz_settings = QuizSettings(**authoring_projection["quiz_settings"])
        assets = authoring_projection["assets"]
        source_kind = "composite"
        source_lineage_key = authoring_projection["model"]["source"]["source_lineage_key"]
        source_scope = "file_set"
    else:
        if args.settings:
            raise ValueError("--settings requires a coursecraft.quiz/1 JSON source.")
        questions, sections = load_workbook_data(source_path)
        quiz_title = args.quiz_title or source_path.stem.replace("_", " ")
        source_digest = sha256_file(source_path)
        quiz_key = f"cc:quiz:workbook:{source_digest[:24]}"
        run_id = f"cc:run:quiz-build:{hashlib.sha256((source_digest + started_at).encode()).hexdigest()[:24]}"
        settings_receipt = default_settings_receipt(quiz_key, run_id)
        workbook_question_keys = {
            question.question_code: f"cc:question:workbook:{source_digest[:16]}:{safe_ident(question.question_code, 'Q').lower()}"
            for question in questions
        }
        for question in questions:
            if not question.randomize_answers:
                continue
            target_key = workbook_question_keys[question.question_code]
            input_id = f"in.randomize_answers.{safe_ident(question.question_code, 'Q').lower()}.question"
            settings_receipt["inputs"].append(
                {
                    "input_id": input_id,
                    "setting": "randomize_answers",
                    "layer": "question_override",
                    "target_entity_key": target_key,
                    "state": "known",
                    "value": True,
                    "source_reference": f"workbook question_code {question.question_code}",
                    "source_evidence_keys": [],
                    "extensions": {},
                }
            )
            settings_receipt["resolutions"].append(
                {
                    "setting": "randomize_answers",
                    "target_entity_key": target_key,
                    "state": "known",
                    "effective_value": True,
                    "winner_input_id": input_id,
                    "considered_input_ids": [input_id],
                    "defaulted": False,
                    "coerced": False,
                    "coercion_note": None,
                    "source_evidence_keys": [],
                    "extensions": {},
                }
            )
        settings_receipt = materialize_effective_settings_receipt(
            settings_receipt,
            quiz_key=quiz_key,
            run_id=run_id,
            question_keys=set(workbook_question_keys.values()),
        )
        quiz_settings = QuizSettings()
        source_kind = "workbook"
        source_lineage_key = f"cc:lineage:workbook:{source_digest[:24]}"
        source_scope = "workbook"

    quiz_ident = args.quiz_ident or f"QUIZ_{safe_ident(quiz_title, 'QUIZ')}"
    quiz_file = args.quiz_file or f"quiz_d2l_{safe_ident(quiz_title, 'quiz').lower()}.xml"
    resource_code = args.resource_code or stable_uuid(f"{quiz_title}:resource")
    condition_set = args.condition_set or stable_uuid(f"{quiz_title}:condition_set")
    question_library_ident = args.question_library_ident or f"QLIB_{safe_ident(quiz_title, 'QUIZ')}"
    manifest_ident = args.manifest_ident or f"MANIFEST_{safe_ident(quiz_title, 'QUIZ')}"

    if len({question.question_code for question in questions}) != len(questions):
        raise ValueError("Question permanent codes/question_code values must be unique in one build.")
    identifier_issues = target_identifier_issues(
        [(row.question_code, f"question-row:{index}") for index, row in enumerate(questions)],
        [(row.bank_id, row.bank_id) for row in questions],
        [(row.section_order, row.source_bank_id, f"section-row:{index}") for index, row in enumerate(sections)],
    )
    if identifier_issues:
        raise ValueError(identifier_issues[0][1])
    output_dir.mkdir(parents=True, exist_ok=True)

    copied_asset_paths: list[str] = []
    for asset in assets:
        destination = output_dir.joinpath(*asset["archive_path"].parts)
        destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(asset["source"], destination)
        copied_digest = sha256_file(destination)
        if copied_digest != asset["sha256"]:
            raise ValueError(f"Copied asset checksum mismatch: {asset['archive_path']}")
        copied_asset_paths.append(asset["package_path"].as_posix())

    write_xml(
        output_dir / "questiondb.xml",
        create_questiondb(
            questions,
            question_library_ident,
            true_false_answer_texttype=args.true_false_answer_texttype,
        ),
    )

    if args.library_only:
        write_xml(
            output_dir / "imsmanifest.xml",
            create_manifest(
                manifest_ident,
                quiz_file,
                quiz_title,
                resource_code,
                library_only=True,
                asset_paths=copied_asset_paths,
            ),
        )
    else:
        grade_resource_code = None
        if args.grade_item:
            grade_resource_code = stable_uuid(f"{quiz_title}:grade")
            out_of = (
                float(args.grade_out_of)
                if args.grade_out_of
                else sum(s.recommended_draw_count * s.points_per_question for s in sections)
            ) or float(len(questions))
            write_xml(
                output_dir / "grades_d2l.xml",
                ET.ElementTree(build_grades_xml(
                    [{"key": "quiz_grade", "name": quiz_title, "points": out_of}],
                    {"quiz_grade": grade_resource_code},
                )),
            )
        write_xml(
            output_dir / quiz_file,
            create_quiz(
                questions,
                sections,
                quiz_title,
                quiz_ident,
                resource_code,
                condition_set,
                grade_resource_code,
                quiz_settings,
            ),
        )
        write_xml(
            output_dir / "imsmanifest.xml",
            create_manifest(
                manifest_ident, quiz_file, quiz_title, resource_code,
                grade_resource_code=grade_resource_code,
                module_title=args.module,
                asset_paths=copied_asset_paths,
            ),
        )

    orgunit_note = write_orgunitconfig(output_dir, Path(args.orgunitconfig).expanduser().resolve() if args.orgunitconfig else None)
    write_readme(output_dir, source_path, quiz_title, quiz_file, len(questions), len(sections), orgunit_note)

    zip_path = Path(args.zip_output).expanduser().resolve() if args.zip_output else None
    if args.zip_output:
        write_zip(output_dir, zip_path)

    receipt_dir = (
        Path(args.receipt_dir).expanduser().resolve()
        if args.receipt_dir
        else output_dir.parent / f"{output_dir.name}__receipts"
    )
    receipt_dir.mkdir(parents=True, exist_ok=True)
    settings_receipt_path = receipt_dir / "quiz_build.settings.json"
    settings_receipt_path.write_text(json.dumps(settings_receipt, indent=2) + "\n", encoding="utf-8")

    package_artifacts = [
        artifact_row(path, "package")
        for path in sorted(item for item in output_dir.rglob("*") if item.is_file())
    ]
    if zip_path:
        package_artifacts.append(artifact_row(zip_path, "package"))
    package_artifacts.append(artifact_row(settings_receipt_path, "receipt"))
    inputs = [artifact_row(source_path, "input", status="read")]
    if args.settings:
        inputs.append(artifact_row(Path(args.settings).expanduser().resolve(), "receipt", status="read"))
    inputs.append(artifact_row(CAPABILITY_REGISTRY_PATH, "input", status="read"))
    if args.promotion_receipt:
        inputs.append(
            artifact_row(
                Path(args.promotion_receipt).expanduser().resolve(),
                "receipt",
                status="read",
            )
        )
    if args.phase5_candidate_authorization:
        authorization_input = artifact_row(
            Path(args.phase5_candidate_authorization).expanduser().resolve(),
            "other",
            status="read",
        )
        authorization_input["extensions"]["coursecraft.artifact_kind"] = (
            "phase5_candidate_authorization"
        )
        inputs.append(authorization_input)
    for asset_source in sorted({asset["source"] for asset in assets}):
        inputs.append(artifact_row(asset_source, "input", status="read"))

    capability_names = sorted({f"question_type:{question.target_question_type}" for question in questions})
    capability_names.extend(["question_library_banks", "itemref_joins"])
    if not args.library_only:
        capability_names.append("random_draw_sections")
    if args.grade_item:
        capability_names.append("grade_item_join")
    if args.module:
        capability_names.append("module_placement")
    if args.library_only:
        capability_names.append("library_only_package")
    if assets:
        capability_names.append("resolved_question_assets_fixture_validated")
    capability_names.append("settings_resolution_projection")
    run_receipt = build_run_receipt(
        run_id=run_id,
        started_at=started_at,
        source_path=source_path,
        source_kind=source_kind,
        source_scope=source_scope,
        inputs=inputs,
        artifacts=package_artifacts,
        capability_names=capability_names,
        settings_receipt_path=str(settings_receipt_path),
        source_lineage_key=source_lineage_key,
    )
    if authoring_projection is not None and authoring_projection.get("promotion_receipt"):
        promotion_receipt = authoring_projection["promotion_receipt"]
        run_receipt["extensions"].update(
            {
                "coursecraft.authoring_route": "strict_model_projection",
                "coursecraft.promotion_token": promotion_receipt["promotion_token"],
                "coursecraft.source_model_sha256": promotion_receipt[
                    "source_model_sha256"
                ],
                "coursecraft.decision_overlay_sha256": promotion_receipt[
                    "decision_overlay_sha256"
                ],
                "coursecraft.promotion_receipt_sha256": sha256_file(
                    Path(args.promotion_receipt).expanduser().resolve()
                ),
            }
        )
    run_receipt["extensions"]["coursecraft.true_false_answer_texttype"] = (
        args.true_false_answer_texttype
    )
    if authoring_projection is not None and authoring_projection.get("route_derivations"):
        run_receipt["extensions"]["coursecraft.source_observed_route_derivations"] = (
            authoring_projection["route_derivations"]
        )
    if authoring_projection is not None and authoring_projection.get(
        "phase5_candidate_authorization"
    ):
        authorization = authoring_projection["phase5_candidate_authorization"]
        run_receipt["extensions"].update(
            {
                "coursecraft.evidence_status": "phase5_candidate_not_evidence",
                "coursecraft.phase5_candidate_authorization_fingerprint": authorization[
                    "authorization_fingerprint"
                ],
                "coursecraft.phase5_candidate_authorization_sha256": sha256_file(
                    Path(args.phase5_candidate_authorization).expanduser().resolve()
                ),
            }
        )
    run_receipt_path = receipt_dir / "quiz_build.run.json"
    run_receipt_path.write_text(json.dumps(run_receipt, indent=2) + "\n", encoding="utf-8")

    print(f"Built quiz package in {output_dir}")
    if args.zip_output:
        print(f"Wrote import zip {zip_path}")
    print(f"Wrote settings receipt {settings_receipt_path}")
    print(f"Wrote run receipt {run_receipt_path}")
    return output_dir


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build a Brightspace quiz package from a normalized workbook or reviewed coursecraft.quiz/1 model."
    )
    parser.add_argument("source", help="Path to a normalized workbook or reviewed coursecraft.quiz/1 JSON model.")
    parser.add_argument("--output-dir", default="workspace/generated/import-ready-components/quizzes/quiz_package")
    parser.add_argument("--quiz-title", default="")
    parser.add_argument("--quiz-ident", default="")
    parser.add_argument("--quiz-file", default="")
    parser.add_argument("--question-library-ident", default="")
    parser.add_argument("--manifest-ident", default="")
    parser.add_argument("--resource-code", default="")
    parser.add_argument("--condition-set", default="")
    parser.add_argument("--orgunitconfig", default="", help="Optional target-shell orgunitconfig.xml to copy into the package.")
    parser.add_argument("--zip-output", default="", help="Optional zip path to write after generating the package folder.")
    parser.add_argument("--grade-item", action="store_true", help="Emit a grade item for the quiz and wire it to the gradebook (assess_procextension link + grades_d2l.xml).")
    parser.add_argument("--grade-out-of", default="", help="Total points for the quiz grade item (default: sum of draw_count x points_per_question across sections).")
    parser.add_argument("--module", default="", help="Place the quiz under a module of this title (adds a manifest organization + quiz quicklink).")
    parser.add_argument("--library-only", action="store_true", help="Emit only the question library (questiondb + manifest); no quiz, grade item, or placement.")
    parser.add_argument("--quiz-entity-key", default="", help="Quiz entity key to build when a model contains more than one quiz.")
    parser.add_argument("--settings", default="", help="Optional validated coursecraft.quiz_settings/1 receipt (model input only).")
    parser.add_argument("--asset-root", action="append", default=[], help="Root for relative model asset source_path values; repeat to supply multiple folders (default: model directory).")
    parser.add_argument("--promotion-receipt", default="", help="Verified Quiz Binder promotion receipt to chain into a strict-model build.")
    parser.add_argument("--phase5-candidate-authorization", default="", help="Exact local-only authorization for an extraction_only Phase 5 candidate build; requires --promotion-receipt.")
    parser.add_argument("--receipt-dir", default="", help="Receipt output directory (default: sibling <package>__receipts; excluded from import ZIP).")
    parser.add_argument(
        "--true-false-answer-texttype",
        choices=("text/html", "text/plain"),
        default="text/plain",
        help=(
            "QTI mattext type for True/False answer labels. text/plain is the "
            "tenant-verified default; text/html is retained only for regression fixtures."
        ),
    )
    parser.add_argument(
        "--allow-nonempty-output",
        action="store_true",
        help="Allow writing into a nonempty folder. Prefer a new folder; strict validation is required to detect stale files.",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv or sys.argv[1:])
    try:
        build_package(args)
    except Exception as exc:  # noqa: BLE001
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
