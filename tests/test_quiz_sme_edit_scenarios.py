"""Reviewer actions against either this checkout or QUIZ_TOOLS_REPO.

Set QUIZ_TOOLS_REPO to exercise the public snapshot with the same synthetic
inputs. These tests use no real course material.
"""
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
import xml.etree.ElementTree as ET

import pytest
from openpyxl import load_workbook

TEST_ROOT = Path(__file__).resolve().parents[1]
TOOL_ROOT = Path(os.environ.get('QUIZ_TOOLS_REPO', TEST_ROOT)).resolve()
sys.path.insert(0, str(TOOL_ROOT / 'scripts'))
from quiz_assessment_review import prepare_review_packet, validate_packet
from quiz_review_workbook_reingest import ReingestError, materialize_workbook_decisions
from quiz_promote_revisions import promote_revisions


@pytest.fixture(scope='module')
def template(tmp_path_factory):
    root = tmp_path_factory.mktemp('sme-edit-scenarios')
    source = root / 'source'
    shutil.copytree(TEST_ROOT / 'tests/fixtures/quiz_source_choice_fidelity', source)
    quiz_path = source / 'quiz_d2l_choices.xml'
    tree = ET.parse(quiz_path)
    extension = ET.SubElement(tree.find('.//assessment'), 'assess_procextension')
    ET.SubElement(extension, '{http://desire2learn.com/xsd/d2lcp_v2p0}time_limit').text = '20'
    tree.write(quiz_path, encoding='utf-8', xml_declaration=True)
    extraction = root / 'extraction'
    result = subprocess.run([sys.executable, str(TOOL_ROOT / 'scripts/extract_quiz_pool_review.py'),
        str(source), '--asset-mode', 'copy', '--output-dir', str(extraction)], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    packet = root / 'packet'
    prepare_review_packet(next(extraction.glob('*reviewer.xlsx')), next(extraction.glob('*.model.json')), packet)
    return packet


@pytest.fixture
def packet(template, tmp_path):
    result = tmp_path / 'packet'
    shutil.copytree(template, result)
    return result


def row_for(packet, title='DUPLICATE'):
    spec = validate_packet(packet)
    model = json.loads((packet / 'model.json').read_text())
    key = next(q['entity_key'] for q in model['questions'] if q['title'] == title)
    book = load_workbook(packet / 'reviewer_working.xlsx')
    for tab in spec['sheets']:
        sheet = book[tab['title']]
        columns = {c.value: c.column for c in sheet[2]}
        for row in tab['rows']:
            if sheet.cell(row['row'], columns['question_entity_key']).value == key:
                return book, sheet, columns, row['row'], key
    raise AssertionError('Missing synthetic question')


def collect(packet):
    return materialize_workbook_decisions(packet / 'model.json', packet / 'reviewer_baseline_DO_NOT_EDIT.xlsx',
                                         packet / 'reviewer_working.xlsx', metadata_policy='optional')


def promote(packet, overlay):
    path = packet / 'decisions.json'
    path.write_text(json.dumps(overlay))
    return promote_revisions(packet / 'model.json', path,
        TOOL_ROOT / 'workspace/reference/schemas/quiz/quiz_build_capabilities.json')


@pytest.mark.parametrize('field,value', [
    pytest.param('revised_question_text', 'A clearly revised SME prompt.',
                 marks=pytest.mark.xfail(strict=True, reason='This snapshot excludes accepted prompt revisions.')),
    pytest.param('revised_response_options', '["First revised option", "Second revised option", "Third revised option"]',
                 marks=pytest.mark.xfail(strict=True, reason='This snapshot excludes accepted response-option revisions.')),
    pytest.param('revised_answer_key', '{"correct_option_keys": ["B"]}',
                 marks=pytest.mark.xfail(strict=True, reason='This snapshot excludes accepted answer-key revisions.')),
])
def test_accepted_plain_question_edits_reach_promoted_content(packet, field, value):
    book, sheet, cols, row, key = row_for(packet)
    sheet.cell(row, cols[field], value)
    sheet.cell(row, cols['approval_status'], 'accepted')
    book.save(packet / 'reviewer_working.xlsx')
    output, receipt = promote(packet, collect(packet))
    assert receipt['summary']['applied_change_count'] == 1, receipt['excluded']
    question = next(q for q in output['questions'] if q['entity_key'] == key)
    assert question['build_support']['level'] == 'extraction_only'
    if field == 'revised_question_text':
        assert question['prompt']['content'] == value
    elif field == 'revised_response_options':
        assert [o['content']['content'] for o in question['type_payload']['options']] == json.loads(value)
    else:
        assert [o['option_key'] for o in question['type_payload']['options'] if o['correct']] == ['B']


@pytest.mark.parametrize('status', ['open', 'rejected'])
def test_unaccepted_edits_do_not_change_content(packet, status):
    book, sheet, cols, row, key = row_for(packet)
    sheet.cell(row, cols['revised_question_text'], 'DO NOT APPLY')
    sheet.cell(row, cols['approval_status'], status)
    book.save(packet / 'reviewer_working.xlsx')
    output, receipt = promote(packet, collect(packet))
    assert receipt['summary']['applied_change_count'] == 0
    question = next(q for q in output['questions'] if q['entity_key'] == key)
    assert 'DO NOT APPLY' not in question['prompt']['content']


def test_replacing_original_image_file_is_rejected(packet):
    spec = validate_packet(packet)
    (packet / spec['assets'][0]['path']).write_text('<svg xmlns="http://www.w3.org/2000/svg"><text>Replacement</text></svg>')
    with pytest.raises(ReingestError, match='missing or changed'):
        collect(packet)


def test_linking_a_new_image_in_source_image_cell_is_rejected(packet):
    book, sheet, cols, row, _ = row_for(packet, 'IMAGES')
    new_image = packet / 'Images/replacement.svg'
    new_image.write_text('<svg xmlns="http://www.w3.org/2000/svg"><text>Replacement</text></svg>')
    cell = sheet.cell(row, cols['image_link'])
    cell.value = 'New image'
    cell.hyperlink = 'Images/replacement.svg'
    book.save(packet / 'reviewer_working.xlsx')
    with pytest.raises(ReingestError):
        collect(packet)


def test_image_replacement_note_records_request_without_replacing_content(packet):
    book, sheet, cols, row, key = row_for(packet, 'IMAGES')
    note = f'New image uploaded to Images/replacement.svg; requested for question in row {row}.'
    (packet / 'Images/replacement.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg"><text>Replacement</text></svg>')
    sheet.cell(row, cols['reviewer_note'], note)
    book.save(packet / 'reviewer_working.xlsx')
    overlay = collect(packet)
    assert overlay['content_minimized_summary']['annotation_count'] == 1
    assert note in json.dumps(overlay)
    output, receipt = promote(packet, overlay)
    original = json.loads((packet / 'model.json').read_text())
    assert output['assets'] == original['assets']
    assert next(q for q in output['questions'] if q['entity_key'] == key)['type_payload'] == next(q for q in original['questions'] if q['entity_key'] == key)['type_payload']
    assert receipt['summary']['applied_change_count'] == 0


@pytest.mark.parametrize('kind', ['question', 'image_instruction'])
def test_new_rows_in_extracted_assessment_are_rejected(packet, kind):
    book, sheet, cols, row, _ = row_for(packet)
    target = sheet.max_row + 1
    if kind == 'question':
        sheet.cell(target, cols['revised_question_text'], 'A newly added question')
        sheet.cell(target, cols['approval_status'], 'accepted')
    else:
        sheet.cell(target, cols['reviewer_note'], 'New image uploaded to Images/replacement.svg')
    book.save(packet / 'reviewer_working.xlsx')
    with pytest.raises(ReingestError):
        collect(packet)


@pytest.mark.parametrize('field,value', [('source_points', 9), ('source_feedback', 'Changed feedback'), ('answer_key', 'B')])
def test_source_fields_are_protected(packet, field, value):
    book, sheet, cols, row, _ = row_for(packet)
    sheet.cell(row, cols[field], value)
    book.save(packet / 'reviewer_working.xlsx')
    with pytest.raises(ReingestError):
        collect(packet)


def test_invalid_option_count_is_reported_instead_of_applied(packet):
    book, sheet, cols, row, _ = row_for(packet)
    sheet.cell(row, cols['revised_response_options'], '["Only one option"]')
    sheet.cell(row, cols['approval_status'], 'accepted')
    book.save(packet / 'reviewer_working.xlsx')
    _, receipt = promote(packet, collect(packet))
    assert receipt['summary']['applied_change_count'] == 0
    assert receipt['summary']['excluded_change_count'] == 1


@pytest.mark.parametrize('status', ['open', 'accepted', 'rejected'])
def test_quiz_setting_proposal_uses_explicit_acceptance(packet, status):
    book = load_workbook(packet / 'reviewer_working.xlsx')
    sheet = book['Quiz Settings']
    cols = {c.value: c.column for c in sheet[1]}
    row = next(r for r in range(2, sheet.max_row + 1)
               if sheet.cell(r, cols['setting']).value == 'time_limit')
    sheet.cell(row, cols['proposed_value'], '45')
    sheet.cell(row, cols['approval_status'], status)
    book.save(packet / 'reviewer_working.xlsx')
    overlay = collect(packet)
    assert len(overlay['settings_decisions']) == 1
    assert len(overlay['settings_inputs']) == (status == 'accepted')
    if status == 'accepted':
        assert overlay['settings_inputs'][0]['value'] == 45


def test_new_question_rows_in_blank_drafting_template_are_ingested(tmp_path):
    from quiz_draft_intake import ingest, read_workbook

    source = TOOL_ROOT / 'workspace/reference/examples/quiz_draft/blank.xlsx'
    before = source.read_bytes()
    book = load_workbook(source)
    config = {'lineage_code': 'sme-synthetic-additions', 'quiz_code': 'new-rows', 'quiz_title': 'Synthetic new question test'}
    for row in book['Config']:
        if row[0].value in config:
            row[1].value = config[row[0].value]
    for code, prompt in [('NEW-Q1', 'First newly drafted question.'), ('NEW-Q2', 'Second newly drafted question.')]:
        book['Questions'].append([code, 'Multiple Choice', prompt, 2, 'NEW-P1', None, None, 'plain_text'])
        book['Responses'].append([code, 'A', 'option', 'Correct', 'plain_text', True])
        book['Responses'].append([code, 'B', 'option', 'Incorrect', 'plain_text', False])
    book['Pools'].append(['NEW-P1', 'New questions'])
    book['Draws'].append(['NEW-D1', 'NEW-P1', 'Two new questions', 2, 2])
    path = tmp_path / 'new-questions.xlsx'
    book.save(path)
    model, report = ingest(read_workbook(path), source_name=path.name, input_format='xlsx')
    assert report['valid_intake'] and not report['ready_to_build']
    assert len(model['questions']) == 2
    assert {q['prompt']['content'] for q in model['questions']} == {'First newly drafted question.', 'Second newly drafted question.'}
    assert {q['identity']['permanent_code'] for q in model['questions']} == {'NEW-Q1', 'NEW-Q2'}
    assert source.read_bytes() == before
