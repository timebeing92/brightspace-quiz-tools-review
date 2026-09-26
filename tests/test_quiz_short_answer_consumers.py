from __future__ import annotations

from copy import deepcopy
import json
from pathlib import Path
import shutil
import subprocess
import sys

from openpyxl import load_workbook
import pytest


REPO_ROOT = Path(__file__).resolve().parents[1]
SCRIPT_ROOT = REPO_ROOT / "scripts"
FIXTURE_ROOT = REPO_ROOT / "tests" / "fixtures"
SHORT_ANSWER_FIXTURE = FIXTURE_ROOT / "quiz_xml" / "short_answer_projection"
AUTHORING_FIXTURE = FIXTURE_ROOT / "quiz_authoring"
REGISTRY = (
    REPO_ROOT
    / "workspace"
    / "reference"
    / "schemas"
    / "quiz"
    / "quiz_build_capabilities.json"
)
sys.path.insert(0, str(SCRIPT_ROOT))

from extract_quiz_pool_review import (  # noqa: E402
    build_payload,
    build_reviewer_question_rows,
    write_reviewer_workbook,
)
from quiz_authoring_readiness import analyze_authoring_readiness  # noqa: E402
from quiz_build_support import load_authoring_projection  # noqa: E402
from quiz_normalization import (  # noqa: E402
    build_normalized_model,
    enrich_payload,
)
from quiz_promote_revisions import promote_revisions  # noqa: E402
from quiz_review_workbook_reingest import (  # noqa: E402
    materialize_workbook_decisions,
)


EXPECTED_VALUES = [
    "[ALPHA]",
    "Crème: β",
    "ratio/term",
    "A & B",
    "left / right?!",
]
EXPECTED_REVIEWER_HEADERS = [
    "quiz",
    "question_number",
    "quiz_section (blank = quiz root)",
    "occurrence_key",
    "question_entity_key",
    "question_occurrence_count",
    "placement",
    "relationship_status",
    "selection_mode",
    "draw_count",
    "candidate_pool_size",
    "pool_context_basis",
    "question_type",
    "proposed_permanent_code",
    "proposed_scoring_mode",
    "question_text",
    "revised_question_text",
    "image_link",
    "response_options",
    "revised_response_options",
    "answer_key",
    "revised_answer_key",
    "question_library_location",
    "revised_question_library_location",
    "library_link_status",
    "image_count",
    "note",
    "revision_reason",
    "proposed_by",
    "proposed_at",
    "approval_status",
    "approved_by",
    "approved_at",
    "reviewer_note",
]


@pytest.fixture
def live_short_answer_projection() -> tuple[
    dict[str, object],
    dict[str, object],
    dict[str, object],
]:
    payload = build_payload(SHORT_ANSWER_FIXTURE)
    enrich_payload(payload, SHORT_ANSWER_FIXTURE)
    model, _source_meta = build_normalized_model(
        payload,
        SHORT_ANSWER_FIXTURE,
        SHORT_ANSWER_FIXTURE,
        "directory",
        source_lineage_key="cc:lineage:fixture:short-answer-consumers",
        run_id="cc:run:fixture:short-answer-consumers",
    )
    question = next(
        row
        for row in model["questions"]
        if [
            answer["value"]
            for answer in row["type_payload"]["accepted_responses"]
        ]
        == EXPECTED_VALUES
    )
    assert question["scoring"]["mode"] == "exact"
    assert question["type_payload"]["blanks"] == []
    return payload, model, question


def write_model(path: Path, model: dict[str, object]) -> None:
    path.write_text(
        json.dumps(model, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def write_review_pair(
    tmp_path: Path,
    payload: dict[str, object],
    model: dict[str, object],
) -> tuple[Path, Path]:
    baseline = tmp_path / "baseline.xlsx"
    edited = tmp_path / "edited.xlsx"
    write_reviewer_workbook(
        baseline,
        payload,
        SHORT_ANSWER_FIXTURE,
        "2026-07-24 00:00:00 UTC",
        SCRIPT_ROOT / "extract_quiz_pool_review.py",
        model,
    )
    shutil.copyfile(baseline, edited)
    return baseline, edited


def accept_answer_revision(workbook_path: Path, question_number: int = 1) -> str:
    revised_value = json.dumps(
        ["[REVISED-ANSWER]"], ensure_ascii=False, separators=(",", ":")
    )
    workbook = load_workbook(workbook_path)
    sheet = workbook["Quiz Questions"]
    headers = [cell.value for cell in sheet[1]]
    row_number = next(
        row
        for row in range(2, sheet.max_row + 1)
        if sheet.cell(row=row, column=headers.index("question_number") + 1).value
        == question_number
    )
    values = {
        "revised_answer_key": revised_value,
        "revision_reason": "Exercise the existing Short Answer review path.",
        "proposed_by": "Synthetic Reviewer",
        "proposed_at": "2026-07-24T10:00:00-04:00",
        "approval_status": "accepted",
        "approved_by": "Synthetic Approver",
        "approved_at": "2026-07-24T10:05:00-04:00",
    }
    for header, value in values.items():
        sheet.cell(
            row=row_number,
            column=headers.index(header) + 1,
            value=value,
        )
    workbook.save(workbook_path)
    return revised_value


def materialize_answer_revision(
    tmp_path: Path,
    payload: dict[str, object],
    model: dict[str, object],
) -> tuple[Path, Path, dict[str, object], str]:
    model_path = tmp_path / "short-answer.model.json"
    write_model(model_path, model)
    baseline, edited = write_review_pair(tmp_path, payload, model)
    revised_value = accept_answer_revision(edited)
    overlay = materialize_workbook_decisions(model_path, baseline, edited)
    overlay_path = tmp_path / "decisions.json"
    overlay_path.write_text(
        json.dumps(overlay, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return model_path, overlay_path, overlay, revised_value


def authoring_model_with_live_short_answer(
    live_question: dict[str, object],
) -> tuple[dict[str, object], str]:
    model = json.loads(
        (AUTHORING_FIXTURE / "buildable_quiz.model.json").read_text(
            encoding="utf-8"
        )
    )
    target = model["questions"][0]
    target_key = target["entity_key"]
    target["kind"] = "short_answer"
    target["source_kind"] = "Short Answer"
    target["prompt"] = deepcopy(live_question["prompt"])
    target["type_payload"] = deepcopy(live_question["type_payload"])
    target["type_payload"]["raw_response_models"] = []
    target["scoring"] = deepcopy(live_question["scoring"])
    target["feedback"] = deepcopy(live_question["feedback"])
    target["build_support"] = deepcopy(live_question["build_support"])
    return model, target_key


def test_enriched_multi_synonym_reviewer_output_uses_pipe_values_not_blank_artifact(
    tmp_path: Path,
    live_short_answer_projection: tuple[
        dict[str, object],
        dict[str, object],
        dict[str, object],
    ],
) -> None:
    payload, model, _question = live_short_answer_projection
    expected_answer = " || ".join(EXPECTED_VALUES)
    blank_artifact = " / ".join(EXPECTED_VALUES)

    reviewer_rows = build_reviewer_question_rows(payload)
    row = next(item for item in reviewer_rows if item["question_number"] == 1)
    assert row["answer_key"] == expected_answer
    assert row["answer_key"] != blank_artifact

    workbook_path = tmp_path / "reviewer.xlsx"
    write_reviewer_workbook(
        workbook_path,
        payload,
        SHORT_ANSWER_FIXTURE,
        "2026-07-24 00:00:00 UTC",
        SCRIPT_ROOT / "extract_quiz_pool_review.py",
        model,
    )
    workbook = load_workbook(workbook_path, read_only=True)
    sheet = workbook["Quiz Questions"]
    headers = [cell.value for cell in sheet[1]]
    assert headers == EXPECTED_REVIEWER_HEADERS
    assert {
        "case_sensitive",
        "weight",
        "profile",
        "source_response_facts",
    }.isdisjoint(headers)
    answer_column = headers.index("answer_key") + 1
    question_number_column = headers.index("question_number") + 1
    workbook_row = next(
        row_number
        for row_number in range(2, sheet.max_row + 1)
        if sheet.cell(row=row_number, column=question_number_column).value == 1
    )
    assert sheet.cell(row=workbook_row, column=answer_column).value == expected_answer


def test_enriched_short_answer_reingest_uses_stable_occurrence_join_and_keeps_annotation_shape(
    tmp_path: Path,
    live_short_answer_projection: tuple[
        dict[str, object],
        dict[str, object],
        dict[str, object],
    ],
) -> None:
    payload, model, question = live_short_answer_projection
    model_path, _overlay_path, overlay, revised_value = materialize_answer_revision(
        tmp_path,
        payload,
        model,
    )
    source_bytes = model_path.read_bytes()
    assert json.loads(model_path.read_text(encoding="utf-8")) == model
    assert "coursecraft.binder.reviewer_row" not in question["extensions"]
    native_records = [
        row
        for row in question["type_payload"]["raw_response_models"]
        if row["source_kind"] == "d2l_qti_response_facts/0"
    ]
    assert native_records
    assert all(
        "quiz_title" not in row["payload"]
        and "quiz_item_number" not in row["payload"]
        for row in native_records
    )

    approvals = [
        row
        for row in overlay["annotations"]
        if row["kind"] == "approved_change"
    ]
    assert len(approvals) == 1
    approval = approvals[0]
    assert approval["target_entity_key"] == question["entity_key"]
    assert approval["field_path"] == "/type_payload/answer_key"
    assert approval["value"] == revised_value
    proposal = next(
        row
        for row in overlay["annotations"]
        if row["kind"] == "proposed_revision"
    )
    assert proposal["extensions"]["coursecraft.binder.source_column"] == (
        "revised_answer_key"
    )
    assert approval["extensions"]["coursecraft.binder.proposal_id"] == proposal[
        "annotation_id"
    ]
    assert overlay["row_diffs"][0]["row_key_headers"] == ["occurrence_key"]
    assert len(overlay["row_diffs"][0]["row_key"]) == 1
    assert overlay["row_diffs"][0]["row_key"][0].startswith("cc:occurrence:")
    serialized = json.dumps(overlay, ensure_ascii=False, sort_keys=True)
    assert "case_sensitive" not in serialized
    assert '"weight"' not in serialized
    assert model_path.read_bytes() == source_bytes


def test_promotion_excludes_enriched_short_answer_answer_edit_without_metadata_loss(
    tmp_path: Path,
    live_short_answer_projection: tuple[
        dict[str, object],
        dict[str, object],
        dict[str, object],
    ],
) -> None:
    payload, model, question = live_short_answer_projection
    original_question = deepcopy(question)
    model_path, overlay_path, _overlay, _revised_value = (
        materialize_answer_revision(tmp_path, payload, model)
    )
    source_bytes = model_path.read_bytes()

    promoted, receipt = promote_revisions(model_path, overlay_path, REGISTRY)

    promoted_question = next(
        row for row in promoted["questions"] if row["entity_key"] == question["entity_key"]
    )
    assert promoted_question == original_question
    assert [
        (answer["value"], answer["case_sensitive"], answer["weight"])
        for answer in promoted_question["type_payload"]["accepted_responses"]
    ] == [
        (value, False, "100.000000000")
        for value in EXPECTED_VALUES
    ]
    assert promoted_question["scoring"]["mode"] == "exact"
    assert promoted_question["type_payload"]["blanks"] == []
    assert receipt["summary"] == {
        "accepted_change_count": 1,
        "applied_change_count": 0,
        "excluded_change_count": 1,
    }
    assert receipt["excluded"][0]["reason"] == "question_kind_extraction_only"
    assert model_path.read_bytes() == source_bytes


def test_readiness_keeps_enriched_short_answer_extraction_only_and_blocked(
    tmp_path: Path,
    live_short_answer_projection: tuple[
        dict[str, object],
        dict[str, object],
        dict[str, object],
    ],
) -> None:
    _payload, _model, question = live_short_answer_projection
    authoring_model, target_key = authoring_model_with_live_short_answer(question)
    model_path = tmp_path / "enriched-short-answer.model.json"
    write_model(model_path, authoring_model)
    source_bytes = model_path.read_bytes()

    report = analyze_authoring_readiness(
        model_path,
        settings_path=AUTHORING_FIXTURE / "buildable_quiz.settings.json",
        asset_root=AUTHORING_FIXTURE,
    )

    target_codes = {
        issue["code"]
        for issue in report["issues"]
        if issue["entity_key"] == target_key
    }
    assert report["ready"] is False
    assert target_codes == {
        "question_kind_not_buildable",
        "question_not_approved_for_build",
    }
    assert report["question_capabilities"] == {
        "extraction_only": 1,
        "roundtrip_verified": 3,
    }
    assert report["summary"]["error_count"] == 2
    assert report["summary"]["warning_count"] == 0
    assert model_path.read_bytes() == source_bytes


def test_authoring_projection_and_builder_reject_enriched_short_answer(
    tmp_path: Path,
    live_short_answer_projection: tuple[
        dict[str, object],
        dict[str, object],
        dict[str, object],
    ],
) -> None:
    _payload, _model, question = live_short_answer_projection
    authoring_model, _target_key = authoring_model_with_live_short_answer(question)
    model_path = tmp_path / "enriched-short-answer.model.json"
    write_model(model_path, authoring_model)
    expected_error = (
        "kind 'short_answer' is extraction_only, not build-supported"
    )

    with pytest.raises(ValueError, match=expected_error):
        load_authoring_projection(
            model_path,
            settings_path=AUTHORING_FIXTURE / "buildable_quiz.settings.json",
            asset_root=AUTHORING_FIXTURE,
        )

    output_dir = tmp_path / "package"
    result = subprocess.run(
        [
            sys.executable,
            "scripts/build_quiz_package_from_workbook.py",
            str(model_path),
            "--settings",
            str(AUTHORING_FIXTURE / "buildable_quiz.settings.json"),
            "--asset-root",
            str(AUTHORING_FIXTURE),
            "--output-dir",
            str(output_dir),
        ],
        cwd=REPO_ROOT,
        text=True,
        capture_output=True,
        check=False,
    )
    assert result.returncode == 2
    assert expected_error in result.stderr
    assert not output_dir.exists()
