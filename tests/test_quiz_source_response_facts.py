from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import sys
import xml.etree.ElementTree as ET

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))
sys.path.insert(0, str(REPO_ROOT / "scripts"))

from scripts.extract_quiz_pool_review import (  # noqa: E402
    build_payload,
    extract_source_response_facts,
    write_workbook,
)
from scripts.quiz_normalization import build_normalized_model  # noqa: E402


SPECIMEN_ROOT = (
    REPO_ROOT
    / "workspace"
    / "review"
    / "quiz_capability_lab_r1"
    / "fixtures"
    / "specimens"
)
OCCURRENCE_FIXTURE = (
    REPO_ROOT
    / "tests"
    / "fixtures"
    / "quiz_xml"
    / "occurrence_scoped_source_facts"
)
D2L_NS = "http://desire2learn.com/xsd/d2lcp_v2p0"


def specimen_item(name: str) -> ET.Element:
    return ET.parse(SPECIMEN_ROOT / name).getroot()


def specimen_facts(name: str) -> dict[str, object]:
    return extract_source_response_facts(specimen_item(name))


def attribute_value(
    attributes: list[dict[str, str | None]],
    name: str,
    *,
    namespace_uri: str | None = None,
) -> str | None:
    for attribute in attributes:
        if (
            attribute["name"] == name
            and attribute["namespace_uri"] == namespace_uri
        ):
            return attribute["raw_value"]
    return None


def iter_fact_nodes(node: dict[str, object]):
    yield node
    for child in node["children"]:
        yield from iter_fact_nodes(child)


def iter_scalar_values(value: object):
    if isinstance(value, dict):
        for child in value.values():
            yield from iter_scalar_values(child)
    elif isinstance(value, list):
        for child in value:
            yield from iter_scalar_values(child)
    else:
        yield value


def condition_pair(condition: dict[str, object]) -> tuple[str, str] | None:
    for predicate in condition["predicate_trees"]:
        for node in iter_fact_nodes(predicate):
            if node["name"] != "varequal":
                continue
            respident = attribute_value(node["attributes"], "respident")
            value = str(node["raw_text"] or "").strip()
            if respident and value:
                return respident, value
    return None


def setvar_value(
    condition: dict[str, object],
    *,
    varname: str | None = None,
    action: str | None = None,
) -> str | None:
    for setvar in condition["setvars"]:
        attributes = setvar["attributes"]
        if varname is not None and attribute_value(attributes, "varname") != varname:
            continue
        if action is not None and attribute_value(attributes, "action") != action:
            continue
        return str(setvar["raw_value"] or "").strip()
    return None


def test_sanitized_native_families_emit_versioned_portable_fact_bundles() -> None:
    expectations = {
        "sa_short_answer.xml": ("Short Answer", 1),
        "msa_multi_short_answer.xml": ("Multi-Short Answer", 1),
        "fib_fill_in_blanks.xml": ("Fill in the Blanks", 2),
        "matching_bijective_repeated_group.xml": ("Matching", 3),
        "ordering_sequence.xml": ("Ordering", 1),
    }

    for filename, (question_type, declaration_count) in expectations.items():
        facts = specimen_facts(filename)
        assert facts["schema"] == "coursecraft.quiz_source_response_facts/0"
        assert facts["source_question_type"] == question_type
        assert len(facts["response_declarations"]) == declaration_count
        assert facts["diagnostics"] == []
        assert not any(
            isinstance(value, str)
            and (
                value.startswith(("/Users/", "/private/", "/tmp/"))
                or (len(value) >= 3 and value[1:3] in {":\\", ":/"})
            )
            for value in iter_scalar_values(facts)
        )


def test_text_response_facts_keep_msa_boxes_and_fib_sequence_independent() -> None:
    msa = specimen_facts("msa_multi_short_answer.xml")
    msa_declaration = msa["response_declarations"][0]
    renderer_attributes = msa_declaration["renderer"]["attributes"]

    assert len(msa["response_declarations"]) == 1
    assert len(msa_declaration["response_labels"]) == 1
    assert attribute_value(
        renderer_attributes,
        "input_boxes",
        namespace_uri=D2L_NS,
    ) == "3"
    assert len(msa["response_conditions"]) == 4
    assert len(msa["outcomes"]) == 0
    assert msa["grading_type"] == {
        "state": "absent",
        "raw_value": None,
        "occurrences": [],
    }

    fib = specimen_facts("fib_fill_in_blanks.xml")
    assert [token["kind"] for token in fib["presentation_sequence"]] == [
        "material",
        "response",
        "material",
        "response",
        "material",
    ]
    response_ids = [
        declaration["response_labels"][0]["source_ident"]
        for declaration in fib["response_declarations"]
    ]
    accepted_by_response = {response_id: [] for response_id in response_ids}
    for condition in fib["response_conditions"]:
        pair = condition_pair(condition)
        if pair and pair[0] in accepted_by_response:
            accepted_by_response[pair[0]].append(pair[1])
    assert {key: len(values) for key, values in accepted_by_response.items()} == {
        response_ids[0]: 2,
        response_ids[1]: 2,
    }
    assert all(
        setvar_value(condition, action="Set") == "50.000000000"
        for condition in fib["response_conditions"]
    )


def test_grading_type_presence_and_decvar_attributes_remain_raw() -> None:
    item = specimen_item("sa_short_answer.xml")
    grading_type = next(
        node for node in item.iter() if node.tag.endswith("grading_type")
    )
    grading_type.text = None

    facts = extract_source_response_facts(item)
    assert facts["grading_type"]["state"] == "present_empty"
    assert facts["grading_type"]["raw_value"] is None
    assert len(facts["grading_type"]["occurrences"]) == 1

    decvar_attributes = facts["outcomes"][0]["attributes"]
    assert attribute_value(decvar_attributes, "varname") == "Blank_1"
    assert attribute_value(decvar_attributes, "vartype") == "Integer"
    assert attribute_value(decvar_attributes, "minvalue") == "0"
    assert attribute_value(decvar_attributes, "maxvalue") == "100"


def test_item_metadata_and_parent_trees_are_occurrence_complete() -> None:
    facts = specimen_facts("sa_short_answer.xml")

    assert facts["source_question_type"] == "Short Answer"
    assert {
        key: (
            value["state"],
            value["raw_value"],
            len(value["occurrences"]),
        )
        for key, value in facts["item_metadata"].items()
    } == {
        "qmd_questiontype": ("present_value", "Short Answer", 1),
        "qmd_computerscored": ("present_value", "yes", 1),
        "qmd_weighting": ("present_value", "1.875000000", 1),
    }
    for field_name, value in facts["item_metadata"].items():
        occurrence = value["occurrences"][0]
        assert occurrence["ordinal"] == 1
        assert occurrence["raw_value"] == value["raw_value"]
        assert occurrence["node"]["name"] == "qti_metadatafield"
        assert [child["name"] for child in occurrence["node"]["children"]] == [
            "fieldlabel",
            "fieldentry",
        ]
        assert occurrence["node"]["children"][0]["raw_text"] == field_name

    assert facts["presentation"]["state"] == "present"
    assert len(facts["presentation"]["occurrences"]) == 1
    presentation = facts["presentation"]["occurrences"][0]["node"]
    assert presentation["name"] == "presentation"
    assert [child["name"] for child in presentation["children"]] == ["flow"]
    assert [child["name"] for child in presentation["children"][0]["children"]] == [
        "material",
        "response_extension",
        "response_str",
    ]

    assert facts["response_processing"]["state"] == "present"
    assert len(facts["response_processing"]["occurrences"]) == 1
    response_processing = facts["response_processing"]["occurrences"][0]["node"]
    assert response_processing["name"] == "resprocessing"
    assert [
        child["name"] for child in response_processing["children"]
    ] == ["outcomes", "respcondition", "respcondition"]


def test_item_metadata_states_retain_missing_empty_repeated_and_malformed() -> None:
    item = specimen_item("sa_short_answer.xml")
    metadata = next(
        node for node in item.iter() if node.tag.endswith("qtimetadata")
    )
    fields = [
        node
        for node in list(metadata)
        if node.tag.endswith("qti_metadatafield")
    ]
    weighting = next(
        field
        for field in fields
        if next(
            child
            for child in list(field)
            if child.tag.endswith("fieldlabel")
        ).text
        == "qmd_weighting"
    )
    metadata.remove(weighting)

    computer_scored = next(
        field
        for field in fields
        if next(
            child
            for child in list(field)
            if child.tag.endswith("fieldlabel")
        ).text
        == "qmd_computerscored"
    )
    computer_entry = next(
        child
        for child in list(computer_scored)
        if child.tag.endswith("fieldentry")
    )
    computer_entry.text = None

    question_type = next(
        field
        for field in fields
        if next(
            child
            for child in list(field)
            if child.tag.endswith("fieldlabel")
        ).text
        == "qmd_questiontype"
    )
    metadata.append(deepcopy(question_type))

    malformed = ET.SubElement(metadata, "qti_metadatafield")
    ET.SubElement(malformed, "fieldlabel").text = "qmd_computerscored"
    ET.SubElement(malformed, "fieldlabel").text = "qmd_computerscored"
    ET.SubElement(malformed, "fieldentry").text = "yes"

    facts = extract_source_response_facts(item)
    assert facts["item_metadata"]["qmd_weighting"] == {
        "state": "absent",
        "raw_value": None,
        "occurrences": [],
    }
    assert facts["item_metadata"]["qmd_questiontype"]["state"] == "present_value"
    assert facts["item_metadata"]["qmd_questiontype"]["raw_value"] is None
    assert len(
        facts["item_metadata"]["qmd_questiontype"]["occurrences"]
    ) == 2
    assert facts["item_metadata"]["qmd_computerscored"]["state"] == "present_empty"
    assert facts["item_metadata"]["qmd_computerscored"]["raw_value"] is None
    computer_occurrences = facts["item_metadata"]["qmd_computerscored"][
        "occurrences"
    ]
    assert len(computer_occurrences) == 2
    assert [occurrence["raw_value"] for occurrence in computer_occurrences] == [
        None,
        None,
    ]
    assert [
        child["name"]
        for child in computer_occurrences[1]["node"]["children"]
    ] == ["fieldlabel", "fieldlabel", "fieldentry"]


def test_matching_conditions_preserve_one_to_many_and_many_to_one_edges() -> None:
    categorization = specimen_facts("matching_categorization_two_premise.xml")
    correct_pairs = [
        condition_pair(condition)
        for condition in categorization["response_conditions"]
        if setvar_value(condition, varname="D2L_Correct", action="Add") == "1"
    ]
    correct_pairs = [pair for pair in correct_pairs if pair is not None]

    first_premise = categorization["response_declarations"][0]["source_respident"]
    assert len([pair for pair in correct_pairs if pair[0] == first_premise]) == 6
    assert len(set(correct_pairs)) == len(correct_pairs)

    reused_option_item = specimen_item("matching_bijective_repeated_group.xml")
    conditions = [
        node
        for node in reused_option_item.iter()
        if node.tag.endswith("respcondition")
    ]
    target_pair = (
        "QUES_1_1_C2",
        "QUES_1_1_M1",
    )
    for condition in conditions:
        varequal = next(
            (node for node in condition.iter() if node.tag.endswith("varequal")),
            None,
        )
        if varequal is None:
            continue
        if (
            varequal.attrib.get("respident"),
            (varequal.text or "").strip(),
        ) != target_pair:
            continue
        setvar = next(
            node for node in condition.iter() if node.tag.endswith("setvar")
        )
        setvar.attrib["varname"] = "D2L_Correct"
        break

    option_reuse = extract_source_response_facts(reused_option_item)
    reuse_pairs = [
        condition_pair(condition)
        for condition in option_reuse["response_conditions"]
        if setvar_value(condition, varname="D2L_Correct", action="Add") == "1"
    ]
    assert ("QUES_1_1_C1", "QUES_1_1_M1") in reuse_pairs
    assert target_pair in reuse_pairs
    assert all(
        len(declaration["response_labels"]) == 5
        for declaration in option_reuse["response_declarations"]
    )


def test_ordering_positions_grading_type_and_feedback_slots_survive() -> None:
    ordering = specimen_facts("ordering_sequence.xml")
    declaration = ordering["response_declarations"][0]
    option_ids = [
        label["source_ident"] for label in declaration["response_labels"]
    ]
    positions = {
        pair[0]: int(pair[1])
        for condition in ordering["response_conditions"]
        if setvar_value(condition, varname="D2L_Correct", action="Add") == "1"
        if (pair := condition_pair(condition)) is not None
    }

    assert declaration["element_kind"] == "response_grp"
    assert declaration["rcardinality"] == "Ordered"
    assert ordering["grading_type"]["state"] == "present_value"
    assert ordering["grading_type"]["raw_value"] == "1"
    assert positions == {
        option_id: position
        for position, option_id in enumerate(option_ids, start=1)
    }
    assert len(ordering["feedback_slots"]) == len(option_ids) == 4
    assert all(slot["has_content"] is False for slot in ordering["feedback_slots"])


def test_unknown_predicate_and_empty_shell_are_preserved_with_diagnostics() -> None:
    item = specimen_item("sa_short_answer.xml")
    conditionvar = next(
        node for node in item.iter() if node.tag.endswith("conditionvar")
    )
    unknown = ET.SubElement(conditionvar, "vendor_future_predicate", {"mode": "x"})
    unknown.text = "opaque-token"

    facts = extract_source_response_facts(item)
    diagnostic_codes = {row["code"] for row in facts["diagnostics"]}
    assert "response_fact_unknown_predicate" in diagnostic_codes
    assert any(
        node["name"] == "vendor_future_predicate"
        and node["raw_text"] == "opaque-token"
        for node in iter_fact_nodes(
            facts["response_conditions"][0]["predicate_trees"][0]
        )
    )

    shell = specimen_item("matching_bijective_repeated_group.xml")
    flow = next(node for node in shell.iter() if node.tag.endswith("flow"))
    for child in list(flow):
        if child.tag.endswith("response_grp"):
            flow.remove(child)
    shell_facts = extract_source_response_facts(shell)
    assert shell_facts["response_declarations"] == []
    assert "response_shape_empty_shell" in {
        row["code"] for row in shell_facts["diagnostics"]
    }


def test_declaration_and_condition_trees_preserve_all_children_in_order() -> None:
    item = specimen_item("sa_short_answer.xml")
    declaration = next(
        node for node in item.iter() if node.tag.endswith("response_str")
    )
    ET.SubElement(declaration, "render_choice", {"shuffle": "yes"})

    condition = next(
        node for node in item.iter() if node.tag.endswith("respcondition")
    )
    condition.insert(1, ET.Element("displayfeedback", {"linkrefid": "FUTURE"}))
    extra_setvar = ET.SubElement(condition, "setvar", {"action": "Add"})
    extra_setvar.text = "25.000000000"

    facts = extract_source_response_facts(item)
    declaration_fact = facts["response_declarations"][0]
    condition_fact = facts["response_conditions"][0]

    assert declaration_fact["renderer"]["name"] == "render_fib"
    assert [
        child["name"] for child in declaration_fact["tree"]["children"]
    ] == ["render_fib", "render_choice"]
    assert "response_fact_multiple_renderers" in {
        diagnostic["code"] for diagnostic in facts["diagnostics"]
    }
    assert [child["name"] for child in condition_fact["tree"]["children"]] == [
        "conditionvar",
        "displayfeedback",
        "setvar",
        "setvar",
    ]
    assert [
        str(setvar["raw_value"] or "").strip()
        for setvar in condition_fact["setvars"]
    ] == ["100.000000000", "25.000000000"]


def test_normalization_requires_occurrence_facts_for_authoritative_choice_projection() -> None:
    payload_with_facts = build_payload(OCCURRENCE_FIXTURE)
    payload_without_facts = deepcopy(payload_with_facts)
    for row in payload_without_facts["quiz_question_rows"]:
        row.pop("source_response_facts")

    model_with_facts, _ = build_normalized_model(
        payload_with_facts,
        OCCURRENCE_FIXTURE,
        OCCURRENCE_FIXTURE,
        "directory",
    )
    model_without_facts, _ = build_normalized_model(
        payload_without_facts,
        OCCURRENCE_FIXTURE,
        OCCURRENCE_FIXTURE,
        "directory",
    )

    question_with = deepcopy(model_with_facts["questions"][0])
    question_without = deepcopy(model_without_facts["questions"][0])
    fact_records = [
        row
        for row in question_with["type_payload"]["raw_response_models"]
        if row["source_kind"] == "d2l_qti_response_facts/0"
    ]
    question_with["type_payload"]["raw_response_models"] = [
        row
        for row in question_with["type_payload"]["raw_response_models"]
        if row["source_kind"] != "d2l_qti_response_facts/0"
    ]

    assert len(fact_records) == 2
    assert [record["payload"] for record in fact_records] == [
        row["source_response_facts"]
        for row in payload_with_facts["quiz_question_rows"]
    ]
    assert all(
        "source_response_facts" not in record["payload"]
        for record in question_with["type_payload"]["raw_response_models"]
    )
    # Choice normalization now uses source IDs/material, never the lossy
    # reviewer strings. Removing those facts must withhold authoritative keys,
    # while preserving identity and the unmodified fact payloads above.
    assert question_with["entity_key"] == question_without["entity_key"]
    assert question_with["identity"]["source_aliases"] == question_without["identity"]["source_aliases"]
    assert question_without["type_payload"]["options"] == []
    assert question_without["scoring"]["state"] == "unresolved"
    assert any(
        diagnostic["code"] == "source_choice_projection_unresolved"
        and question_without["entity_key"] in diagnostic["entity_keys"]
        for diagnostic in model_without_facts["diagnostics"]
    )


def test_native_itemref_and_inline_occurrences_do_not_first_match_collapse() -> None:
    payload = build_payload(OCCURRENCE_FIXTURE)
    assert len(payload["quiz_question_rows"]) == 2
    assert payload["quiz_question_rows"][0]["quiz_item_ident"] == ""
    assert payload["quiz_question_rows"][1]["quiz_item_ident"] == "OBJ_INLINE_SHARED"
    model, _ = build_normalized_model(
        payload,
        OCCURRENCE_FIXTURE,
        OCCURRENCE_FIXTURE,
        "directory",
    )
    assert len(model["questions"]) == 1
    raw_models = model["questions"][0]["type_payload"]["raw_response_models"]
    fact_records = [
        row
        for row in raw_models
        if row["source_kind"] == "d2l_qti_response_facts/0"
    ]
    legacy_records = [
        row
        for row in raw_models
        if row["source_kind"].startswith("canonical-extractor-row:")
    ]
    assert len(fact_records) == 2
    assert len(legacy_records) == 2

    occurrence_evidence = [
        next(
            key
            for key in record["source_evidence_keys"]
            if key.startswith("ev.occurrence.")
        )
        for record in fact_records
    ]
    assert len(set(occurrence_evidence)) == 2
    assert [
        record["payload"]["grading_type"]["raw_value"]
        for record in fact_records
    ] == ["0", "1"]
    assert [
        attribute_value(
            record["payload"]["response_declarations"][0]["renderer"]["attributes"],
            "shuffle",
        )
        for record in fact_records
    ] == ["no", "yes"]

    occurrence_rows = {
        row["evidence_key"]: row
        for row in model["evidence"]
        if row["evidence_key"] in occurrence_evidence
    }
    assert {
        (row["source_ref"], row["locator"])
        for row in occurrence_rows.values()
    } == {
        ("quiz_d2l_occurrences.xml", "quiz item 1"),
        ("quiz_d2l_occurrences.xml", "quiz item 2"),
    }

    expected_legacy_payloads = [
        {
            key: value
            for key, value in row.items()
            if "absolute_path" not in key and key != "source_response_facts"
        }
        for row in payload["quiz_question_rows"]
    ]
    assert [record["payload"] for record in legacy_records] == expected_legacy_payloads
    assert [record["extensions"] for record in legacy_records] == [
        {},
        {"additional_occurrence": True},
    ]


def test_nested_facts_stay_in_json_rows_but_not_scalar_worksheets(
    tmp_path: Path,
) -> None:
    payload = build_payload(OCCURRENCE_FIXTURE)
    assert {
        row["source_response_facts"]["schema"]
        for row in payload["quiz_question_rows"]
    } == {
        "coursecraft.quiz_source_response_facts/0"
    }
    assert {
        row["source_response_facts"]["grading_type"]["raw_value"]
        for row in payload["quiz_question_rows"]
    } == {
        "0",
        "1",
    }

    output_path = tmp_path / "review.xlsx"
    write_workbook(
        output_path,
        payload,
        OCCURRENCE_FIXTURE,
        "2026-07-24 00:00:00 UTC",
        REPO_ROOT / "scripts" / "extract_quiz_pool_review.py",
    )
    workbook = load_workbook(output_path, read_only=True)
    headers = [cell.value for cell in next(workbook["Quiz Questions"].iter_rows())]
    assert "source_response_facts" not in headers
