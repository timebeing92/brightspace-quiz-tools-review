"""Synthetic coverage of target identity, encoding, feedback and asset fidelity."""
from copy import deepcopy
import html
import json
from pathlib import Path
import subprocess
import sys
import xml.etree.ElementTree as ET

import pytest

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests/fixtures/quiz_authoring"
sys.path.insert(0, str(ROOT / "scripts"))

from quiz_authoring_readiness import analyze_authoring_readiness
from quiz_build_support import load_authoring_projection, target_identifier_issues
from build_quiz_package_from_workbook import html_fragment


def model():
    return json.loads((FIXTURE / "buildable_quiz.model.json").read_text())


def save(tmp_path, value):
    path = tmp_path / "source.model.json"
    path.write_text(json.dumps(value))
    return path


def readiness(path):
    return analyze_authoring_readiness(path, asset_root=FIXTURE)


def build(path, tmp_path):
    target = tmp_path / "package"
    result = subprocess.run(
        [sys.executable, str(ROOT / "scripts/build_quiz_package_from_workbook.py"),
         str(path), "--output-dir", str(target), "--asset-root", str(FIXTURE)],
        text=True, capture_output=True, check=False,
    )
    return result, target


@pytest.mark.parametrize("codes", [("Q-a", "q-A"), ("Q-A", "Q_A"), ("Q.A", "Q-A"), ("1", "Q_1")])
def test_question_target_collisions_block_readiness_and_direct_build(tmp_path, codes):
    value = model()
    for question, code in zip(value["questions"], codes):
        question["identity"]["permanent_code"] = code
    path = save(tmp_path, value)
    before = path.read_bytes()
    report = readiness(path)
    assert not report["ready"]
    assert "target_question_identifier_collision" in {row["code"] for row in report["issues"]}
    with pytest.raises(ValueError, match="Target question identifier"):
        load_authoring_projection(path, asset_root=FIXTURE)
    result, target = build(path, tmp_path)
    assert result.returncode != 0 and "Target question identifier" in result.stderr
    assert not target.exists()
    assert path.read_bytes() == before


def split_pools(value, codes):
    pool, draw = value["structures"]
    second_pool, second_draw = deepcopy(pool), deepcopy(draw)
    second_pool["entity_key"] += "-second"
    second_draw["entity_key"] += "-second"
    second_draw["ordinal"] = 2
    second_draw["identity"]["permanent_code"] += "-second"
    pool["identity"]["permanent_code"], second_pool["identity"]["permanent_code"] = codes
    for row in (pool, draw, second_pool, second_draw):
        row["selection"]["available_count"] = 2
    value["structures"].extend([second_pool, second_draw])
    for kind in ("contains", "draws_from"):
        original = next(row for row in value["relationships"] if row["kind"] == kind)
        duplicate = deepcopy(original)
        duplicate["relationship_key"] += "-second"
        if kind == "contains":
            duplicate["to_entity_key"] = second_draw["entity_key"]
        else:
            duplicate["from_entity_key"] = second_draw["entity_key"]
            duplicate["to_entity_key"] = second_pool["entity_key"]
        value["relationships"].append(duplicate)
    moved = {row["entity_key"] for row in value["questions"][2:]}
    for row in value["relationships"]:
        if row["kind"] == "member_of" and row["from_entity_key"] in moved:
            row["to_entity_key"] = second_pool["entity_key"]


@pytest.mark.parametrize("codes", [("P-A", "P_A"), ("P.A", "p-a"), ("P A", "P_A")])
def test_pool_collisions_cannot_merge_memberships(tmp_path, codes):
    value = model()
    split_pools(value, codes)
    path = save(tmp_path, value)
    report = readiness(path)
    assert "target_pool_identifier_collision" in {row["code"] for row in report["issues"]}
    result, target = build(path, tmp_path)
    assert result.returncode != 0 and "Target pool identifier" in result.stderr
    assert not target.exists()


def test_same_bank_reference_is_allowed_but_duplicate_draw_target_is_not():
    assert not target_identifier_issues([], [("P", "same"), ("P", "same")], [])
    issues = target_identifier_issues([], [], [(1, "P-A", "draw-a"), (1, "P_A", "draw-b")])
    assert issues[0][0] == "target_draw_identifier_collision"


def test_populated_approved_keys_cannot_bypass_unresolved_source_evidence(tmp_path):
    value = model()
    question = value["questions"][0]
    assert question["build_support"]["level"] == "roundtrip_verified"
    assert sum(option["correct"] is True for option in question["type_payload"]["options"]) == 1
    question["type_payload"]["extensions"]["coursecraft.source_choice_projection"] = {
        "state": "unresolved", "reason": "Synthetic contradictory native scoring evidence."
    }
    path = save(tmp_path, value)
    before = path.read_bytes()
    report = readiness(path)
    assert "source_choice_projection_unresolved" in {row["code"] for row in report["issues"]}
    assert not report["ready"]
    result, target = build(path, tmp_path)
    assert result.returncode != 0 and "unresolved source choice evidence" in result.stderr
    assert not target.exists()
    assert path.read_bytes() == before


@pytest.mark.parametrize("fmt", ["plain_text", "html", "xhtml"])
def test_explicit_encoding_survives_projection_to_serialized_fields(tmp_path, fmt):
    value = model()
    literal = '  <b>literal &amp; \"quoted\"</b> & x < y > z\nnext  '
    for question in value["questions"]:
        question["prompt"] = {"format": fmt, "content": literal, "extensions": {}}
        if question["kind"] in {"multiple_choice", "multi_select"}:
            for option in question["type_payload"]["options"]:
                option["content"] = {"format": fmt, "content": literal, "extensions": {}}
        if question["kind"] == "long_answer":
            question["type_payload"]["manual_answer_key"] = {"format": fmt, "content": literal, "extensions": {}}
    path = save(tmp_path, value)
    assert readiness(path)["ready"]
    result, target = build(path, tmp_path)
    assert result.returncode == 0, result.stdout + result.stderr
    root = ET.parse(target / "questiondb.xml").getroot()
    expected = f"<p>{html.escape(literal, quote=False)}</p>" if fmt == "plain_text" else literal
    items = root.findall(".//item")
    assert len(items) == 4
    for item in items:
        assert item.find("./presentation/flow/material/mattext").text == expected
    choices = [item for item in items if item.find("./presentation/flow/response_lid") is not None]
    for item in (choices[0], choices[2]):
        assert all(node.text == expected for node in item.findall(".//response_label/flow_mat/material/mattext"))
    assert items[-1].find("./answer_key/answer_key_material/flow_mat/material/mattext").text == expected
    assert path.read_text() == json.dumps(value)


def feedback(channel="answer_key", text="Synthetic evaluator guidance"):
    return {"channel": channel, "source_kind": "synthetic", "content": {"format": "plain_text", "content": text, "extensions": {}},
            "source_evidence_keys": [], "extensions": {}}


@pytest.mark.parametrize("case,expected_code", [
    ("general", "feedback_projection_not_supported"),
    ("duplicate", "answer_key_feedback_projection_not_supported"),
    ("shadowed", "answer_key_feedback_projection_not_supported"),
    ("nonwritten", "answer_key_feedback_projection_not_supported"),
    ("nonwritten_manual", "manual_answer_key_projection_not_supported"),
    ("empty", "empty_evaluator_answer_key"),
])
def test_unrepresented_feedback_blocks_both_entrypoints(tmp_path, case, expected_code):
    value = model()
    written = value["questions"][-1]
    written["type_payload"].pop("manual_answer_key")
    written["feedback"] = [feedback()]
    if case == "general":
        written["feedback"].append(feedback("general"))
    elif case == "duplicate":
        written["feedback"].append(feedback())
    elif case == "shadowed":
        written["type_payload"]["manual_answer_key"] = deepcopy(written["feedback"][0]["content"])
    elif case == "nonwritten":
        value["questions"][0]["feedback"] = [feedback()]
    elif case == "nonwritten_manual":
        value["questions"][0]["type_payload"]["manual_answer_key"] = deepcopy(written["feedback"][0]["content"])
    elif case == "empty":
        written["feedback"][0]["content"]["content"] = ""
    path = save(tmp_path, value)
    report = readiness(path)
    assert not report["ready"]
    issue = next(row for row in report["issues"] if row["code"] == expected_code)
    result, target = build(path, tmp_path)
    assert result.returncode != 0 and issue["message"] in result.stderr
    assert not target.exists()


def test_one_written_feedback_key_is_preserved_and_escaped(tmp_path):
    value = model()
    written = value["questions"][-1]
    written["type_payload"].pop("manual_answer_key")
    written["feedback"] = [feedback(text="Compare <x> & <y> literally")]
    path = save(tmp_path, value)
    assert readiness(path)["ready"]
    result, target = build(path, tmp_path)
    assert result.returncode == 0, result.stderr
    root = ET.parse(target / "questiondb.xml").getroot()
    assert root.find(".//answer_key/answer_key_material/flow_mat/material/mattext").text == "<p>Compare &lt;x&gt; &amp; &lt;y&gt; literally</p>"


@pytest.mark.parametrize("change", [
    lambda q: q["type_payload"]["options"][0]["content"].update(format="html", content="<b>True</b>"),
    lambda q: q["type_payload"]["options"].reverse(),
])
def test_true_false_custom_labels_are_not_silently_replaced(tmp_path, change):
    value = model()
    change(value["questions"][1])
    path = save(tmp_path, value)
    report = readiness(path)
    assert "true_false_option_projection_not_supported" in {row["code"] for row in report["issues"]}
    result, target = build(path, tmp_path)
    assert result.returncode != 0 and not target.exists()


def test_legacy_workbook_collisions_are_refused_before_output(tmp_path):
    from test_quiz_package_builder_validator import write_source_workbook
    from openpyxl import load_workbook
    source = tmp_path / "legacy.xlsx"
    write_source_workbook(source)
    wb = load_workbook(source)
    ws = wb["Bank_1"]
    ws.cell(2, 4, "Q-A")
    ws.cell(3, 4, "Q_A")
    wb.save(source)
    result, target = build(source, tmp_path)
    assert result.returncode != 0 and "Target question identifier" in result.stderr
    assert not target.exists()


def test_legacy_workbook_html_heuristic_remains_unchanged():
    assert html_fragment("  <b>legacy</b> ") == "<b>legacy</b>"
    assert html_fragment(" ordinary ") == "<p>ordinary</p>"


def put_html(value, location, markup):
    question = value["questions"][0 if location in {"prompt", "option"} else -1]
    content = {"format": "html", "content": markup, "extensions": {}}
    if location == "prompt":
        question["prompt"] = content
    elif location == "option":
        question["type_payload"]["options"][0]["content"] = content
    elif location == "key":
        question["type_payload"]["manual_answer_key"] = content
    else:
        question["type_payload"].pop("manual_answer_key")
        question["feedback"] = [feedback()]
        question["feedback"][0]["content"] = content
    return question


def bind_asset(value, question, asset_key=None):
    asset_key = asset_key or value["assets"][0]["entity_key"]
    template = next(row for row in value["relationships"] if row["kind"] == "uses_asset")
    if template["from_entity_key"] == question["entity_key"] and template["to_entity_key"] == asset_key:
        return
    relation = deepcopy(template)
    relation["relationship_key"] += f"-additional-{len(value['relationships'])}"
    relation["from_entity_key"] = question["entity_key"]
    relation["to_entity_key"] = asset_key
    value["relationships"].append(relation)


@pytest.mark.parametrize("location", ["prompt", "option", "key", "feedback"])
@pytest.mark.parametrize("markup", [
    '<p>Illustration <IMG SRC="missing.svg"></p>',
    '<math altimg="missing.svg"><mi>x</mi></math>',
])
def test_unbound_html_media_blocks_readiness_and_direct_builder(tmp_path, location, markup):
    value = model()
    put_html(value, location, markup)
    path = save(tmp_path, value)
    original = path.read_bytes()
    report = readiness(path)
    assert not report["ready"]
    assert "html_asset_reference_not_bound" in {row["code"] for row in report["issues"]}
    result, target = build(path, tmp_path)
    assert result.returncode != 0 and "no matching selected asset" in result.stderr
    assert not target.exists() and path.read_bytes() == original


@pytest.mark.parametrize("location", ["prompt", "option", "key", "feedback"])
@pytest.mark.parametrize("markup", [
    '<p>Illustration <img src="quiz-assets/concept-map.svg?view=1&amp;mode=2#diagram"></p>',
    '<math altimg="quiz-assets/concept-map.svg"><mi>x</mi></math>',
])
def test_bound_media_is_copied_without_rewriting_html(tmp_path, location, markup):
    value = model()
    question = put_html(value, location, markup)
    bind_asset(value, question)
    path = save(tmp_path, value)
    original = path.read_bytes()
    assert readiness(path)["ready"]
    result, target = build(path, tmp_path)
    assert result.returncode == 0, result.stdout + result.stderr
    assert (target / "quiz-assets/concept-map.svg").read_bytes() == (FIXTURE / "assets/concept-map.svg").read_bytes()
    assert markup in [node.text for node in ET.parse(target / "questiondb.xml").iter("mattext")]
    assert path.read_bytes() == original


@pytest.mark.parametrize("ref", [
    '\u00a0quiz-assets/concept-map.svg', 'quiz-assets/concept-map.svg\u00a0',
    '&nbsp;quiz-assets/concept-map.svg', 'quiz-assets/concept-map.svg&nbsp;',
])
def test_unicode_whitespace_cannot_alias_a_clean_asset_path(tmp_path, ref):
    value = model()
    put_html(value, "prompt", f'<img src="{ref}">')
    path = save(tmp_path, value)
    original = path.read_bytes()
    report = readiness(path)
    assert not report["ready"]
    assert "html_asset_reference_not_bound" in {row["code"] for row in report["issues"]}
    result, target = build(path, tmp_path)
    assert result.returncode != 0 and not target.exists()
    assert path.read_bytes() == original


def test_ascii_space_url_trimming_preserves_authored_html(tmp_path):
    value = model()
    markup = '<img src="  quiz-assets/concept-map.svg  ">'
    put_html(value, "prompt", markup)
    path = save(tmp_path, value)
    assert readiness(path)["ready"]
    result, target = build(path, tmp_path)
    assert result.returncode == 0, result.stderr
    assert (target / "quiz-assets/concept-map.svg").is_file()
    assert markup in [node.text for node in ET.parse(target / "questiondb.xml").iter("mattext")]


def test_percent_encoded_local_media_uses_existing_archive_member_mapping(tmp_path):
    value = model()
    value["assets"][0]["package_path"] = "quiz-assets/concept%20map.svg"
    markup = '<img src="quiz-assets/concept%20map.svg">'
    put_html(value, "prompt", markup)
    path = save(tmp_path, value)
    assert readiness(path)["ready"]
    result, target = build(path, tmp_path)
    assert result.returncode == 0, result.stderr
    assert (target / "quiz-assets/concept map.svg").is_file()
    assert markup in [node.text for node in ET.parse(target / "questiondb.xml").iter("mattext")]


@pytest.mark.parametrize("ref", [
    "../outside.svg", "%2e%2e/outside.svg", "quiz-assets%2fconcept-map.svg",
    "file:///tmp/a.svg", "javascript:alert(1)", "/content/enforced/a.svg",
    "data:image/svg+xml,<svg/>", "quiz-assets\\concept-map.svg", "broken%ZZ.svg", "",
    "https:///figure.svg", "https://host:invalid/figure.svg", "//host:65536/figure.svg",
    "https://bad host/figure.svg", "https://user:password@host/figure.svg",
])
def test_unsafe_or_unsupported_references_are_refused(tmp_path, ref):
    value = model()
    put_html(value, "prompt", f'<img src="{html.escape(ref, quote=True)}">')
    path = save(tmp_path, value)
    assert "unsafe_or_unsupported_html_reference" in {row["code"] for row in readiness(path)["issues"]}
    result, target = build(path, tmp_path)
    assert result.returncode != 0 and not target.exists()


@pytest.mark.parametrize("markup", [
    '<img srcset="missing.svg 1x, other.svg 2x">',
    '<p style="background: url(missing.svg)">Figure</p>',
    '<style>p { background: url(missing.svg); }</style>',
    '<base href="https://example.test/"><img src="quiz-assets/concept-map.svg">',
    '<img src="quiz-assets/concept-map.svg" src="missing.svg">',
    '<object codebase="nested/" data="quiz-assets/concept-map.svg"></object>',
    '<svg xml:base="nested/"><image href="quiz-assets/concept-map.svg"/></svg>',
])
def test_ambiguous_reference_syntax_is_not_silently_skipped(tmp_path, markup):
    value = model()
    put_html(value, "prompt", markup)
    path = save(tmp_path, value)
    assert "html_reference_projection_not_supported" in {row["code"] for row in readiness(path)["issues"]}
    result, target = build(path, tmp_path)
    assert result.returncode != 0 and not target.exists()


@pytest.mark.parametrize("case", ["unselected", "wrong_question", "wrong_package_path", "ambiguous"])
def test_asset_record_presence_does_not_replace_an_exact_binding(tmp_path, case):
    value = model()
    if case == "unselected":
        value["relationships"] = [row for row in value["relationships"] if row["kind"] != "uses_asset"]
    elif case == "wrong_question":
        put_html(value, "key", '<img src="quiz-assets/concept-map.svg">')
    elif case == "wrong_package_path":
        put_html(value, "prompt", '<img src="assets/concept-map.svg">')
    else:
        asset = deepcopy(value["assets"][0])
        asset["entity_key"] += "-alias"
        value["assets"].append(asset)
        bind_asset(value, value["questions"][0], asset["entity_key"])
    path = save(tmp_path, value)
    expected = "html_asset_reference_ambiguous" if case == "ambiguous" else "html_asset_reference_not_bound"
    assert expected in {row["code"] for row in readiness(path)["issues"]}
    result, target = build(path, tmp_path)
    assert result.returncode != 0 and not target.exists()


def test_existing_external_link_semantics_and_literal_plain_text_are_preserved(tmp_path):
    value = model()
    markup = '<p data="metadata"><img src="https://example.test/figure.svg"><img src="//example.test/figure.svg"><a href="mailto:reviewer@example.test">Contact</a><a href="#note">Note</a></p>'
    put_html(value, "prompt", markup)
    value["questions"][1]["prompt"]["content"] = '<img src="unbound.svg"> is literal text'
    path = save(tmp_path, value)
    assert readiness(path)["ready"]
    result, target = build(path, tmp_path)
    assert result.returncode == 0, result.stderr
    texts = [node.text for node in ET.parse(target / "questiondb.xml").iter("mattext")]
    assert markup in texts
    assert '<p>&lt;img src="unbound.svg"&gt; is literal text</p>' in texts
