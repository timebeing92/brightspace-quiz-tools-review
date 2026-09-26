from __future__ import annotations

import json
from pathlib import Path
import shutil
import sys
import zipfile

import pytest
from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
SCRIPT_ROOT = REPO_ROOT / "scripts"
FIXTURE_ROOT = REPO_ROOT / "tests" / "fixtures" / "quiz_xml"
if str(SCRIPT_ROOT) not in sys.path:
    sys.path.insert(0, str(SCRIPT_ROOT))

import quiz_unbind
from quiz_unbind_receipt import SCHEMA_ID, sha256_file, validate_unbind_record
from verify_quiz_unbind_run import UnbindVerificationError, verify_unbind_run


MIXED = FIXTURE_ROOT / "mixed_inline_itemref_and_root_bank"
QUIZ_ONLY = FIXTURE_ROOT / "quiz_only_inline_tf"
DUPLICATE_MATCH = FIXTURE_ROOT / "duplicate_pool_and_similarity"
SENTINEL = "QUIZUNBINDCANARYPROMPT"


def _run(source: Path, output: Path, **kwargs):
    return quiz_unbind.run_unbind(
        export_path=source,
        output_dir=output,
        source_lineage_key=kwargs.pop(
            "source_lineage_key", "cc:lineage:fixture:quiz-unbind"
        ),
        **kwargs,
    )


def _record(output: Path) -> dict:
    return json.loads((output / quiz_unbind.UNBIND_RECORD_NAME).read_text(encoding="utf-8"))


def _extraction_receipt(output: Path) -> tuple[Path, dict]:
    path = next((output / "extraction").glob("*.run.json"))
    return path, json.loads(path.read_text(encoding="utf-8"))


def _staging_dirs(output: Path) -> list[Path]:
    return list(output.parent.glob(f".{output.name}.unbind-staging-*"))


def _refresh_direct_artifact(output: Path, artifact_key: str) -> None:
    record = _record(output)
    ref = record["artifacts"][artifact_key]
    artifact = output / ref["path"]
    ref["bytes"] = artifact.stat().st_size
    ref["sha256"] = sha256_file(artifact)
    (output / quiz_unbind.UNBIND_RECORD_NAME).write_text(
        json.dumps(record, indent=2) + "\n",
        encoding="utf-8",
    )


def _zip_tree(source: Path, archive: Path) -> None:
    with zipfile.ZipFile(archive, "w") as handle:
        for path in sorted(source.rglob("*")):
            if path.is_file():
                handle.write(path, path.relative_to(source).as_posix())


def _build_asset_export(tmp_path: Path) -> Path:
    export = tmp_path / "asset_export"
    shutil.copytree(QUIZ_ONLY, export)
    images = export / "images"
    images.mkdir()
    (images / "alpha.png").write_bytes(b"synthetic-png-for-unbind")
    quiz = export / "quiz_d2l_only.xml"
    raw = quiz.read_text(encoding="utf-8").replace(
        "is this synthetic prompt true or false?",
        "is this synthetic prompt true or false? "
        "&lt;img src=&quot;images/alpha.png&quot; alt=&quot;alpha&quot; /&gt;",
    )
    quiz.write_text(raw, encoding="utf-8")
    return export


def test_full_unbind_writes_draft_receipt_station_and_recursive_verified_close(
    tmp_path: Path,
) -> None:
    output = tmp_path / "run"
    events: list[dict] = []
    result = quiz_unbind.run_unbind(
        export_path=MIXED,
        output_dir=output,
        source_lineage_key="cc:lineage:fixture:mixed-unbind",
        event_sinks=(events.append,),
    )

    record = result["record"]
    assert record["schema"] == SCHEMA_ID == quiz_unbind.UNBIND_SCHEMA_ID
    assert record["contract_status"] == "draft_unratified"
    assert validate_unbind_record(record) == []
    assert record["verification"]["status"] == "verified"
    assert record["evidence_level"] == "extraction_only"
    assert record["brightspace"] == {
        "roundtrip": "not_performed",
        "network_operations": 0,
        "brightspace_operations": 0,
    }
    assert (output / "station" / "station.json").is_file()
    assert (output / "station" / "index.html").is_file()
    assert (output / "review" / "quiz_review_projection.json").is_file()
    assert (output / "reading_room" / "station.json").is_file()
    assert (output / "reading_room" / "index.html").is_file()
    assert (output / "UNBIND_SUMMARY.md").is_file()
    assert record["source_export"]["scope"] == {
        "coverage_kind": "question_library_with_itemrefs",
        "questiondb_present": True,
        "inline_occurrence_count": 1,
        "itemref_occurrence_count": 2,
        "resolved_itemref_count": 2,
        "unresolved_itemref_count": 0,
    }
    assert record["model_summary"]["question_occurrence_count"] == 3

    stable_stages = [
        "source_check",
        "extract_normalize",
        "project_review",
        "readiness",
        "verified_close",
    ]
    assert [(row["stage"], row["status"]) for row in events] == [
        (stage, status) for stage in stable_stages for status in ("started", "completed")
    ]
    assert events == record["events"]
    verified = verify_unbind_run(output)
    reading_room_manifest = json.loads(
        (output / "reading_room" / "station.json").read_text(encoding="utf-8")
    )
    shard_count = len(reading_room_manifest["shards"])
    assert shard_count > 0
    assert verified["direct_artifact_count"] == 10
    assert verified["extraction_artifact_count"] == 5
    assert verified["reading_room_shard_count"] == shard_count
    assert verified["total_reference_count"] == 15 + shard_count
    assert record["verification"]["total_reference_count"] == 15 + shard_count
    assert verified["close_scope"] == "local_extraction_evidence_only"


def test_quiz_only_run_is_content_and_path_minimized(tmp_path: Path, capsys) -> None:
    output = tmp_path / "run"
    assert quiz_unbind.main(
        [
            str(QUIZ_ONLY),
            "--output-dir",
            str(output),
            "--source-lineage-key",
            "cc:lineage:fixture:quiz-only-unbind",
            "--json-events",
        ]
    ) == 0
    captured = capsys.readouterr()
    assert SENTINEL in (QUIZ_ONLY / "quiz_d2l_only.xml").read_text(encoding="utf-8")
    minimized = "\n".join(
        [
            (output / "quiz_unbind_run.json").read_text(encoding="utf-8"),
            (output / "UNBIND_SUMMARY.md").read_text(encoding="utf-8"),
            captured.out,
            captured.err,
        ]
    )
    assert SENTINEL not in minimized
    assert str(tmp_path) not in minimized
    receipt_path, receipt = _extraction_receipt(output)
    capability = next(
        row for row in receipt["capabilities"] if row["name"] == "question_library_resolution"
    )
    assert capability["status"] == "skipped"
    assert receipt_path.is_file()
    record = _record(output)
    assert record["source_export"]["scope"] == {
        "coverage_kind": "quiz_only_inline",
        "questiondb_present": False,
        "inline_occurrence_count": 1,
        "itemref_occurrence_count": 0,
        "resolved_itemref_count": 0,
        "unresolved_itemref_count": 0,
    }
    assert "verified close -- local extraction evidence only" in captured.out


def test_occupied_output_and_symlinked_export_are_refused_before_work(
    tmp_path: Path,
) -> None:
    occupied = tmp_path / "occupied"
    occupied.mkdir()
    (occupied / "keep.txt").write_text("user data", encoding="utf-8")
    assert quiz_unbind.main([str(QUIZ_ONLY), "--output-dir", str(occupied)]) == 2
    assert not (occupied / "extraction").exists()
    assert (occupied / "keep.txt").read_text(encoding="utf-8") == "user data"

    link = tmp_path / "export-link"
    link.symlink_to(QUIZ_ONLY, target_is_directory=True)
    assert quiz_unbind.main([str(link), "--output-dir", str(tmp_path / "linked")]) == 2
    assert not (tmp_path / "linked").exists()


def test_success_is_verified_in_sibling_staging_then_atomically_promoted(
    tmp_path: Path,
) -> None:
    output = tmp_path / "run"
    final_visibility: list[bool] = []

    result = quiz_unbind.run_unbind(
        export_path=QUIZ_ONLY,
        output_dir=output,
        source_lineage_key="cc:lineage:fixture:atomic-unbind",
        event_sinks=(lambda _event: final_visibility.append(output.exists()),),
    )

    assert final_visibility
    assert not any(final_visibility[:-1])
    assert final_visibility[-1] is True
    assert output.is_dir()
    assert not _staging_dirs(output)
    assert result["record_path"] == output / quiz_unbind.UNBIND_RECORD_NAME
    assert result["summary_path"] == output / "UNBIND_SUMMARY.md"
    assert verify_unbind_run(output)["status"] == "verified"


def test_failed_run_removes_only_owned_staging_and_publishes_no_partial_output(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys
) -> None:
    output = tmp_path / "run"

    def fail_station(_model: Path, _registry: Path, station: Path) -> None:
        station.mkdir()
        (station / "partial.txt").write_text("partial", encoding="utf-8")
        raise quiz_unbind.UnbindFailed("synthetic station failure")

    monkeypatch.setattr(quiz_unbind, "_render_station", fail_station)
    assert quiz_unbind.main([str(QUIZ_ONLY), "--output-dir", str(output)]) == 2

    captured = capsys.readouterr()
    assert captured.err.strip() == "ERROR: synthetic station failure"
    assert not output.exists()
    assert not _staging_dirs(output)


def test_ctrl_c_is_concise_cleans_staging_and_preserves_empty_output(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys
) -> None:
    output = tmp_path / "run"
    output.mkdir()

    def cancel_station(_model: Path, _registry: Path, station: Path) -> None:
        station.mkdir()
        (station / "partial.txt").write_text(SENTINEL, encoding="utf-8")
        raise KeyboardInterrupt

    monkeypatch.setattr(quiz_unbind, "_render_station", cancel_station)
    assert quiz_unbind.main([str(QUIZ_ONLY), "--output-dir", str(output)]) == 130

    captured = capsys.readouterr()
    assert captured.out == ""
    assert captured.err.strip() == (
        "CANCELLED: Unbind cancelled; no output was promoted"
    )
    assert SENTINEL not in captured.err
    assert str(tmp_path) not in captured.err
    assert output.is_dir()
    assert not any(output.iterdir())
    assert not _staging_dirs(output)


def test_output_occupied_during_run_is_preserved_and_blocks_promotion(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    output = tmp_path / "run"
    events: list[dict] = []
    original = quiz_unbind._render_station

    def occupy_output(model: Path, registry: Path, station: Path) -> None:
        original(model, registry, station)
        output.mkdir()
        (output / "keep.txt").write_text("user data", encoding="utf-8")

    monkeypatch.setattr(quiz_unbind, "_render_station", occupy_output)
    with pytest.raises(quiz_unbind.UnbindRefused, match="became occupied"):
        quiz_unbind.run_unbind(
            export_path=QUIZ_ONLY,
            output_dir=output,
            source_lineage_key="cc:lineage:fixture:mid-run-occupied",
            event_sinks=(events.append,),
        )

    assert (output / "keep.txt").read_text(encoding="utf-8") == "user data"
    assert not (output / "extraction").exists()
    assert not _staging_dirs(output)
    assert not any(
        row["stage"] == "verified_close" and row["status"] == "completed"
        for row in events
    )


def test_unsafe_archive_member_is_refused(tmp_path: Path) -> None:
    archive = tmp_path / "unsafe.zip"
    with zipfile.ZipFile(archive, "w") as handle:
        handle.writestr("../escape.xml", "<unsafe />")
        handle.writestr("quiz_d2l_safe.xml", "<questestinterop />")
    output = tmp_path / "run"
    assert quiz_unbind.main([str(archive), "--output-dir", str(output)]) == 2
    assert not (tmp_path / "escape.xml").exists()


def test_recursive_close_detects_tamper_of_every_extraction_review_artifact(
    tmp_path: Path,
) -> None:
    original = tmp_path / "original"
    _run(MIXED, original)
    _receipt_path, receipt = _extraction_receipt(original)
    review_rows = [row for row in receipt["artifacts"] if row["role"] == "review"]
    assert len(review_rows) == 4

    for index, row in enumerate(review_rows):
        variant = tmp_path / f"tampered-{index}"
        shutil.copytree(original, variant)
        artifact = variant / "extraction" / row["path"]
        with artifact.open("ab") as handle:
            handle.write(b"\ntampered")
        with pytest.raises(UnbindVerificationError, match="mismatch"):
            verify_unbind_run(variant)


def test_recursive_close_hash_checks_every_declared_reading_room_shard(
    tmp_path: Path,
) -> None:
    original = tmp_path / "original"
    _run(MIXED, original)
    manifest = json.loads(
        (original / "reading_room" / "station.json").read_text(encoding="utf-8")
    )
    assert manifest["shards"]

    for index, row in enumerate(manifest["shards"]):
        variant = tmp_path / f"shard-tampered-{index}"
        shutil.copytree(original, variant)
        shard = variant / "reading_room" / row["path"]
        payload = bytearray(shard.read_bytes())
        payload[0] ^= 1
        shard.write_bytes(payload)
        with pytest.raises(UnbindVerificationError, match="hash mismatch"):
            verify_unbind_run(variant)


def test_recursive_close_size_checks_reading_room_shards(tmp_path: Path) -> None:
    output = tmp_path / "run"
    _run(QUIZ_ONLY, output)
    manifest = json.loads(
        (output / "reading_room" / "station.json").read_text(encoding="utf-8")
    )
    shard = output / "reading_room" / manifest["shards"][0]["path"]
    with shard.open("ab") as handle:
        handle.write(b"\nextra")
    with pytest.raises(UnbindVerificationError, match="byte-size mismatch"):
        verify_unbind_run(output)


def test_recursive_close_rejects_unsafe_reading_room_shard_path(
    tmp_path: Path,
) -> None:
    output = tmp_path / "run"
    _run(QUIZ_ONLY, output)
    manifest_path = output / "reading_room" / "station.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    manifest["shards"][0]["path"] = "../outside.html"
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    _refresh_direct_artifact(output, "reading_room_data")

    with pytest.raises(UnbindVerificationError, match="unsafe Reading Room shard path"):
        verify_unbind_run(output)


def test_recursive_close_rejects_symlinked_reading_room_shard(
    tmp_path: Path,
) -> None:
    output = tmp_path / "run"
    _run(QUIZ_ONLY, output)
    manifest = json.loads(
        (output / "reading_room" / "station.json").read_text(encoding="utf-8")
    )
    shard = output / "reading_room" / manifest["shards"][0]["path"]
    target = shard.with_name("symlink-target.html")
    target.write_bytes(shard.read_bytes())
    shard.unlink()
    shard.symlink_to(target.name)

    with pytest.raises(UnbindVerificationError, match="symlink"):
        verify_unbind_run(output)


def test_recursive_close_rejects_stale_undeclared_reading_room_file(
    tmp_path: Path,
) -> None:
    output = tmp_path / "run"
    _run(QUIZ_ONLY, output)
    (output / "reading_room" / "stale-extra.html").write_text(
        "stale",
        encoding="utf-8",
    )

    with pytest.raises(UnbindVerificationError, match="inventory is not exact"):
        verify_unbind_run(output)


def test_recursive_close_rejects_undeclared_reading_room_symlink(
    tmp_path: Path,
) -> None:
    output = tmp_path / "run"
    _run(QUIZ_ONLY, output)
    (output / "reading_room" / "undeclared-link.html").symlink_to("index.html")

    with pytest.raises(UnbindVerificationError, match="contains a symlink"):
        verify_unbind_run(output)


@pytest.mark.parametrize(
    "artifact_key",
    [
        "readiness_report",
        "readiness_markdown",
        "station_data",
        "station_html",
        "review_projection",
        "reading_room_data",
        "reading_room_html",
        "summary_markdown",
    ],
)
def test_close_detects_tamper_of_direct_review_artifacts(
    artifact_key: str, tmp_path: Path
) -> None:
    output = tmp_path / "run"
    _run(QUIZ_ONLY, output)
    record = _record(output)
    artifact = output / record["artifacts"][artifact_key]["path"]
    with artifact.open("ab") as handle:
        handle.write(b"\ntampered")
    with pytest.raises(UnbindVerificationError, match="mismatch"):
        verify_unbind_run(output)


def test_close_rejects_self_consistent_top_level_model_swap(tmp_path: Path) -> None:
    output = tmp_path / "run"
    _run(MIXED, output)
    record = _record(output)
    model = output / record["artifacts"]["model"]["path"]
    payload = json.loads(model.read_text(encoding="utf-8"))
    payload["extensions"]["test_swap"] = True
    model.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")

    record["artifacts"]["model"]["sha256"] = sha256_file(model)
    record["artifacts"]["model"]["bytes"] = model.stat().st_size
    (output / "quiz_unbind_run.json").write_text(
        json.dumps(record, indent=2) + "\n", encoding="utf-8"
    )
    with pytest.raises(UnbindVerificationError, match="extraction artifact|chain mismatch"):
        verify_unbind_run(output)


def test_asset_copy_default_and_reference_opt_out_for_folder_and_zip(
    tmp_path: Path,
) -> None:
    export = _build_asset_export(tmp_path)
    copied = tmp_path / "copied"
    referenced = tmp_path / "referenced"
    zipped_reference = tmp_path / "zipped-reference"
    archive = tmp_path / "asset_export.zip"
    _zip_tree(export, archive)

    copied_record = _run(export, copied)["record"]
    reference_record = _run(export, referenced, asset_mode="reference")["record"]
    zip_record = _run(archive, zipped_reference, asset_mode="reference")["record"]
    assert copied_record["source_export"]["asset_mode"] == "self_contained"
    assert copied_record["source_export"]["copied_asset_count"] == 1
    assert next((copied / "extraction").glob("*_assets/images/alpha.png")).is_file()
    assert reference_record["source_export"]["copied_asset_count"] == 0
    assert not list((referenced / "extraction").glob("*_assets"))
    assert zip_record["source_export"]["intake_kind"] == "zip"
    assert zip_record["source_export"]["copied_asset_count"] == 0
    assert not list((zipped_reference / "extraction").glob("*_assets"))
    assert verify_unbind_run(zipped_reference)["status"] == "verified"


def test_folder_zip_parity_and_explicit_lineage(tmp_path: Path) -> None:
    archive = tmp_path / "quiz_only.zip"
    _zip_tree(QUIZ_ONLY, archive)
    lineage = "cc:lineage:fixture:folder-zip-parity"
    folder = _run(QUIZ_ONLY, tmp_path / "folder", source_lineage_key=lineage)["record"]
    zipped = _run(archive, tmp_path / "zip", source_lineage_key=lineage)["record"]
    assert folder["source_export"]["intake_kind"] == "directory"
    assert zipped["source_export"]["intake_kind"] == "zip"
    assert folder["source_export"]["source_lineage_key"] == lineage
    assert zipped["source_export"]["source_lineage_key"] == lineage
    assert folder["model_summary"] == zipped["model_summary"]


def test_review_counts_match_across_receipt_workbook_status_station_and_reading_room(
    tmp_path: Path,
) -> None:
    output = tmp_path / "run"
    result = _run(DUPLICATE_MATCH, output)
    record = result["record"]
    projection = json.loads(
        (output / "review" / "quiz_review_projection.json").read_text(
            encoding="utf-8"
        )
    )
    status_station = json.loads(
        (output / "station" / "station.json").read_text(encoding="utf-8")
    )
    reading_room = json.loads(
        (output / "reading_room" / "station.json").read_text(encoding="utf-8")
    )
    reviewer_workbook_path = next(
        (output / "extraction").glob("*__quiz_pool_review_reviewer.xlsx")
    )
    workbook = load_workbook(reviewer_workbook_path, read_only=True)
    occurrence_sheet = workbook["Quiz Occurrences"]
    entity_sheet = workbook["Question Entities"]
    occurrence_headers = [cell.value for cell in occurrence_sheet[1]]
    occurrence_rows = [
        dict(zip(occurrence_headers, values))
        for values in occurrence_sheet.iter_rows(min_row=2, values_only=True)
    ]

    assert record["model_summary"]["question_entity_count"] == 2
    assert record["model_summary"]["question_occurrence_count"] == 2
    assert projection["summary"]["unique_question_count"] == 2
    assert projection["summary"]["question_occurrence_count"] == 2
    assert reading_room["summary"]["unique_question_count"] == 2
    assert reading_room["summary"]["occurrence_count"] == 2
    assert len(status_station["questions"]) == 2
    assert entity_sheet.max_row - 1 == 2
    assert occurrence_sheet.max_row - 1 == 2

    workbook_labels = {row["library_match_status"] for row in occurrence_rows}
    room_labels = {
        row["match_status"] for row in reading_room["question_occurrences"]
    }
    assert workbook_labels == room_labels == {
        "Linked by identical content",
        "Possible link based on similar content",
    }
    assert status_station["queues"]["summary"]["probable_match_count"] == 2
    assert reading_room["queue_summary"]["probable_match_count"] == 2
    exact = next(
        row for row in occurrence_rows if row["library_match_code"] == "inferred_exact"
    )
    assert "dup-1 (score 1.0)" in exact["library_match_candidates"]
    assert "dup-2 (score 1.0)" in exact["library_match_candidates"]
    assert "Binder will not guess" in exact["library_match_explanation"]

    # Readiness is a separate build-support claim; it does not invalidate the
    # completed extraction-only close.
    assert record["readiness"]["ready"] is False
    assert record["readiness"]["blocker_codes"]
    assert result["verification"]["status"] == "verified"
    assert result["verification"]["evidence_level"] == "extraction_only"


def test_multi_quiz_run_preserves_aggregate_stop_and_allows_selection(
    tmp_path: Path,
) -> None:
    export = tmp_path / "multi_quiz"
    shutil.copytree(QUIZ_ONLY, export)
    original = (export / "quiz_d2l_only.xml").read_text(encoding="utf-8")
    second = original
    for before, after in (
        ("QUIZ_ONLY", "QUIZ_SECOND"),
        ("ONLY_SECTION", "SECOND_SECTION"),
        ("ONLY_OBJECT", "SECOND_OBJECT"),
        ("ONLY_Q", "SECOND_Q"),
        ("ONLY_LID", "SECOND_LID"),
        ("ONLY_TRUE", "SECOND_TRUE"),
        ("ONLY_FALSE", "SECOND_FALSE"),
        ("ONLY001", "SECOND001"),
    ):
        second = second.replace(before, after)
    (export / "quiz_d2l_second.xml").write_text(second, encoding="utf-8")

    aggregate_output = tmp_path / "aggregate"
    aggregate = _run(export, aggregate_output)["record"]
    assert aggregate["model_summary"]["quiz_count"] == 2
    assert aggregate["selection"]["selected_quiz_entity_key"] is None
    assert "quiz_selection_required" in aggregate["readiness"]["blocker_codes"]
    model = json.loads(
        next((aggregate_output / "extraction").glob("*.model.json")).read_text(
            encoding="utf-8"
        )
    )
    selected_key = model["quizzes"][0]["entity_key"]

    selected = _run(
        export,
        tmp_path / "selected",
        quiz_entity_key=selected_key,
    )["record"]
    assert selected["selection"] == {
        "requested_quiz_entity_key": selected_key,
        "selected_quiz_entity_key": selected_key,
    }
    assert "quiz_selection_required" not in selected["readiness"]["blocker_codes"]
