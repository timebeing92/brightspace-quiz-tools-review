from __future__ import annotations

from copy import deepcopy
import hashlib
import json
from pathlib import Path
import sys
import xml.etree.ElementTree as ET

from openpyxl import load_workbook
import pytest


REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT))
sys.path.insert(0, str(REPO_ROOT / "scripts"))

from scripts.extract_quiz_pool_review import (  # noqa: E402
    build_payload,
    extract_correct_answer,
    extract_fill_in_blank_structure,
    extract_source_response_facts,
    build_reviewer_question_rows,
    write_reviewer_workbook,
)
from scripts.quiz_authoring_readiness import analyze_authoring_readiness  # noqa: E402
from scripts.quiz_build_support import load_authoring_projection  # noqa: E402
from scripts.quiz_normalization import (  # noqa: E402
    SHORT_ANSWER_AWARD,
    _enrich_short_answer_question,
    _recognize_short_answer_occurrence,
    _type_payload,
    build_normalized_model,
)
from scripts.quiz_promote_revisions import promote_revisions  # noqa: E402
from scripts.quiz_review_workbook_reingest import (  # noqa: E402
    FORMAT as OVERLAY_FORMAT,
    materialize_workbook_decisions,
)


FIXTURE = (
    REPO_ROOT
    / "tests"
    / "fixtures"
    / "quiz_xml"
    / "short_answer_projection"
)
SPECIMEN = (
    REPO_ROOT
    / "workspace"
    / "review"
    / "quiz_capability_lab_r1"
    / "fixtures"
    / "specimens"
    / "sa_short_answer.xml"
)
AUTHORING_FIXTURE = REPO_ROOT / "tests" / "fixtures" / "quiz_authoring"
REGISTRY = (
    REPO_ROOT
    / "workspace"
    / "reference"
    / "schemas"
    / "quiz"
    / "quiz_build_capabilities.json"
)


def local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def descendants(element: ET.Element, name: str) -> list[ET.Element]:
    return [node for node in element.iter() if local_name(node.tag) == name]


def build_fixture() -> tuple[dict[str, object], dict[str, object]]:
    payload = build_payload(FIXTURE)
    model, _receipt = build_normalized_model(
        payload,
        FIXTURE,
        FIXTURE,
        "directory",
        source_lineage_key="cc:lineage:fixture:short-answer-projection",
    )
    return payload, model


def alias_value(question: dict[str, object], namespace: str) -> str | None:
    return next(
        (
            str(alias["value"])
            for alias in question["identity"]["source_aliases"]
            if alias["namespace"] == namespace
        ),
        None,
    )


def question_from_item(
    item: ET.Element,
    *,
    question_weight: object = "1.875000000",
) -> tuple[dict[str, object], list[dict[str, object]]]:
    answer = extract_correct_answer(item)
    blank = extract_fill_in_blank_structure(item)
    row = {
        "quiz_file": "quiz_d2l_manual_short_answer.xml",
        "quiz_item_number": 1,
        "question_type": "Short Answer",
        "question_weight": question_weight,
        "correct_answer": answer["correct_answer"],
        "correct_response_ids": answer["correct_response_ids"],
        "correct_answer_basis": answer["correct_answer_basis"],
        "choices_text": "",
        "fill_in_blank_count": blank["blank_count"],
        "fill_in_blank_answers_text": blank["blank_answers_text"],
        "source_response_facts": extract_source_response_facts(item),
    }
    evidence_key = "ev.occurrence.manual-short-answer"
    question = {
        "kind": "short_answer",
        "source_kind": "Short Answer",
        "type_payload": _type_payload(row, [evidence_key]),
        "scoring": {
            "state": "known",
            "mode": "unknown",
            "maximum_points": row["question_weight"],
            "rules": [],
            "extensions": {},
        },
        "feedback": [],
    }
    evidence = [
        {
            "kind": "xml_element",
            "evidence_key": evidence_key,
            "source_ref": row["quiz_file"],
            "locator": "quiz item 1",
        }
    ]
    return question, evidence


def manual_question(
    values: list[str],
) -> tuple[dict[str, object], list[dict[str, object]]]:
    item = ET.parse(SPECIMEN).getroot()
    response_processing = descendants(item, "resprocessing")[0]
    conditions = [
        child
        for child in list(response_processing)
        if local_name(child.tag) == "respcondition"
    ]
    accepted_template = conditions[0]
    aggregator = conditions[-1]
    descendants(accepted_template, "varequal")[0].text = values[0]
    for value in values[1:]:
        condition = deepcopy(accepted_template)
        descendants(condition, "varequal")[0].text = value
        response_processing.insert(list(response_processing).index(aggregator), condition)
    return question_from_item(item)


def fact_record(question: dict[str, object]) -> dict[str, object]:
    return next(
        record
        for record in question["type_payload"]["raw_response_models"]
        if record["source_kind"] == "d2l_qti_response_facts/0"
    )


def source_facts(question: dict[str, object]) -> dict[str, object]:
    return fact_record(question)["payload"]


def fact_nodes(value: object):
    if isinstance(value, dict):
        if set(value) == {
            "qualified_name",
            "namespace_uri",
            "name",
            "attributes",
            "raw_text",
            "raw_tail",
            "children",
        }:
            yield value
        for child in value.values():
            yield from fact_nodes(child)
    elif isinstance(value, list):
        for child in value:
            yield from fact_nodes(child)


def named_node(value: object, name: str, occurrence: int = 0) -> dict[str, object]:
    return [
        node
        for node in fact_nodes(value)
        if node["name"] == name
    ][occurrence]


def add_attribute(
    node: dict[str, object],
    name: str,
    value: str,
    namespace_uri: str | None = None,
) -> None:
    qualified_name = (
        f"{{{namespace_uri}}}{name}" if namespace_uri is not None else name
    )
    node["attributes"].append(
        {
            "qualified_name": qualified_name,
            "namespace_uri": namespace_uri,
            "name": name,
            "raw_value": value,
        }
    )


def remove_attribute(node: dict[str, object], name: str) -> None:
    node["attributes"] = [
        attribute
        for attribute in node["attributes"]
        if attribute["name"] != name
    ]


def rename_node(node: dict[str, object], name: str) -> None:
    node["qualified_name"] = name
    node["namespace_uri"] = None
    node["name"] = name


def assert_atomic_refusal(
    question: dict[str, object],
    evidence: list[dict[str, object]],
) -> None:
    before = deepcopy(question)
    assert _enrich_short_answer_question(question, evidence) is False
    assert question == before


def two_occurrence_question(
    second_values: list[str] | None = None,
) -> tuple[
    dict[str, object],
    list[dict[str, object]],
]:
    question, evidence = manual_question(["alpha"])
    second_question, _second_evidence = manual_question(
        second_values or ["alpha"]
    )
    second_records = deepcopy(
        second_question["type_payload"]["raw_response_models"]
    )
    second_key = "ev.occurrence.manual-short-answer-second"
    second_records[0]["payload"]["quiz_item_number"] = 2
    for record in second_records:
        record["source_evidence_keys"] = [second_key]
    question["type_payload"]["raw_response_models"].extend(second_records)
    evidence.append(
        {
            "kind": "xml_element",
            "evidence_key": second_key,
            "source_ref": "quiz_d2l_manual_short_answer.xml",
            "locator": "quiz item 2",
        }
    )
    return question, evidence


def mutate_fact_boundary(
    question: dict[str, object],
    mutation: str,
) -> None:
    facts = source_facts(question)
    presentation = facts["presentation"]["occurrences"][0]["node"]
    flow = named_node(presentation, "flow")
    response_str = named_node(presentation, "response_str")
    renderer = named_node(response_str, "render_fib")
    response_extension = named_node(presentation, "response_extension")
    grading = named_node(response_extension, "grading_type")
    processing = facts["response_processing"]["occurrences"][0]["node"]
    outcomes = named_node(processing, "outcomes")
    decvar = named_node(outcomes, "decvar")
    conditions = [
        child
        for child in processing["children"]
        if child["name"] == "respcondition"
    ]
    accepted = conditions[0]
    accepted_conditionvar = named_node(accepted, "conditionvar")
    accepted_varequal = named_node(accepted_conditionvar, "varequal")
    accepted_setvar = named_node(accepted, "setvar")
    aggregator = conditions[-1]
    aggregator_setvar = named_node(aggregator, "setvar")

    if mutation == "case_yes":
        next(
            attribute
            for attribute in accepted_varequal["attributes"]
            if attribute["name"] == "case"
        )["raw_value"] = "yes"
    elif mutation == "case_missing":
        remove_attribute(accepted_varequal, "case")
    elif mutation == "award_token":
        accepted_setvar["raw_text"] = "100"
    elif mutation == "award_action":
        remove_attribute(accepted_setvar, "action")
    elif mutation == "award_multiple":
        accepted["children"].append(deepcopy(accepted_setvar))
    elif mutation == "award_missing":
        accepted["children"].remove(accepted_setvar)
    elif mutation == "grading_missing":
        response_extension["children"].remove(grading)
    elif mutation == "grading_repeated":
        response_extension["children"].append(deepcopy(grading))
    elif mutation == "grading_nonzero":
        grading["raw_text"] = "1"
    elif mutation == "compound_predicate":
        accepted_conditionvar["children"].append(deepcopy(accepted_varequal))
    elif mutation == "negated_predicate":
        accepted_conditionvar["children"].remove(accepted_varequal)
        wrapper = deepcopy(accepted_varequal)
        rename_node(wrapper, "not")
        wrapper["raw_text"] = None
        wrapper["attributes"] = []
        wrapper["children"] = [accepted_varequal]
        accepted_conditionvar["children"].append(wrapper)
    elif mutation == "cardinality":
        next(
            attribute
            for attribute in response_str["attributes"]
            if attribute["name"] == "rcardinality"
        )["raw_value"] = "Multiple"
    elif mutation == "declaration_missing":
        flow["children"].remove(response_str)
    elif mutation == "declaration_repeated":
        flow["children"].append(deepcopy(response_str))
    elif mutation == "response_label_repeated":
        renderer["children"].append(deepcopy(renderer["children"][0]))
    elif mutation == "input_boxes":
        add_attribute(
            renderer,
            "input_boxes",
            "2",
            "http://desire2learn.com/xsd/d2lcp_v2p0",
        )
    elif mutation == "renderer":
        next(
            attribute
            for attribute in renderer["attributes"]
            if attribute["name"] == "rows"
        )["raw_value"] = "2"
    elif mutation == "outcome_missing":
        processing["children"].remove(outcomes)
    elif mutation == "outcome_different":
        next(
            attribute
            for attribute in decvar["attributes"]
            if attribute["name"] == "varname"
        )["raw_value"] = "Different"
    elif mutation == "aggregator_missing":
        processing["children"].remove(aggregator)
    elif mutation == "aggregator_repeated":
        processing["children"].append(deepcopy(aggregator))
    elif mutation == "aggregator_different":
        aggregator_setvar["raw_text"] = "Different"
    elif mutation == "diagnostic":
        facts["diagnostics"].append(
            {"code": "synthetic_refusal_oracle", "severity": "warning"}
        )
    elif mutation == "feedback_slot":
        facts["feedback_slots"].append(
            {
                "ordinal": 1,
                "source_ident": "synthetic-feedback",
                "has_content": True,
            }
        )
    elif mutation == "metadata_member_missing":
        facts["item_metadata"].pop("qmd_weighting")
    elif mutation == "metadata_state_missing":
        facts["item_metadata"]["qmd_weighting"]["state"] = "absent"
    elif mutation == "questiontype_repeated":
        field = facts["item_metadata"]["qmd_questiontype"]
        duplicate = deepcopy(field["occurrences"][0])
        duplicate["ordinal"] = 2
        field["occurrences"].append(duplicate)
    elif mutation == "questiontype_conflicting":
        field = facts["item_metadata"]["qmd_questiontype"]
        duplicate = deepcopy(field["occurrences"][0])
        duplicate["ordinal"] = 2
        duplicate["raw_value"] = "Long Answer"
        named_node(duplicate["node"], "fieldentry")["raw_text"] = "Long Answer"
        field["occurrences"].append(duplicate)
        field["raw_value"] = None
    elif mutation == "presentation_attr":
        add_attribute(presentation, "future", "1")
    elif mutation == "flow_attr":
        add_attribute(flow, "future", "1")
    elif mutation == "second_presentation":
        duplicate = deepcopy(facts["presentation"]["occurrences"][0])
        duplicate["ordinal"] = 2
        facts["presentation"]["occurrences"].append(duplicate)
    elif mutation == "second_flow":
        presentation["children"].append(deepcopy(flow))
    elif mutation == "processing_attr":
        add_attribute(processing, "future", "1")
    elif mutation == "second_processing":
        duplicate = deepcopy(facts["response_processing"]["occurrences"][0])
        duplicate["ordinal"] = 2
        duplicate["node"]["children"] = []
        facts["response_processing"]["occurrences"].append(duplicate)
    elif mutation == "second_outcomes":
        processing["children"].insert(1, deepcopy(outcomes))
    elif mutation == "outcomes_attr":
        add_attribute(outcomes, "future", "1")
    elif mutation == "processing_reordered":
        processing["children"][0], processing["children"][1] = (
            processing["children"][1],
            processing["children"][0],
        )
    elif mutation == "processing_extra_child":
        extra = deepcopy(outcomes)
        rename_node(extra, "future_processing_child")
        processing["children"].append(extra)
    elif mutation == "regex":
        extension = deepcopy(accepted_varequal)
        rename_node(extension, "var_extension")
        extension["raw_text"] = None
        extension["attributes"] = []
        child = deepcopy(accepted_varequal)
        rename_node(child, "answer_is_regexp")
        child["raw_text"] = "1"
        child["attributes"] = []
        child["children"] = []
        extension["children"] = [child]
        accepted_conditionvar["children"].append(extension)
    elif mutation == "zero_accepted":
        processing["children"] = [outcomes, aggregator]
    elif mutation == "six_accepted":
        processing["children"][-1:-1] = [
            deepcopy(accepted)
            for _index in range(5)
        ]
    else:
        raise AssertionError(f"Unknown mutation: {mutation}")


def metadata_field(item: ET.Element, field_name: str) -> ET.Element:
    return next(
        field
        for field in descendants(item, "qti_metadatafield")
        if any(
            local_name(child.tag) == "fieldlabel"
            and (child.text or "").strip() == field_name
            for child in list(field)
        )
    )


def mutate_native_item(item: ET.Element, mutation: str) -> None:
    processing = descendants(item, "resprocessing")[0]
    conditions = [
        child
        for child in list(processing)
        if local_name(child.tag) == "respcondition"
    ]
    accepted = conditions[0]
    aggregator = conditions[-1]
    conditionvar = descendants(accepted, "conditionvar")[0]
    varequal = descendants(accepted, "varequal")[0]
    award = descendants(accepted, "setvar")[0]
    grading = descendants(item, "grading_type")[0]
    outcomes = descendants(processing, "outcomes")[0]
    decvar = descendants(outcomes, "decvar")[0]
    response_str = descendants(item, "response_str")[0]
    response_label = descendants(response_str, "response_label")[0]

    if mutation == "case_yes":
        varequal.attrib["case"] = "yes"
    elif mutation == "case_missing":
        varequal.attrib.pop("case")
    elif mutation == "award_token":
        award.text = "100"
    elif mutation == "award_whitespace":
        award.text = " 100.000000000 "
    elif mutation == "award_action":
        award.attrib["action"] = "Add"
    elif mutation == "grading_missing":
        parent = descendants(item, "response_extension")[0]
        parent.remove(grading)
    elif mutation == "grading_nonzero":
        grading.text = "1"
    elif mutation == "grading_whitespace":
        grading.text = " 0 "
    elif mutation == "compound_predicate":
        conditionvar.append(deepcopy(varequal))
    elif mutation == "negated_predicate":
        conditionvar.remove(varequal)
        wrapper = ET.Element("not")
        wrapper.append(varequal)
        conditionvar.append(wrapper)
    elif mutation == "outcome_missing":
        processing.remove(outcomes)
    elif mutation == "outcome_different":
        decvar.attrib["varname"] = "Different"
    elif mutation == "aggregator_missing":
        processing.remove(aggregator)
    elif mutation == "aggregator_different":
        descendants(aggregator, "setvar")[0].text = "Different"
    elif mutation == "wrong_respident":
        varequal.attrib["respident"] = "WRONG_RESPONSE"
    elif mutation == "missing_declaration_id":
        response_str.attrib.pop("ident")
    elif mutation == "missing_label_id":
        response_label.attrib.pop("ident")
    elif mutation == "extra_setvar_varname":
        award.attrib["varname"] = "Unexpected"
    elif mutation == "decvar_extra_attribute":
        decvar.attrib["future"] = "1"
    elif mutation == "decvar_text":
        decvar.text = "unexpected"
    elif mutation == "feedback":
        feedback = ET.SubElement(item, "itemfeedback", {"ident": "FEEDBACK"})
        material = ET.SubElement(feedback, "material")
        ET.SubElement(material, "mattext", {"texttype": "text/html"}).text = (
            "<p>Feedback.</p>"
        )
    elif mutation == "metadata_duplicate":
        metadata = descendants(item, "qtimetadata")[0]
        metadata.append(deepcopy(metadata_field(item, "qmd_questiontype")))
    elif mutation == "metadata_malformed":
        field = metadata_field(item, "qmd_computerscored")
        ET.SubElement(field, "fieldlabel").text = "qmd_computerscored"
    elif mutation == "weight_whitespace":
        field = metadata_field(item, "qmd_weighting")
        next(
            child
            for child in list(field)
            if local_name(child.tag) == "fieldentry"
        ).text = " 1.875000000 "
    else:
        raise AssertionError(f"Unknown native mutation: {mutation}")


def test_full_export_agreement_enriches_and_occurrence_conflict_refuses() -> None:
    payload, model = build_fixture()
    questions = {
        alias_value(question, "d2l.qmd_globalid"): question
        for question in model["questions"]
    }

    agreeing = questions["fixture-sa-agree-global"]
    assert [row["value"] for row in agreeing["type_payload"]["accepted_responses"]] == [
        "[ALPHA]",
        "Crème: β",
        "ratio/term",
        "A & B",
        "left / right?!",
    ]
    assert {
        row["case_sensitive"]
        for row in agreeing["type_payload"]["accepted_responses"]
    } == {False}
    assert {
        row["weight"]
        for row in agreeing["type_payload"]["accepted_responses"]
    } == {SHORT_ANSWER_AWARD}
    assert agreeing["type_payload"]["blanks"] == []
    assert agreeing["scoring"]["mode"] == "exact"
    assert agreeing["scoring"]["maximum_points"] == "2.500000000"

    compact = questions["fixture-sa-compact-global"]
    assert compact["scoring"]["mode"] == "exact"
    assert compact["type_payload"]["blanks"] == []

    conflict = questions["fixture-sa-conflict-global"]
    assert conflict["scoring"]["mode"] == "unknown"
    assert conflict["type_payload"]["accepted_responses"][0][
        "case_sensitive"
    ] is None
    assert conflict["type_payload"]["accepted_responses"][0]["weight"] is None
    assert len(conflict["type_payload"]["blanks"]) == 1

    agreeing_rows = payload["quiz_question_rows"][:2]
    agreeing_models = agreeing["type_payload"]["raw_response_models"]
    legacy = [
        record
        for record in agreeing_models
        if record["source_kind"] == "canonical-extractor-row:Short Answer"
    ]
    facts = [
        record
        for record in agreeing_models
        if record["source_kind"] == "d2l_qti_response_facts/0"
    ]
    assert [record["payload"]["correct_answer"] for record in legacy] == [
        row["correct_answer"] for row in agreeing_rows
    ]
    assert [record["payload"] for record in facts] == [
        row["source_response_facts"] for row in agreeing_rows
    ]
    assert all(
        "source_response_facts" not in record["payload"]
        for record in legacy
    )
    assert len(
        [
            key
            for key in legacy[0]["source_evidence_keys"]
            if key.startswith("ev.occurrence.")
        ]
    ) == 2
    assert all(
        len(
            [
                key
                for key in record["source_evidence_keys"]
                if key.startswith("ev.occurrence.")
            ]
        )
        == 1
        for record in facts
    )
    assert not any(
        isinstance(value, str)
        and value.startswith(("/Users/", "/private/", "/tmp/"))
        for record in agreeing_models
        for value in record["payload"].values()
    )
    assert all(record["payload"]["diagnostics"] == [] for record in facts)
    assert agreeing["diagnostic_ids"] == []


@pytest.mark.parametrize("count", [1, 2, 3, 4, 5])
def test_one_to_five_native_synonyms_project_without_text_loss(count: int) -> None:
    source_values = [
        "[ONE]",
        "Crème: β",
        "ratio/term",
        "A & B",
        "left / right?!",
    ][:count]
    question, evidence = manual_question(source_values)

    assert _enrich_short_answer_question(question, evidence) is True
    assert [row["value"] for row in question["type_payload"]["accepted_responses"]] == (
        source_values
    )
    assert all(
        row["case_sensitive"] is False
        and row["weight"] == SHORT_ANSWER_AWARD
        for row in question["type_payload"]["accepted_responses"]
    )
    assert question["type_payload"]["blanks"] == []
    assert question["scoring"]["mode"] == "exact"


def test_spaced_slash_is_literal_but_legacy_double_pipe_is_ambiguous() -> None:
    slash, slash_evidence = manual_question(["left / right"])
    assert _enrich_short_answer_question(slash, slash_evidence) is True
    assert slash["type_payload"]["accepted_responses"][0]["value"] == "left / right"

    ambiguous, ambiguous_evidence = manual_question(["left || right"])
    assert_atomic_refusal(ambiguous, ambiguous_evidence)


def test_projection_only_fills_null_fields_and_clears_proven_legacy_blank() -> None:
    question, evidence = manual_question(["alpha", "beta"])
    accepted = question["type_payload"]["accepted_responses"]
    for response in accepted:
        response["case_sensitive"] = False
        response["weight"] = SHORT_ANSWER_AWARD
    question["scoring"]["mode"] = "exact"
    assert _enrich_short_answer_question(question, evidence) is True
    assert question["type_payload"]["blanks"] == []

    mismatch_mutations = [
        lambda row: row["type_payload"]["accepted_responses"][0].__setitem__(
            "case_sensitive", True
        ),
        lambda row: row["type_payload"]["accepted_responses"][0].__setitem__(
            "weight", "100"
        ),
        lambda row: row["type_payload"]["blanks"][0]["accepted_responses"][0].__setitem__(
            "value", "alpha"
        ),
    ]
    for mutate in mismatch_mutations:
        candidate, candidate_evidence = manual_question(["alpha", "beta"])
        mutate(candidate)
        assert_atomic_refusal(candidate, candidate_evidence)


def test_effective_question_points_are_preserved_independently_of_qmd_weighting() -> None:
    question, evidence = manual_question(["alpha"])
    question["scoring"]["maximum_points"] = "9.500000000"

    assert _enrich_short_answer_question(question, evidence) is True
    assert question["scoring"]["maximum_points"] == "9.500000000"
    assert source_facts(question)["item_metadata"]["qmd_weighting"][
        "raw_value"
    ] == "1.875000000"


@pytest.mark.parametrize(
    "mutation",
    [
        "case_yes",
        "case_missing",
        "award_token",
        "award_action",
        "award_multiple",
        "award_missing",
        "grading_missing",
        "grading_repeated",
        "grading_nonzero",
        "compound_predicate",
        "negated_predicate",
        "cardinality",
        "declaration_missing",
        "declaration_repeated",
        "response_label_repeated",
        "input_boxes",
        "renderer",
        "outcome_missing",
        "outcome_different",
        "aggregator_missing",
        "aggregator_repeated",
        "aggregator_different",
        "diagnostic",
        "feedback_slot",
        "metadata_member_missing",
        "metadata_state_missing",
        "questiontype_repeated",
        "questiontype_conflicting",
        "presentation_attr",
        "flow_attr",
        "second_presentation",
        "second_flow",
        "processing_attr",
        "second_processing",
        "second_outcomes",
        "outcomes_attr",
        "processing_reordered",
        "processing_extra_child",
        "regex",
        "zero_accepted",
        "six_accepted",
    ],
)
def test_full_parent_fact_boundary_refuses_atomically(mutation: str) -> None:
    candidate, evidence = manual_question(["alpha"])
    mutate_fact_boundary(candidate, mutation)
    assert _recognize_short_answer_occurrence(source_facts(candidate)) is None
    assert_atomic_refusal(candidate, evidence)


@pytest.mark.parametrize(
    "mutation",
    [
        "case_yes",
        "case_missing",
        "award_token",
        "award_whitespace",
        "award_action",
        "grading_missing",
        "grading_nonzero",
        "grading_whitespace",
        "compound_predicate",
        "negated_predicate",
        "outcome_missing",
        "outcome_different",
        "aggregator_missing",
        "aggregator_different",
        "wrong_respident",
        "missing_declaration_id",
        "missing_label_id",
        "extra_setvar_varname",
        "decvar_extra_attribute",
        "decvar_text",
        "feedback",
        "metadata_duplicate",
        "metadata_malformed",
        "weight_whitespace",
    ],
)
def test_coherent_native_source_mutations_refuse_atomically(
    mutation: str,
) -> None:
    item = ET.parse(SPECIMEN).getroot()
    mutate_native_item(item, mutation)
    candidate, evidence = question_from_item(item)

    assert _recognize_short_answer_occurrence(source_facts(candidate)) is None
    assert_atomic_refusal(candidate, evidence)


@pytest.mark.parametrize(
    "values",
    [
        ["duplicate", "duplicate"],
        [" edge-padded "],
        ["one", "two", "three", "four", "five", "six"],
    ],
)
def test_coherent_literal_duplicate_whitespace_and_six_values_refuse(
    values: list[str],
) -> None:
    candidate, evidence = manual_question(values)

    assert _recognize_short_answer_occurrence(source_facts(candidate)) is None
    assert_atomic_refusal(candidate, evidence)


@pytest.mark.parametrize(
    "mutation",
    [
        "accepted_value",
        "accepted_order",
        "accepted_count",
        "scoring_mode",
        "scoring_rules",
        "scoring_shape",
        "typed_feedback",
        "blank_key",
        "blank_metadata",
        "blank_extensions",
        "blank_evidence",
        "blank_response_extensions",
    ],
)
def test_projection_boundary_refuses_prepopulated_conflicts_atomically(
    mutation: str,
) -> None:
    candidate, evidence = manual_question(["alpha", "beta"])
    accepted = candidate["type_payload"]["accepted_responses"]
    blank = candidate["type_payload"]["blanks"][0]
    if mutation == "accepted_value":
        accepted[0]["value"] = "normalized-alpha"
    elif mutation == "accepted_order":
        accepted.reverse()
    elif mutation == "accepted_count":
        accepted.pop()
    elif mutation == "scoring_mode":
        candidate["scoring"]["mode"] = "all_or_nothing"
    elif mutation == "scoring_rules":
        candidate["scoring"]["rules"] = [{"kind": "unrelated"}]
    elif mutation == "scoring_shape":
        candidate["scoring"]["future"] = True
    elif mutation == "typed_feedback":
        candidate["feedback"] = [{"channel": "general"}]
    elif mutation == "blank_key":
        blank["blank_key"] = "blank-2"
    elif mutation == "blank_metadata":
        blank["accepted_responses"][0]["case_sensitive"] = False
    elif mutation == "blank_extensions":
        blank["extensions"]["future"] = True
    elif mutation == "blank_evidence":
        blank["source_evidence_keys"] = ["ev.occurrence.wrong"]
    elif mutation == "blank_response_extensions":
        blank["accepted_responses"][0]["extensions"]["future"] = True
    assert_atomic_refusal(candidate, evidence)


@pytest.mark.parametrize(
    ("target", "value"),
    [
        ("case", 0),
        ("case", []),
        ("case", {}),
        ("weight", []),
        ("weight", {}),
        ("mode", []),
        ("mode", {}),
        ("maximum_points", []),
        ("maximum_points", {}),
    ],
)
def test_untrusted_typed_values_fail_closed_without_type_errors(
    target: str,
    value: object,
) -> None:
    candidate, evidence = manual_question(["alpha"])
    if target == "case":
        candidate["type_payload"]["accepted_responses"][0][
            "case_sensitive"
        ] = value
    elif target == "weight":
        candidate["type_payload"]["accepted_responses"][0]["weight"] = value
    elif target == "mode":
        candidate["scoring"]["mode"] = value
    else:
        candidate["scoring"]["maximum_points"] = value

    assert_atomic_refusal(candidate, evidence)


@pytest.mark.parametrize("malformed_record", [[], {}])
def test_malformed_raw_records_fail_closed_without_type_errors(
    malformed_record: object,
) -> None:
    candidate, evidence = manual_question(["alpha"])
    candidate["type_payload"]["raw_response_models"].append(malformed_record)

    assert_atomic_refusal(candidate, evidence)


@pytest.mark.parametrize("malformed_key", [[], {}])
def test_malformed_raw_evidence_keys_fail_closed_without_type_errors(
    malformed_key: object,
) -> None:
    candidate, evidence = manual_question(["alpha"])
    fact_record(candidate)["source_evidence_keys"].append(malformed_key)

    assert_atomic_refusal(candidate, evidence)


@pytest.mark.parametrize(
    "fact_path",
    [
        "metadata",
        "presentation",
        "grading",
        "outcome",
        "condition",
        "response_label",
    ],
)
def test_boolean_true_never_satisfies_exact_occurrence_ordinal(
    fact_path: str,
) -> None:
    candidate, evidence = manual_question(["alpha"])
    facts = source_facts(candidate)
    if fact_path == "metadata":
        facts["item_metadata"]["qmd_weighting"]["occurrences"][0][
            "ordinal"
        ] = True
    elif fact_path == "presentation":
        facts["presentation"]["occurrences"][0]["ordinal"] = True
    elif fact_path == "grading":
        facts["grading_type"]["occurrences"][0]["ordinal"] = True
    elif fact_path == "outcome":
        facts["outcomes"][0]["ordinal"] = True
    elif fact_path == "condition":
        facts["response_conditions"][0]["ordinal"] = True
    else:
        facts["response_declarations"][0]["response_labels"][0][
            "ordinal"
        ] = True

    assert _recognize_short_answer_occurrence(facts) is None
    assert_atomic_refusal(candidate, evidence)


@pytest.mark.parametrize(
    ("field_name", "raw_value"),
    [
        ("qmd_computerscored", "no"),
        ("qmd_weighting", "NaN"),
        ("qmd_weighting", "0"),
        ("qmd_weighting", "-1"),
    ],
)
def test_invalid_item_metadata_values_refuse_atomically(
    field_name: str,
    raw_value: str,
) -> None:
    candidate, evidence = manual_question(["alpha"])
    field = source_facts(candidate)["item_metadata"][field_name]
    field["raw_value"] = raw_value
    field["occurrences"][0]["raw_value"] = raw_value
    named_node(field["occurrences"][0]["node"], "fieldentry")[
        "raw_text"
    ] = raw_value
    assert _recognize_short_answer_occurrence(source_facts(candidate)) is None
    assert_atomic_refusal(candidate, evidence)


def test_occurrence_pairing_refuses_missing_duplicate_and_mislocated_facts() -> None:
    candidate, evidence = manual_question(["alpha"])
    raw_models = candidate["type_payload"]["raw_response_models"]
    raw_models.remove(fact_record(candidate))
    assert_atomic_refusal(candidate, evidence)

    candidate, evidence = manual_question(["alpha"])
    raw_models = candidate["type_payload"]["raw_response_models"]
    raw_models.append(deepcopy(fact_record(candidate)))
    assert_atomic_refusal(candidate, evidence)

    candidate, evidence = manual_question(["alpha"])
    fact_record(candidate)["source_evidence_keys"] = ["ev.occurrence.unmatched"]
    assert_atomic_refusal(candidate, evidence)


@pytest.mark.parametrize(
    "mutation",
    [
        "missing_legacy",
        "duplicate_legacy",
        "duplicate_evidence_locator",
        "partial_facts",
        "cross_spliced_fact_evidence",
        "cross_spliced_fact_payload",
        "semantic_weight_conflict",
    ],
)
def test_two_occurrence_pairing_failures_refuse_atomically(
    mutation: str,
) -> None:
    candidate, evidence = two_occurrence_question(
        ["beta"]
        if mutation.startswith("cross_spliced")
        else None
    )
    raw_models = candidate["type_payload"]["raw_response_models"]
    if mutation == "missing_legacy":
        raw_models.pop(2)
    elif mutation == "duplicate_legacy":
        raw_models.insert(2, deepcopy(raw_models[2]))
    elif mutation == "duplicate_evidence_locator":
        evidence.append(
            {
                "kind": "xml_element",
                "evidence_key": "ev.occurrence.duplicate-locator",
                "source_ref": "quiz_d2l_manual_short_answer.xml",
                "locator": "quiz item 2",
            }
        )
    elif mutation == "partial_facts":
        raw_models.pop(3)
    elif mutation == "cross_spliced_fact_evidence":
        (
            raw_models[1]["source_evidence_keys"],
            raw_models[3]["source_evidence_keys"],
        ) = (
            raw_models[3]["source_evidence_keys"],
            raw_models[1]["source_evidence_keys"],
        )
    elif mutation == "cross_spliced_fact_payload":
        raw_models[1]["payload"], raw_models[3]["payload"] = (
            raw_models[3]["payload"],
            raw_models[1]["payload"],
        )
    elif mutation == "semantic_weight_conflict":
        facts = raw_models[3]["payload"]
        field = facts["item_metadata"]["qmd_weighting"]
        field["raw_value"] = "2.000000000"
        field["occurrences"][0]["raw_value"] = "2.000000000"
        named_node(field["occurrences"][0]["node"], "fieldentry")[
            "raw_text"
        ] = "2.000000000"

    assert_atomic_refusal(candidate, evidence)


def test_fact_record_order_and_unrelated_raw_records_are_not_pairing_authority() -> None:
    candidate, evidence = two_occurrence_question()
    raw_models = candidate["type_payload"]["raw_response_models"]
    raw_models[:] = [
        raw_models[0],
        raw_models[2],
        {
            "source_kind": "unrelated-observation",
            "payload": {"safe": True},
            "source_evidence_keys": [],
            "extensions": {},
        },
        raw_models[3],
        raw_models[1],
    ]

    assert _enrich_short_answer_question(candidate, evidence) is True
    assert candidate["scoring"]["mode"] == "exact"


def jq_style_hash(value: object) -> str:
    rendered = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        indent=2,
        separators=(",", ": "),
    )
    return hashlib.sha256(f"{rendered}\n".encode("utf-8")).hexdigest()


def test_mixed_fixture_remains_the_pinned_safe_refusal_oracle() -> None:
    fixture = (
        REPO_ROOT
        / "tests"
        / "fixtures"
        / "quiz_xml"
        / "mixed_inline_itemref_and_root_bank"
    )
    payload = build_payload(fixture)
    model, _receipt = build_normalized_model(
        payload,
        fixture,
        fixture,
        "directory",
    )
    payload_without_facts = deepcopy(payload)
    for row in payload_without_facts["quiz_question_rows"]:
        row.pop("source_response_facts", None)
    model_without_facts, _receipt = build_normalized_model(
        payload_without_facts,
        fixture,
        fixture,
        "directory",
    )

    short_answer = next(
        question for question in model["questions"] if question["kind"] == "short_answer"
    )
    short_answer_without_facts = next(
        question
        for question in model_without_facts["questions"]
        if question["kind"] == "short_answer"
    )
    fact_records = [
        record
        for record in short_answer["type_payload"]["raw_response_models"]
        if record["source_kind"] == "d2l_qti_response_facts/0"
    ]
    stripped_fact_records = deepcopy(fact_records)
    for record in stripped_fact_records:
        for key in ("item_metadata", "presentation", "response_processing"):
            record["payload"].pop(key)
    assert jq_style_hash(stripped_fact_records) == (
        "24e99c1170aae6fd8680d4fce6b4e07e9b635790bbd05f836b2aac16c2571b24"
    )

    fact_stripped_question = deepcopy(short_answer)
    fact_stripped_question["type_payload"]["raw_response_models"] = [
        record
        for record in fact_stripped_question["type_payload"]["raw_response_models"]
        if record["source_kind"] != "d2l_qti_response_facts/0"
    ]
    assert fact_stripped_question == short_answer_without_facts
    assert short_answer["scoring"]["mode"] == "unknown"
    assert short_answer["type_payload"]["accepted_responses"] == [
        {
            "value": "accepted",
            "case_sensitive": None,
            "weight": None,
            "extensions": {},
        }
    ]
    assert short_answer["type_payload"]["blanks"] == []

    projection = [
        {
            "entity_key": short_answer["entity_key"],
            "scoring": short_answer["scoring"],
            "type_payload": {
                "accepted_responses": short_answer["type_payload"][
                    "accepted_responses"
                ],
                "blanks": short_answer["type_payload"]["blanks"],
            },
        }
    ]
    # Repinned by the Unbind fidelity correction. The refusal-bearing fields are
    # asserted individually above and are byte-unchanged; only the entity key --
    # now minted from the item's own identity rather than the matched library
    # record -- and the new placement-scoped scoring extensions moved.
    assert jq_style_hash(projection) == (
        "936084dd6c02920bad230a5c063d4168e69f130d47846319a1797f4b70baac6d"
    )
    assert short_answer["identity"] == short_answer_without_facts["identity"]
    assert model["relationships"] == model_without_facts["relationships"]
    assert [
        question["build_support"] for question in model["questions"]
    ] == [
        question["build_support"] for question in model_without_facts["questions"]
    ]


def test_reviewer_and_reingest_keep_legacy_answer_join_unchanged(
    tmp_path: Path,
) -> None:
    payload, model = build_fixture()
    reviewer_rows = build_reviewer_question_rows(payload)
    source_row = payload["quiz_question_rows"][0]
    assert reviewer_rows[0]["answer_key"] == source_row["correct_answer"]
    assert reviewer_rows[0]["answer_key"] == (
        "[ALPHA] || Crème: β || ratio/term || A & B || left / right?!"
    )
    assert not str(reviewer_rows[0]["answer_key"]).startswith("Blank 1:")

    model_path = tmp_path / "short-answer.model.json"
    baseline = tmp_path / "baseline.xlsx"
    edited = tmp_path / "edited.xlsx"
    model_path.write_text(
        json.dumps(model, ensure_ascii=False, sort_keys=True),
        encoding="utf-8",
    )
    write_reviewer_workbook(
        baseline,
        payload,
        FIXTURE,
        "2026-07-24T00:00:00Z",
        REPO_ROOT / "scripts" / "extract_quiz_pool_review.py",
        model,
    )
    edited.write_bytes(baseline.read_bytes())
    workbook = load_workbook(edited)
    sheet = workbook["Quiz Questions"]
    headers = {
        str(cell.value): index
        for index, cell in enumerate(sheet[1], start=1)
    }
    edits = {
        "revised_answer_key": "alpha || beta",
        "revision_reason": "Synthetic legacy-join regression",
        "proposed_by": "Test Reviewer",
        "proposed_at": "2026-07-24T12:00:00Z",
        "approval_status": "open",
    }
    for header, value in edits.items():
        sheet.cell(row=2, column=headers[header], value=value)
    workbook.save(edited)
    before = model_path.read_bytes()

    overlay = materialize_workbook_decisions(model_path, baseline, edited)

    proposal = next(
        annotation
        for annotation in overlay["annotations"]
        if annotation["field_path"] == "/type_payload/answer_key"
    )
    assert proposal["value"] == "alpha || beta"
    assert model_path.read_bytes() == before


def test_promotion_excludes_enriched_short_answer_without_metadata_loss(
    tmp_path: Path,
) -> None:
    _payload, model = build_fixture()
    question = next(
        row
        for row in model["questions"]
        if alias_value(row, "d2l.qmd_globalid") == "fixture-sa-agree-global"
    )
    model_path = tmp_path / "source.model.json"
    overlay_path = tmp_path / "overlay.json"
    model_path.write_text(
        json.dumps(model, ensure_ascii=False, sort_keys=True),
        encoding="utf-8",
    )
    before = model_path.read_bytes()
    annotation = {
        "annotation_id": "ann.short-answer.excluded",
        "target_entity_key": question["entity_key"],
        "field_path": "/prompt/content",
        "kind": "approved_change",
        "value": "PROMOTION_MUST_NOT_APPLY",
        "actor": "Synthetic Approver",
        "timestamp": "2026-07-24T12:00:00Z",
        "source_evidence_keys": [],
        "status": "accepted",
        "extensions": {
            "coursecraft.binder.proposal_id": "proposal.short-answer.excluded"
        },
    }
    overlay_path.write_text(
        json.dumps(
            {
                "format": OVERLAY_FORMAT,
                "source_model": {
                    "model_id": model["model_id"],
                    "sha256": hashlib.sha256(before).hexdigest(),
                    "source_fingerprint": model["source"]["fingerprint"]["digest"],
                },
                "workbooks": {
                    "baseline_sha256": "b" * 64,
                    "edited_sha256": "c" * 64,
                },
                "materialized_view_fingerprint": "d" * 64,
                "row_diffs": [],
                "annotations": [annotation],
                "settings_inputs": [],
                "content_minimized_summary": {
                    "changed_row_count": 1,
                    "annotation_count": 1,
                    "accepted_change_count": 1,
                },
            },
            sort_keys=True,
        ),
        encoding="utf-8",
    )

    promoted, receipt = promote_revisions(model_path, overlay_path, REGISTRY)

    promoted_question = next(
        row for row in promoted["questions"] if row["entity_key"] == question["entity_key"]
    )
    assert receipt["excluded"] == [
        {
            "annotation_id": annotation["annotation_id"],
            "target_entity_key": question["entity_key"],
            "field_path": "/prompt/content",
            "reason": "question_kind_extraction_only",
        }
    ]
    for key in ("type_payload", "scoring", "build_support"):
        assert promoted_question[key] == question[key]
    assert promoted_question["prompt"] == question["prompt"]
    assert model_path.read_bytes() == before


def test_builder_and_readiness_reject_enriched_short_answer_without_mutation(
    tmp_path: Path,
) -> None:
    _payload, normalized = build_fixture()
    enriched = next(
        row
        for row in normalized["questions"]
        if alias_value(row, "d2l.qmd_globalid") == "fixture-sa-agree-global"
    )
    authoring = json.loads(
        (AUTHORING_FIXTURE / "buildable_quiz.model.json").read_text(
            encoding="utf-8"
        )
    )
    target = authoring["questions"][0]
    for key in (
        "kind",
        "source_kind",
        "title",
        "prompt",
        "type_payload",
        "scoring",
        "feedback",
        "build_support",
    ):
        target[key] = deepcopy(enriched[key])
    for record in target["type_payload"]["raw_response_models"]:
        record["source_evidence_keys"] = []
    model_path = tmp_path / "enriched-short-answer.model.json"
    model_path.write_text(
        json.dumps(authoring, ensure_ascii=False, sort_keys=True),
        encoding="utf-8",
    )
    before = model_path.read_bytes()

    with pytest.raises(
        ValueError,
        match="kind 'short_answer' is extraction_only, not build-supported",
    ):
        load_authoring_projection(
            model_path,
            settings_path=AUTHORING_FIXTURE / "buildable_quiz.settings.json",
            asset_root=AUTHORING_FIXTURE,
        )
    report = analyze_authoring_readiness(
        model_path,
        settings_path=AUTHORING_FIXTURE / "buildable_quiz.settings.json",
        asset_root=AUTHORING_FIXTURE,
    )

    assert {issue["code"] for issue in report["issues"]} == {
        "question_kind_not_buildable",
        "question_not_approved_for_build",
    }
    assert report["question_capabilities"] == {
        "extraction_only": 1,
        "roundtrip_verified": 3,
    }
    assert target["type_payload"]["accepted_responses"] == enriched[
        "type_payload"
    ]["accepted_responses"]
    assert target["scoring"] == enriched["scoring"]
    assert model_path.read_bytes() == before
