"""Synthetic packet tests: portability, source protection and actual reingest."""
import json
from pathlib import Path
import shutil
import subprocess
import sys

import pytest
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from quiz_assessment_review import BASELINE, WORKING, LAYOUT, prepare_review_packet, validate_packet
from quiz_review_workbook_reingest import ReingestError, materialize_workbook_decisions
from quiz_review_text import readable_html
from test_extract_quiz_pool_review import build_image_export


@pytest.fixture(scope="module")
def template(tmp_path_factory):
    root = tmp_path_factory.mktemp("assessment-review")
    export = build_image_export(root)
    output = root / "extracted"
    result = subprocess.run([sys.executable, str(ROOT / "scripts/extract_quiz_pool_review.py"),
        str(export), "--output-dir", str(output), "--asset-mode", "copy"], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    packet = root / "packet"
    prepare_review_packet(next(output.glob("*reviewer.xlsx")), next(output.glob("*.model.json")), packet)
    return packet


@pytest.fixture
def packet(template, tmp_path):
    dest = tmp_path / "moved packet with spaces"
    shutil.copytree(template, dest)
    return dest


def collect(packet, policy="optional"):
    return materialize_workbook_decisions(packet / "model.json", packet / BASELINE,
                                         packet / WORKING, metadata_policy=policy)


def edit_row(packet):
    spec = validate_packet(packet)
    book = load_workbook(packet / WORKING)
    tab = spec["sheets"][0]
    sheet = book[tab["title"]]
    columns = {c.value: c.column for c in sheet[2]}
    return book, sheet, columns, tab["rows"][0]["row"]


def test_relocated_packet_links_and_zero_change_import(packet):
    spec = validate_packet(packet)
    assert len(spec["assets"]) == 2  # occurrence and library-only diagrams
    for filename in (WORKING, BASELINE, "native_source_DO_NOT_EDIT.xlsx"):
        book = load_workbook(packet / filename)
        for sheet in book:
            for row in sheet:
                for cell in row:
                    if cell.hyperlink and cell.hyperlink.target and not cell.hyperlink.target.startswith("https:"):
                        assert (packet / cell.hyperlink.target).is_file()
    overlay = collect(packet)
    assert overlay["content_minimized_summary"]["changed_row_count"] == 0
    assert overlay["assessment_review"]["question_count"] == spec["question_count"]


@pytest.mark.parametrize("status,accepted", [(None, 0), ("open", 0), ("accepted", 1), ("rejected", 0)])
def test_grouped_revision_reaches_native_importer(packet, status, accepted):
    book, sheet, cols, row = edit_row(packet)
    sheet.cell(row, cols["proposed_permanent_code"], "SYNTHETIC-REVIEW-001")
    sheet.cell(row, cols["approval_status"], status)
    book.save(packet / WORKING)
    result = collect(packet)
    assert result["content_minimized_summary"]["changed_row_count"] == 1
    assert result["content_minimized_summary"]["accepted_change_count"] == accepted
    assert all(a["actor"] is None for a in result["annotations"])


@pytest.mark.parametrize("damage", ["question_text", "occurrence_key", "group_heading", "native_snapshot", "markup", "formula", "missing_row", "extra_tab"])
def test_invalid_review_changes_are_rejected(packet, damage):
    book, sheet, cols, row = edit_row(packet)
    if damage in {"question_text", "occurrence_key"}:
        sheet.cell(row, cols[damage], "changed source")
    elif damage == "group_heading":
        sheet.cell(row-1, cols["reviewer_note"], "stray review")
    elif damage == "native_snapshot":
        book["Quiz Questions"].cell(2, 2, 999)
    elif damage == "markup":
        sheet.cell(row, cols["response_options_markup"], "lost raw content")
    elif damage == "formula":
        sheet.cell(row, cols["reviewer_note"], "=1+1")
    elif damage == "missing_row":
        sheet.delete_rows(row)
    else:
        book.create_sheet("Unexpected")
    book.save(packet / WORKING)
    with pytest.raises(ReingestError):
        collect(packet)


def test_hidden_layout_and_review_notes_survive_collection(packet):
    book, sheet, cols, row = edit_row(packet)
    sheet.row_dimensions[row].hidden = True
    sheet.column_dimensions["C"].hidden = True
    sheet.cell(row, cols["reviewer_note"], "Review this diagram")
    book.save(packet / WORKING)
    before = (packet / WORKING).read_bytes()
    result = collect(packet)
    assert result["content_minimized_summary"]["annotation_count"] == 1
    assert (packet / WORKING).read_bytes() == before


@pytest.mark.parametrize("companion", ["image", "model", "baseline"])
def test_changed_companions_are_rejected(packet, companion):
    spec = validate_packet(packet)
    ref = spec["assets"][0] if companion == "image" else spec[companion]
    (packet / ref["path"]).write_bytes(b"changed")
    with pytest.raises(ReingestError, match="missing or changed"):
        collect(packet)


def test_rich_content_drafts_collect_but_accepted_edits_hold(packet):
    book, sheet, cols, row = edit_row(packet)
    sheet.cell(row, cols["revised_question_text"], "Proposed diagram question")
    book.save(packet / WORKING)
    assert collect(packet)["content_minimized_summary"]["changed_row_count"] == 1
    sheet.cell(row, cols["approval_status"], "accepted")
    book.save(packet / WORKING)
    with pytest.raises(ReingestError, match="rich-content"):
        collect(packet)


def test_native_baseline_cannot_silently_miss_grouped_edits(packet):
    with pytest.raises(ReingestError, match="packet baseline"):
        materialize_workbook_decisions(packet / "model.json", packet / "native_source_DO_NOT_EDIT.xlsx", packet / WORKING)


def test_existing_packet_is_not_overwritten(packet):
    before = (packet / WORKING).read_bytes()
    with pytest.raises(ValueError, match="existing reviewer edits"):
        prepare_review_packet(packet / "native_source_DO_NOT_EDIT.xlsx", packet / "model.json", packet)
    assert (packet / WORKING).read_bytes() == before


@pytest.mark.parametrize("damage", ["native_row", "occurrence_key", "coverage", "hold", "image"])
def test_manifest_cannot_redirect_or_drop_review_inputs(packet, damage):
    path = packet / LAYOUT
    spec = json.loads(path.read_text())
    item = spec["sheets"][0]["rows"][0]
    if damage == "native_row":
        item["native_row"] = 999
    elif damage == "occurrence_key":
        item["occurrence_key"] = "wrong-key"
    elif damage == "coverage":
        spec["sheets"][0]["rows"] = []
    elif damage == "hold":
        item["rich_revision_hold"] = False
    else:
        spec["assets"] = []
    path.write_text(json.dumps(spec))
    with pytest.raises(ReingestError):
        collect(packet)


def test_math_display_preserves_declared_structure_and_hides_annotations():
    raw = '<p>Value <math><semantics><msup><mi>x</mi><mn>2</mn></msup><annotation>{"math":"noise"}</annotation></semantics></math>.</p>'
    assert readable_html(raw) == "Value x²."
    assert readable_html("H<sub>2</sub>O") == "H₂O"
    assert readable_html('<m:math xmlns:m="http://www.w3.org/1998/Math/MathML"><m:msub><m:mi>H</m:mi><m:mn>2</m:mn></m:msub></m:math>') == "H₂"
    assert readable_html("<math><mfrac><mi>a</mi><mi>b</mi></mfrac></math>") == "(a)/(b)"
    with pytest.raises(ValueError, match="Unsupported MathML"):
        readable_html("<math><mtable><mtr /></mtable></math>")


def test_library_only_export_keeps_inventory_without_inventing_assessment(tmp_path):
    source = ROOT / "tests/fixtures/quiz_xml/answer_key_types"
    output = tmp_path / "extracted"
    result = subprocess.run([sys.executable, str(ROOT / "scripts/extract_quiz_pool_review.py"),
        str(source), "--output-dir", str(output), "--asset-mode", "copy"], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    packet = tmp_path / "packet"
    spec = prepare_review_packet(next(output.glob("*reviewer.xlsx")), next(output.glob("*.model.json")), packet)
    assert spec["profile"] == "inventory_only"
    assert spec["sheets"] == []
    assert spec["inventory_question_count"] > 0
    assert validate_packet(packet)["question_count"] == 0
