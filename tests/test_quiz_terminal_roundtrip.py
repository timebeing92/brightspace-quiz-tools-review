"""Whole terminal workflow using synthetic content; no network or tenant access."""
from copy import deepcopy
from datetime import datetime, timedelta, timezone
import hashlib
import json
from pathlib import Path
import struct
import sys
import zlib

from openpyxl import load_workbook
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
import build_quiz_package_from_workbook as builder
import quiz_terminal as terminal
from quiz_phase5_authorization import create_candidate_authorization
from quiz_assessment_review import validate_packet
from quiz_normalization import source_choice_projection


def png(rgb):
    """Small valid RGB test image, generated without external image dependencies."""
    def chunk(kind, data):
        return struct.pack('>I', len(data)) + kind + data + struct.pack('>I', zlib.crc32(kind + data))
    return (b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>2I5B', 64, 48, 8, 2, 0, 0, 0))
            + chunk(b'IDAT', zlib.compress((b'\0' + bytes(rgb) * 64) * 48)) + chunk(b'IEND', b''))


def make_source(root, title='CC QA 20260926 A Baseline'):
    root.mkdir(parents=True, exist_ok=True)
    fixture = ROOT / 'tests/fixtures/quiz_authoring'
    model = json.loads((fixture / 'buildable_quiz.model.json').read_text())
    model['quizzes'][0]['title'] = title
    plain = deepcopy(model['questions'][0])
    plain['entity_key'] = 'cc:question:fixture:plain-edit'
    plain['identity']['permanent_code'] = 'CCQA-PLAIN'
    plain['title'] = 'CCQA Plain Choice'
    plain['prompt'] = {'format': 'plain_text', 'content': 'Which token is blue?', 'extensions': {}}
    plain['type_payload']['options'][0]['content']['content'] = 'Blue token'
    plain['type_payload']['options'][1]['content']['content'] = 'Red token'
    model['questions'].append(plain)
    membership = deepcopy(next(r for r in model['relationships'] if r['kind'] == 'member_of'))
    membership.update(relationship_key='cc:relationship:fixture:plain-pool', from_entity_key=plain['entity_key'], ordinal=5)
    model['relationships'].append(membership)
    for s in model['structures']:
        s['selection']['available_count'] = 5
        if s['kind'] == 'draw': s['selection']['requested_count'] = 5
    for q in model['questions']:
        if q is not plain: q['identity']['permanent_code'] = 'CCQA-' + q['identity']['permanent_code']
    image = root / 'original.png'
    image.write_bytes(png((30, 100, 220)))
    asset = model['assets'][0]
    asset.update(source_path='original.png', package_path='quiz-assets/original.png', media_type='image/png')
    asset['fingerprint']['digest'] = hashlib.sha256(image.read_bytes()).hexdigest()
    model['questions'][0]['prompt']['content'] = '<p>Look at the blue test rectangle.</p><img src="quiz-assets/original.png" alt="Synthetic test rectangle"><p>Which option is correct?</p>'
    model_path = root / 'baseline.model.json'
    terminal.write_json(model_path, model)
    settings = json.loads((fixture / 'buildable_quiz.settings.json').read_text())
    for row in settings['inputs']:
        if row['setting'] == 'time_limit': row['value'] = 12
        if row['setting'] == 'is_forward_only': row['value'] = False
    for row in settings['resolutions']:
        if row['setting'] == 'time_limit': row['effective_value'] = 12
        if row['setting'] == 'is_forward_only': row['effective_value'] = False
    settings_path = root / 'baseline.settings.json'
    terminal.write_json(settings_path, settings)
    builder.build_package(builder.parse_args([str(model_path), '--settings', str(settings_path),
        '--asset-root', str(root), '--output-dir', str(root / 'package'),
        '--zip-output', str(root / 'baseline.zip')]))
    return root / 'baseline.zip'


def edit_workbook(workspace):
    packet = workspace / 'compose'
    spec = validate_packet(packet)
    model = json.loads((packet / 'model.json').read_text())
    plain_key = next(q['entity_key'] for q in model['questions'] if q['title'] == 'CCQA Plain Choice')
    book = load_workbook(packet / 'reviewer_working.xlsx')
    edited_count = 0
    code_index = 0
    for info in spec['sheets']:
        sheet = book[info['title']]
        cols = {c.value: c.column for c in sheet[2]}
        for entry in info['rows']:
            row = entry['row']
            if sheet.cell(row, cols['question_type']).value == 'Multi-Select':
                sheet.cell(row, cols['proposed_scoring_mode'], 'all_or_nothing')
            code_index += 1
            sheet.cell(row, cols['proposed_permanent_code'], f'CCQA-REVIEW-{code_index:03d}')
            sheet.cell(row, cols['approval_status'], 'accepted')
            if sheet.cell(row, cols['question_entity_key']).value == plain_key:
                for name, value in {
                    'revised_question_text': 'Which revised token is green?',
                    'revised_response_options': '["Red revised token", "Green revised token"]',
                    'revised_answer_key': '{"correct_option_keys": ["B"]}',
                    'approval_status': 'accepted',
                }.items(): sheet.cell(row, cols[name], value)
                edited_count += 1
    assert edited_count == 1
    image_row = spec['image_replacements']['rows'][0]
    sheet = book['Image Replacements']; cols = {c.value: c.column for c in sheet[1]}
    replacement = packet / 'Replacement Images/green.png'
    replacement.parent.mkdir(exist_ok=True)
    replacement.write_bytes(png((30, 190, 80)))
    sheet.cell(image_row['row'], cols['replacement_file'], 'Replacement Images/green.png')
    sheet.cell(image_row['row'], cols['approval_status'], 'accepted')
    pool = spec['question_additions']['target_pools'][0]
    sheet = book['New Questions']; cols = {c.value: c.column for c in sheet[1]}
    values = {'question_code': 'CCQA-NEW', 'question_type': 'Multiple Choice',
              'question_text': 'Which new token is yellow?', 'points': 1,
              'target_quiz_entity_key': pool['quiz_key'], 'target_pool_entity_key': pool['pool_key'],
              'content_format': 'plain_text', 'approval_status': 'accepted'}
    for name, value in values.items(): sheet.cell(2, cols[name], value)
    sheet = book['New Responses']; cols = {c.value: c.column for c in sheet[1]}
    for row, (key, text, correct) in enumerate([('A', 'Yellow new token', True), ('B', 'Purple new token', False)], 2):
        for name, value in dict(question_code='CCQA-NEW', response_key=key, role='option', text=text, content_format='plain_text', correct=correct).items():
            sheet.cell(row, cols[name], value)
    sheet = book['Quiz Settings']; cols = {c.value: c.column for c in sheet[1]}
    settings = {'time_limit': '17', 'attempts_allowed': '3', 'is_active': 'false', 'enforce_time_limit': 'true'}
    found = set()
    for row in range(2, sheet.max_row + 1):
        setting = sheet.cell(row, cols['setting']).value
        if setting in settings:
            sheet.cell(row, cols['proposed_value'], settings[setting])
            sheet.cell(row, cols['approval_status'], 'accepted')
            found.add(setting)
    assert found == set(settings)
    book.save(packet / 'reviewer_working.xlsx')
    return pool['quiz_key']


def compose_candidate(workspace, quiz_key, approved_by='Synthetic regression test'):
    terminal.compose_workspace(workspace, quiz_entity_keys=[quiz_key])
    generated = workspace / 'compose/generated'
    model_path = generated / 'promoted.model.json'
    receipt_path = generated / 'promotion.receipt.json'
    model = json.loads(model_path.read_text()); receipt = json.loads(receipt_path.read_text())
    assert receipt['summary']['applied_change_count'] == 11, receipt
    assert not receipt['excluded']
    now = datetime.now(timezone.utc)
    auth = create_candidate_authorization(model_path=model_path, promotion_receipt_path=receipt_path,
        registry_path=terminal.REGISTRY, quiz_entity_key=quiz_key,
        question_entity_keys=[q['entity_key'] for q in model['questions']],
        approved_by=approved_by, approved_at=now.isoformat(), valid_until=(now+timedelta(hours=24)).isoformat(),
        activation_base_commit=receipt['activation_base_commit'])
    auth_path = workspace / 'candidate.authorization.json'
    terminal.write_json(auth_path, auth)
    state = terminal.compose_workspace(workspace, quiz_entity_keys=[quiz_key], phase5_candidate_authorization=auth_path)
    assert state['compose']['ready_count'] == 1, state['compose']['results']
    assert len(state['compose']['asset_roots']) == 2
    return terminal.rebind_workspace(workspace, quiz_entity_key=quiz_key)


def assert_returned(model):
    assert len(model['questions']) == 6
    plain = next(q for q in model['questions'] if q['title'] == 'CCQA Plain Choice')
    assert 'Which revised token is green?' in plain['prompt']['content']
    options = plain['type_payload']['options']
    assert [o['correct'] for o in options] == [False, True]
    assert 'Green revised token' in options[1]['content']['content']
    expected_keys = {
        'Read a concept map': [True, False],
        'Contracts preserve identity': [True, False],
        'Identify contract features': [True, True, False],
        'CCQA-NEW': [True, False],
    }
    for title, expected in expected_keys.items():
        question = next(q for q in model['questions'] if q['title'] == title)
        assert [o['correct'] for o in question['type_payload']['options']] == expected, title
    multi = next(q for q in model['questions'] if q['title'] == 'Identify contract features')
    # Extraction retains observed grading in source facts. The authoring-mode
    # decision remains explicit on the next review, rather than auto-approved.
    facts = [r['payload'] for r in multi['type_payload']['raw_response_models']
             if r['source_kind'] == 'd2l_qti_response_facts/0']
    assert facts
    assert all(source_choice_projection(f)['scoring_mode'] == 'all_or_nothing' for f in facts)
    digest = hashlib.sha256(png((30, 190, 80))).hexdigest()
    assert any(a['fingerprint']['digest'] == digest for a in model['assets'])


def test_terminal_workbook_edits_rebind_and_reextract(tmp_path):
    source = make_source(tmp_path / 'source')
    workspace = tmp_path / 'review'
    terminal.unbind_workspace(source, workspace)
    quiz_key = edit_workbook(workspace)
    compose_candidate(workspace, quiz_key)
    zip_path = next((workspace / 'rebind').glob('*/quiz-import.zip'))
    returned = tmp_path / 'returned'
    terminal.unbind_workspace(zip_path, returned)
    assert_returned(json.loads((returned / 'compose/model.json').read_text()))
    # Source image and protected baseline survive the entire workflow.
    assert (tmp_path / 'source/original.png').read_bytes() == png((30, 100, 220))
    terminal.verify_workspace(*terminal.load_state(workspace))


def test_terminal_refuses_excluded_accepted_revision(tmp_path, monkeypatch):
    source = make_source(tmp_path / 'source')
    workspace = tmp_path / 'review'
    state = terminal.unbind_workspace(source, workspace)
    model = json.loads((workspace / 'compose/model.json').read_text())
    original = terminal.promote_revisions
    def excluded(*args, **kwargs):
        model, receipt = original(*args, **kwargs)
        receipt['excluded'] = [{'reason': 'synthetic unsupported revision'}]
        return model, receipt
    monkeypatch.setattr(terminal, 'promote_revisions', excluded)
    with pytest.raises(terminal.WorkflowError, match='could not be applied'):
        terminal.compose_workspace(workspace, quiz_entity_keys=[model['quizzes'][0]['entity_key']])
    assert terminal.load_state(workspace)[1]['compose'] is None
