from __future__ import annotations

import json
import subprocess
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook


REPO_ROOT = Path(__file__).resolve().parents[1]
FIXTURE_ROOT = REPO_ROOT / "tests" / "fixtures" / "quiz_xml"


def run_script(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, *args],
        cwd=REPO_ROOT,
        text=True,
        capture_output=True,
        check=False,
    )


def local_name(tag: str) -> str:
    return tag.split("}", 1)[-1] if "}" in tag else tag


def first_item(path: Path) -> ET.Element:
    root = ET.parse(path).getroot()
    for elem in root.iter():
        if local_name(elem.tag) == "item":
            return elem
    raise AssertionError(f"No <item> found in {path}")


def first_assessment_item(path: Path) -> ET.Element:
    return first_item(path)


def render_fixture_payload(fixture_name: str, tmp_path: Path) -> dict[str, object]:
    fixture_dir = FIXTURE_ROOT / fixture_name
    output_dir = tmp_path / "review"
    result = run_script(
        "scripts/extract_quiz_pool_review.py",
        str(fixture_dir),
        "--output-dir",
        str(output_dir),
    )
    assert result.returncode == 0, result.stderr
    return json.loads((output_dir / f"{fixture_name}__quiz_pool_review.json").read_text(encoding="utf-8"))


def test_matching_response_groups_fixture_preserves_group_structure_and_matches_pool(tmp_path: Path) -> None:
    quiz_item = first_assessment_item(FIXTURE_ROOT / "matching_response_groups" / "quiz_d2l_5166.xml")
    response_groups = [elem for elem in quiz_item.iter() if local_name(elem.tag) == "response_grp"]
    assert len(response_groups) == 5
    assert all(group.attrib.get("rcardinality") == "Single" for group in response_groups)

    payload = render_fixture_payload("matching_response_groups", tmp_path)
    row = payload["quiz_question_rows"][0]
    assert row["question_type"] == "Matching"
    assert row["evidence_level"] == "inferred_exact"
    assert row["pool_title"] == "Final Exam Obj 5"
    assert row["correct_answer_basis"] == "matching_pairs"
    assert row["matching_prompts_text"] == (
        "shape one || shape two || shape three || shape four || shape five"
    )
    assert row["matching_options_text"] == (
        "Circle, a round shape || "
        "Square, a four sided shape || "
        "Triangle, a three sided shape || "
        "Pentagon, a five sided shape || "
        "Hexagon, a six sided shape"
    )
    assert "shape one -> Circle, a round shape" in row["correct_answer"]
    assert "shape five -> Hexagon, a six sided shape" in row["correct_answer"]

    reviewer_workbook = load_workbook(tmp_path / "review" / "matching_response_groups__quiz_pool_review_reviewer.xlsx", read_only=True)
    reviewer_sheet = reviewer_workbook["Quiz Questions"]
    headers = [cell.value for cell in reviewer_sheet[1]]
    reviewer_row = dict(zip(headers, next(reviewer_sheet.iter_rows(min_row=2, max_row=2, values_only=True))))
    assert reviewer_row["response_options"] == (
        "Prompt set:\n"
        "- shape one\n"
        "- shape two\n"
        "- shape three\n"
        "- shape four\n"
        "- shape five\n\n"
        "Match options:\n"
        "- Circle, a round shape\n"
        "- Square, a four sided shape\n"
        "- Triangle, a three sided shape\n"
        "- Pentagon, a five sided shape\n"
        "- Hexagon, a six sided shape"
    )
    assert reviewer_row["answer_key"] == (
        "shape one -> Circle, a round shape\n"
        "shape two -> Square, a four sided shape\n"
        "shape three -> Triangle, a three sided shape\n"
        "shape four -> Pentagon, a five sided shape\n"
        "shape five -> Hexagon, a six sided shape"
    )


def test_ordering_response_group_fixture_extracts_ordered_sequence(tmp_path: Path) -> None:
    quiz_item = first_assessment_item(FIXTURE_ROOT / "ordering_response_group" / "quiz_d2l_5166.xml")
    response_groups = [elem for elem in quiz_item.iter() if local_name(elem.tag) == "response_grp"]
    assert len(response_groups) == 1
    assert response_groups[0].attrib.get("rcardinality") == "Ordered"

    payload = render_fixture_payload("ordering_response_group", tmp_path)
    row = payload["quiz_question_rows"][0]
    assert row["question_type"] == "Ordering"
    assert row["evidence_level"] == "inferred_exact"
    assert row["pool_title"] == "Essay Pool 3"
    assert row["correct_answer_basis"] == "ordering_sequence"
    assert row["correct_answer"] == (
        "Open the empty box || "
        "Place a blue token in the box || "
        "Close the box lid || "
        "Label the closed box"
    )

    reviewer_workbook = load_workbook(tmp_path / "review" / "ordering_response_group__quiz_pool_review_reviewer.xlsx", read_only=True)
    reviewer_sheet = reviewer_workbook["Quiz Questions"]
    headers = [cell.value for cell in reviewer_sheet[1]]
    reviewer_row = dict(zip(headers, next(reviewer_sheet.iter_rows(min_row=2, max_row=2, values_only=True))))
    assert reviewer_row["response_options"] == (
        "Items to order:\n"
        "- Open the empty box\n"
        "- Place a blue token in the box\n"
        "- Close the box lid\n"
        "- Label the closed box"
    )
    assert reviewer_row["answer_key"] == (
        "1. Open the empty box\n"
        "2. Place a blue token in the box\n"
        "3. Close the box lid\n"
        "4. Label the closed box"
    )


def test_fill_in_the_blanks_multiblank_fixture_preserves_two_blanks_and_quiz_local_status(tmp_path: Path) -> None:
    quiz_item = first_assessment_item(FIXTURE_ROOT / "fill_in_the_blanks_multiblank" / "quiz_d2l_5156.xml")
    response_strs = [elem for elem in quiz_item.iter() if local_name(elem.tag) == "response_str"]
    assert len(response_strs) == 2

    payload = render_fixture_payload("fill_in_the_blanks_multiblank", tmp_path)
    assert payload["summary"]["pool_count"] == 0
    row = payload["quiz_question_rows"][0]
    assert row["question_type"] == "Fill in the Blanks"
    assert row["evidence_level"] == "unmatched"
    assert row["correct_answer_basis"] == "text_entry"
    assert row["correct_answer"] == "10 || K"
    assert "Number: [Blank 1] Letter: [Blank 2]" in row["question_text"]
    assert row["fill_in_blank_count"] == 2
    assert row["fill_in_blank_answers_text"] == "Blank 1: 10 || Blank 2: K"

    reviewer_workbook = load_workbook(tmp_path / "review" / "fill_in_the_blanks_multiblank__quiz_pool_review_reviewer.xlsx", read_only=True)
    reviewer_sheet = reviewer_workbook["Quiz Questions"]
    headers = [cell.value for cell in reviewer_sheet[1]]
    reviewer_row = dict(zip(headers, next(reviewer_sheet.iter_rows(min_row=2, max_row=2, values_only=True))))
    assert "Number: [Blank 1] Letter: [Blank 2]" in reviewer_row["question_text"]
    assert reviewer_row["answer_key"] == "Blank 1: 10\nBlank 2: K"


def test_long_answer_with_answer_key_fixture_preserves_answer_key_and_surfaces_it(tmp_path: Path) -> None:
    library_item = first_item(FIXTURE_ROOT / "long_answer_with_answer_key" / "questiondb.xml")
    answer_keys = [elem for elem in library_item.iter() if local_name(elem.tag) == "answer_key"]
    assert len(answer_keys) == 1

    payload = render_fixture_payload("long_answer_with_answer_key", tmp_path)
    row = payload["quiz_question_rows"][0]
    assert row["question_type"] == "Long Answer"
    assert row["evidence_level"] == "inferred_exact"
    assert row["pool_title"] == "Essay Pool 1"
    assert row["correct_answer_basis"] == "answer_key_material"
    assert "A synthetic answer describes a blue token inside a closed box." in row["correct_answer"]

    output_dir = tmp_path / "review"
    model = json.loads(
        (
            output_dir
            / "long_answer_with_answer_key__quiz_pool_review.model.json"
        ).read_text(encoding="utf-8")
    )
    placed = [
        question
        for question in model["questions"]
        if not (question.get("extensions") or {}).get("library_only")
    ]
    assert len(placed) == 1
    manual = placed[0]["type_payload"]["manual_answer_key"]
    assert "A synthetic answer describes" in manual["content"]

    reviewer = load_workbook(
        output_dir
        / "long_answer_with_answer_key__quiz_pool_review_reviewer.xlsx",
        read_only=True,
    )
    entity_sheet = reviewer["Question Entities"]
    headers = [cell.value for cell in entity_sheet[1]]
    entity_row = dict(
        zip(
            headers,
            next(entity_sheet.iter_rows(min_row=2, max_row=2, values_only=True)),
        )
    )
    assert "A synthetic answer describes" in entity_row["answer_key"]


def test_mixed_inline_itemref_fixture_preserves_occurrences_and_root_bank(tmp_path: Path) -> None:
    payload = render_fixture_payload("mixed_inline_itemref_and_root_bank", tmp_path)

    assert payload["summary"]["quiz_count"] == 1
    assert payload["summary"]["quiz_question_count"] == 3
    assert payload["summary"]["source_evidence_count"] == 2
    assert payload["summary"]["unmatched_count"] == 1

    rows = {row["quiz_item_label"]: row for row in payload["quiz_question_rows"]}
    assert rows["ROOT_Q"]["evidence_level"] == "source_evidence"
    assert rows["ROOT_Q"]["pool_title"] == "Question Library Root Items"
    assert rows["SECTION_Q"]["evidence_level"] == "source_evidence"
    assert rows["SECTION_Q"]["pool_title"] == "Section Bank"
    assert rows["SECTION_Q"]["draw_count"] == "1"
    assert rows["INLINE_Q"]["evidence_level"] == "unmatched"
    assert rows["INLINE_Q"]["question_type"] == "True/False"
