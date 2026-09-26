from __future__ import annotations

import json
from pathlib import Path
import shutil
import subprocess
import sys

from openpyxl import Workbook, load_workbook
import pytest


REPO_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO_ROOT / "scripts"))

from quiz_contracts import validate_contract  # noqa: E402
from quiz_review_workbook_reingest import (  # noqa: E402
    ReingestError,
    materialize_workbook_decisions,
)
from quiz_review_projection import build_quiz_review_projection  # noqa: E402


FIXTURE = REPO_ROOT / "tests" / "fixtures" / "quiz_binder" / "golden_course.model.json"
HEADERS = [
    "quiz",
    "question_number",
    "question_type",
    "proposed_permanent_code",
    "proposed_scoring_mode",
    "question_text",
    "revised_question_text",
    "response_options",
    "revised_response_options",
    "answer_key",
    "revised_answer_key",
    "question_library_location",
    "revised_question_library_location",
    "library_link_status",
    "revision_reason",
    "proposed_by",
    "proposed_at",
    "approval_status",
    "approved_by",
    "approved_at",
    "reviewer_note",
]
SETTINGS_HEADERS = [
    "quiz",
    "target_entity_key",
    "setting",
    "observed_state",
    "observed_value",
    "proposed_value",
    "revision_reason",
    "proposed_by",
    "proposed_at",
    "approval_status",
    "approved_by",
    "approved_at",
    "reviewer_note",
]
UNRESOLVED_HEADERS = [
    "quiz",
    "question_number",
    "quiz_section (blank = quiz root)",
    "question_type",
    "proposed_permanent_code",
    "proposed_scoring_mode",
    "question_text",
    "revised_question_text",
    "response_options",
    "revised_response_options",
    "answer_key",
    "revised_answer_key",
    "primary_image",
    "why_it_needs_review",
    "revision_reason",
    "proposed_by",
    "proposed_at",
    "approval_status",
    "approved_by",
    "approved_at",
    "reviewer_note",
]


def make_workbook(path: Path) -> None:
    model = json.loads(FIXTURE.read_text(encoding="utf-8"))
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quiz Questions"
    sheet.append(HEADERS)
    for number, question in enumerate(model["questions"], start=1):
        sheet.append(
            [
                "Synthetic Review-State Quiz",
                number,
                question["source_kind"],
                "",
                "",
                question["prompt"]["content"],
                "",
                "source options",
                "",
                "source answer",
                "",
                "Synthetic Pool",
                "",
                question["extensions"]["coursecraft.binder.match_state"],
                "",
                "",
                "",
                "",
                "",
                "",
                "",
            ]
        )
    settings = workbook.create_sheet("Quiz Settings")
    settings.append(SETTINGS_HEADERS)
    settings.append(
        [
            "Synthetic Review-State Quiz",
            "cc:quiz:fixture:binder-golden",
            "time_limit",
            "unresolved",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
        ]
    )
    unresolved = workbook.create_sheet("No Safe Library Match")
    unresolved.append(UNRESOLVED_HEADERS)
    unresolved.append(
        [
            "Synthetic Review-State Quiz",
            4,
            "Quiz root",
            "Short Answer",
            "",
            "",
            "SENTINEL_STEM_NO_SAFE",
            "",
            "",
            "",
            "SENTINEL_ACCEPTED_RESPONSE",
            "",
            "",
            "No safe question-library location was found automatically.",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
        ]
    )
    workbook.save(path)


def set_row_values(
    path: Path,
    row: int,
    values: dict[str, object],
    *,
    sheet_name: str = "Quiz Questions",
) -> None:
    workbook = load_workbook(path)
    sheet = workbook[sheet_name]
    headers = [cell.value for cell in sheet[1]]
    for header, value in values.items():
        sheet.cell(row=row, column=headers.index(header) + 1, value=value)
    workbook.save(path)


def prepare_files(tmp_path: Path) -> tuple[Path, Path, Path]:
    model = tmp_path / "source.model.json"
    baseline = tmp_path / "baseline.xlsx"
    edited = tmp_path / "edited.xlsx"
    shutil.copyfile(FIXTURE, model)
    make_workbook(baseline)
    shutil.copyfile(baseline, edited)
    return model, baseline, edited


def test_materializes_open_accepted_and_rejected_marginalia_using_existing_annotations(
    tmp_path: Path,
) -> None:
    model_path, baseline, edited = prepare_files(tmp_path)
    shared = {
        "revision_reason": "Synthetic reason",
        "proposed_by": "Reviewer A",
        "proposed_at": "2026-07-20T14:00:00-04:00",
    }
    set_row_values(
        edited,
        2,
        {**shared, "revised_question_text": "Approved prompt", "approval_status": "open"},
    )
    set_row_values(
        edited,
        3,
        {
            **shared,
            "revised_answer_key": '{"correct_option_keys":["F"]}',
            "approval_status": "accepted",
            "approved_by": "Approver B",
            "approved_at": "2026-07-20T15:00:00-04:00",
        },
    )
    set_row_values(
        edited,
        4,
        {
            **shared,
            "revised_question_text": "Rejected prompt",
            "approval_status": "rejected",
            "approved_by": "Approver B",
            "approved_at": "2026-07-20T15:05:00-04:00",
        },
    )

    overlay = materialize_workbook_decisions(model_path, baseline, edited)

    proposals = [row for row in overlay["annotations"] if row["kind"] == "proposed_revision"]
    approvals = [row for row in overlay["annotations"] if row["kind"] == "approved_change"]
    assert {row["status"] for row in proposals} == {"open", "accepted", "rejected"}
    assert len(approvals) == 1
    assert approvals[0]["actor"] == "Approver B"
    assert approvals[0]["extensions"]["coursecraft.binder.proposal_id"]
    assert len(overlay["row_diffs"]) == 3

    model = json.loads(model_path.read_text(encoding="utf-8"))
    model["annotations"].extend(overlay["annotations"])
    assert validate_contract(model, mode="transform") == []


def test_reingest_rejects_any_source_field_edit(tmp_path: Path) -> None:
    model, baseline, edited = prepare_files(tmp_path)
    set_row_values(edited, 2, {"question_text": "Changed extracted source"})

    with pytest.raises(ReingestError, match="Source field edit rejected"):
        materialize_workbook_decisions(model, baseline, edited)


def test_reingest_accepts_unchanged_optional_placeholder_sheet(tmp_path: Path) -> None:
    model, baseline, edited = prepare_files(tmp_path)
    for path in (baseline, edited):
        workbook = load_workbook(path)
        del workbook["No Safe Library Match"]
        sheet = workbook.create_sheet("No Safe Library Match")
        sheet.append(["note"])
        sheet.append(["No unresolved rows in this extraction."])
        workbook.save(path)

    overlay = materialize_workbook_decisions(model, baseline, edited)

    assert overlay["content_minimized_summary"]["changed_row_count"] == 0

    workbook = load_workbook(edited)
    workbook["No Safe Library Match"]["A2"] = "Changed placeholder source copy"
    workbook.save(edited)
    with pytest.raises(ReingestError, match="Source field edit rejected"):
        materialize_workbook_decisions(model, baseline, edited)


def test_accepted_permanent_code_uses_existing_annotation_shape(tmp_path: Path) -> None:
    model, baseline, edited = prepare_files(tmp_path)
    set_row_values(
        edited,
        2,
        {
            "proposed_permanent_code": "P5-TWIN-M-Q001",
            "revision_reason": "Assign an explicit synthetic Phase 5 code",
            "proposed_by": "Reviewer A",
            "proposed_at": "2026-07-20T14:00:00Z",
            "approval_status": "accepted",
            "approved_by": "Approver B",
            "approved_at": "2026-07-20T15:00:00Z",
        },
    )

    overlay = materialize_workbook_decisions(model, baseline, edited)

    proposal = next(
        row
        for row in overlay["annotations"]
        if row["kind"] == "proposed_revision"
        and row["field_path"] == "/identity/permanent_code"
    )
    approval = next(
        row
        for row in overlay["annotations"]
        if row["kind"] == "approved_change"
        and row["field_path"] == "/identity/permanent_code"
    )
    assert proposal["value"] == "P5-TWIN-M-Q001"
    assert approval["extensions"]["coursecraft.binder.proposal_id"] == proposal[
        "annotation_id"
    ]


def test_accepted_scoring_mode_uses_existing_annotation_shape(tmp_path: Path) -> None:
    model, baseline, edited = prepare_files(tmp_path)
    set_row_values(
        edited,
        2,
        {
            "proposed_scoring_mode": "all_or_nothing",
            "revision_reason": "Resolve the synthetic multi-select scoring policy",
            "proposed_by": "Reviewer A",
            "proposed_at": "2026-07-20T14:00:00Z",
            "approval_status": "accepted",
            "approved_by": "Approver B",
            "approved_at": "2026-07-20T15:00:00Z",
        },
    )

    overlay = materialize_workbook_decisions(model, baseline, edited)

    approval = next(
        row
        for row in overlay["annotations"]
        if row["kind"] == "approved_change"
        and row["field_path"] == "/scoring/mode"
    )
    assert approval["value"] == "all_or_nothing"


def test_reingest_requires_separate_proposal_and_approval_attribution(tmp_path: Path) -> None:
    model, baseline, edited = prepare_files(tmp_path)
    set_row_values(
        edited,
        2,
        {
            "revised_question_text": "Approved prompt",
            "revision_reason": "Synthetic reason",
            "proposed_by": "Reviewer A",
            "proposed_at": "2026-07-20T14:00:00Z",
            "approval_status": "accepted",
        },
    )

    with pytest.raises(ReingestError, match="requires approved_by"):
        materialize_workbook_decisions(model, baseline, edited)


def test_reingest_is_deterministic_for_identical_inputs(tmp_path: Path) -> None:
    model, baseline, edited = prepare_files(tmp_path)
    set_row_values(
        edited,
        2,
        {
            "revised_question_text": "Open prompt",
            "revision_reason": "Synthetic reason",
            "proposed_by": "Reviewer A",
            "proposed_at": "2026-07-20T14:00:00Z",
            "approval_status": "open",
        },
    )
    assert materialize_workbook_decisions(model, baseline, edited) == materialize_workbook_decisions(
        model, baseline, edited
    )


@pytest.mark.parametrize("status", ["", "open", "accepted", "rejected"])
def test_optional_metadata_preserves_unknowns_and_explicit_decisions(tmp_path: Path, status: str) -> None:
    model, baseline, edited = prepare_files(tmp_path)
    set_row_values(edited, 2, {"revised_question_text": "Minimal review", "approval_status": status})
    result = materialize_workbook_decisions(model, baseline, edited, metadata_policy="optional")
    assert result == materialize_workbook_decisions(model, baseline, edited, metadata_policy="optional")
    proposals = [a for a in result["annotations"] if a["kind"] == "proposed_revision"]
    assert len(proposals) == 1
    assert proposals[0]["status"] == (status or "open")
    assert proposals[0]["extensions"]["coursecraft.binder.reason"] is None
    assert all(a["actor"] is None and a["timestamp"] is None for a in result["annotations"])
    assert result["content_minimized_summary"]["accepted_change_count"] == (status == "accepted")
    canonical = json.loads(model.read_text())
    canonical["annotations"].extend(result["annotations"])
    assert validate_contract(canonical, mode="transform") == []


def test_optional_metadata_keeps_populated_output_identical(tmp_path: Path) -> None:
    model, baseline, edited = prepare_files(tmp_path)
    set_row_values(edited, 2, {"revised_question_text": "Accepted review", "approval_status": "accepted",
        "revision_reason": "Supplied reason", "proposed_by": "Reviewer", "proposed_at": "2026-09-15T10:00:00-04:00",
        "approved_by": "Approver", "approved_at": "2026-09-15T15:00:00Z"})
    assert materialize_workbook_decisions(model, baseline, edited, metadata_policy="optional") == materialize_workbook_decisions(model, baseline, edited)


@pytest.mark.parametrize("field,value", [
    ("proposed_at", "not a date"), ("proposed_at", "2026-09-15T12:00:00"),
    ("approved_at", "not a date"), ("approved_at", "2026-09-15T12:00:00")])
def test_optional_metadata_still_validates_supplied_timestamps(tmp_path: Path, field: str, value: str) -> None:
    model, baseline, edited = prepare_files(tmp_path)
    set_row_values(edited, 2, {"revised_question_text": "Edited", "approval_status": "accepted", field: value})
    with pytest.raises(ReingestError, match=field):
        materialize_workbook_decisions(model, baseline, edited, metadata_policy="optional")


@pytest.mark.parametrize("status", ["open", "accepted", "rejected"])
def test_optional_settings_metadata_retains_setting_contract(tmp_path: Path, status: str) -> None:
    model, baseline, edited = prepare_files(tmp_path)
    set_row_values(edited, 2, {"proposed_value": "1800", "approval_status": status}, sheet_name="Quiz Settings")
    result = materialize_workbook_decisions(model, baseline, edited, metadata_policy="optional")
    assert len(result["settings_decisions"]) == 1
    decision = result["settings_decisions"][0]
    assert decision["status"] == status and decision["value"] == 1800
    assert all(decision[f] is None for f in ("proposed_by", "proposed_at", "approved_by", "approved_at", "reason"))
    assert len(result["settings_inputs"]) == (status == "accepted")


def test_optional_metadata_cli_and_revision_promotion(tmp_path: Path) -> None:
    from quiz_promote_revisions import promote_revisions
    model, baseline, edited = prepare_files(tmp_path)
    set_row_values(edited, 2, {"revised_question_text": "Minimal accepted prompt", "approval_status": "accepted"})
    output = tmp_path / "decisions.json"
    result = subprocess.run([sys.executable, str(REPO_ROOT / "scripts/quiz_review_workbook_reingest.py"),
        "--model", str(model), "--baseline-workbook", str(baseline), "--edited-workbook", str(edited),
        "--output", str(output), "--metadata-policy", "optional"], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    promoted, receipt = promote_revisions(model, output,
        REPO_ROOT / "workspace/reference/schemas/quiz/quiz_build_capabilities.json")
    assert any(q["prompt"]["content"] == "Minimal accepted prompt" for q in promoted["questions"])
    assert validate_contract(promoted, mode="transform") == []


def test_default_still_requires_metadata_and_invalid_policy_is_rejected(tmp_path: Path) -> None:
    model, baseline, edited = prepare_files(tmp_path)
    set_row_values(edited, 2, {"revised_question_text": "Edited"})
    with pytest.raises(ReingestError, match="requires revision_reason"):
        materialize_workbook_decisions(model, baseline, edited)
    with pytest.raises(ReingestError, match="Unknown review metadata policy"):
        materialize_workbook_decisions(model, baseline, edited, metadata_policy="typo")


@pytest.mark.parametrize("sheet", ["Quiz Questions", "Quiz Settings"])
def test_optional_metadata_cannot_supply_a_decision_without_a_proposal(tmp_path: Path, sheet: str) -> None:
    model, baseline, edited = prepare_files(tmp_path)
    set_row_values(edited, 2, {"approval_status": "accepted"}, sheet_name=sheet)
    with pytest.raises(ReingestError, match="decision but no proposed"):
        materialize_workbook_decisions(model, baseline, edited, metadata_policy="optional")


def test_stable_occurrence_key_replaces_fragile_quiz_and_row_number_identity(
    tmp_path: Path,
) -> None:
    model_path, baseline, edited = prepare_files(tmp_path)
    model = json.loads(model_path.read_text(encoding="utf-8"))
    occurrences = build_quiz_review_projection(model)["question_occurrences"]
    occurrence_by_order = {
        row["quiz_projection_order"]: row["occurrence_key"] for row in occurrences
    }

    for path in (baseline, edited):
        workbook = load_workbook(path)
        questions = workbook["Quiz Questions"]
        questions.insert_cols(1)
        questions["A1"] = "occurrence_key"
        for row_number in range(2, questions.max_row + 1):
            source_number = row_number - 1
            questions.cell(row=row_number, column=1, value=occurrence_by_order[source_number])
            # These legacy display values are deliberately unsuitable as model
            # identity. The stable occurrence key must carry the join instead.
            questions.cell(row=row_number, column=2, value="Changed review label")
            questions.cell(row=row_number, column=3, value=900 + source_number)

        unresolved = workbook["No Safe Library Match"]
        unresolved.insert_cols(1)
        unresolved["A1"] = "occurrence_key"
        unresolved["A2"] = occurrence_by_order[4]
        unresolved["B2"] = "Changed review label"
        unresolved["C2"] = 904
        workbook.save(path)

    set_row_values(
        edited,
        2,
        {
            "revised_question_text": "Stable-key prompt revision",
            "revision_reason": "Exercise occurrence-backed workbook identity",
            "proposed_by": "Reviewer A",
            "proposed_at": "2026-08-03T15:00:00-04:00",
            "approval_status": "open",
        },
    )

    overlay = materialize_workbook_decisions(model_path, baseline, edited)

    assert overlay["row_diffs"][0]["row_key_headers"] == ["occurrence_key"]
    assert overlay["row_diffs"][0]["row_key"] == [occurrence_by_order[1]]
    assert overlay["row_diffs"][0]["target_entity_key"] == (
        "cc:question:fixture:binder-direct"
    )


def test_accepted_setting_decision_materializes_existing_settings_input_shape(
    tmp_path: Path,
) -> None:
    model, baseline, edited = prepare_files(tmp_path)
    set_row_values(
        edited,
        2,
        {
            "proposed_value": 45,
            "revision_reason": "Approved synthetic duration",
            "proposed_by": "Reviewer A",
            "proposed_at": "2026-07-20T14:00:00Z",
            "approval_status": "accepted",
            "approved_by": "Approver B",
            "approved_at": "2026-07-20T15:00:00Z",
        },
        sheet_name="Quiz Settings",
    )

    overlay = materialize_workbook_decisions(model, baseline, edited)

    assert len(overlay["settings_decisions"]) == 1
    assert overlay["settings_decisions"][0]["status"] == "accepted"
    assert len(overlay["settings_inputs"]) == 1
    setting_input = overlay["settings_inputs"][0]
    assert setting_input["setting"] == "time_limit"
    assert setting_input["layer"] == "quiz_decision"
    assert setting_input["target_entity_key"] == "cc:quiz:fixture:binder-golden"
    assert setting_input["state"] == "known"
    assert setting_input["value"] == 45
    assert overlay["content_minimized_summary"]["accepted_setting_input_count"] == 1


def test_reingest_distinguishes_repeated_setting_observations_by_stable_key(
    tmp_path: Path,
) -> None:
    model, baseline, edited = prepare_files(tmp_path)
    for path in (baseline, edited):
        workbook = load_workbook(path)
        settings = workbook["Quiz Settings"]
        settings.insert_cols(3)
        settings["C1"] = "setting_observation_key"
        settings["C2"] = "setting.observation.synthetic.001"
        values = [cell.value for cell in settings[2]]
        values[2] = "setting.observation.synthetic.002"
        settings.append(values)
        workbook.save(path)

    unchanged = materialize_workbook_decisions(model, baseline, edited)
    assert unchanged["settings_row_diffs"] == []

    proposal = {
        "proposed_value": 45,
        "revision_reason": "Resolve one repeated source observation",
        "proposed_by": "Reviewer A",
        "proposed_at": "2026-09-04T14:00:00Z",
        "approval_status": "open",
    }
    set_row_values(edited, 3, proposal, sheet_name="Quiz Settings")
    overlay = materialize_workbook_decisions(model, baseline, edited)
    assert overlay["settings_row_diffs"][0]["row_key_headers"] == [
        "setting_observation_key"
    ]
    assert overlay["settings_row_diffs"][0]["row_key"] == [
        "setting.observation.synthetic.002"
    ]

    set_row_values(edited, 2, proposal, sheet_name="Quiz Settings")
    with pytest.raises(ReingestError, match="more than one proposal"):
        materialize_workbook_decisions(model, baseline, edited)


def test_no_safe_sheet_marginalia_is_ingested_and_duplicate_sheet_edits_fail(
    tmp_path: Path,
) -> None:
    model, baseline, edited = prepare_files(tmp_path)
    marginalia = {
        "revised_question_text": "Review-only no-safe revision",
        "revision_reason": "Synthetic reason",
        "proposed_by": "Reviewer A",
        "proposed_at": "2026-07-20T14:00:00Z",
        "approval_status": "open",
    }
    set_row_values(
        edited,
        2,
        marginalia,
        sheet_name="No Safe Library Match",
    )

    overlay = materialize_workbook_decisions(model, baseline, edited)
    proposal = next(row for row in overlay["annotations"] if row["kind"] == "proposed_revision")
    assert proposal["target_entity_key"] == "cc:question:fixture:binder-no-safe"
    assert proposal["extensions"]["coursecraft.binder.workbook_sheet"] == "No Safe Library Match"

    set_row_values(edited, 5, marginalia, sheet_name="Quiz Questions")
    with pytest.raises(ReingestError, match="more than one question sheet"):
        materialize_workbook_decisions(model, baseline, edited)
